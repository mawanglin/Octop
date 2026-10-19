"""Contract tests for the bundled developer experts (仓库导读 / SQL / Git / 接口测试 / 安全扫描 / 网页工具 / 正则脚本)."""

from __future__ import annotations

import http.server
import importlib.util
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import threading
from collections.abc import Iterator
from pathlib import Path
from types import ModuleType

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
# Captured at import: per-test fixtures swap HOME, which would hide Playwright's browser cache from subprocesses.
ORIGINAL_ENV = dict(os.environ)
REPO_SCAN = LIBRARY / "repo-navigator" / "skills" / "repo-scan" / "scripts" / "repo_scan.py"
SQL = LIBRARY / "sql-expert" / "skills" / "sql-toolkit"
GIT = LIBRARY / "git-helper" / "skills" / "git-report" / "scripts" / "git_report.py"
API = LIBRARY / "api-tester" / "skills" / "api-toolkit"
SECRET = LIBRARY / "code-security" / "skills" / "secret-scan" / "scripts" / "secret_scan.py"
HTML = LIBRARY / "web-tool-builder" / "skills" / "html-check"
REGEX = LIBRARY / "regex-script" / "skills" / "regex-lab" / "scripts" / "regex_test.py"

EXPERTS = {
    "repo-navigator": ("代码仓库导读", {"repo-scan", "onboarding-doc"}),
    "sql-expert": ("SQL 数据库专家", {"sql-toolkit", "sql-optimization"}),
    "git-helper": ("Git 版本管理助手", {"git-report", "git-rescue"}),
    "api-tester": ("接口测试专家", {"api-toolkit", "api-test-design"}),
    "code-security": ("代码安全扫描", {"secret-scan", "secure-coding"}),
    "web-tool-builder": ("网页小工具专家", {"html-check", "web-tool-patterns"}),
    "regex-script": ("正则与脚本助手", {"regex-lab", "script-recipes"}),
}


def _run(
    script: Path, *args: str | Path, cwd: Path | None = None, base_env: dict[str, str] | None = None
) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **(base_env if base_env is not None else os.environ),
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
        cwd=cwd,
    )


def _load(path: Path) -> ModuleType:
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location(f"_dev_{path.stem}", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


# --------------------------------------------------------------------- repo


def _demo_repo(root: Path) -> None:
    (root / "src" / "shop" / "api").mkdir(parents=True)
    (root / "src" / "shop" / "db").mkdir(parents=True)
    (root / "tests").mkdir()
    (root / "web" / "components").mkdir(parents=True)
    (root / "src" / "shop" / "__init__.py").write_text("", encoding="utf-8")
    (root / "src" / "shop" / "api" / "__init__.py").write_text(
        "from shop.db import repo\n", encoding="utf-8"
    )
    (root / "src" / "shop" / "db" / "__init__.py").write_text("", encoding="utf-8")
    (root / "src" / "shop" / "db" / "repo.py").write_text(
        "# 数据访问\nfrom ..api import x  # TODO: 去掉反向依赖\n\n\ndef get():\n    return 1\n",
        encoding="utf-8",
    )
    (root / "src" / "shop" / "main.py").write_text(
        "from shop.api import app\nimport fastapi\n", encoding="utf-8"
    )
    (root / "tests" / "test_repo.py").write_text(
        "def test_x():\n    assert True\n", encoding="utf-8"
    )
    (root / "web" / "index.ts").write_text(
        "import { Btn } from './components/btn'\n", encoding="utf-8"
    )
    (root / "web" / "components" / "btn.ts").write_text("export const Btn = 1\n", encoding="utf-8")
    (root / "pyproject.toml").write_text(
        '[project]\nname = "shop"\ndependencies = ["fastapi>=0.110", "sqlalchemy"]\n[project.scripts]\nshop = "shop.main:run"\n',
        encoding="utf-8",
    )
    (root / "package.json").write_text(
        json.dumps({"name": "web", "dependencies": {"react": "^18"}, "scripts": {"dev": "vite"}}),
        encoding="utf-8",
    )
    (root / "README.md").write_text("# Shop\n", encoding="utf-8")
    (root / "node_modules" / "x").mkdir(parents=True)
    (root / "node_modules" / "x" / "huge.js").write_text("x\n" * 10000, encoding="utf-8")


def test_repo_scan_overview_tree_and_graph(tmp_path: Path) -> None:
    root = tmp_path / "proj"
    _demo_repo(root)
    data = json.loads(_run(REPO_SCAN, "scan", root, "--json").stdout)
    assert data["文件总数"] == 11  # node_modules is skipped
    assert data["语言"]["Python"]["文件"] == 6 and data["语言"]["TypeScript"]["文件"] == 2
    assert {"FastAPI", "SQLAlchemy", "React"} <= set(data["框架和主要库"])
    assert "src/shop/main.py" in data["入口和构建文件"] and "web/index.ts" in data["入口和构建文件"]
    assert (
        data["测试文件数"] == 1
        and data["TODO/FIXME"] == {"TODO": 1}
        and data["README"] == "README.md"
    )
    py = next(m for m in data["依赖清单"] if m["类型"] == "Python")
    assert py["依赖"] == ["fastapi", "sqlalchemy"] and py["脚本"] == {"shop": "shop.main:run"}

    tree = _run(REPO_SCAN, "tree", root, "--depth", "2").stdout
    assert "src/  (5 个文件)" in tree and "node_modules" not in tree
    assert "│   └── …1 个文件" in tree  # files below the depth limit are summarised, not dropped

    out = tmp_path / "g.mmd"
    result = _run(REPO_SCAN, "graph", root, "--lang", "python", "--depth", "2", "-o", out)
    assert result.returncode == 0, result.stderr
    assert "shop.api ⇄ shop.db" in result.stdout or "shop.db ⇄ shop.api" in result.stdout
    mmd = out.read_text(encoding="utf-8")
    assert mmd.startswith("graph LR") and '["shop.api"]' in mmd and '["shop.db"]' in mmd

    pkg = _run(REPO_SCAN, "graph", root / "src" / "shop", "--depth", "2")
    assert pkg.returncode == 0 and "shop.db" in pkg.stdout  # package root gets its name prefixed
    js = _run(REPO_SCAN, "graph", root / "web", "--lang", "js", "--depth", "1")
    assert js.returncode == 0 and '["components"]' in js.stdout


# ---------------------------------------------------------------------- sql


def test_sql_lint_rules() -> None:
    script = SQL / "scripts" / "sql_check.py"
    result = _run(script, "lint", SQL / "references" / "example-queries.sql")
    assert result.returncode == 3
    out = result.stdout
    assert "共 9 条语句：错误 2，警告 6，提示 2" in out
    blocks = {
        int(m.group(1)): m.group(2)
        for m in re.finditer(r"^\[(\d+)\] (.*?)(?=^\[\d+\]|^共)", out, re.M | re.S)
    }
    assert "SELECT *" in blocks[1] and "YEAR()" in blocks[1]
    assert "隐式连接" in blocks[2]
    assert "LIKE '%…'" in blocks[3] and "OR 连接了不同列" in blocks[3]
    assert "NOT EXISTS" in blocks[4]
    assert "[错误] UPDATE 没有 WHERE" in blocks[5]
    assert "IS NULL" in blocks[6]
    assert "没有写列名" in blocks[7]
    assert "深分页" in blocks[8]
    assert "没有发现问题" in blocks[9]

    clean = _run(script, "lint", "SELECT name FROM users WHERE note = 'SELECT * FROM x' LIMIT 1")
    assert clean.returncode == 0  # text inside string literals is ignored
    assert "[错误] DELETE 没有 WHERE" in _run(script, "lint", "-- 清空\nDELETE FROM t").stdout


def test_sql_er_and_explain(tmp_path: Path) -> None:
    script = SQL / "scripts" / "sql_check.py"
    schema = SQL / "references" / "example-schema.sql"
    result = _run(script, "er", schema)
    assert (
        "5 张表，3 个外键关系" in result.stdout and "没有主键的表：operation_log" in result.stdout
    )
    assert (
        "INTEGER order_id PK, FK" in result.stdout
        and 'customers ||--o{ orders : "customer_id"' in result.stdout
    )

    db = tmp_path / "shop.db"
    conn = sqlite3.connect(db)
    conn.executescript(schema.read_text(encoding="utf-8"))
    conn.close()
    plan = _run(
        script,
        "explain",
        db,
        "SELECT o.id FROM orders o JOIN customers c ON c.id = o.customer_id WHERE c.city = 'x'",
    )
    assert plan.returncode == 0, plan.stderr
    assert (
        "USING COVERING INDEX idx_orders_customer" in plan.stdout and "全表扫描：c" in plan.stdout
    )
    assert _run(script, "explain", db, "DELETE FROM orders").returncode == 1
    from_db = _run(script, "er", db)
    assert "orders ||--o{ order_items" in from_db.stdout


# ---------------------------------------------------------------------- git

needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed")


def _git(repo: Path, *args: str, date: str | None = None) -> None:
    env = {**os.environ, "GIT_AUTHOR_NAME": "张三", "GIT_AUTHOR_EMAIL": "a@example.com",
           "GIT_COMMITTER_NAME": "张三", "GIT_COMMITTER_EMAIL": "a@example.com"}  # fmt: skip
    if date:
        env |= {"GIT_AUTHOR_DATE": date, "GIT_COMMITTER_DATE": date}
    subprocess.run(["git", "-C", str(repo), *args], check=True, capture_output=True, env=env)


@needs_git
def test_git_report_changelog_branches_and_staged(tmp_path: Path) -> None:
    repo = tmp_path / "r"
    repo.mkdir()
    _git(repo, "init", "-q", "-b", "main")
    _git(repo, "config", "commit.gpgsign", "false")

    def commit(msg: str, name: str = "a.py", date: str | None = None) -> None:
        with (repo / name).open("a", encoding="utf-8") as f:
            f.write("x\n")
        _git(repo, "add", name)
        _git(repo, "commit", "-q", "-m", msg, date=date)

    commit("feat: 初始化项目")
    _git(repo, "tag", "v1.0.0")
    commit("fix(login): 修复验证码过期判断 (#12)")
    commit("docs: 补充安装说明", "README.md")
    commit("feat(api)!: 订单接口改为分页返回")
    commit("改了点东西")
    _git(repo, "checkout", "-q", "-b", "feature/old")
    commit("feat: 旧功能", "old.py", date="2025-01-01T10:00:00")
    _git(repo, "checkout", "-q", "main")
    _git(repo, "branch", "feature/done")

    log = _run(GIT, "--repo", repo, "changelog", "--from", "v1.0.0").stdout
    sections = re.findall(r"^### (.+)$", log, re.M)
    assert sections == ["⚠️ 破坏性变更", "新增", "修复", "文档", "其他"]
    assert "- **login**：修复验证码过期判断 (#12)" in log and "1 条提交没有按" in log

    stats = _run(GIT, "--repo", repo, "stats").stdout
    assert (
        "5 次提交" in stats and "| 张三 | 5 |" in stats
    )  # current branch only; feature/old isn't counted
    branches = _run(GIT, "--repo", repo, "branches", "--days", "90").stdout
    assert "git branch -d feature/done" in branches
    assert re.search(r"\| feature/old \| 2025-01-01 \| 张三 \| 否 \| 长期未动", branches)

    (repo / "test_new.py").write_text("x\n", encoding="utf-8")
    _git(repo, "add", "test_new.py")
    staged = _run(GIT, "--repo", repo, "staged").stdout
    assert "| test_new.py | 新增 | 1 | 0 |" in staged and "建议类型：test" in staged

    bad = _run(GIT, "--repo", tmp_path, "stats")
    assert bad.returncode == 1 and "不是 git 仓库" in bad.stderr


# ---------------------------------------------------------------------- api


class _FakeApi(http.server.BaseHTTPRequestHandler):
    def do_GET(self) -> None:  # noqa: N802 - http.server API
        path = self.path.split("?")[0]
        if path == "/books" and self.headers.get("Authorization") != "Bearer t":
            code = 401
        else:
            code = {"/health": 200, "/books": 200, "/books/1": 500}.get(path, 404)
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(b"{}")

    def log_message(self, *args: object) -> None:
        pass


@pytest.fixture
def fake_api() -> Iterator[str]:
    server = http.server.HTTPServer(("127.0.0.1", 0), _FakeApi)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


def test_api_list_cases_and_curl(tmp_path: Path) -> None:
    script = API / "scripts" / "api_check.py"
    spec = API / "references" / "example-openapi.yaml"
    out = _run(script, "list", spec).stdout
    assert "示例图书接口 1.0.0：6 个接口" in out
    assert "| GET | /health | 健康检查 | - | - | 200 | 否 |" in out
    assert "| POST | /books | 新增图书 | - | 有（必填） | 201,422 | 是 |" in out

    xlsx = tmp_path / "cases.xlsx"
    result = _run(script, "cases", spec, "-o", xlsx)
    assert "6 个接口，26 条用例" in result.stdout
    rows = list(load_workbook(xlsx)["测试用例"].iter_rows(min_row=2, values_only=True))
    post_ok = next(r for r in rows if r[1] == "POST /books" and r[2] == "正常请求")
    assert json.loads(post_ok[4]) == {"title": "三体", "price": 23.5}
    assert any(r[1] == "POST /books" and r[2] == "请求体缺少必填字段 title" for r in rows)
    assert any(
        r[1] == "GET /books/{book_id}" and r[2] == "资源不存在" and r[5] == "404" for r in rows
    )

    code = _run(script, "curl2py", "curl -X POST 'https://api.example.com/v1/items' -H 'Content-Type: application/json' "
                "-H 'Authorization: Bearer abc' -d '{\"name\":\"true story\",\"ok\":true,\"n\":null}' -k").stdout  # fmt: skip
    assert "payload = {'name': 'true story', 'ok': True, 'n': None}" in code
    assert (
        "httpx.Client(verify=False)" in code
        and "json=payload" in code
        and '"Content-Type"' not in code
    )
    compile(code, "generated", "exec")


def test_api_smoke_is_read_only(tmp_path: Path, fake_api: str) -> None:
    script = API / "scripts" / "api_check.py"
    spec = API / "references" / "example-openapi.yaml"
    result = _run(script, "smoke", spec, "--base", fake_api, "-o", tmp_path / "smoke.xlsx")
    assert result.returncode == 3  # /books/1 returns 500
    out = result.stdout
    assert "通过 1，失败 1，未通过鉴权 1，跳过 1（只测试 GET 接口）" in out
    assert "[未通过鉴权] GET /books 401" in out and "[失败] GET /books/{book_id} 500" in out
    assert (
        "[跳过] GET /authors/{author_id}/books" in out and "POST" not in out and "DELETE" not in out
    )
    authed = _run(script, "smoke", spec, "--base", fake_api, "--header", "Authorization: Bearer t")
    assert "通过 2，失败 1，未通过鉴权 0" in authed.stdout


# ----------------------------------------------------------------- security


def _leaky_project(root: Path) -> None:
    (root / "app").mkdir(parents=True)
    fake_aws = (
        "AKIA" + "Q" * 16
    )  # built at runtime so this source never contains a key-shaped literal
    fake_gh = "ghp" + "_" + "a1B2" * 9
    (root / "app" / "config.py").write_text(
        f'AWS_KEY = "{fake_aws}"\nDB_URL = "postgres://admin:S3cretPass@db.internal:5432/app"\n'
        'password = "hunter2hunter2"\napi_key = os.environ.get("API_KEY")\ntoken = "your_token_here"\n'
        f'DEBUG = True\nGH = "{fake_gh}"  # nosecret\n',
        encoding="utf-8",
    )
    (root / "app" / "views.py").write_text(
        'def f(cur, q):\n    cur.execute(f"SELECT * FROM u WHERE n = \'{q}\'")\n    cur.execute("SELECT 1 WHERE a = %s", (q,))\n'
        "    subprocess.run('ls ' + q, shell=True)\n    yaml.load(q)\n    yaml.load(q, Loader=yaml.SafeLoader)\n"
        "    requests.get(q, verify=False)\n    return eval(q)\n",
        encoding="utf-8",
    )
    (root / "web.js").write_text("el.innerHTML = userInput;\n", encoding="utf-8")
    (root / ".env").write_text("SECRET_KEY=abc\n", encoding="utf-8")
    (root / ".env.example").write_text("SECRET_KEY=\n", encoding="utf-8")


def test_secret_scan_findings_and_masking(tmp_path: Path) -> None:
    root = tmp_path / "proj"
    _leaky_project(root)
    result = _run(SECRET, root, "--json")
    assert result.returncode == 3
    findings = json.loads(result.stdout)
    by_rule = {(f["rule"], f["file"], f["line"]) for f in findings}
    assert ("AWS Access Key", "app/config.py", 1) in by_rule
    assert ("带密码的连接串", "app/config.py", 2) in by_rule
    assert ("写死的密码 / 密钥", "app/config.py", 3) in by_rule
    assert not any(
        f["line"] in (4, 5, 7) and f["file"] == "app/config.py" for f in findings
    )  # env read, placeholder, nosecret
    assert ("环境变量文件没有被 .gitignore 忽略", ".env", 0) in by_rule
    assert not any(f["file"] == ".env.example" for f in findings)
    views = {f["rule"] for f in findings if f["file"] == "app/views.py"}
    assert views == {"拼接 SQL", "shell=True", "不安全的反序列化", "关闭证书校验", "eval / exec"}
    assert (
        sum(1 for f in findings if f["rule"] == "不安全的反序列化") == 1
    )  # safe_load-style call is not flagged
    text = json.dumps(findings, ensure_ascii=False)
    assert (
        "Q" * 16 not in text and "S3cretPass" not in text and "hunter2hunter2" not in text
    )  # masked
    assert "AKIA****QQ" in text

    (root / ".gitignore").write_text(".env\n", encoding="utf-8")
    report = tmp_path / "r.xlsx"
    again = _run(SECRET, root, "--min-level", "高", "-o", report)
    assert ".env" not in again.stdout.split("处理建议")[0].replace(".env.example", "")
    assert load_workbook(report)["扫描结果"].max_row == 4  # header + 3 high findings


@needs_git
def test_secret_scan_flags_tracked_env(tmp_path: Path) -> None:
    root = tmp_path / "proj"
    root.mkdir()
    (root / ".env").write_text("A=1\n", encoding="utf-8")
    (root / ".gitignore").write_text(".env\n", encoding="utf-8")
    _git(root, "init", "-q")
    _git(root, "add", "-f", ".env")
    findings = json.loads(_run(SECRET, root, "--json").stdout)
    assert findings[0]["level"] == "高" and findings[0]["rule"] == "环境变量文件已被 git 跟踪"


# --------------------------------------------------------------------- html


def _chromium_ok() -> bool:
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as pw:
            pw.chromium.launch().close()
        return True
    except Exception:  # noqa: BLE001 - any failure means the browser isn't usable here
        return False


@pytest.mark.skipif(not _chromium_ok(), reason="Playwright Chromium is not installed")
def test_html_check_example_and_broken_page(tmp_path: Path) -> None:
    script = HTML / "scripts" / "html_check.py"
    ok = _run(
        script,
        HTML / "references" / "example-tool.html",
        "-o",
        tmp_path / "shots",
        base_env=ORIGINAL_ENV,
    )
    assert ok.returncode == 0, ok.stdout + ok.stderr
    assert "错误 0，警告 0" in ok.stdout
    assert (tmp_path / "shots" / "desktop.png").is_file() and (
        tmp_path / "shots" / "mobile.png"
    ).is_file()

    bad = tmp_path / "bad.html"
    bad.write_text(
        '<html><head><meta name="viewport" content="width=device-width, initial-scale=1"></head><body>'
        '<div style="width:900px">wide</div><img src="missing.png"><input name="q"><script>undefinedFn();</script></body></html>',
        encoding="utf-8",
    )
    result = _run(script, bad, "--viewports", "desktop,mobile", base_env=ORIGINAL_ENV)
    assert result.returncode == 3
    out = result.stdout
    assert (
        "脚本报错：undefinedFn is not defined" in out
        and "资源加载失败" in out
        and "missing.png" in out
    )
    assert "Failed to load resource" not in out  # duplicates of requestfailed are dropped
    assert "手机宽度（390px）下内容超出" in out and "缺少 <title>" in out
    assert "1 张图片没有 alt" in out and "1 个输入框没有关联 label：q" in out
    assert _run(script, bad, "--viewports", "watch").returncode == 1


# -------------------------------------------------------------------- regex


def test_regex_expectations_find_replace_and_errors() -> None:
    ok = _run(REGEX, r"1[3-9]\d{9}", "-M", "13800138000", "-N", "12345678901", "-N", "138001380000")
    assert ok.returncode == 0, ok.stdout + ok.stderr
    assert "模式：fullmatch" in ok.stdout and "预期检查：3/3 符合" in ok.stdout

    wrong = _run(REGEX, r"\d{11}", "-M", "13800138000", "-N", "12345678901")
    assert wrong.returncode == 3 and "← 预期不应该匹配" in wrong.stdout

    found = _run(REGEX, r"(?P<y>\d{4})-(?P<m>\d{2})-(?P<d>\d{2})", "-s", "日期 2026-10-19 和 2025-01-02",
                 "--mode", "findall", "--replace", r"\g<y>年\g<m>月\g<d>日")  # fmt: skip
    assert "全部结果：['2026-10-19', '2025-01-02']" in found.stdout
    assert "替换后：'日期 2026年10月19日 和 2025年01月02日'" in found.stdout

    js_style = _run(REGEX, r"(?<year>\d+)", "-s", "x")
    assert js_style.returncode == 1 and "(?P<name>…)" in js_style.stderr

    mod = _load(REGEX)
    tips = "\n".join(
        mod.lint(r"\d+.\d+.\d+.\d+", validating=False) + mod.lint("[A-z]+", validating=False)
    )
    assert "没有转义" in tips and "[A-Za-z]" in tips
    assert any("嵌套量词" in t for t in mod.lint(r"(a+)+$", validating=False))
    assert not mod.lint(r"^\d{3}\.\d{2}$", validating=True)


def test_regex_catastrophic_backtracking_times_out() -> None:
    result = _run(REGEX, r"(a+)+$", "-s", "a" * 40 + "b", "--timeout", "1")
    assert result.returncode == 3 and "[超时]" in result.stdout
