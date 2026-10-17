"""Contract tests for the bundled office experts (会议 / 方案 / 数据分析)."""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from docx import Document
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import Workbook, load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
MINUTES = LIBRARY / "meeting-assistant" / "skills" / "meeting-minutes"
ACTIONS = LIBRARY / "meeting-assistant" / "skills" / "action-tracking"
PLAN_EXPORT = LIBRARY / "plan-writer" / "skills" / "plan-export"
TOOLKIT = LIBRARY / "data-analyst" / "skills" / "data-toolkit"
SALES = TOOLKIT / "references" / "example-sales.csv"

EXPERTS = {
    "meeting-assistant": ("会议助手", {"meeting-prep", "meeting-minutes", "action-tracking"}),
    "plan-writer": ("方案策划专家", {"event-plan", "project-plan", "plan-export"}),
    "data-analyst": ("数据分析报告专家", {"data-toolkit", "analysis-methods", "analysis-report"}),
}


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


@pytest.mark.parametrize("name", ["docx_build.py", "budget_xlsx.py"])
def test_shared_script_copies_are_identical(name: str) -> None:
    copies = sorted(LIBRARY.glob(f"*/skills/*/scripts/{name}"))
    if len(copies) < 2:
        pytest.skip(f"only {len(copies)} copy of {name} in this tree")
    first = copies[0].read_bytes()
    for copy in copies[1:]:
        assert copy.read_bytes() == first, f"{copy} drifted from {copies[0]}"


# ------------------------------------------------------------------ meetings


def test_minutes_example_exports_to_word(tmp_path: Path) -> None:
    out = tmp_path / "minutes.docx"
    result = _run(
        MINUTES / "scripts" / "docx_build.py",
        MINUTES / "references" / "example-minutes.md",
        "-o",
        out,
        "--style",
        "report",
    )
    assert result.returncode == 0, result.stderr
    doc = Document(str(out))
    assert doc.paragraphs[0].text == "经营例会纪要"
    assert len(doc.tables) == 2
    assert doc.tables[0].cell(1, 0).text == "提交 Q4 预算调整方案"


def test_actions_tracker_flags_overdue_and_round_trips(tmp_path: Path) -> None:
    script = ACTIONS / "scripts" / "actions_xlsx.py"
    out = tmp_path / "actions.xlsx"
    result = _run(
        script, ACTIONS / "references" / "example-actions.md", "-o", out, "--today", "2026-10-18"
    )
    assert result.returncode == 0, result.stderr
    assert "共 5 项，已完成 1，逾期 1，3 天内到期 1" in result.stdout
    assert "提交 Q4 预算调整方案｜张三｜截止 2026-10-15（已逾期 3 天）" in result.stdout
    assert "1 项没有负责人" in result.stdout and "1 项没有截止日期" in result.stdout
    wb = load_workbook(out)
    ws = wb["待办"]
    assert ws.cell(4, 2).value == "提交 Q4 预算调整方案"  # overdue sorted first
    assert ws.cell(4, 2).fill.fgColor.rgb.endswith("FEE2E2")
    assert [c.value for c in wb["按负责人"][2]] == ["张三", 1, 1, 0]

    refused = _run(script, out, "-o", out, "--today", "2026-10-18")
    assert refused.returncode == 1 and "--overwrite" in refused.stderr
    again = _run(script, out, "-o", out, "--today", "2026-10-18", "--overwrite")
    assert again.returncode == 0
    assert "共 5 项，已完成 1，逾期 1" in again.stdout

    bad = tmp_path / "bad.md"
    bad.write_text(
        "| 待办 | 负责人 | 截止 |\n|---|---|---|\n| 写周报 | 张三 | 下周 |\n", encoding="utf-8"
    )
    result = _run(script, bad, "-o", tmp_path / "x.xlsx")
    assert result.returncode == 1 and "无法识别" in result.stderr


# --------------------------------------------------------------------- plans


def test_plan_budget_example(tmp_path: Path) -> None:
    out = tmp_path / "budget.xlsx"
    result = _run(
        PLAN_EXPORT / "scripts" / "budget_xlsx.py",
        PLAN_EXPORT / "references" / "example-budget.md",
        "-o",
        out,
    )
    assert result.returncode == 0, result.stderr
    assert "估算总计 200,340.00 元" in result.stdout
    labels = [load_workbook(out)["预算表"].cell(r, 2).value for r in range(1, 30)]
    assert "预备费（5%）" in labels and "税费（6%）" in labels and "管理费（10%）" not in labels


# ---------------------------------------------------------------------- data


def test_profile_data_reports_quality_issues() -> None:
    result = _run(TOOLKIT / "scripts" / "profile_data.py", SALES)
    assert result.returncode == 0, result.stderr
    out = result.stdout
    assert "行数 75，列数 6" in out
    assert re.search(r"\| 订单号 \| 文本 .*唯一标识", out)
    assert re.search(r"\| 日期 \| 日期 \|", out)
    assert re.search(r"\| 销售额 \| 数值 .*异常值 1 个，非数字 1 个", out)
    assert "唯一标识" not in out.split("| 销售额 |")[1].split("\n")[0]
    assert "「区域」存在大小写/空格不同" in out

    as_json = json.loads(_run(TOOLKIT / "scripts" / "profile_data.py", SALES, "--json").stdout)
    assert as_json["行数"] == 75
    assert {c["字段"]: c["类型"] for c in as_json["字段"]}["数量"] == "数值"


def test_profile_data_reads_xlsx_with_offset_header(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["2026 年销售明细"])
    ws.append(["客户", "金额"])
    ws.append(["甲", 100])
    ws.append(["乙", "1,200"])
    path = tmp_path / "t.xlsx"
    wb.save(path)
    result = _run(TOOLKIT / "scripts" / "profile_data.py", path, "--header-row", "2")
    assert result.returncode == 0, result.stderr
    assert "| 金额 | 数值 |" in result.stdout and "100 ~ 1200" in result.stdout


def test_pivot_chart_category_time_and_cross(tmp_path: Path) -> None:
    script = TOOLKIT / "scripts" / "pivot_chart.py"
    by_region = tmp_path / "region.xlsx"
    result = _run(script, SALES, "-o", by_region, "--rows", "区域", "--values", "销售额")
    assert result.returncode == 0, result.stderr
    assert "图表：pie" in result.stdout
    assert "1 行因字段为空" in result.stdout  # the "待定" amount is excluded and reported
    ws = load_workbook(by_region)["汇总"]
    assert [c.value for c in ws[1]] == ["区域", "销售额（合计）", "占比"]
    regions = [ws.cell(r, 1).value for r in range(2, 6)]
    assert sorted(regions) == ["华东", "华北", "华南", "西南"]  # " 华东" merged after strip
    assert str(ws.cell(6, 2).value).startswith("=SUM(")
    assert len(ws._charts) == 1

    by_quarter = tmp_path / "q.xlsx"
    result = _run(
        script,
        SALES,
        "-o",
        by_quarter,
        "--date",
        "日期",
        "--period",
        "quarter",
        "--values",
        "销售额",
    )
    assert result.returncode == 0, result.stderr
    ws = load_workbook(by_quarter)["汇总"]
    assert [c.value for c in ws[1]] == ["日期", "销售额（合计）", "环比", "同比"]
    assert [ws.cell(r, 1).value for r in range(2, 8)] == [
        "2025-Q1", "2025-Q2", "2025-Q3", "2025-Q4", "2026-Q1", "2026-Q2",
    ]  # fmt: skip
    assert ws.cell(2, 3).value is None and ws.cell(3, 3).value is not None
    assert ws.cell(6, 4).value is not None  # 2026-Q1 vs 2025-Q1

    cross = tmp_path / "cross.xlsx"
    result = _run(
        script,
        SALES,
        "-o",
        cross,
        "--rows",
        "产品",
        "--columns",
        "区域",
        "--agg",
        "count",
        "--chart",
        "column",
    )
    assert result.returncode == 0, result.stderr
    header = [c.value for c in load_workbook(cross)["汇总"][1]]
    assert header[0] == "产品" and "总计" in header and "华东" in header

    missing = _run(
        script, SALES, "-o", tmp_path / "x.xlsx", "--rows", "不存在", "--values", "销售额"
    )
    assert missing.returncode == 1 and "可用字段" in missing.stderr
