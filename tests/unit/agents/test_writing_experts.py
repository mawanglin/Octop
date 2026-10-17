"""Contract tests for the bundled writing experts (公文 / 汇报 / 公众号 / 讲话稿)."""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from docx import Document
from octop_harness.backends.workspace import BackendWorkspace

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
GOV = LIBRARY / "gov-document" / "skills" / "gov-format"
REPORT = LIBRARY / "work-report" / "skills" / "report-export"
SPEECH = LIBRARY / "speech-writer" / "skills" / "speech-export"
WECHAT = LIBRARY / "wechat-writer" / "skills"

EXPERTS = {
    "gov-document": ("政务公文专家", {"gov-doc-types", "gov-report-material", "gov-format"}),
    "work-report": (
        "工作汇报专家",
        {"daily-weekly-report", "periodic-report", "annual-review", "report-export"},
    ),
    "wechat-writer": (
        "公众号写作专家",
        {"topic-planning", "headline", "article-writing", "wechat-layout", "platform-adapt"},
    ),
    "speech-writer": ("讲话稿专家", {"speech-writing", "host-script", "speech-export"}),
}
DOCX_BUILD_COPIES = [
    GOV / "scripts" / "docx_build.py",
    REPORT / "scripts" / "docx_build.py",
    SPEECH / "scripts" / "docx_build.py",
]


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


# ------------------------------------------------------------------- catalog


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert all(
        q["title"]["zh"] and q["title"]["en"] and q["prompt"]["zh"] and q["prompt"]["en"]
        for q in manifest["quick_prompts"]
    )
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        # Every referenced bundled file must exist.
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    for rel in paths:
        assert await workspace.adownload_bytes(rel) == (expert_dir / rel).read_bytes()
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


def test_docx_build_copies_are_identical() -> None:
    first = DOCX_BUILD_COPIES[0].read_bytes()
    for copy in DOCX_BUILD_COPIES[1:]:
        assert copy.read_bytes() == first, f"{copy} drifted from {DOCX_BUILD_COPIES[0]}"


# ---------------------------------------------------------------- docx_build


def test_gov_style_follows_gb_t_9704(tmp_path: Path) -> None:
    out = tmp_path / "notice.docx"
    result = _run(
        DOCX_BUILD_COPIES[0], GOV / "references" / "example-notice.md", "-o", out, "--style", "gov"
    )
    assert result.returncode == 0, result.stderr

    doc = Document(str(out))
    sec = doc.sections[0]
    assert round(sec.top_margin.mm) == 37 and round(sec.left_margin.mm) == 28
    emu_per_mm = 36000
    assert round((sec.page_width - sec.left_margin - sec.right_margin) / emu_per_mm) == 156
    assert round((sec.page_height - sec.top_margin - sec.bottom_margin) / emu_per_mm) == 225
    assert doc.settings.odd_and_even_pages_header_footer

    texts = [p.text for p in doc.paragraphs]
    header = next(p for p in doc.paragraphs if p.text == "示例市人民政府办公室文件")
    assert header.runs[0].font.color.rgb == (0xFF, 0x00, 0x00)
    number = next(p for p in doc.paragraphs if "〔2026〕12号" in p.text)
    assert 'w:color="FF0000"' in number._p.xml  # 版头红色分隔线
    # Title split on \n into two centered lines.
    assert "关于开展2026年秋冬季安全生产" in texts and "大检查的通知" in texts
    assert "各区人民政府，市政府各部门：" in texts
    first_level = next(p for p in doc.paragraphs if p.text == "一、检查时间")
    assert 'w:eastAsia="黑体"' in first_level._p.xml
    second = next(p for p in doc.paragraphs if p.text.startswith("（一）道路交通安全"))
    assert 'w:eastAsia="楷体_GB2312"' in second.runs[0]._r.xml
    assert 'w:eastAsia="仿宋_GB2312"' in second.runs[-1]._r.xml
    assert any(t.startswith("附件：1．") for t in texts)
    assert "（此件公开发布）" in texts
    colophon = doc.tables[-1]
    assert colophon.cell(0, 0).text.startswith("抄送：")
    assert colophon.cell(1, 0).text.endswith("印发")
    assert "— " in sec.footer.paragraphs[0].text or "PAGE" in sec.footer.paragraphs[0]._p.xml


def test_report_and_speech_styles(tmp_path: Path) -> None:
    report = tmp_path / "report.docx"
    result = _run(
        DOCX_BUILD_COPIES[1],
        REPORT / "references" / "example-report.md",
        "-o",
        report,
        "--style",
        "report",
    )
    assert result.returncode == 0, result.stderr
    doc = Document(str(report))
    assert doc.paragraphs[0].text == "运营中心季度工作汇报"
    assert len(doc.tables) == 1
    assert any(p.text.startswith("• ") for p in doc.paragraphs)

    speech = tmp_path / "speech.docx"
    result = _run(
        DOCX_BUILD_COPIES[2],
        SPEECH / "references" / "example-speech.md",
        "-o",
        speech,
        "--style",
        "speech",
    )
    assert result.returncode == 0, result.stderr
    doc = Document(str(speech))
    salutation = next(p for p in doc.paragraphs if p.text == "同志们：")
    assert salutation.paragraph_format.first_line_indent is None
    body = next(p for p in doc.paragraphs if p.text.startswith("今天召开"))
    assert body.runs[0].font.size.pt == 18


def test_docx_build_rejects_bad_input(tmp_path: Path) -> None:
    src = tmp_path / "bad.md"
    src.write_text("---\n标题 没有冒号\n---\n正文\n", encoding="utf-8")
    result = _run(DOCX_BUILD_COPIES[0], src, "-o", tmp_path / "x.docx", "--style", "gov")
    assert result.returncode == 1
    assert "冒号" in result.stderr


# ------------------------------------------------------------------ gov_lint


def test_gov_lint_example_is_clean_and_bad_draft_is_flagged(tmp_path: Path) -> None:
    lint = GOV / "scripts" / "gov_lint.py"
    clean = _run(lint, GOV / "references" / "example-notice.md")
    assert clean.returncode == 0, clean.stdout

    bad = tmp_path / "bad.md"
    bad.write_text(
        "---\n发文字号: 示政发[2026]12号\n日期: 2026年01月05日\n---\n"
        "# 关于申请专项经费的报告。\n\n一、基本情况\n\n三、经费需求\n\n"
        "需要经费100万元,时间为2026.11.01。\n\n妥否，请批示。\n",
        encoding="utf-8",
    )
    result = _run(lint, bad)
    assert result.returncode == 3
    for expected in (
        "标题中除书名号",
        "此处应为“二、”",
        "“2026.11.01”",
        "月、日不编虚位",
        "半角标点“,”",
        "六角括号",
        "报告中不得夹带请示事项",
    ):
        assert expected in result.stdout, expected


# ----------------------------------------------------------- wechat / speech


def test_wechat_layout_inlines_styles_and_footnotes_links(tmp_path: Path) -> None:
    script = WECHAT / "wechat-layout" / "scripts" / "wechat_layout.py"
    example = WECHAT / "wechat-layout" / "references" / "example-article.md"
    out = tmp_path / "a.html"
    result = _run(script, example, "-o", out, "--theme", "blue", "--fragment")
    assert result.returncode == 0, result.stderr
    assert "提示" not in result.stdout
    html = out.read_text(encoding="utf-8")
    assert html.startswith('<section style="')
    assert "<style" not in html and "class=" not in html  # inline styles only
    assert "#1e5bd8" in html
    assert "<ul" not in html and "<ol" not in html  # lists become paragraphs
    assert "参考链接" in html and "https://github.com/TencentCloud/Octop" in html
    assert 'href="https://github.com' not in html  # external links are not clickable

    long_title = tmp_path / "t.md"
    long_title.write_text("---\n标题: " + "长" * 70 + "\n---\n正文\n", encoding="utf-8")
    stats = _run(script, long_title, "--stats")
    assert "上限 64 字" in stats.stdout


def test_content_check_flags_risky_claims(tmp_path: Path) -> None:
    script = WECHAT / "article-writing" / "scripts" / "content_check.py"
    risky = tmp_path / "ad.md"
    risky.write_text(
        "全网最低价，7天见效，根治痘痘！转发到朋友圈即可领取。投资稳赚。第一次用。",
        encoding="utf-8",
    )
    result = _run(script, risky)
    assert result.returncode == 3
    for category in ("广告法极限用语", "医疗健康功效", "诱导分享/关注", "金融投资承诺"):
        assert category in result.stdout
    assert "“第一”" not in result.stdout  # "第一次" is not a superlative
    clean = _run(script, WECHAT / "wechat-layout" / "references" / "example-article.md")
    assert clean.returncode == 0, clean.stdout


def test_speech_timer_estimates_and_targets() -> None:
    script = SPEECH / "scripts" / "speech_timer.py"
    result = _run(script, SPEECH / "references" / "example-speech.md", "--target", "3")
    assert result.returncode == 0, result.stderr
    assert "合计：" in result.stdout
    assert "约需增加" in result.stdout
    assert "一、提高政治站位" in result.stdout
