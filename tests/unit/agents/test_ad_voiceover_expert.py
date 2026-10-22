"""Contract tests for the bundled 投流口播专家 (ad-voiceover) expert."""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path

from deepagents.backends.local_shell import LocalShellBackend
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import Workbook, load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
EXPERT_DIR = LIBRARY / "ad-voiceover"
SKILLS = {"agent-product-pitch", "ad-script-writing", "ad-compliance", "ad-test-review"}
CHECK = EXPERT_DIR / "skills" / "ad-compliance" / "scripts" / "ad_script_check.py"
AD_TEST = EXPERT_DIR / "skills" / "ad-test-review" / "scripts" / "ad_test.py"
EXAMPLE_SCRIPTS = EXPERT_DIR / "skills" / "ad-script-writing" / "references" / "example-scripts.md"
EXAMPLE_REPORT = EXPERT_DIR / "skills" / "ad-test-review" / "references" / "example-report.csv"


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


def test_expert_catalog_manifest_and_skills() -> None:
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get("ad-voiceover")
    assert expert is not None
    assert expert.summary.label_zh == "投流口播专家"
    assert expert.summary.icon_url == "/experts/avatars/ad-voiceover.svg"

    manifest = json.loads((EXPERT_DIR / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == "ad-voiceover"
    for rel in manifest["prompt_files"]:
        assert (EXPERT_DIR / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (EXPERT_DIR / "skills").iterdir() if p.is_dir()} == SKILLS
    soul = (EXPERT_DIR / "SOUL.md").read_text(encoding="utf-8")
    for name in SKILLS:
        raw = (EXPERT_DIR / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (EXPERT_DIR / "skills" / name / rel).is_file(), f"{name}: {rel}"


async def test_expert_seeds_into_workspace(tmp_path: Path) -> None:
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(EXPERT_DIR)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=EXPERT_DIR, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (EXPERT_DIR / "manifest.json").read_bytes()
    )


# ------------------------------------------------------------------- check


def test_example_scripts_good_versions_pass_bad_one_flagged() -> None:
    proc = _run(CHECK, EXAMPLE_SCRIPTS, "--json")
    assert proc.returncode == 3, proc.stderr
    data = json.loads(proc.stdout)
    assert data["meta"]["时长"] == "30"
    a, b, c = data["variants"]
    for good in (a, b):
        assert good["issues"] == [] and good["compliance"] == [], good
        assert good["has_cta"] and good["hook_signals"]
    assert a["hook"] == "还在每天加班整理销售日报"
    assert "痛点" in a["hook_signals"] and "数字" in b["hook_signals"]
    categories = {hit["类别"] for hit in c["compliance"]}
    assert categories == {
        "广告法极限用语",
        "冒用品牌或名人",
        "AI 能力夸大",
        "效果和收益承诺",
        "虚假紧迫和诱导",
    }
    words = {hit["命中"] for hit in c["compliance"]}
    assert {"最强", "完全取代你的运营团队", "100% 准确", "马上涨价"} <= words
    assert not c["has_cta"]


def test_check_flags_structure_and_digital_human(tmp_path: Path) -> None:
    script = tmp_path / "稿.md"
    script.write_text(
        "---\n时长: 10\n数字人: 是\n---\n"
        "## 慢开头\n"
        "口播：大家好今天我想跟大家分享一个我最近一直在用而且觉得非常不错的人工智能产品\n"
        "口播：它基于 RAG 和 Workflow 技术打造。数据绝不泄露。\n",
        encoding="utf-8",
    )
    proc = _run(CHECK, script)
    assert proc.returncode == 3
    out = proc.stdout
    assert "开头第一句约" in out and "超过 4 秒" in out
    assert "结尾没有明确的行动引导" in out
    assert "超过目标 10 秒" in out
    assert "一口气超过 18 字没有停顿" in out
    assert "术语“RAG”没有解释" in out
    assert "[合规·隐私和安全承诺] “绝不泄露”" in out
    assert "AI 生成" in out

    plain = tmp_path / "无前缀.md"
    plain.write_text("画面：录屏\n你还在手动回复客户吗？点击下方链接免费试用。\n", encoding="utf-8")
    ok = _run(CHECK, plain, "--target", "5")
    assert ok.returncode == 0, ok.stdout
    assert "✓ 全文" in ok.stdout

    assert _run(CHECK, tmp_path / "missing.md").returncode == 1


# ------------------------------------------------------------------ matrix


def test_matrix_is_balanced_and_named(tmp_path: Path) -> None:
    out = tmp_path / "矩阵.xlsx"
    proc = _run(
        AD_TEST,
        "matrix",
        "--product",
        "小智",
        "--hooks",
        "痛点,数字,反常识,对比",
        "--points",
        "省时，降本,零门槛",
        "--ctas",
        "免费试用,预约演示",
        "--max",
        "8",
        "--date",
        "1022",
        "-o",
        out,
    )
    assert proc.returncode == 0, proc.stderr
    ws = load_workbook(out)["测试矩阵"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 8
    names = [r[1] for r in rows]
    assert names[0] == "1022_小智_痛点_省时_免费试用_30s_01"
    assert len(set(names)) == 8
    hooks = [r[2] for r in rows]
    assert {hooks.count(h) for h in set(hooks)} == {2}  # 4 hooks x 2 each
    assert {r[3] for r in rows} == {"省时", "降本", "零门槛"}
    assert {r[4] for r in rows} == {"免费试用", "预约演示"}

    assert (
        _run(
            AD_TEST, "matrix", "--product", "a_b", "--hooks", "x", "--points", "y", "--ctas", "z"
        ).returncode
        == 1
    )


# ------------------------------------------------------------------ review


def test_review_example_report() -> None:
    proc = _run(AD_TEST, "review", EXAMPLE_REPORT, "--json")
    assert proc.returncode == 0, proc.stderr
    data = json.loads(proc.stdout)
    verdicts = {m["素材"].rsplit("_", 1)[-1]: m["结论"] for m in data["materials"]}
    assert verdicts == {
        "01": "放量",
        "05": "放量",
        "02": "改开头",
        "03": "改承接",
        "07": "关停",
        "04": "关停",
        "06": "数据不足",
    }
    first = data["materials"][0]
    assert first["展示"] == 100000 and first["转化"] == 50  # two days merged
    assert data["overall"]["转化成本"] == 158.52
    hooks = data["dimensions"]["钩子"]
    assert hooks[0]["取值"] == "痛点"
    assert set(data["dimensions"]) == {"钩子", "卖点", "引导"}


def test_review_xlsx_with_other_headers(tmp_path: Path) -> None:
    src = tmp_path / "腾讯广告导出.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["广告报表 2026-10-21"])
    ws.append(["创意名称", "曝光量", "点击量", "花费（元）", "表单数"])
    ws.append(["A版", 10000, 300, "1,200.00", 10])
    ws.append(["B版", 10000, 100, "1,200.00", 1])
    wb.save(src)
    out = tmp_path / "复盘.xlsx"
    proc = _run(AD_TEST, "review", src, "-o", out)
    assert proc.returncode == 0, proc.stderr
    assert "[放量] A版" in proc.stdout
    assert load_workbook(out).sheetnames == ["素材排行", "维度汇总", "整体"]

    bad = tmp_path / "bad.csv"
    bad.write_text("名字,金额\nx,1\n", encoding="utf-8")
    proc = _run(AD_TEST, "review", bad)
    assert proc.returncode == 1
    assert "没有识别到这些列" in proc.stderr
