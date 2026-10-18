"""Contract tests for the bundled NAS / file experts (NAS 管家 / 文件整理 / 电子书整理)."""

from __future__ import annotations

import csv
import datetime as dt
import importlib.util
import json
import os
import re
import struct
import subprocess
import sys
import zipfile
from pathlib import Path
from types import ModuleType

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
HEALTH = LIBRARY / "nas-manager" / "skills" / "nas-health"
COMPOSE = LIBRARY / "nas-manager" / "skills" / "docker-compose"
FILES = LIBRARY / "file-organizer" / "skills" / "file-toolkit" / "scripts"
EBOOKS = LIBRARY / "ebook-librarian" / "skills" / "ebook-toolkit" / "scripts"

EXPERTS = {
    "nas-manager": ("NAS 管家", {"nas-health", "docker-compose", "backup-strategy"}),
    "file-organizer": ("文件整理专家", {"file-toolkit", "organize-rules"}),
    "ebook-librarian": ("电子书整理专家", {"ebook-toolkit", "reading-list"}),
}


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


def _load(path: Path) -> ModuleType:
    sys.dont_write_bytecode = True
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))  # the tools import their sibling media_ops.py
    spec = importlib.util.spec_from_file_location(f"_nas_{path.stem}", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _plan(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


def test_media_ops_copies_are_identical() -> None:
    copies = sorted(LIBRARY.glob("*/skills/*/scripts/media_ops.py"))
    assert len(copies) >= 2
    first = copies[0].read_bytes()
    for copy in copies[1:]:
        assert copy.read_bytes() == first, f"{copy} drifted from {copies[0]}"


# ----------------------------------------------------------------------- nas


def test_smart_grades_hdd_and_nvme() -> None:
    script = HEALTH / "scripts" / "nas_check.py"
    refs = HEALTH / "references"
    result = _run(script, "smart", refs / "example-smart-hdd.txt", refs / "example-smart-nvme.json")
    assert result.returncode == 3, result.stderr
    out = result.stdout
    assert (
        "| example-smart-hdd | WDC WD40EFRX-68N32N0（****0001） | 4.00 TB | ATA | 37°C | 45210 小时 | 关注 |"
        in out
    )
    assert (
        "| example-smart-nvme | Samsung SSD 980 PRO 1TB（****0001） | 1.00 TB | NVMe | 41°C | 5432 小时 | 正常 |"
        in out
    )
    assert "重映射扇区 8 个" in out and "UDMA CRC 错误 3 次" in out and "约 5.2 年" in out
    assert "WD-WCC7K0000001" not in out  # serials are masked

    as_json = json.loads(_run(script, "smart", refs / "example-smart-nvme.json", "--json").stdout)
    assert as_json[0]["nvme"]["Percentage Used"] == 3 and as_json[0]["level"] == "正常"


def test_smart_failed_and_worn_nvme(tmp_path: Path) -> None:
    hdd = tmp_path / "bad.txt"
    hdd.write_text(
        (HEALTH / "references" / "example-smart-hdd.txt")
        .read_text(encoding="utf-8")
        .replace("result: PASSED", "result: FAILED!")
        .replace(
            "Current_Pending_Sector  0x0032   200   200   000    Old_age   Always       -       0",
            "Current_Pending_Sector  0x0032   200   200   000    Old_age   Always       -       16",
        ),  # fmt: skip
        encoding="utf-8",
    )
    nvme = tmp_path / "worn.json"
    data = json.loads(
        (HEALTH / "references" / "example-smart-nvme.json").read_text(encoding="utf-8")
    )
    data["nvme_smart_health_information_log"].update(percentage_used=85, available_spare=5)
    nvme.write_text(json.dumps(data), encoding="utf-8")
    out = _run(HEALTH / "scripts" / "nas_check.py", "smart", hdd, nvme).stdout
    assert "bad（危险）" in out and "待映射扇区 16 个" in out and "FAILED" in out
    assert "worn（危险）" in out and "已用寿命 85%" in out and "可用备用块 5%" in out


def test_capacity_forecast_and_record(tmp_path: Path) -> None:
    mod = _load(HEALTH / "scripts" / "nas_check.py")
    d0 = dt.date(2026, 9, 1)
    pts = [(d0, 500), (d0 + dt.timedelta(days=10), 600), (d0 + dt.timedelta(days=20), 700)]
    slope, d90, d100 = mod.forecast(pts, 1000)
    assert slope == pytest.approx(10.0) and d90 == 20 and d100 == 30
    assert mod.forecast(pts[:1], 1000) is None
    assert mod.forecast([(d0, 500), (d0 + dt.timedelta(days=3), 600)], 1000) is None  # < 7 days

    history = tmp_path / "hist.csv"
    script = HEALTH / "scripts" / "nas_check.py"
    first = _run(
        script, "capacity", tmp_path, "--history", history, "--record", "--today", "2026-10-01"
    )
    assert first.returncode in (0, 3), first.stderr
    assert "历史记录不足 7 天" in first.stdout
    again = _run(
        script, "capacity", tmp_path, "--history", history, "--record", "--today", "2026-10-01"
    )
    assert again.returncode in (0, 3)
    rows = list(csv.DictReader(history.open(encoding="utf-8-sig")))
    assert len(rows) == 1 and rows[0]["日期"] == "2026-10-01"  # same day isn't recorded twice
    assert _run(script, "capacity", tmp_path / "missing").returncode == 1


def test_compose_check_finds_common_problems(tmp_path: Path) -> None:
    script = COMPOSE / "scripts" / "compose_check.py"
    result = _run(script, COMPOSE / "references" / "example-compose.yml", "--json")
    assert result.returncode == 3
    issues = json.loads(result.stdout)
    text = "\n".join(f"[{i['等级']}] {i['服务']}：{i['问题']}" for i in issues)
    assert "[错误] portainer、qbittorrent：宿主机端口 8080 被多个服务占用" in text
    assert "[警告] mysql：镜像「mysql」没有写版本标签" in text
    assert "[警告] mysql：端口 3306:3306 把 MySQL 暴露在所有网卡上" in text
    assert "[警告] mysql：环境变量 MYSQL_ROOT_PASSWORD 直接写了明文值" in text
    assert (
        "[警告] portainer：privileged: true" in text
        and "[警告] portainer：挂载了 /var/run/docker.sock" in text
    )
    assert "[警告] qbittorrent：没有设置重启策略" in text
    assert (
        "[建议] qbittorrent：该镜像支持 PUID/PGID" in text
        and "[建议] qbittorrent：相对路径挂载" in text
    )
    assert not [i for i in issues if i["服务"] == "jellyfin"]  # the well-formed service is clean

    good = tmp_path / "ok.yml"
    good.write_text(
        "services:\n  a:\n    image: nginx:1.27\n    restart: unless-stopped\n    environment:\n      TZ: Asia/Shanghai\n"
        "      DB_PASSWORD: ${DB_PASSWORD}\n    ports:\n      - '127.0.0.1:3306:3306'\n      - '8080:80'\n"
        "    healthcheck:\n      test: ['CMD', 'true']\n  b:\n    image: nginx:1.27\n    restart: always\n"
        "    environment: [TZ=Asia/Shanghai]\n    ports: ['127.0.0.2:8080:80']\n    healthcheck:\n      test: ['CMD', 'true']\n",
        encoding="utf-8",
    )
    result = _run(script, good)
    assert "端口 8080 被多个服务占用" in result.stdout  # 0.0.0.0:8080 clashes with 127.0.0.2:8080
    assert "明文" not in result.stdout and "MySQL" not in result.stdout

    broken = tmp_path / "broken.yml"
    broken.write_text("services:\n  a:\n   image: x\n  - bad\n", encoding="utf-8")
    assert "YAML 无法解析" in _run(script, broken).stdout


# --------------------------------------------------------------------- files


def test_file_organizer_scan_dupes_plan_apply_undo(tmp_path: Path) -> None:
    root = tmp_path / "下载"
    (root / "sub").mkdir(parents=True)
    (root / "报告.docx").write_text("a", encoding="utf-8")
    (root / "sub" / "报告副本.docx").write_text("a", encoding="utf-8")
    (root / "数据.xlsx").write_text("bb", encoding="utf-8")
    (root / "sub" / "数据.xlsx").write_text("cc", encoding="utf-8")
    (root / "setup.exe").write_text("d", encoding="utf-8")
    (root / "照片.JPG").write_text("e", encoding="utf-8")
    (root / "~$报告.docx").write_text("f", encoding="utf-8")
    old = dt.datetime(2024, 1, 5).timestamp()
    for name in ("数据.xlsx", "setup.exe"):
        os.utime(root / name, (old, old))
    tool, ops = FILES / "file_tool.py", FILES / "media_ops.py"

    scan = _run(tool, "scan", root, "--today", "2026-10-18")
    assert scan.returncode == 0, scan.stderr
    assert "7 个文件" in scan.stdout and "| 文档 | 3 |" in scan.stdout
    assert "一年以上未修改：2 个" in scan.stdout and "临时或系统文件：1 个" in scan.stdout

    dupes = tmp_path / "d.csv"
    assert _run(tool, "dupes", root, "-o", dupes).returncode == 0
    rows = _plan(dupes)
    assert len(rows) == 1 and rows[0]["源路径"].endswith("报告副本.docx")
    assert Path(rows[0]["目标路径"]) == root / "_重复文件待确认" / "sub" / "报告副本.docx"

    plan = tmp_path / "p.csv"
    result = _run(
        tool,
        "plan",
        root,
        "--dest",
        root / "已整理",
        "-o",
        plan,
        "--by",
        "type-date",
        "--today",
        "2026-10-18",
    )
    assert result.returncode == 0, result.stderr
    assert "跳过 1 个临时或系统文件" in result.stdout
    targets = {
        Path(r["源路径"]).relative_to(root).as_posix(): Path(r["目标路径"]) for r in _plan(plan)
    }
    assert targets["setup.exe"] == root / "已整理" / "安装包" / "2024" / "setup.exe"
    assert targets["照片.JPG"].name == "照片.jpg"
    assert {targets["数据.xlsx"].parent.name, targets["sub/数据.xlsx"].parent.name} == {
        "2024",
        "2026",
    }

    applied = _run(ops, "apply", plan, "--root", root, "--prune-empty")
    assert applied.returncode == 0, applied.stderr
    assert (root / "已整理" / "表格" / "2024" / "数据.xlsx").read_text(encoding="utf-8") == "bb"
    assert (root / "~$报告.docx").exists() and not (root / "sub").exists()
    again = tmp_path / "p2.csv"
    _run(tool, "plan", root, "--dest", root / "已整理", "-o", again)
    assert _plan(again) == []  # the destination is never re-planned

    log = next((root / ".octop-media-undo").glob("*.csv"))
    assert _run(ops, "undo", log).returncode == 0
    assert (root / "sub" / "数据.xlsx").read_text(encoding="utf-8") == "cc"

    recent = tmp_path / "p3.csv"
    result = _run(
        tool,
        "plan",
        root,
        "--dest",
        root / "归档",
        "-o",
        recent,
        "--older-than",
        "90",
        "--today",
        "2026-10-18",
    )
    assert {Path(r["源路径"]).name for r in _plan(recent)} == {"数据.xlsx", "setup.exe"}
    assert "近 90 天修改过" in result.stdout


# -------------------------------------------------------------------- ebooks


def _epub(path: Path, title: str, creator: str, lang: str = "zh") -> None:
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        z.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            f"<dc:title>{title}</dc:title><dc:creator>{creator}</dc:creator><dc:language>{lang}</dc:language>"
            "<dc:date>2019-05-01</dc:date><dc:identifier>urn:isbn:9787000000001</dc:identifier></metadata></package>",
        )


def _mobi(path: Path, title: str, author: str) -> None:
    t, a = title.encode(), author.encode()
    exth = b"EXTH" + struct.pack(">II", 12 + 8 + len(a) + 8 + len(t), 2)
    exth += struct.pack(">II", 100, 8 + len(a)) + a + struct.pack(">II", 503, 8 + len(t)) + t
    hdr_len = 232
    rec = bytearray(16) + b"MOBI" + struct.pack(">III", hdr_len, 2, 65001) + bytes(hdr_len - 16)
    struct.pack_into(">II", rec, 84, 16 + hdr_len + len(exth), len(t))
    struct.pack_into(">I", rec, 128, 0x40)
    rec += exth + t
    head = (
        b"x" * 32
        + bytes(28)
        + b"BOOKMOBI"
        + bytes(8)
        + struct.pack(">H", 1)
        + struct.pack(">II", 86, 0)
    )
    path.write_bytes(head + bytes(rec))


def _library(root: Path) -> None:
    from pypdf import PdfWriter

    (root / "杂").mkdir(parents=True)
    _epub(root / "abc123.epub", "三体", "[中] 刘慈欣 著")
    _epub(root / "杂" / "book.epub", "Sapiens", "Harari, Yuval Noah", "en")
    (root / "三体副本.epub").write_bytes((root / "abc123.epub").read_bytes())
    _mobi(root / "santi2.azw3", "三体Ⅱ·黑暗森林", "刘慈欣")
    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.add_metadata({"/Title": "Microsoft Word - 草稿.docx", "/Author": "Administrator"})
    with (root / "余华 - 活着.pdf").open("wb") as f:
        w.write(f)
    (root / "《平凡的世界》路遥.txt").write_bytes("平凡的世界".encode("gb18030"))
    (root / "杂" / "broken.epub").write_bytes(b"not a zip")


def test_ebook_metadata_and_filename_rules() -> None:
    mod = _load(EBOOKS / "ebook_tool.py")
    cases = {
        "活着 - 余华": ("活着", "余华"),
        "余华 - 活着": ("活着", "余华"),
        "钱钟书 - 围城": ("围城", "钱钟书"),
        "百年孤独 - 加西亚·马尔克斯": ("百年孤独", "加西亚·马尔克斯"),
        "George Orwell - 1984": ("1984", "George Orwell"),
        "明朝那些事儿 - 当年明月": ("明朝那些事儿", "当年明月"),
        "《围城》钱钟书": ("围城", "钱钟书"),
        "[东野圭吾]白夜行": ("白夜行", "东野圭吾"),
        "人类简史(尤瓦尔·赫拉利 著)": ("人类简史", "尤瓦尔·赫拉利"),
        "z-lib.org 原则 (epub)": ("原则", ""),
    }
    assert {k: mod.from_filename(k) for k in cases} == cases
    assert mod.norm_author("[美] 卡尔·萨根 著；李四 译") == "卡尔·萨根"
    assert mod.norm_author("Harari, Yuval Noah") == "Yuval Noah Harari"
    assert mod.norm_author("刘慈欣、王晋康") == "刘慈欣 等"
    assert mod.clean("Microsoft Word - 草稿.docx") == "" and mod.clean("Administrator") == ""


def test_ebook_scan_plan_dupes_catalog(tmp_path: Path) -> None:
    root = tmp_path / "books"
    _library(root)
    tool, ops = EBOOKS / "ebook_tool.py", EBOOKS / "media_ops.py"

    scan = _run(tool, "scan", root)
    assert scan.returncode == 0, scan.stderr
    assert "7 本" in scan.stdout and "epub 4" in scan.stdout
    assert "同一本书同一格式有多个文件 1 组" in scan.stdout
    assert "GBK 编码的 txt 1 本" in scan.stdout

    plan = tmp_path / "plan.csv"
    dest = root / "书库"
    assert _run(tool, "plan", root, "--dest", dest, "-o", plan).returncode == 0
    rows = {Path(r["源路径"]).name: r for r in _plan(plan)}
    assert Path(rows["abc123.epub"]["目标路径"]) == dest / "刘慈欣" / "刘慈欣 - 三体.epub"
    assert Path(rows["三体副本.epub"]["目标路径"]).name == "刘慈欣 - 三体 (2).epub"
    assert Path(rows["santi2.azw3"]["目标路径"]) == dest / "刘慈欣" / "刘慈欣 - 三体Ⅱ·黑暗森林.azw3"
    assert (
        Path(rows["book.epub"]["目标路径"])
        == dest / "Yuval Noah Harari" / "Yuval Noah Harari - Sapiens.epub"
    )
    assert Path(rows["余华 - 活着.pdf"]["目标路径"]) == dest / "余华" / "余华 - 活着.pdf"
    assert rows["余华 - 活着.pdf"]["说明"].startswith("文件名")  # junk PDF metadata ignored
    assert Path(rows["broken.epub"]["目标路径"]).parent.name == "未知作者"

    assert _run(ops, "apply", plan, "--root", root, "--prune-empty").returncode == 0
    assert (dest / "刘慈欣" / "刘慈欣 - 三体.epub").exists() and not (root / "杂").exists()

    dupes = tmp_path / "d.csv"
    assert _run(tool, "dupes", root, "-o", dupes).returncode == 0
    assert [Path(r["源路径"]).name for r in _plan(dupes)] == ["刘慈欣 - 三体 (2).epub"]

    cat = tmp_path / "list.xlsx"
    result = _run(tool, "catalog", root, "-o", cat)
    assert result.returncode == 0, result.stderr
    ws = load_workbook(cat)["书单"]
    assert [c.value for c in ws[1]][:3] == ["书名", "作者", "格式"]
    titles = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert titles["三体"][7] == "9787000000001" and titles["三体"][6] == "2019"

    custom = tmp_path / "flat.csv"
    assert (
        _run(
            tool,
            "plan",
            root,
            "--dest",
            tmp_path / "flat",
            "-o",
            custom,
            "--layout",
            "{lang}/{title}",
        ).returncode
        == 0
    )
    assert (
        _run(
            tool, "plan", root, "--dest", tmp_path / "x", "-o", custom, "--layout", "{isbn}"
        ).returncode
        == 1
    )
