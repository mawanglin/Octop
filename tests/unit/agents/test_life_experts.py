"""Contract tests for the bundled life experts (记账 / 旅行 / 健身 / 作业辅导)."""

from __future__ import annotations

import importlib.util
import json
import os
import re
import subprocess
import sys
from collections import Counter
from fractions import Fraction
from pathlib import Path
from types import ModuleType

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from docx import Document
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
BILLS = LIBRARY / "home-bookkeeper" / "skills" / "bill-import"
TRIP = LIBRARY / "travel-planner" / "skills" / "trip-sheets"
FITNESS = LIBRARY / "fitness-coach" / "skills" / "nutrition" / "scripts" / "fitness.py"
PLAN = LIBRARY / "fitness-coach" / "skills" / "training-plan" / "references" / "example-plan.md"
MISTAKES = LIBRARY / "homework-tutor" / "skills" / "mistake-book"
DRILL = LIBRARY / "homework-tutor" / "skills" / "practice-sheets" / "scripts" / "math_drill.py"

EXPERTS = {
    "home-bookkeeper": ("家庭记账专家", {"bill-import", "budget-review"}),
    "travel-planner": ("旅行规划专家", {"itinerary", "trip-sheets"}),
    "fitness-coach": ("健身教练", {"training-plan", "nutrition"}),
    "homework-tutor": ("作业辅导老师", {"step-by-step", "mistake-book", "practice-sheets"}),
}


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


def _load(path: Path) -> ModuleType:
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location(f"_life_{path.stem}", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


# ------------------------------------------------------------------- ledger


def test_ledger_imports_alipay_and_wechat_and_reports(tmp_path: Path) -> None:
    script = BILLS / "scripts" / "ledger.py"
    refs = BILLS / "references"
    ledger = tmp_path / "ledger.xlsx"
    result = _run(
        script,
        "import",
        refs / "example-alipay.csv",  # GBK with a preamble before the header
        refs / "example-wechat.csv",
        "-o",
        ledger,
        "--rules",
        refs / "example-rules.csv",
    )
    assert result.returncode == 0, result.stderr
    assert "识别为支付宝账单，读取 24 笔" in result.stdout
    assert "识别为微信账单，读取 15 笔" in result.stdout
    assert "跳过 3 笔（状态：交易关闭）" in result.stdout
    assert "共 39 笔" in result.stdout and "“其他”类 0 笔" in result.stdout
    ws = load_workbook(ledger)["账本"]
    assert [c.value for c in ws[1]][:6] == ["日期", "时间", "来源", "收支", "金额", "类别"]
    cats = Counter(r[5] for r in ws.iter_rows(min_row=2, values_only=True))
    assert cats["工资"] == 3 and cats["转账"] == 6  # red packets / transfers are not spending
    assert cats["运动"] == 3  # from the custom rules file

    again = _run(
        script, "import", refs / "example-alipay.csv", "-o", tmp_path / "l2.xlsx", "--merge", ledger
    )
    assert again.returncode == 0, again.stderr
    assert "共 39 笔" in again.stdout and "去重 24 笔" in again.stdout

    report = tmp_path / "report.xlsx"
    result = _run(script, "report", ledger, "-o", report, "--budget", refs / "example-budget.csv")
    assert result.returncode == 0, result.stderr
    out = result.stdout
    assert "收入 45,000.00 元，消费支出 4,337.70 元" in out
    assert "支出前三：餐饮 2,179、运动 597、居住 469" in out  # ranked by amount, not count
    assert "疑似订阅扣费 3 项" in out and "大额异常支出 1 笔" in out
    assert "超预算：餐饮 月均 726 元" in out
    wb = load_workbook(report)
    assert wb.sheetnames == ["月度汇总", "分类×月份", "Top商户", "疑似订阅", "大额支出", "预算对比"]
    subs = {r[0] for r in wb["疑似订阅"].iter_rows(min_row=2, values_only=True)}
    assert subs == {"某健身房", "中国移动", "爱奇艺"}
    assert len(wb["月度汇总"]._charts) == 1


def test_ledger_rejects_unknown_layout(tmp_path: Path) -> None:
    bad = tmp_path / "bad.csv"
    bad.write_text("名称,数量\n苹果,3\n", encoding="utf-8")
    result = _run(BILLS / "scripts" / "ledger.py", "import", bad, "-o", tmp_path / "x.xlsx")
    assert result.returncode == 1 and "表头" in result.stderr


# --------------------------------------------------------------------- trip


def test_trip_sheet_flags_overlap_and_budget(tmp_path: Path) -> None:
    out = tmp_path / "trip.xlsx"
    result = _run(
        TRIP / "scripts" / "trip.py", "sheet", TRIP / "references" / "example-trip.md", "-o", out
    )
    assert result.returncode == 0, result.stderr
    assert "预计花费 4,070 元（人均 2,035 元）" in result.stdout
    assert "「看熊猫」与「午餐」时间重叠" in result.stdout
    assert "超出预算 4,000 元 70 元" in result.stdout
    assert "没有住宿" not in result.stdout  # "两晚" covers both nights
    wb = load_workbook(out)
    assert wb.sheetnames == ["行程", "每日费用", "分类预算"]
    assert len(wb["分类预算"]._charts) == 1


def test_trip_packing_list(tmp_path: Path) -> None:
    out = tmp_path / "pack.xlsx"
    result = _run(
        TRIP / "scripts" / "trip.py",
        "packing",
        "--days",
        "10",
        "--climate",
        "cold,rain",
        "--type",
        "abroad",
        "-o",
        out,
    )
    assert result.returncode == 0, result.stderr
    assert "护照" in result.stdout and "内衣袜子 × 7" in result.stdout
    ws = load_workbook(out)["行李清单"]
    assert [c.value for c in ws[1]] == ["分类", "物品", "已准备"]


# ------------------------------------------------------------------ fitness


def test_fitness_tdee_and_low_bmi_warning() -> None:
    result = _run(FITNESS, "tdee", "--sex", "m", "--age", "30", "--height", "175", "--weight", "75",
                  "--activity", "moderate", "--goal", "gain")  # fmt: skip
    assert result.returncode == 0, result.stderr
    assert "BMI 24.5（超重" in result.stdout
    assert "基础代谢 约 1699 千卡" in result.stdout and "建议每日摄入 约 2896 千卡" in result.stdout
    assert "蛋白质 135 克" in result.stdout

    low = _run(
        FITNESS,
        "tdee",
        "--sex",
        "f",
        "--age",
        "25",
        "--height",
        "165",
        "--weight",
        "50",
        "--goal",
        "lose",
    )
    assert low.returncode == 0 and "不建议继续减重" in low.stdout

    child = _run(FITNESS, "tdee", "--sex", "f", "--age", "10", "--height", "140", "--weight", "35")
    assert child.returncode == 1 and "咨询医生" in child.stderr


def test_fitness_meal_and_plan(tmp_path: Path) -> None:
    result = _run(
        FITNESS, "meal", "米饭 1碗, 鸡胸肉 150g, 鸡蛋 2个, 西兰花 200g, 食用油 10g, 榴莲酥 1个"
    )
    assert result.returncode == 0, result.stderr
    assert "| **合计** | | **679** | **62.3** |" in result.stdout
    assert "未计入（不在内置数据中）：榴莲酥" in result.stdout

    out = tmp_path / "log.xlsx"
    result = _run(FITNESS, "plan", PLAN, "-o", out, "--weeks", "6")
    assert result.returncode == 0, result.stderr
    header = [c.value for c in load_workbook(out)["训练打卡"][1]]
    assert header[:2] == ["训练日", "动作"] and header[-1] == "第6周"


# ------------------------------------------------------------------- tutor


def test_mistake_book_build_merge_and_due(tmp_path: Path) -> None:
    script = MISTAKES / "scripts" / "mistakes.py"
    book = tmp_path / "book.xlsx"
    result = _run(script, "build", MISTAKES / "references" / "example-mistakes.md", "-o", book)
    assert result.returncode == 0, result.stderr
    assert "共 6 道错题（本次新增 6 道）" in result.stdout
    assert "有余数除法 2 道" in result.stdout
    wb = load_workbook(book)
    ws = wb["错题"]
    header = [c.value for c in ws[1]]
    assert header[-6:] == ["复习1", "复习2", "复习3", "复习4", "复习5", "已掌握"]
    first = [c.value for c in ws[2]]
    assert first[1].strftime("%m-%d") == "10-03"
    assert [d.strftime("%m-%d") for d in first[9:14]] == [
        "10-04",
        "10-05",
        "10-07",
        "10-10",
        "10-18",
    ]

    ws.cell(2, header.index("已掌握") + 1).value = "是"
    wb.save(book)
    merged = tmp_path / "merged.xlsx"
    result = _run(
        script,
        "build",
        MISTAKES / "references" / "example-mistakes.md",
        "-o",
        merged,
        "--merge",
        book,
    )
    assert result.returncode == 0, result.stderr
    assert "共 6 道错题（本次新增 0 道）" in result.stdout

    due = _run(script, "due", merged, "--today", "2026-10-18")
    assert due.returncode == 0, due.stderr
    assert "待复习 5 道" in due.stdout  # the mastered one is skipped
    assert "36 × 24" not in due.stdout
    assert "[数学·有余数除法] 第 4 轮：59 ÷ 7" in due.stdout

    later = _run(script, "due", merged, "--today", "2026-10-21")
    assert "[语文·形近字] 第 3 轮（已过 1 天）" in later.stdout  # missed by a day, still listed

    bad = tmp_path / "bad.md"
    bad.write_text("| 学科 | 题目 |\n|---|---|\n| 数学 | 1+1 |\n", encoding="utf-8")
    result = _run(script, "build", bad, "-o", tmp_path / "x.xlsx")
    assert result.returncode == 1 and "缺少「日期」" in result.stderr


def test_math_drill_answers_are_correct() -> None:
    drill = _load(DRILL)
    for kind in drill.TYPES:
        for seed in range(5):
            problems = drill.generate(kind, 30, 100, "any", seed)
            assert len(problems) == 30, kind
            for question, answer in problems:
                expr = question.removesuffix(" =")
                if kind == "equation":
                    x = int(answer.removeprefix("x = ").replace("−", "-"))
                    left, right = expr.replace("−", "-").split(" = ")
                    value = eval(re.sub(r"(\d)x", r"\1*x", left), {"x": x})  # noqa: S307
                    assert value == int(right), question
                elif kind == "divr":
                    a, b = map(int, expr.split(" ÷ "))
                    q, r = map(int, answer.split("……"))
                    assert a == b * q + r and 0 < r < b, question
                elif kind == "fraction":
                    assert drill._eval(expr.replace("/", "÷")) == Fraction(answer.split("（")[0]), (
                        question
                    )
                elif kind == "decimal":
                    assert abs(eval(expr.replace("−", "-")) - float(answer)) < 1e-9, question  # noqa: S307
                else:
                    value = drill._eval(expr)
                    assert value == int(answer) and value >= 0, question

    assert all(drill.has_carry(*map(int, q.removesuffix(" =").split(" + ")))
               for q, _ in drill.generate("add", 30, 100, "require", 1))  # fmt: skip
    assert not any(drill.has_borrow(*map(int, q.removesuffix(" =").split(" − ")))
                   for q, _ in drill.generate("sub", 30, 100, "avoid", 1))  # fmt: skip


def test_math_drill_exports_word_with_answer_page(tmp_path: Path) -> None:
    out = tmp_path / "drill.docx"
    result = _run(
        DRILL,
        "--type",
        "mixed",
        "-n",
        "20",
        "--seed",
        "3",
        "--cols",
        "2",
        "--title",
        "混合运算",
        "-o",
        out,
    )
    assert result.returncode == 0, result.stderr
    doc = Document(str(out))
    assert doc.paragraphs[0].text == "混合运算"
    assert len(doc.tables) == 2
    questions = [c.text for row in doc.tables[0].rows for c in row.cells if c.text]
    answers = [c.text for row in doc.tables[1].rows for c in row.cells if c.text]
    assert len(questions) == len(answers) == 20
    assert "w:br" in doc.element.xml and 'w:type="page"' in doc.element.xml

    again = _run(DRILL, "--type", "mixed", "-n", "20", "--seed", "3")
    assert again.stdout.splitlines()[0].startswith("1. ")
    assert again.stdout == _run(DRILL, "--type", "mixed", "-n", "20", "--seed", "3").stdout

    bad = _run(DRILL, "--type", "add", "-n", "500")
    assert bad.returncode == 1
