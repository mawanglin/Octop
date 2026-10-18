"""Contract tests for the bundled health / learning / pet experts (体检 / 英语 / 宠物)."""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from docx import Document
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
CHECKUP = LIBRARY / "health-report" / "skills" / "checkup-compare"
VOCAB = LIBRARY / "english-coach" / "skills" / "vocab-book"
PET = LIBRARY / "pet-care" / "skills" / "care-schedule" / "scripts" / "pet_care.py"

EXPERTS = {
    "health-report": ("体检报告解读", {"checkup-compare", "health-habits"}),
    "english-coach": ("英语学习教练", {"vocab-book", "english-practice"}),
    "pet-care": ("宠物照护", {"care-schedule", "pet-health"}),
}


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


# ------------------------------------------------------------------- checkup


def test_checkup_wide_table_flags_and_trends(tmp_path: Path) -> None:
    script = CHECKUP / "scripts" / "checkup_compare.py"
    out = tmp_path / "ck.xlsx"
    result = _run(script, CHECKUP / "references" / "example-checkup.csv", "-o", out)
    assert result.returncode == 0, result.stderr
    text = result.stdout
    assert "10 项指标，3 次体检（2024、2025、2026）" in text
    assert "## 最近一次异常（6）" in text and "## 连续两次及以上异常（3）" in text
    assert "## 连续三次越来越偏离参考范围（2）" in text
    assert "- 高密度脂蛋白：↓ 1.0 mmol/L（参考 >1.04）" in text
    assert "- 尿酸：450 → 470 → 440 μmol/L" in text
    assert (
        "幽门螺杆菌" not in text.split("## 最近一次异常")[1].split("## 连续")[0]
    )  # positive in 2025 only
    assert "不能代替医生诊断" in text
    ws = load_workbook(out)["历年对比"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["总胆固醇"][5] == "5.9 ↑" and rows["幽门螺杆菌"][4] == "阳性 阳性"
    assert ws["F2"].fill.fgColor.rgb.endswith("FECACA")

    data = json.loads(_run(script, CHECKUP / "references" / "example-checkup.csv", "--json").stdout)
    assert {r["指标"] for r in data["持续变差"]} == {"总胆固醇", "甘油三酯"}


def test_checkup_long_table_and_bad_range(tmp_path: Path) -> None:
    script = CHECKUP / "scripts" / "checkup_compare.py"
    long = tmp_path / "long.csv"
    long.write_text(
        "年份,指标,结果,单位,参考范围\n2025,ALT,45,U/L,9-50\n2026,ALT,72,U/L,9-50\n2026,尿蛋白,+,,阴性\n2026,血红蛋白,138,g/L,130～175\n",
        encoding="utf-8",
    )
    text = _run(script, long).stdout
    assert "- ALT：↑ 72 U/L（参考 9-50） —— 肝细胞损伤的敏感指标" in text
    assert (
        "- 尿蛋白：阳性 +" in text
        and "血红蛋白" not in text.split("## 最近一次异常")[1].split("## 连续")[0]
    )

    bad = tmp_path / "bad.csv"
    bad.write_text("指标,单位,参考范围,2026\n血糖,mmol/L,正常偏低,5\n", encoding="utf-8")
    result = _run(script, bad)
    assert result.returncode == 1 and "无法识别" in result.stderr


# --------------------------------------------------------------------- vocab


def test_vocab_build_merge_due_and_quiz(tmp_path: Path) -> None:
    script = VOCAB / "scripts" / "vocab.py"
    book = tmp_path / "vb.xlsx"
    result = _run(
        script,
        "build",
        VOCAB / "references" / "example-words.md",
        "-o",
        book,
        "--today",
        "2026-10-11",
    )
    assert result.returncode == 0, result.stderr
    assert "共 5 个单词（本次新增 5 个）" in result.stdout
    ws = load_workbook(book)["单词本"]
    header = [c.value for c in ws[1]]
    assert header[-7:] == ["复习1", "复习2", "复习3", "复习4", "复习5", "复习6", "已掌握"]
    ws.cell(2, header.index("已掌握") + 1).value = "是"  # ambiguous (first alphabetically)
    ws.parent.save(book)

    extra = tmp_path / "w2.md"
    extra.write_text(
        "| 单词 | 释义 |\n|---|---|\n| Deadline | n. 期限 |\n| itinerary | n. 行程 |\n",
        encoding="utf-8",
    )
    merged = tmp_path / "vb2.xlsx"
    result = _run(script, "build", extra, "-o", merged, "--merge", book, "--today", "2026-10-17")
    assert (
        "共 6 个单词（本次新增 1 个）" in result.stdout
    )  # "Deadline" is a case-insensitive duplicate

    due = _run(script, "due", merged, "--today", "2026-10-18").stdout
    assert "待复习 5 个（单词本共 6 个，已掌握 1 个）" in due
    assert "ambiguous" not in due and "deadline /ˈdedlaɪn/ n. 截止日期（第 4 轮）" in due
    assert "itinerary n. 行程（第 1 轮）" in due

    quiz = tmp_path / "q.docx"
    result = _run(
        script,
        "quiz",
        merged,
        "-o",
        quiz,
        "-n",
        "4",
        "--mode",
        "zh2en",
        "--seed",
        "1",
        "--today",
        "2026-10-18",
    )
    assert result.returncode == 0, result.stderr
    paras = [p.text for p in Document(str(quiz)).paragraphs]
    assert paras[0] == "单词默写（4 题）"
    answers = paras[paras.index("答案") + 1 :]
    assert len(answers) == 4 and all("ambiguous" not in a for a in answers)
    assert all(re.match(r"\d\. .+：[A-Za-z]+$", a) for a in answers)


# ----------------------------------------------------------------------- pet


def test_pet_schedule_for_puppy(tmp_path: Path) -> None:
    out, ics = tmp_path / "pet.xlsx", tmp_path / "pet.ics"
    result = _run(PET, "plan", "--species", "dog", "--name", "旺财", "--birth", "2026-08-20", "--today", "2026-10-18",
                  "--days", "45", "-o", out, "--ics", ics)  # fmt: skip
    assert result.returncode == 0, result.stderr
    text = result.stdout
    assert "旺财（狗，8 周龄）" in text
    assert "- 2026-11-12：犬多联疫苗（核心苗）（首免第 2 针（12 周龄））" in text
    assert "另有 6 项日期已过" in text
    rows = list(load_workbook(out)["日程"].iter_rows(min_row=2, values_only=True))
    rabies = [r for r in rows if r[2] == "狂犬疫苗"]
    assert rabies[0][0].date().isoformat() == "2026-12-10" and rabies[0][3] == "首针（约 4 月龄）"
    week8 = {
        r[2]: r[4] for r in rows if r[0].date().isoformat() == "2026-10-15"
    }  # 8 weeks old: three items due
    assert week8 == dict.fromkeys(
        ("犬多联疫苗（核心苗）", "体内驱虫", "体外驱虫"), "已过（请确认是否完成）"
    )
    cal = ics.read_bytes().decode()
    assert cal.count("BEGIN:VEVENT") == sum(1 for r in rows if r[4] == "待办")
    assert "TRIGGER:-P1D" in cal


def test_pet_schedule_for_adult_with_gaps() -> None:
    text = _run(PET, "plan", "--species", "cat", "--name", "咪咪", "--birth", "2023-05-01", "--last-core", "2026-03-10",
                "--last-deworm-in", "2026-09-01", "--today", "2026-10-18", "--days", "60", "--months", "6").stdout  # fmt: skip
    assert "咪咪（猫，约 3.5 岁）" in text
    assert "- 2026-10-18：狂犬疫苗（没有上次接种记录：请尽快咨询兽医补种）" in text
    assert "- 2026-12-01：体内驱虫（每 3 个月）" in text
    assert "猫三联" not in text  # next booster (2027-03) is outside the 60-day window
    assert _run(PET, "plan", "--species", "cat").returncode == 1
    assert (
        _run(
            PET, "plan", "--species", "dog", "--birth", "2030-01-01", "--today", "2026-10-18"
        ).returncode
        == 1
    )
