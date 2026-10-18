"""Contract tests for the bundled work experts (项目管理 / 培训课程 / 电商运营)."""

from __future__ import annotations

import csv
import importlib.util
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from types import ModuleType

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from docx import Document
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
SCHEDULE = LIBRARY / "project-manager" / "skills" / "project-schedule"
EXAM = LIBRARY / "training-designer" / "skills" / "exam-builder"
LISTING = LIBRARY / "ecommerce-ops" / "skills" / "listing-check"

EXPERTS = {
    "project-manager": ("项目管理专家", {"project-schedule", "project-control"}),
    "training-designer": ("培训课程设计", {"course-design", "exam-builder"}),
    "ecommerce-ops": ("电商运营专家", {"listing-check", "store-operations"}),
}


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


def _load(path: Path) -> ModuleType:
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location(f"_work2_{path.stem}", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


# ------------------------------------------------------------------ schedule


def test_gantt_critical_path_holidays_and_progress(tmp_path: Path) -> None:
    script = SCHEDULE / "scripts" / "gantt.py"
    holidays = tmp_path / "hol.txt"
    holidays.write_text("# 节假日\n2026-11-11\n", encoding="utf-8")
    out = tmp_path / "plan.xlsx"
    result = _run(script, SCHEDULE / "references" / "example-tasks.md", "--start", "2026-11-02", "-o", out,
                  "--today", "2026-11-20", "--holidays", holidays)  # fmt: skip
    assert result.returncode == 0, result.stderr
    text = result.stdout
    assert "共 8 个任务，29 个工作日，2026-11-02 开始，预计 2026-12-11 完成" in text
    assert (
        "关键路径：1 需求调研 → 2 方案设计 → 4 后端开发 → 5 接口联调 → 7 系统测试 → 8 上线" in text
    )
    assert "3 前端开发（2 天）" in text and "6 测试用例编写（9 天）" in text
    assert "截至 2026-11-20（第 13 个工作日）：1 个任务落后" in text
    assert "6 测试用例编写（钱七）：计划应完成 50%，实际 0%，还有 9 天浮动时间" in text

    wb = load_workbook(out)
    rows = {r[0]: r for r in wb["排期"].iter_rows(min_row=2, values_only=True)}
    assert rows["6"][5].date().isoformat() == "2026-11-18"  # 2+2 lag, skipping the 11-11 holiday
    assert (
        rows["8"][5].date().isoformat() == rows["8"][6].date().isoformat() == "2026-12-11"
    )  # milestone
    assert rows["4"][10] == "是" and not rows["3"][10] and rows["3"][9] == 2
    g = wb["甘特图"]
    header = [c.value for c in g[1]]
    assert header[3] == "11-02" and header[-1] == "12-11"
    ms_row = next(r for r in g.iter_rows(min_row=3) if r[0].value == "8")
    assert ms_row[header.index("12-11")].value == "◆"
    holiday = g.cell(3, header.index("11-11") + 1)
    assert holiday.fill.fgColor.rgb.endswith("F3F4F6")


def test_gantt_rejects_bad_dependencies(tmp_path: Path) -> None:
    script = SCHEDULE / "scripts" / "gantt.py"
    cyc = tmp_path / "cyc.md"
    cyc.write_text(
        "| 编号 | 任务 | 工期 | 前置任务 |\n|---|---|---|---|\n| 1 | A | 2 | 2 |\n| 2 | B | 1 | 1 |\n",
        encoding="utf-8",
    )
    result = _run(script, cyc, "--start", "2026-11-02")
    assert result.returncode == 1 and "循环依赖：1 → 2 → 1" in result.stderr
    missing = tmp_path / "miss.md"
    missing.write_text(
        "| 编号 | 任务 | 工期 | 前置任务 |\n|---|---|---|---|\n| 1 | A | 2 | 9 |\n",
        encoding="utf-8",
    )
    assert "前置任务 9 不存在" in _run(script, missing, "--start", "2026-11-02").stderr


# ---------------------------------------------------------------------- exam


def test_exam_bank_check(tmp_path: Path) -> None:
    script = EXAM / "scripts" / "exam.py"
    ok = _run(script, "check", EXAM / "references" / "example-bank.csv")
    assert ok.returncode == 0, ok.stdout
    assert "题库共 18 题：单选 8、多选 4、判断 4、填空 1、简答 1" in ok.stdout

    bad = tmp_path / "bad.csv"
    bad.write_text(
        "题型,题目,选项,答案\n单选,一加一,1|2|3,D\n多选,选两个,a|b|c,A\n判断,天是蓝的,,也许\n填空,没有空的题,,x\n单选,一加一,1|2|3,B\n",
        encoding="utf-8",
    )
    out = _run(script, "check", bad).stdout
    for msg in ("第 2 行：答案「D」不在选项 ABC 内", "第 3 行：多选题答案只有一个", "第 4 行：判断题答案应为 对/错",
                "第 5 行：填空题题目里没有空", "第 6 行：与第 2 行题目重复"):  # fmt: skip
        assert msg in out


def test_exam_paper_shuffles_options_with_correct_answers(tmp_path: Path) -> None:
    script = EXAM / "scripts" / "exam.py"
    bank_path = EXAM / "references" / "example-bank.csv"
    out = tmp_path / "paper.docx"
    result = _run(script, "paper", bank_path, "-o", out, "--single", "5", "--multi", "3", "--judge", "4",
                  "--blank", "1", "--short", "1", "--title", "安全生产培训考试", "--seed", "7")  # fmt: skip
    assert result.returncode == 0, result.stderr
    assert "14 题，满分 35 分" in result.stdout

    with bank_path.open(encoding="utf-8-sig") as f:
        bank = {r["题目"]: r for r in csv.DictReader(f)}
    paras = [p.text for p in Document(str(out)).paragraphs]
    assert paras[0] == "安全生产培训考试" and "满分 35 分" in paras[1]
    split = paras.index("安全生产培训考试（参考答案）")
    questions, answers = paras[:split], paras[split + 1 :]
    key = {int(m.group(1)): m.group(2) for a in answers if (m := re.match(r"(\d+)\. (\S+)", a))}
    checked = 0
    for i, line in enumerate(questions):
        m = re.match(r"(\d+)\. (.+?)（　）$", line)
        if not m or not questions[i + 1].startswith("A. "):
            continue
        stem = m.group(2)
        options = dict(re.findall(r"([A-H])\. (.+?)(?=　　[A-H]\. |$)", questions[i + 1]))
        source = bank[stem]
        correct = {source["选项"].split("|")["ABCDEFGH".index(c)] for c in source["答案"]}
        assert {options[c] for c in key[int(m.group(1))]} == correct, stem
        checked += 1
    assert (
        checked == 8
    )  # every single- and multi-choice question still points at the right option text

    again = tmp_path / "again.docx"
    _run(script, "paper", bank_path, "-o", again, "--single", "5", "--multi", "3", "--judge", "4", "--seed", "7",
         "--title", "安全生产培训考试")  # fmt: skip
    first = [p.text for p in Document(str(again)).paragraphs]
    _run(script, "paper", bank_path, "-o", again, "--single", "5", "--multi", "3", "--judge", "4", "--seed", "7",
         "--title", "安全生产培训考试")  # fmt: skip
    assert [p.text for p in Document(str(again)).paragraphs] == first  # same seed, same paper

    short = _run(
        script, "paper", bank_path, "-o", tmp_path / "x.docx", "--single", "3", "--topic", "用电"
    )
    assert short.returncode == 1 and "只有 2 道" in short.stderr


def test_exam_shuffle_keeps_all_of_the_above(tmp_path: Path) -> None:
    mod = _load(EXAM / "scripts" / "exam.py")
    import random

    q = mod.Question(2, "单选", "题", ["甲", "乙", "以上都对"], "C", 2.0, "", "", "")
    for seed in range(10):
        assert mod.shuffled(q, random.Random(seed), True) == (["甲", "乙", "以上都对"], "C")
    q2 = mod.Question(3, "多选", "题", ["甲", "乙", "丙", "丁"], "BD", 3.0, "", "", "")
    opts, ans = mod.shuffled(q2, random.Random(3), True)
    assert {opts["ABCD".index(c)] for c in ans} == {"乙", "丁"}


# ------------------------------------------------------------------- listing


def test_listing_check_levels_and_export(tmp_path: Path) -> None:
    script = LISTING / "scripts" / "listing_check.py"
    out = tmp_path / "lc.xlsx"
    result = _run(script, LISTING / "references" / "example-items.csv", "-o", out)
    assert result.returncode == 3
    text = result.stdout
    blocks = re.split(r"\n(?=\d+\. )", text)
    assert "违规 0，风险 1，建议 1" in blocks[0] and "「100%」需要检测报告" in blocks[0]
    assert "违规 5，风险 4" in blocks[1]
    for word in ("最低价", "顶级", "根治", "好评返现", "纯天然", "祛斑"):
        assert f"「{word}」" in blocks[1]
    assert "普通食品不得宣称「养胃」" in blocks[2] and "标题中重复的词：小米" in blocks[2]
    assert "与其他 1 个商品标题完全相同" in blocks[3]
    rows = list(load_workbook(out)["检查结果"].iter_rows(min_row=2, values_only=True))
    assert sum(1 for r in rows if r[2] == "违规") == 8

    clean = _run(
        script,
        "纯棉短袖T恤男 圆领宽松 夏季休闲百搭打底衫",
        "--platform",
        "douyin",
        "--keywords",
        "T恤,纯棉,宽松",
    )
    assert clean.returncode == 0 and "违规 0，风险 0" in clean.stdout
    long = _run(script, "超长标题" * 20, "--platform", "taobao").stdout
    assert "标题长度 160/60（中文算 2）" in long
