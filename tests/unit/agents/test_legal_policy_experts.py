"""Contract tests for the bundled contract, policy, and business writing experts."""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from docx import Document
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
ANNOTATE = LIBRARY / "contract-reviewer" / "skills" / "contract-annotate"
COMPARE = LIBRARY / "policy-interpreter" / "skills" / "policy-compare"

EXPERTS = {
    "contract-reviewer": (
        "合同审阅专家",
        {"contract-review", "contract-annotate", "contract-compare"},
    ),
    "policy-interpreter": (
        "政策解读专家",
        {"policy-reading", "policy-compare", "policy-implementation"},
    ),
    "business-writing": (
        "商务沟通专家",
        {"business-email", "workplace-messages", "tricky-conversations"},
    ),
}

CONTRACT_LINES = [
    "甲方：示例科技有限公司",
    "乙方：示例软件服务有限公司",
    "第一条 服务内容：乙方为甲方提供客服系统的部署与运维服务。",
    "第二条 合同价款：总价人民币 30 万元，甲方应在合同签订后 3 日内一次性支付全部款项。",
    "第三条 验收：系统部署完成后，视为验收合格。",
    "第四条 违约责任：甲方逾期付款的，每日按合同总价的 1% 支付违约金；乙方延期交付的，每日按合同总价的 0.01% 支付违约金。",
    "第五条 知识产权：本合同项下开发的全部成果归乙方所有。",
    "第六条 争议解决：双方协商不成的，提交乙方所在地人民法院诉讼解决。",
]


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


def _make_contract(path: Path) -> None:
    doc = Document()
    doc.add_heading("技术服务合同", 1)
    for line in CONTRACT_LINES:
        p = doc.add_paragraph()
        mid = len(line) // 2
        p.add_run(line[:mid])
        p.add_run(line[mid:]).bold = True  # snippets must work across run boundaries
    doc.save(str(path))


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


def test_contract_reviewer_states_it_is_not_legal_advice() -> None:
    soul = (LIBRARY / "contract-reviewer" / "SOUL.md").read_text(encoding="utf-8")
    review = (LIBRARY / "contract-reviewer" / "skills" / "contract-review" / "SKILL.md").read_text(
        encoding="utf-8"
    )
    assert "不构成法律意见" in soul and "不构成法律意见" in review
    assert "不编造法条" in soul


def test_review_docx_extract_and_annotate(tmp_path: Path) -> None:
    contract = tmp_path / "contract.docx"
    _make_contract(contract)
    script = ANNOTATE / "scripts" / "review_docx.py"
    extract = _run(script, "extract", contract)
    assert extract.returncode == 0, extract.stderr
    assert "[P5] 第二条 合同价款" in extract.stdout

    out, report = tmp_path / "reviewed.docx", tmp_path / "risk.xlsx"
    result = _run(
        script,
        "annotate",
        contract,
        ANNOTATE / "references" / "example-review.md",
        "-o",
        out,
        "--report",
        report,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "5/5 条批注（高 3、中 2、低 0），原合同未修改" in result.stdout

    reviewed = Document(str(out))
    comments = list(reviewed.comments)
    assert len(comments) == 5
    assert comments[0].text.startswith("【高风险】预付全款")
    # The contract text itself is untouched even though runs were split for anchoring.
    assert [p.text for p in reviewed.paragraphs][1:] == CONTRACT_LINES
    assert [p.text for p in Document(str(contract)).paragraphs][1:] == CONTRACT_LINES
    assert len(list(Document(str(contract)).comments)) == 0
    # The comment range covers exactly the quoted snippet.
    xml = reviewed.paragraphs[4]._p.xml
    inside = re.search(r"<w:commentRangeStart[^>]*/>(.*?)<w:commentRangeEnd", xml, re.S)
    assert inside is not None
    assert (
        "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", inside.group(1)))
        == "合同签订后 3 日内一次性支付全部款项"
    )

    ws = load_workbook(report)["风险清单"]
    assert [ws.cell(r, 2).value for r in range(2, 7)] == ["高", "高", "高", "中", "中"]


def test_review_docx_reports_missing_snippets_and_guards(tmp_path: Path) -> None:
    contract = tmp_path / "contract.docx"
    _make_contract(contract)
    script = ANNOTATE / "scripts" / "review_docx.py"
    review = tmp_path / "review.md"
    review.write_text(
        "| 段落 | 原文片段 | 风险等级 | 问题 | 修改建议 |\n|---|---|---|---|---|\n"
        "| | 合同 里 没有 的话 | 低 | 测试 | 测试 |\n"
        "| | 视为验收合格 | 高 | 无验收标准 | 补充验收条款 |\n",
        encoding="utf-8",
    )
    result = _run(script, "annotate", contract, review, "-o", tmp_path / "r.docx")
    assert result.returncode == 3
    assert "1/2 条批注" in result.stdout and "未标注：第 1 条" in result.stdout

    same = _run(script, "annotate", contract, review, "-o", contract)
    assert same.returncode == 1 and "不能覆盖原合同" in same.stderr

    bad = tmp_path / "bad.md"
    bad.write_text(
        "| 原文片段 | 风险等级 | 问题 | 修改建议 |\n|---|---|---|---|\n| 验收 | 很高 | x | y |\n",
        encoding="utf-8",
    )
    result = _run(script, "annotate", contract, bad, "-o", tmp_path / "b.docx")
    assert result.returncode == 1 and "高 / 中 / 低" in result.stderr


def test_text_diff_marks_changes(tmp_path: Path) -> None:
    script = COMPARE / "scripts" / "text_diff.py"
    out, xlsx = tmp_path / "diff.md", tmp_path / "diff.xlsx"
    result = _run(
        script,
        COMPARE / "references" / "example-old.md",
        COMPARE / "references" / "example-new.md",
        "-o",
        out,
        "--xlsx",
        xlsx,
    )
    assert result.returncode == 0, result.stderr
    md = out.read_text(encoding="utf-8")
    assert "新增 0 处，删除 0 处，修改 5 处" in md
    assert "【-5000-】【+8000+】" in md  # numbers change as a whole
    assert "【+、退役军人+】" in md
    assert "| 第三条 |" in md
    ws = load_workbook(xlsx)["条文对比"]
    assert ws.cell(1, 6).value == "解读（人工填写）"

    old, new = tmp_path / "a.txt", tmp_path / "b.txt"
    old.write_text("第一条 甲。\n第二条 乙。\n", encoding="utf-8")
    new.write_text("第一条 甲。\n第三条 全新的条款内容。\n", encoding="utf-8")
    result = _run(script, old, new)
    assert "删除" in result.stdout and "新增" in result.stdout
