"""Contract tests for the bundled work experts (简历 / 招聘 / 翻译 / 论文 / 投标)."""

from __future__ import annotations

import importlib.util
import json
import os
import re
import subprocess
import sys
from decimal import Decimal
from pathlib import Path
from types import ModuleType

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from docx import Document
from docx.shared import RGBColor
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
RESUME = LIBRARY / "resume-writer" / "skills" / "resume-build"
JD = LIBRARY / "resume-writer" / "skills" / "jd-match"
SCREEN = LIBRARY / "hr-recruiter" / "skills" / "screening"
INTERVIEW = LIBRARY / "hr-recruiter" / "skills" / "interview-kit"
TRANSLATE = LIBRARY / "doc-translator" / "skills" / "docx-translate"
REFS = LIBRARY / "thesis-writer" / "skills" / "references-format"
THESIS = LIBRARY / "thesis-writer" / "skills" / "format-check" / "scripts" / "thesis_check.py"
TENDER = LIBRARY / "bid-writer" / "skills" / "tender-analysis"
BID_CHECK = LIBRARY / "bid-writer" / "skills" / "bid-response" / "scripts" / "bid_check.py"

EXPERTS = {
    "resume-writer": ("求职简历专家", {"resume-build", "jd-match", "interview-prep"}),
    "hr-recruiter": ("招聘专家", {"job-posting", "screening", "interview-kit"}),
    "doc-translator": ("文档翻译专家", {"docx-translate", "translation-style"}),
    "thesis-writer": ("论文写作专家", {"thesis-structure", "references-format", "format-check"}),
    "bid-writer": ("投标文件专家", {"tender-analysis", "bid-response"}),
}


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


def _load(path: Path) -> ModuleType:
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location(f"_work_{path.stem}", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


# ------------------------------------------------------------------- resume


def test_resume_check_and_three_layouts(tmp_path: Path) -> None:
    script = RESUME / "scripts" / "resume_docx.py"
    example = RESUME / "references" / "example-resume.md"
    result = _run(script, "check", example, "--today", "2026-10-18")
    assert result.returncode == 0, result.stderr
    out = result.stdout
    assert "「负责日常取数需求」没有数字" in out
    assert "「参与 8 个零售客户" in out and "弱动词" in out
    assert "之间有 7 个月空档" in out
    assert "预计约" in out and "联系方式" not in out

    for style in ("classic", "modern", "compact"):
        docx = tmp_path / f"{style}.docx"
        result = _run(script, "build", example, "-o", docx, "--style", style)
        assert result.returncode == 0, result.stderr
        doc = Document(str(docx))
        assert doc.paragraphs[0].text == "张小明"
        texts = [p.text for p in doc.paragraphs]
        assert any(t.startswith("某某零售科技有限公司　数据分析师\t2022.03") for t in texts)
        assert sum(t.startswith("• ") for t in texts) == 10

    bad = tmp_path / "bad.md"
    bad.write_text("- 没有标题的条目\n", encoding="utf-8")
    assert _run(script, "check", bad).returncode == 1


def test_resume_check_flags_sensitive_info_and_order(tmp_path: Path) -> None:
    md = tmp_path / "r.md"
    md.write_text(
        "# 李四\n身份证：110101199001011234 · 已婚\n\n## 工作经历\n"
        "### 甲公司 | 专员 | 2019.01 – 2020.06\n- 整理 30 份合同\n"
        "### 乙公司 | 主管 | 2020.03 – 至今\n- 带领 5 人团队\n",
        encoding="utf-8",
    )
    out = _run(RESUME / "scripts" / "resume_docx.py", "check", md, "--today", "2026-10-18").stdout
    assert "没有找到手机号" in out and "没有找到邮箱" in out
    assert "身份证号" in out and "「已婚」" in out
    assert "请按时间倒序排列" in out
    assert "时间重叠 3 个月" in out


def test_jd_match_reports_coverage() -> None:
    result = _run(
        JD / "scripts" / "jd_match.py",
        JD / "references" / "example-jd.txt",
        RESUME / "references" / "example-resume.md",
        "--today",
        "2026-10-18",
        "--json",
    )
    assert result.returncode == 0, result.stderr
    res = json.loads(result.stdout)
    assert {"sql", "python", "tableau", "数据可视化"} <= set(res["必备_已覆盖"])
    assert {"ab测试", "指标体系"} <= set(res["必备_未覆盖"])
    assert "ab" not in res["必备_已覆盖"] + res["必备_未覆盖"]
    assert "机器学习" in res["加分_未覆盖"] and "零售" in res["加分_已覆盖"]
    assert res["学历"] == "要求本科及以上，简历为本科"
    assert res["工作年限"] == "要求 3 年以上，简历约 7.3 年"
    assert 0 < res["必备覆盖率"] < 1


# ----------------------------------------------------------------------- hr


def test_screening_ranks_and_refuses_discriminatory_criteria(tmp_path: Path) -> None:
    script = SCREEN / "scripts" / "screen.py"
    refs = SCREEN / "references"
    out = tmp_path / "screen.xlsx"
    result = _run(
        script,
        refs / "example-criteria.csv",
        refs / "example-resumes",
        "-o",
        out,
        "--today",
        "2026-10-18",
    )
    assert result.returncode == 0, result.stderr
    assert "3 份简历，必备项全部满足 2 份" in result.stdout
    assert "1. 候选人A.md：100/100" in result.stdout
    assert "2. 候选人C.txt：90/100，缺：实验设计" in result.stdout
    assert "3. 候选人B.md：45/100（必备项未满足）" in result.stdout
    ws = load_workbook(out)["初筛结果"]
    header = [c.value for c in ws[1]]
    assert header[:4] == ["排名", "简历文件", "总分（满分 100）", "必备项"]
    row = {header[i]: c.value for i, c in enumerate(ws[3])}
    assert (
        row["简历文件"] == "候选人C.txt"
        and row["SQL"] == "SQL、MySQL、Hive"
        and row["学历"] == "本科"
    )

    bad = tmp_path / "bad.csv"
    bad.write_text("维度,权重,关键词,必备\n年龄,10,35岁以下,是\n", encoding="utf-8")
    refused = _run(script, bad, refs / "example-resumes", "-o", tmp_path / "x.xlsx")
    assert refused.returncode == 1 and "就业歧视" in refused.stderr


def test_interview_form_and_aggregate(tmp_path: Path) -> None:
    script = INTERVIEW / "scripts" / "interview_sheet.py"
    form = tmp_path / "form.xlsx"
    result = _run(
        script,
        "form",
        INTERVIEW / "references" / "example-questions.md",
        "-o",
        form,
        "--position",
        "数据分析师",
    )
    assert result.returncode == 0, result.stderr
    assert "5 道题，4 个维度" in result.stdout
    ws = load_workbook(form)["评估表"]
    assert ws["A1"].value == "结构化面试评估表（数据分析师）"
    assert str(ws["F14"].value).startswith("=IFERROR(SUMPRODUCT(E8:E12,F8:F12)")

    filled = []
    for who, cand, scores in (("王", "候选人A", [5, 4, 4, 3, 4]), ("李", "候选人A", [4, 4, 5, 5, 3]),
                              ("王", "候选人C", [3, 2, 3, 3, 3])):  # fmt: skip
        wb = load_workbook(form)
        sheet = wb["评估表"]
        sheet["B2"], sheet["B3"] = cand, who
        for i, s in enumerate(scores):
            sheet.cell(8 + i, 6, s)
        path = tmp_path / f"{who}-{cand}.xlsx"
        wb.save(path)
        filled.append(path)
    agg = tmp_path / "agg.xlsx"
    result = _run(script, "aggregate", *filled, "-o", agg)
    assert result.returncode == 0, result.stderr
    assert "候选人A：2 位面试官，平均 84.0 分；分歧较大：沟通协作" in result.stdout
    assert "候选人C：1 位面试官，平均 56.0 分" in result.stdout
    rows = list(load_workbook(agg)["候选人汇总"].iter_rows(min_row=2, values_only=True))
    assert rows[0][:5] == ("候选人A", 2, 84.0, 86.0, 82.0)

    wb = load_workbook(filled[0])
    wb["评估表"].cell(8, 6, 9)
    wb.save(filled[0])
    bad = _run(script, "aggregate", filled[0], "-o", agg)
    assert bad.returncode == 1 and "超出 1–5" in bad.stderr


# --------------------------------------------------------------- translate


def _sample_docx(path: Path) -> None:
    d = Document()
    d.sections[0].header.paragraphs[0].text = "某某公司内部资料"
    d.add_heading("2026 年第三季度运营报告", 1)
    p = d.add_paragraph("本季度")
    p.add_run("数据中心").bold = True
    p.add_run("可用率达到 99.95%，")
    p.add_run("用户").font.color.rgb = RGBColor(0xC0, 0, 0)
    p.add_run("投诉下降 12%。")
    d.add_paragraph("2026-09-30")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "指标", "数值"
    t.cell(1, 0).text, t.cell(1, 1).text = "新增用户", "1,200"
    d.add_paragraph("指标")
    d.save(str(path))


def test_docx_translate_round_trip_keeps_formatting(tmp_path: Path) -> None:
    script = TRANSLATE / "scripts" / "docx_segments.py"
    src = tmp_path / "src.docx"
    _sample_docx(src)
    seg = tmp_path / "seg.json"
    result = _run(script, "extract", src, "-o", seg)
    assert result.returncode == 0, result.stderr
    assert "7 个片段，其中需要翻译 6 个" in result.stdout
    segs = json.loads(seg.read_text(encoding="utf-8"))
    by_id = {s["id"]: s for s in segs}
    assert (
        by_id["p2"]["text"] == "本季度<2>数据中心</2>可用率达到 99.95%，<4>用户</4>投诉下降 12%。"
    )
    assert by_id["p8"]["dup_of"] == "p4" and by_id["p9"]["where"] == "页眉"
    assert "2026-09-30" not in {s["text"] for s in segs}  # digits-only paragraphs are left alone

    tr = {
        "p1": "Q3 2026 Operations Report",
        "p2": "This quarter, <2>data center</2> availability reached 99.95%, and <4>user</4> complaints fell by 12%.",
        "p4": "Metric",
        "p5": "Value",
        "p6": "New users",
        "p9": "Company Internal",
    }
    for s in segs:
        s["translation"] = tr.get(s["id"], "")
    seg.write_text(json.dumps(segs, ensure_ascii=False), encoding="utf-8")
    glossary = TRANSLATE / "references" / "example-glossary.csv"
    check = _run(script, "check", seg, "--glossary", glossary)
    assert (
        "发现 1 个问题" in check.stdout and "术语「季度」" in check.stdout
    )  # "Q3" doesn't contain "quarter"

    out = tmp_path / "en.docx"
    result = _run(script, "fill", src, seg, "-o", out)
    assert result.returncode == 0, result.stderr
    assert "回填 7 个片段，未翻译 0 个" in result.stdout
    doc = Document(str(out))
    p = doc.paragraphs[1]
    assert p.text == tr["p2"].replace("<2>", "").replace("</2>", "").replace("<4>", "").replace(
        "</4>", ""
    )
    runs = {r.text: r for r in p.runs}
    assert runs["data center"].bold and not runs["This quarter, "].bold
    assert runs["user"].font.color.rgb == RGBColor(0xC0, 0, 0)
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert [c.text for row in doc.tables[0].rows for c in row.cells] == [
        "Metric",
        "Value",
        "New users",
        "1,200",
    ]
    assert (
        doc.paragraphs[-1].text == "Metric"
        and doc.sections[0].header.paragraphs[0].text == "Company Internal"
    )

    segs[1]["translation"] = "Availability reached 99.95%."
    segs[4]["translation"] = ""
    seg.write_text(json.dumps(segs, ensure_ascii=False), encoding="utf-8")
    check = _run(script, "check", seg)
    assert check.returncode == 1
    assert "未翻译 1 个：p6" in check.stdout
    assert "p2：格式标签与原文不一致" in check.stdout and "原文数字 12" in check.stdout

    assert _run(script, "fill", src, seg, "-o", src).returncode == 1  # never overwrite the source


# ------------------------------------------------------------------- thesis


def test_gbt7714_format_and_check(tmp_path: Path) -> None:
    script = REFS / "scripts" / "gbt7714.py"
    out = tmp_path / "refs.txt"
    result = _run(script, "format", REFS / "references" / "example.bib", "-o", out)
    assert result.returncode == 0, result.stderr
    lines = out.read_text(encoding="utf-8").splitlines()
    assert lines == [
        "[1] 王小明, 李华, 张伟, 等. 基于深度学习的城市交通流量预测方法研究[J]. 交通运输工程学报, 2023, 23(4): 112-125.",
        "[2] SMITH J A, LEE M, BROWN D. Graph neural networks for traffic forecasting[J]. "
        "Transportation Research Part C, 2021, 125: 103-118. DOI:10.1016/j.trc.2021.000000.",
        "[3] 陈某某. 机器学习[M]. 2版. 北京: 某某大学出版社, 2019.",
        "[4] 刘某. 城市路网交通状态感知与预测关键技术研究[D]. 上海: 某某大学, 2022.",
        "[5] KIM H, PARK J W. Spatio-temporal attention for traffic prediction[C]//"
        "Proceedings of the 2020 International Conference on Data Mining. New York: IEEE, 2020: 201-210.",
        "[6] 某某市交通委员会. 2024 年城市交通发展年度报告[EB/OL]. (2024-03-15)[2026-10-01]. https://example.gov.cn/report/2024",
    ]
    assert _run(script, "check", out).returncode == 0

    bad = tmp_path / "bad.txt"
    bad.write_text(
        "[1] 王小明，李华. 城市交通研究[J]. 交通学报，2023，23(4)：112–125.\n"
        "[3] Smith J, Lee M, Brown D, Green K. Graph networks. Transportation Research, 2021.\n",
        encoding="utf-8",
    )
    result = _run(script, "check", bad)
    assert result.returncode == 1
    for msg in (
        "[1] 使用了中文标点",
        "[1] 页码范围请用",
        "[3] 序号不连续",
        "[3] 缺少文献类型标识",
        "[3] 作者超过 3 人",
    ):
        assert msg in result.stdout

    csv = tmp_path / "refs.csv"
    csv.write_text(
        "类型,作者,题名,刊名,年,卷,期,页码\nJ,张三;李四,某研究,某学报,2024,,,\n", encoding="utf-8"
    )
    result = _run(script, "format", csv)
    assert "[1] 张三, 李四. 某研究[J]. 某学报, 2024." in result.stdout
    assert "缺少：页码" in result.stderr


def test_edition_and_author_rules() -> None:
    mod = _load(REFS / "scripts" / "gbt7714.py")
    assert [mod.edition_text(e, c) for e, c in (("1", True), ("第3版", True), ("2nd", False), ("11", False))] == [
        "", "3版", "2nd ed", "11th ed",
    ]  # fmt: skip
    assert mod.format_author("van der Berg, Anna-Maria") == "VAN DER BERG A M"
    assert mod.authors("A, B;C, D;E, F;G, H", chinese=False) == "A B, C D, E F, et al"


def test_thesis_check_finds_numbering_and_citation_issues(tmp_path: Path) -> None:
    d = Document()
    for text in ("摘  要", "本文研究城市交通流量预测问题。", "关键词：交通预测；深度学习；图神经网络", "Abstract",
                 "This thesis studies traffic forecasting.", "Key words: traffic; deep learning"):  # fmt: skip
        d.add_paragraph(text)
    d.add_heading("第一章 绪论", 1)
    d.add_heading("1.1 研究背景", 2)
    d.add_paragraph("城市交通拥堵日益严重[1]，已有研究[2-3]提出了多种方法。如图 1-1 所示。")
    d.add_paragraph("图 1-1 研究框架")
    d.add_heading("1.3 研究内容", 2)
    d.add_paragraph("本文采用文献[5]的方法，并参考[4]。")
    d.add_heading("第二章 方法", 1)
    d.add_heading("2.1 模型", 2)
    d.add_paragraph("表 2-1 数据集统计")
    d.add_paragraph("图 2-2 模型结构")
    d.add_heading("参考文献", 1)
    for i in range(1, 7):
        d.add_paragraph(f"[{i}] 作者{i}. 题名{i}[J]. 刊名, 2023, 1(1): 1-2.")
    d.add_heading("致谢", 1)
    d.add_paragraph("感谢导师。")
    path = tmp_path / "thesis.docx"
    d.save(str(path))

    result = _run(THESIS, path, "--json")
    assert result.returncode == 1
    res = json.loads(result.stdout)
    assert (res["标题数"], res["图"], res["表"], res["参考文献条数"]) == (5, 2, 1, 6)
    issues = "\n".join(res["问题"])
    assert "「1.3 研究内容」不连续，应为 1.2" in issues
    assert "第 2 章的第一个图应为 图 2-1" in issues
    assert "表 2-1 在正文中没有被提到" in issues
    assert "图 1-1" not in issues
    assert "未被正文引用：[6]" in issues
    assert "[4] 首次出现在 [5] 之后" in issues
    assert "没有找到" not in issues  # both abstracts and keyword lines exist


# ---------------------------------------------------------------------- bid


def test_tender_scan_builds_response_matrix(tmp_path: Path) -> None:
    out = tmp_path / "matrix.xlsx"
    result = _run(
        TENDER / "scripts" / "tender_scan.py",
        TENDER / "references" / "example-tender.md",
        "-o",
        out,
    )
    assert result.returncode == 0, result.stderr
    stdout = result.stdout
    assert "- 项目编号：ZB-2026-0815" in stdout
    assert "- 投标截止时间：2026年11月20日 09:30（北京时间）" in stdout
    assert "- 质保期：未找到" in stdout
    assert "★ 实质性 2 条，▲ 重要 1 条，否决/无效投标 4 条，资格要求 5 条" in stdout
    assert "评分标准：5 项，合计 100 分" in stdout
    wb = load_workbook(out)
    kinds = [r[1] for r in wb["实质性条款"].iter_rows(min_row=2, values_only=True)]
    assert kinds.count("资格") == 5 and kinds.count("否决/无效") == 4
    texts = [r[2] for r in wb["实质性条款"].iter_rows(min_row=2, values_only=True)]
    assert not any(t.startswith("#") for t in texts)  # headings are context, not clauses
    sc = list(wb["评分标准"].iter_rows(min_row=2, values_only=True))
    assert [r[1] for r in sc[:5]] == ["投标报价", "运维方案", "人员配置", "类似业绩", "服务承诺"]
    assert sc[-1][2] == "=SUM(C2:C6)"


def test_bid_check_matrix_and_document(tmp_path: Path) -> None:
    matrix = tmp_path / "matrix.xlsx"
    assert (
        _run(
            TENDER / "scripts" / "tender_scan.py",
            TENDER / "references" / "example-tender.md",
            "-o",
            matrix,
        ).returncode
        == 0
    )
    result = _run(BID_CHECK, "matrix", matrix)
    assert result.returncode == 1
    assert "第 1 条（资格）未填写响应情况" in result.stdout

    wb = load_workbook(matrix)
    ws = wb["实质性条款"]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 5, "满足")
        ws.cell(r, 6, "无偏离")
        ws.cell(r, 7, f"P{r}")
    ws.cell(7, 6, "负偏离")
    for r in range(2, 7):
        wb["评分标准"].cell(r, 5, 20)
    wb.save(matrix)
    result = _run(BID_CHECK, "matrix", matrix)
    assert "第 6 条（★实质性）为负偏离" in result.stdout and "发现 1 个问题" in result.stdout

    d = Document()
    for text in (
        "项目名称：某某市智慧园区运维服务项目  项目编号：ZB-2026-0815",
        "我方某某科技有限公司愿以投标总报价：壹佰柒拾捌万元整（¥1,780,000.00）承担本项目。",
        "保证金：叁万伍仟元整（¥30,000.00）",
        "服务承诺详见第 XX 页。",
        "授权代表：【姓名】",
        "本项目由某某科技公司负责实施，原某某信息有限公司项目团队提供支持。",
        "招标编号：ZB-2026-0816",
        "日期：2025年11月18日",
        "承诺服务期至 2026年12月01日",
    ):
        d.add_paragraph(text)
    bid = tmp_path / "bid.docx"
    d.save(str(bid))
    result = _run(BID_CHECK, "doc", bid, "--bidder", "某某科技有限公司", "--project-no", "ZB-2026-0815",
                  "--other", "某某信息有限公司", "--deadline", "2026-11-20")  # fmt: skip
    assert result.returncode == 1
    out = result.stdout
    assert "金额大小写不一致：大写「叁万伍仟元整」，小写 ¥30,000.00 应写作「叁万元整」" in out
    assert "壹佰柒拾捌万" not in out  # the correct amount is not flagged
    assert "残留占位符「XX」" in out and "残留占位符「【姓名】」" in out
    assert "其他公司名称「某某信息有限公司」" in out
    assert "投标人名称有不同写法：某某科技公司" in out
    assert "项目编号写成了「ZB-2026-0816」" in out
    assert "落款日期 2025-11-18 比投标截止日早 367 天" in out
    assert "2026-12-01" not in out  # service-period dates are not sign-off dates


def test_amount_in_words() -> None:
    mod = _load(BID_CHECK)
    cases = {
        "1800000": "壹佰捌拾万元整",
        "100500.05": "壹拾万零伍佰元零伍分",
        "1001": "壹仟零壹元整",
        "105000000": "壹亿零伍佰万元整",
        "1234567.89": "壹佰贰拾叁万肆仟伍佰陆拾柒元捌角玖分",
        "0.5": "伍角",
    }
    assert {k: mod.to_upper(Decimal(k)) for k in cases} == cases
