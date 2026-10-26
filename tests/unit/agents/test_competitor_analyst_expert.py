"""Contract tests for the bundled 竞品分析专家 (competitor-analyst) expert."""

from __future__ import annotations

import http.server
import json
import os
import re
import socket
import subprocess
import sys
import threading
import urllib.parse
from collections.abc import Iterator
from pathlib import Path

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)
from octop.infra.agents.experts.skillhub_market import _SCENE_ORDER

LIBRARY = default_library_root()
EXPERT_DIR = LIBRARY / "competitor-analyst"
SKILLS = {"product-search", "competitor-framework", "competitive-matrix", "analysis-report"}
SEARCH = EXPERT_DIR / "skills" / "product-search" / "scripts" / "product_search.py"
MATRIX = EXPERT_DIR / "skills" / "competitive-matrix" / "scripts" / "compare_matrix.py"
REPORT = EXPERT_DIR / "skills" / "analysis-report"
CHECK = REPORT / "scripts" / "report_check.py"


def _run(
    script: Path, *args: str | Path, env: dict[str, str] | None = None
) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    full_env = {
        **os.environ,
        **(env or {}),
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=full_env,
        timeout=120,
    )


def test_expert_catalog_manifest_and_skills() -> None:
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get("competitor-analyst")
    assert expert is not None
    assert expert.summary.label_zh == "竞品分析专家"
    assert expert.summary.icon_url == "/experts/avatars/competitor-analyst.svg"
    assert expert.summary.scene in _SCENE_ORDER

    manifest = json.loads((EXPERT_DIR / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == "competitor-analyst"
    for rel in manifest["prompt_files"]:
        assert (EXPERT_DIR / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (EXPERT_DIR / "skills").iterdir() if p.is_dir()} == SKILLS
    soul = (EXPERT_DIR / "SOUL.md").read_text(encoding="utf-8")
    for name in SKILLS:
        raw = (EXPERT_DIR / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (EXPERT_DIR / "skills" / name / rel).is_file(), f"{name}: {rel}"


async def test_expert_seeds_into_workspace(tmp_path: Path) -> None:
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(EXPERT_DIR)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=EXPERT_DIR, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (EXPERT_DIR / "manifest.json").read_bytes()
    )


def test_docx_build_is_the_shared_copy() -> None:
    shared = LIBRARY / "data-analyst" / "skills" / "analysis-report" / "scripts" / "docx_build.py"
    assert (REPORT / "scripts" / "docx_build.py").read_bytes() == shared.read_bytes()


# ------------------------------------------------------------ product search

APP = {
    "trackId": 111,
    "trackName": "速笔",
    "sellerName": "示例科技有限公司",
    "averageUserRating": 4.71,
    "userRatingCount": 123000,
    "price": 0,
    "formattedPrice": "免费",
    "primaryGenreName": "Productivity",
    "version": "3.2.0",
    "currentVersionReleaseDate": "2026-10-18T07:00:00Z",
    "releaseDate": "2023-05-01T07:00:00Z",
    "fileSizeBytes": "104857600",
    "trackViewUrl": "https://apps.apple.com/cn/app/id111?uo=4",
    "releaseNotes": "修复已知问题",
    "description": "录音转写和 AI 整理",
    "minimumOsVersion": "15.0",
    "languageCodesISO2A": ["ZH", "EN"],
}


def _review(rating: int, title: str, content: str, version: str = "3.2.0") -> dict[str, object]:
    return {
        "im:rating": {"label": str(rating)},
        "title": {"label": title},
        "content": {"label": content},
        "im:version": {"label": version},
        "updated": {"label": "2026-10-20T08:00:00-07:00"},
        "author": {"name": {"label": "用户"}},
    }


REVIEWS = [
    {"im:name": {"label": "速笔"}},  # app-info entry that old feeds put first
    _review(1, "转写不准", "人名转写不准，会员价格太贵"),
    _review(2, "太贵了", "会员价格涨了，转写不准的问题也没修"),
    _review(1, "导出", "导出格式太少，会员价格不值"),
    _review(5, "好用", "录音整理很方便，摘要很准确", "3.1.0"),
    _review(4, "不错", "摘要很准确，整理很方便"),
]
SITE = """<!doctype html><html><head><meta charset="utf-8"><title>速笔 - AI 会议记录</title>
<meta name="description" content="录音自动转写，AI 整理会议纪要">
<meta property="og:site_name" content="速笔"></head><body>
<a href="/pricing">价格</a><a href="https://help.example.com/docs">帮助中心</a>
<a href="/careers">加入我们</a><a href="/changelog">更新日志</a>
<a href="https://apps.apple.com/cn/app/subi/id1234567890">App Store</a>
<script id="__NEXT_DATA__">{}</script><script src="https://hm.baidu.com/hm.js"></script>
<footer>© 2026 示例科技有限公司 <a href="https://beian.miit.gov.cn">京ICP备12345678号-2</a></footer>
</body></html>"""


ASSET_SITE = """<!doctype html><html><head><meta charset="utf-8"><title>速笔 - AI 会议记录</title>
<meta name="description" content="录音自动转写，AI 整理会议纪要">
<meta property="og:image" content="/img/og.png">
<link rel="icon" href="/img/favicon.png"><link rel="apple-touch-icon" href="/img/icon.png"></head><body>
<a href="/product">产品功能</a><a href="/pricing">价格</a>
<img src="/img/hero.png" alt="产品主图"><img src="/img/tiny.png"><img src="/img/logo.svg">
<img data-src="/img/feature.png"><img src="/img/hero.png" alt="重复">
<footer>© 2026 示例科技有限公司 京ICP备12345678号-2</footer></body></html>"""
IMAGE_SIZES = {
    "og.png": (1200, 630),
    "favicon.png": (32, 32),
    "icon.png": (180, 180),
    "hero.png": (1600, 900),
    "tiny.png": (60, 40),
    "feature.png": (800, 500),
    "icon-512.png": (512, 512),
}


def _png(size: tuple[int, int], fmt: str = "PNG") -> bytes:
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, (size[0] % 255, 90, 160)).save(buf, fmt)
    return buf.getvalue()


class _Handler(http.server.BaseHTTPRequestHandler):
    def log_message(self, *args: object) -> None:
        pass

    def _send(self, body: object, ctype: str = "application/json") -> None:
        if isinstance(body, bytes):
            data = body
        else:
            data = body.encode() if isinstance(body, str) else json.dumps(body).encode()
        self.send_response(200)
        self.send_header(
            "Content-Type", ctype if ctype.startswith("image/") else f"{ctype}; charset=utf-8"
        )
        self.send_header("Server", "nginx")
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self) -> None:  # noqa: N802 - http.server API
        url = urllib.parse.urlparse(self.path)
        q = urllib.parse.parse_qs(url.query)
        if url.path == "/search":
            assert q["entity"] == ["software"]
            self._send({"resultCount": 1, "results": [APP]})
        elif url.path == "/lookup":
            host = f"http://{self.headers['Host']}"
            app = {
                **APP,
                "artworkUrl512": f"{host}/img/icon-512.png",
                "screenshotUrls": [f"{host}/img/s1/320x480bb.jpg", f"{host}/img/s2/320x480bb.jpg"],
            }
            self._send({"results": [app] if q["id"] == ["111"] else []})
        elif url.path.startswith("/img/s") and url.path.endswith(".jpg"):
            size = (600, 1300) if url.path.endswith("/600x0w.jpg") else (320, 480)
            self._send(_png((size[0] + int(url.path[6]), size[1]), "JPEG"), "image/jpeg")
        elif url.path.startswith("/img/") and url.path[5:] in IMAGE_SIZES:
            self._send(_png(IMAGE_SIZES[url.path[5:]]), "image/png")
        elif url.path == "/img/logo.svg":
            self._send("<svg xmlns='http://www.w3.org/2000/svg'/>", "image/svg+xml")
        elif url.path == "/asset-site":
            self._send(ASSET_SITE, "text/html")
        elif url.path.startswith("/cn/rss/customerreviews/page=1/"):
            self._send({"feed": {"entry": REVIEWS}})
        elif url.path.startswith("/cn/rss/customerreviews/"):
            self._send({"feed": {}})
        elif url.path == "/search/repositories":
            self._send(
                {
                    "total_count": 1,
                    "items": [
                        {
                            "full_name": "acme/notes",
                            "description": "AI notes",
                            "stargazers_count": 1200,
                            "forks_count": 80,
                            "open_issues_count": 5,
                            "pushed_at": "2026-10-20T00:00:00Z",
                            "created_at": "2024-01-01T00:00:00Z",
                            "language": "TypeScript",
                            "license": {"spdx_id": "MIT"},
                            "topics": ["ai", "notes"],
                            "html_url": "https://github.com/acme/notes",
                        }
                    ],
                }
            )
        elif url.path == "/cdx/search/cdx":
            self._send([["timestamp", "original"], ["20200102030405", "http://subi.example/"]])
        elif url.path == "/site":
            self._send(SITE, "text/html")
        else:
            self.send_error(404)


@pytest.fixture
def fake_sources() -> Iterator[dict[str, str]]:
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_address[1]}"
    yield {
        "PRODUCT_SEARCH_ITUNES": base,
        "PRODUCT_SEARCH_GITHUB": base,
        "PRODUCT_SEARCH_WAYBACK": base,
        "BASE": base,
    }
    server.shutdown()
    server.server_close()


def test_appstore_search_lookup_and_excel(fake_sources: dict[str, str], tmp_path: Path) -> None:
    out = tmp_path / "商店.xlsx"
    proc = _run(SEARCH, "appstore", "AI 笔记", "-o", out, env=fake_sources)
    assert proc.returncode == 0, proc.stderr
    assert "名称 速笔，开发者 示例科技有限公司，评分 4.71，评分人数 123000" in proc.stdout
    wb = load_workbook(out)
    assert wb.sheetnames == ["App Store 搜索", "来源"]
    assert wb["来源"]["A1"].value == "来源"

    proc = _run(SEARCH, "app", "111", "--json", env=fake_sources)
    assert proc.returncode == 0, proc.stderr
    row = json.loads(proc.stdout)["结果"][0]
    assert row["上架时间"] == "2023-05-01" and row["最近更新"] == "2026-10-18"
    assert row["大小MB"] == 100.0 and row["链接"] == "https://apps.apple.com/cn/app/id111"
    assert row["更新说明"] == "修复已知问题" and row["支持语言"] == "ZH、EN"

    assert _run(SEARCH, "app", "abc", env=fake_sources).returncode == 1
    assert _run(SEARCH, "app", "999", env=fake_sources).returncode == 1


def test_reviews_summarise_ratings_and_complaints(fake_sources: dict[str, str]) -> None:
    proc = _run(SEARCH, "reviews", "111", "--pages", "3", "--json", env=fake_sources)
    assert proc.returncode == 0, proc.stderr
    data = json.loads(proc.stdout)
    assert data["评论数"] == 5  # app-info entry skipped, empty page 2 stops paging
    assert data["评分分布"] == {"5星": 1, "4星": 1, "3星": 0, "2星": 1, "1星": 2}
    assert data["平均分"] == 2.6
    assert data["各版本"]["3.1.0"] == {"评论数": 1, "平均分": 5.0}
    negative = dict(data["差评高频词"])
    assert negative.get("会员价格") == 3 and negative.get("转写不准") == 2
    assert "摘要很准确" in dict(data["好评高频词"]) or "很准确" in dict(data["好评高频词"])
    assert len(data["来源"]) == 2

    text = _run(SEARCH, "reviews", "111", env=fake_sources).stdout
    assert "差评高频词：会员价格(3)" in text and "[1星 3.2.0 2026-10-20] 转写不准" in text


def test_github_and_site(fake_sources: dict[str, str]) -> None:
    proc = _run(SEARCH, "github", "ai notes", "--json", env=fake_sources)
    assert proc.returncode == 0, proc.stderr
    repo = json.loads(proc.stdout)["结果"][0]
    assert repo["项目"] == "acme/notes" and repo["星标"] == 1200 and repo["许可证"] == "MIT"

    site = fake_sources["BASE"] + "/site"
    proc = _run(SEARCH, "site", site, "--history", "--json", env=fake_sources)
    assert proc.returncode == 0, proc.stderr
    info = json.loads(proc.stdout)["结果"][0]
    assert info["标题"] == "速笔 - AI 会议记录"
    assert info["描述"] == "录音自动转写，AI 整理会议纪要" and info["站点名"] == "速笔"
    assert info["关键页面"]["定价"] == [fake_sources["BASE"] + "/pricing"]
    assert info["关键页面"]["文档"] == ["https://help.example.com/docs"]
    assert {"招聘", "更新日志"} <= set(info["关键页面"])
    assert info["ICP备案"] == ["京ICP备12345678号-2"]
    assert info["页脚主体"] == ["示例科技有限公司"]
    assert {"Next.js", "百度统计", "server: nginx"} <= set(info["技术线索"])
    assert info["AppStore应用ID"] == ["1234567890"]
    assert info["网页时光机"]["最早收录"] == "2020-01-02"
    assert "不等于产品上线时间" in info["网页时光机"]["说明"]


def test_network_errors_exit_2() -> None:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    env = {"PRODUCT_SEARCH_ITUNES": f"http://127.0.0.1:{port}"}
    proc = _run(SEARCH, "appstore", "x", env=env)
    assert proc.returncode == 2
    assert "无法访问" in proc.stderr


# ------------------------------------------------------------------- matrix


def _filled_matrix(tmp_path: Path) -> Path:
    path = tmp_path / "矩阵.xlsx"
    proc = _run(
        MATRIX,
        "template",
        "--type",
        "ai",
        "--competitors",
        "小记,速笔,云签",
        "--self",
        "小记",
        "-o",
        path,
    )
    assert proc.returncode == 0, proc.stderr
    wb = load_workbook(path)
    ws, ev = wb["评分"], wb["证据"]
    assert [c.value for c in ws[1]] == ["维度", "说明", "权重", "小记", "速笔", "云签"]
    assert ws.max_row == 9 and ev.max_row == 1 + 8 * 3
    scores = {
        "小记": [4, 5, 3, 4, 4, 3, 5, 2],
        "速笔": [5, 4, 5, 3, 4, 4, 4, 5],
        "云签": [4, 3, 4, 3, 5, 5, 5, None],
    }
    for i, row in enumerate(ws.iter_rows(min_row=2)):
        for j, name in enumerate(("小记", "速笔", "云签")):
            row[3 + j].value = scores[name][i]
    for r in ev.iter_rows(min_row=2):
        if r[0].value != "小记":
            r[2].value, r[3].value, r[4].value, r[5].value = (
                "官网",
                "https://example.com",
                "2026-10-20",
                "高",
            )
        elif r[1].value != "定价与额度":
            r[2].value = "内部测试"
    wb.save(path)
    return path


def test_matrix_scores_gaps_and_missing_evidence(tmp_path: Path) -> None:
    path = _filled_matrix(tmp_path)
    out = tmp_path / "结果.xlsx"
    proc = _run(MATRIX, "score", path, "--self", "小记", "-o", out, "--json")
    assert proc.returncode == 3, proc.stderr
    res = json.loads(proc.stdout)
    ranking = [(r["竞品"], r["加权得分"], r["已评维度"]) for r in res["排名"]]
    assert ranking == [("速笔", 4.35, "8/8"), ("云签", 4.06, "7/8"), ("小记", 3.85, "8/8")]
    assert res["我方差距"][0]["维度"] == "口碑与增长" and res["我方差距"][0]["差距"] == -3
    assert {"维度": "任务完成度", "领先": "小记", "最高分": 5} in res["维度领先"]
    assert res["缺分"] == ["云签·口碑与增长"]
    assert res["缺证据"] == ["小记·定价与额度"]
    assert len(res["缺来源"]) == 7
    wb = load_workbook(out)
    assert wb.sheetnames == ["排名", "维度对比", "维度领先", "我方差距", "待补充"]
    assert len(wb["排名"]._charts) == 1 and len(wb["维度对比"]._charts) == 1


def test_matrix_input_errors(tmp_path: Path) -> None:
    csv_path = tmp_path / "评分.csv"
    csv_path.write_text("维度,说明,权重,A,B\n功能,,1,6,3\n", encoding="utf-8")
    proc = _run(MATRIX, "score", csv_path)
    assert proc.returncode == 1 and "不在 1–5 之间" in proc.stderr
    csv_path.write_text("名称,A,B\n功能,1,2\n", encoding="utf-8")
    assert _run(MATRIX, "score", csv_path).returncode == 1
    csv_path.write_text("维度,说明,权重,A,B\n功能,,1,4,3\n体验,,1,5,5\n", encoding="utf-8")
    proc = _run(MATRIX, "score", csv_path)
    assert proc.returncode == 0, proc.stdout
    assert "没有证据表" in proc.stdout
    out = tmp_path / "x.xlsx"
    assert _run(MATRIX, "template", "--competitors", "只有一个", "-o", out).returncode == 1


# ------------------------------------------------------------------- report


def test_example_report_passes_quality_check() -> None:
    proc = _run(
        CHECK, REPORT / "references" / "example-report.md", "--competitors", "小记,速笔,云签"
    )
    assert proc.returncode == 0, proc.stdout
    assert "检查都通过" in proc.stdout


def test_weak_report_is_flagged(tmp_path: Path) -> None:
    report = tmp_path / "弱报告.md"
    report.write_text(
        "# 竞品分析\n\n## 产品介绍\nKimi 是国内最强的 AI 助手，用户超过 1 亿。据说豆包增长很快。\n\n"
        "| 产品 | 月活 |\n|---|---|\n| Kimi | 5000 万 |\n\n## 建议\n- 加强推广\n",
        encoding="utf-8",
    )
    proc = _run(CHECK, report, "--competitors", "Kimi,豆包", "--json")
    assert proc.returncode == 3
    res = json.loads(proc.stdout)
    issues = "\n".join(res["问题"])
    for needle in ("核心结论", "建议部分只有 1 条", "信息来源", "截至时间", "没有任何来源", "豆包"):
        assert needle in issues, needle
    kinds = {n["类型"] for n in res["待确认"]}
    assert kinds == {"数据无来源", "判断无依据", "低可信表述", "表格无来源"}


def test_report_exports_to_docx(tmp_path: Path) -> None:
    out = tmp_path / "报告.docx"
    proc = _run(
        REPORT / "scripts" / "docx_build.py",
        REPORT / "references" / "example-report.md",
        "-o",
        out,
        "--style",
        "report",
    )
    assert proc.returncode == 0, proc.stderr
    from docx import Document

    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "小记 vs 速笔、云签" in text and "七、信息来源" in text


# ------------------------------------------------------------- product assets

ASSETS = EXPERT_DIR / "skills" / "product-search" / "scripts" / "product_assets.py"


def test_product_assets_collects_images_and_profile(
    fake_sources: dict[str, str], tmp_path: Path
) -> None:
    site = fake_sources["BASE"] + "/asset-site"
    proc = _run(
        ASSETS, "速笔", "--site", site, "--appstore", "111", "-o", tmp_path, env=fake_sources
    )
    assert proc.returncode == 0, proc.stderr + proc.stdout
    folder = tmp_path / "assets" / "速笔"
    files = sorted(p.name for p in folder.iterdir())
    assert files == [
        "appstore-icon.png",
        "appstore-shot-1.jpg",
        "appstore-shot-2.jpg",
        "assets.json",
        "profile.md",
        "site-image-1.png",
        "site-image-2.png",
        "site-logo.png",
        "site-og.png",
    ]
    data = json.loads((folder / "assets.json").read_text(encoding="utf-8"))
    sizes = {Path(r["文件"]).name: r["尺寸"] for r in data["素材"]}
    assert sizes["appstore-shot-1.jpg"] == "601×1300"  # larger rendition, not the thumbnail
    assert sizes["site-logo.png"] == "180×180"  # apple-touch-icon preferred over favicon
    assert sizes["site-image-1.png"] == "1600×900" and sizes["site-image-2.png"] == "800×500"
    assert all(r["获取日期"] and r["来源"].startswith("http") for r in data["素材"])

    profile = (folder / "profile.md").read_text(encoding="utf-8")
    assert profile.startswith("### 速笔\n")
    assert f"- 官网：[{site}]({site})" in profile
    assert "- 产品介绍页：[" in profile and "/product" in profile
    assert "- 定价页：[" in profile and "/pricing" in profile
    assert "- 运营主体：示例科技有限公司（京ICP备12345678号-2）" in profile
    assert "评分 4.71（123000 人评分），版本 3.2.0" in profile
    assert "![速笔 App Store 截图 1（来源：App Store，" in profile
    assert "](assets/速笔/appstore-shot-1.jpg)" in profile
    assert "![速笔 官网分享图（来源：官网，" in profile


def test_product_assets_input_and_network_errors(tmp_path: Path) -> None:
    assert _run(ASSETS, "x", "-o", tmp_path).returncode == 1
    assert _run(ASSETS, "x", "--appstore", "abc", "-o", tmp_path).returncode == 1
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    proc = _run(ASSETS, "x", "--site", f"http://127.0.0.1:{port}/", "-o", tmp_path)
    assert proc.returncode == 2 and "官网获取失败" in proc.stderr


def _chromium_available() -> bool:
    probe = "from playwright.sync_api import sync_playwright\nwith sync_playwright() as p:\n    p.chromium.launch().close()"
    try:
        return (
            subprocess.run(
                [sys.executable, "-c", probe], env=ORIGINAL_ENV, capture_output=True, timeout=60
            ).returncode
            == 0
        )
    except (OSError, subprocess.TimeoutExpired):
        return False


ORIGINAL_ENV = dict(os.environ)


def test_product_assets_page_screenshots(fake_sources: dict[str, str], tmp_path: Path) -> None:
    if not _chromium_available():
        pytest.skip("Chromium is not installed for Playwright")
    site = fake_sources["BASE"] + "/asset-site"
    env = {**ORIGINAL_ENV, **fake_sources}
    proc = subprocess.run(
        [
            sys.executable,
            str(ASSETS),
            "速笔",
            "--site",
            site,
            "--screenshot",
            "--max-images",
            "0",
            "-o",
            str(tmp_path),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=180,
        env={**env, "PYTHONIOENCODING": "utf-8", "PYTHONDONTWRITEBYTECODE": "1"},
    )
    assert proc.returncode == 0, proc.stderr + proc.stdout
    folder = tmp_path / "assets" / "速笔"
    from PIL import Image

    with Image.open(folder / "page-home.png") as im:
        assert im.width == 1440
    assert (folder / "page-pricing.png").is_file()  # pricing link found on the page
    assert "![速笔 官网首页截图（来源：127.0.0.1" in (folder / "profile.md").read_text(
        encoding="utf-8"
    )


# ------------------------------------------------------- report specificity


def test_report_check_flags_vague_and_missing_profiles(tmp_path: Path) -> None:
    report = tmp_path / "报告.md"
    report.write_text(
        "---\n署名行: 数据截至 2026 年 10 月 20 日\n---\n# 分析\n\n## 一、核心结论\n"
        "1. 速笔功能丰富，体验较好 [1]\n2. 云签各有优劣\n3. 小记值得关注\n\n"
        "## 二、竞品档案\n### 速笔\n- 官网：[https://subi.example.com](https://subi.example.com)\n"
        "![速笔官网首页](assets/missing.png)\n\n## 三、建议\n- 做 A\n- 做 B\n- 做 C\n\n"
        "## 四、信息来源\n1. https://subi.example.com\n",
        encoding="utf-8",
    )
    proc = _run(CHECK, report, "--competitors", "速笔,云签", "--json")
    assert proc.returncode == 3
    res = json.loads(proc.stdout)
    issues = "\n".join(res["问题"])
    assert "宽泛表述" in issues
    assert "竞品“云签”没有官网或产品介绍链接" in issues
    assert "竞品“云签”没有产品截图或介绍图" in issues
    assert "关于“云签”的具体事实只有" in issues
    assert "竞品“速笔”没有官网" not in issues and "竞品“速笔”没有产品截图" not in issues
    notes = {(n["类型"], n["内容"]) for n in res["待确认"]}
    assert ("图片缺失", "assets/missing.png") in notes
    assert any(t == "表述宽泛" and "各有优劣" in c for t, c in notes)
    assert res["竞品档案"]["速笔"]["链接"] and res["竞品档案"]["速笔"]["图片"]


def test_docx_build_embeds_images_and_shows_urls(tmp_path: Path) -> None:
    (tmp_path / "assets").mkdir()
    from PIL import Image

    Image.new("RGB", (1600, 900), "navy").save(tmp_path / "assets" / "home.png")
    Image.new("RGB", (300, 600), "teal").save(tmp_path / "assets" / "app.webp")
    md = tmp_path / "报告.md"
    md.write_text(
        "# 报告\n\n官网：[速笔](https://subi.example.com)\n\n![速笔首页](assets/home.png)\n\n"
        "![速笔 App](assets/app.webp)\n\n![缺图](assets/none.png)\n",
        encoding="utf-8",
    )
    out = tmp_path / "报告.docx"
    proc = _run(REPORT / "scripts" / "docx_build.py", md, "-o", out, "--style", "report")
    assert proc.returncode == 0, proc.stderr
    assert "找不到图片：assets/none.png" in proc.stderr
    from docx import Document

    doc = Document(str(out))
    assert len(doc.inline_shapes) == 2
    widths = sorted(round(s.width.mm) for s in doc.inline_shapes)
    assert widths[0] == 79 and widths[1] <= 170  # 300px at 96 dpi; wide image capped at text width
    text = [p.text for p in doc.paragraphs]
    assert "官网：速笔（https://subi.example.com）" in text
    assert "速笔首页" in text and "〔图片缺失：assets/none.png〕" in text
