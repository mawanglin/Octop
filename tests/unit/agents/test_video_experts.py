"""Contract tests for the bundled video experts (脚本 / 分镜 / 剪辑 / 宣传片)."""

from __future__ import annotations

import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from types import ModuleType

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
SCRIPT_CHECK = LIBRARY / "video-scriptwriter" / "skills" / "script-check" / "scripts"
EXAMPLE_SCRIPT = (
    LIBRARY
    / "video-scriptwriter"
    / "skills"
    / "platform-scripts"
    / "references"
    / "example-script.md"
)
STORYBOARD = LIBRARY / "storyboard-artist" / "skills" / "storyboard-sheet"
MAKE_SRT = LIBRARY / "video-editing-coach" / "skills" / "subtitle-srt" / "scripts" / "make_srt.py"
VIDEO_TOOLS = (
    LIBRARY / "video-editing-coach" / "skills" / "video-processing" / "scripts" / "video_tools.py"
)
BUDGET = LIBRARY / "promo-video-planner" / "skills" / "production-budget"

EXPERTS = {
    "video-scriptwriter": (
        "短视频脚本专家",
        {"platform-scripts", "hook-and-title", "script-check"},
    ),
    "storyboard-artist": ("分镜头脚本专家", {"shot-design", "storyboard-sheet", "shooting-plan"}),
    "video-editing-coach": (
        "剪辑指导专家",
        {"editing-techniques", "edit-plan", "subtitle-srt", "video-processing"},
    ),
    "promo-video-planner": (
        "宣传片策划专家",
        {"creative-concept", "narration", "production-budget"},
    ),
}


def _run(
    script: Path, *args: str | Path, env: dict[str, str] | None = None
) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    full_env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
        **(env or {}),
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=full_env,
        timeout=180,
    )


def _load(script: Path) -> ModuleType:
    spec = importlib.util.spec_from_file_location(f"_test_{script.stem}", script)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # dataclasses look the module up by name
    previous = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = previous
    return module


# ------------------------------------------------------------------- catalog


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


@pytest.mark.parametrize("name", ["content_check.py", "docx_build.py", "budget_xlsx.py"])
def test_shared_script_copies_are_identical(name: str) -> None:
    copies = sorted(LIBRARY.glob(f"*/skills/*/scripts/{name}"))
    if len(copies) < 2:
        pytest.skip(f"only {len(copies)} copy of {name} in this tree")
    first = copies[0].read_bytes()
    for copy in copies[1:]:
        assert copy.read_bytes() == first, f"{copy} drifted from {copies[0]}"


# -------------------------------------------------------------- script timer


def test_script_timer_accepts_example_and_flags_problems(tmp_path: Path) -> None:
    ok = _run(SCRIPT_CHECK / "script_timer.py", EXAMPLE_SCRIPT)
    assert ok.returncode == 0, ok.stdout
    assert "总时长 45.0 秒" in ok.stdout

    bad = tmp_path / "bad.md"
    bad.write_text(
        "## [0:00-0:05] 开场\n口播：大家好我是某某今天给大家分享一个非常非常实用的小技巧\n"
        "## [0:06-0:08] 正文\n画面：屏幕\n口播：" + "很长的口播" * 10 + "\n",
        encoding="utf-8",
    )
    result = _run(SCRIPT_CHECK / "script_timer.py", bad, "--target", "30")
    assert result.returncode == 3
    for expected in ("超过 3 秒", "偏长", "空档", "没有写“画面：”", "删减约", "相差较大"):
        assert expected in result.stdout, expected


# ------------------------------------------------------------------ subtitles


def test_make_srt_aligns_script_segments(tmp_path: Path) -> None:
    out = tmp_path / "a.srt"
    result = _run(MAKE_SRT, EXAMPLE_SCRIPT, "-o", out, "--max-chars", "16")
    assert result.returncode == 0, result.stderr
    module = _load(MAKE_SRT)
    cues = module.parse_srt(out.read_text(encoding="utf-8"))
    assert cues[0].text == "周报写了一小时还是一片空白"
    assert cues[0].end <= 3.0  # first cue stays inside the 0:00-0:03 hook segment
    assert cues[-1].end <= 45.0
    assert all(module.units(c.text) <= 16 for c in cues)
    assert not any(c.text[-1] in "。，！？" for c in cues)
    assert all(b.start >= a.end for a, b in zip(cues, cues[1:], strict=False))
    assert _run(MAKE_SRT, "--check", out).returncode == 0

    broken = tmp_path / "b.srt"
    broken.write_text(
        "1\n00:00:00,000 --> 00:00:02,000\n第一句。\n\n2\n00:00:01,500 --> 00:00:01,900\n"
        "这一句字幕非常非常长而且显示时间特别短根本看不完\n",
        encoding="utf-8",
    )
    check = _run(MAKE_SRT, "--check", broken)
    assert check.returncode == 3
    for expected in ("重叠", "超过 16 字", "来不及看", "阅读压力大", "末尾"):
        assert expected in check.stdout, expected


# ----------------------------------------------------------------- storyboard


def test_storyboard_xlsx_builds_three_sheets(tmp_path: Path) -> None:
    out = tmp_path / "sb.xlsx"
    result = _run(
        STORYBOARD / "scripts" / "storyboard_xlsx.py",
        STORYBOARD / "references" / "example-storyboard.md",
        "-o",
        out,
        "--prompts",
    )
    assert result.returncode == 0, result.stderr
    assert "6 个镜头，2 个场景，总时长 29 秒" in result.stdout
    assert "提示" not in result.stdout
    wb = load_workbook(out)
    assert wb.sheetnames == ["分镜表", "场景汇总", "拍摄顺序"]
    ws = wb["分镜表"]
    header = [c.value for c in ws[4]]
    assert header[header.index("时长") + 1] == "累计"
    assert header[-1] == "AI 绘图提示词"
    assert str(ws.cell(6, header.index("累计") + 1).value).startswith("=SUM(")
    assert "medium shot" in str(ws.cell(6, len(header)).value)
    summary = [[c.value for c in r] for r in wb["场景汇总"].iter_rows(min_row=2)]
    assert summary == [["工位", 4, 20, "1、2、4、5"], ["会议室", 2, 9, "3、6"]]
    order = [r[1].value for r in wb["拍摄顺序"].iter_rows(min_row=2)]
    assert order == ["工位", "工位", "工位", "工位", "会议室", "会议室"]


def test_storyboard_xlsx_validates_input(tmp_path: Path) -> None:
    bad = tmp_path / "bad.md"
    bad.write_text("| 镜号 | 景别 |\n|---|---|\n| 1 | 特写 |\n", encoding="utf-8")
    result = _run(STORYBOARD / "scripts" / "storyboard_xlsx.py", bad, "-o", tmp_path / "x.xlsx")
    assert result.returncode == 1
    assert "画面" in result.stderr

    odd = tmp_path / "odd.md"
    odd.write_text(
        "| 镜号 | 画面 | 景别 | 运镜 | 时长 |\n|---|---|---|---|---|\n| 1 | 街景 | 半身 | 乱晃 | 20 |\n",
        encoding="utf-8",
    )
    result = _run(STORYBOARD / "scripts" / "storyboard_xlsx.py", odd, "-o", tmp_path / "y.xlsx")
    assert result.returncode == 0
    for expected in ("景别「半身」", "运镜「乱晃」", "偏长"):
        assert expected in result.stdout, expected


# --------------------------------------------------------------------- budget


def test_budget_xlsx_uses_formulas(tmp_path: Path) -> None:
    out = tmp_path / "budget.xlsx"
    result = _run(
        BUDGET / "scripts" / "budget_xlsx.py",
        BUDGET / "references" / "example-budget.md",
        "-o",
        out,
    )
    assert result.returncode == 0, result.stderr
    assert "估算总计 229,172.00 元" in result.stdout
    ws = load_workbook(out)["预算表"]
    values = {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}
    assert any(v.startswith("=D") and "*E" in v and "*F" in v for v in values)
    labels = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
    for label in ("小计合计", "管理费（10%）", "预备费（5%）", "税费（6%）", "总计"):
        assert label in labels
    total = ws.cell(labels.index("总计") + 1, 7).value
    assert str(total).startswith("=")

    bad = tmp_path / "bad.md"
    bad.write_text("| 项目 | 单价 | 数量 |\n|---|---|---|\n| 摄影 | 三千 | 1 |\n", encoding="utf-8")
    result = _run(BUDGET / "scripts" / "budget_xlsx.py", bad, "-o", tmp_path / "x.xlsx")
    assert result.returncode == 1
    assert "不是数字" in result.stderr


# ---------------------------------------------------------------- video tools


def test_video_tools_dry_run_and_guards(tmp_path: Path) -> None:
    src = tmp_path / "in.mp4"
    src.write_bytes(b"not really a video")
    dry = _run(
        VIDEO_TOOLS,
        "cut",
        src,
        "-o",
        tmp_path / "out.mp4",
        "--keep",
        "0:01-0:03",
        "--keep",
        "5-7.5",
        "--dry-run",
    )
    assert dry.returncode in (0, 2)  # 2 when the runner has no ffmpeg
    assert "trim=start=1.0:end=3.0" in dry.stdout
    assert "concat=n=2:v=1:a=1" in dry.stdout
    assert not (tmp_path / "out.mp4").exists()

    same = _run(VIDEO_TOOLS, "export", src, "--preset", "douyin", "-o", src, "--overwrite")
    assert same.returncode == 1
    assert "原素材" in same.stderr

    bad_range = _run(
        VIDEO_TOOLS, "cut", src, "-o", tmp_path / "o.mp4", "--keep", "0:05-0:02", "--dry-run"
    )
    assert bad_range.returncode == 1


def test_video_tools_reports_missing_ffmpeg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    module = _load(VIDEO_TOOLS)
    monkeypatch.setattr(module, "find_binary", lambda name, explicit: None)
    src = tmp_path / "in.mp4"
    src.write_bytes(b"x")
    code = module.main(["export", str(src), "--preset", "bilibili", "-o", str(tmp_path / "o.mp4")])
    out = capsys.readouterr()
    assert code == 2
    assert "scale=1920:1080" in out.out  # the command is still shown for manual use
    assert "未找到 ffmpeg" in out.err
    assert module.main(["which"]) == 2


@pytest.mark.skipif(
    shutil.which("ffmpeg") is None or shutil.which("ffprobe") is None, reason="needs ffmpeg"
)
def test_video_tools_real_pipeline(tmp_path: Path) -> None:
    src = tmp_path / "src.mp4"
    subprocess.run(
        ["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "testsrc=size=320x180:rate=25",
         "-f", "lavfi", "-i", "sine=frequency=440", "-t", "6", "-pix_fmt", "yuv420p",
         "-c:v", "libx264", "-c:a", "aac", str(src)],
        check=True,
    )  # fmt: skip
    cut = tmp_path / "cut.mp4"
    assert (
        _run(
            VIDEO_TOOLS, "cut", src, "-o", cut, "--keep", "0:01-0:02", "--keep", "3-4.5"
        ).returncode
        == 0
    )
    vertical = tmp_path / "v.mp4"
    assert _run(VIDEO_TOOLS, "export", cut, "--preset", "douyin", "-o", vertical).returncode == 0
    probe = _run(VIDEO_TOOLS, "probe", vertical)
    assert probe.returncode == 0, probe.stderr
    assert "1080×1920 竖屏" in probe.stdout
    assert re.search(r"时长：0:02\.[45]", probe.stdout), probe.stdout
    cover = tmp_path / "c.jpg"
    assert _run(VIDEO_TOOLS, "cover", vertical, "--at", "0:01", "-o", cover).returncode == 0
    assert cover.stat().st_size > 0
    again = _run(VIDEO_TOOLS, "cover", vertical, "--at", "0:01", "-o", cover)
    assert again.returncode == 1 and "--overwrite" in again.stderr
