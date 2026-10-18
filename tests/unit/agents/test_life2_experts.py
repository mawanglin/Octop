"""Contract tests for the bundled household experts (法律 / 财税 / 菜谱 / 证件)."""

from __future__ import annotations

import importlib.util
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from types import ModuleType

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
LEGAL = LIBRARY / "life-legal" / "skills" / "legal-calc" / "scripts" / "legal_calc.py"
FINANCE = LIBRARY / "family-finance" / "skills" / "finance-calc" / "scripts" / "finance_calc.py"
MENU = LIBRARY / "kitchen-chef" / "skills" / "menu-plan" / "scripts" / "menu_plan.py"
DOCS = LIBRARY / "home-docs" / "skills" / "doc-tracker"

EXPERTS = {
    "life-legal": ("生活法律助手", {"legal-calc", "legal-documents"}),
    "family-finance": ("家庭财税计算", {"finance-calc", "tax-guide"}),
    "kitchen-chef": ("厨房菜谱", {"menu-plan", "cooking-guide"}),
    "home-docs": ("家庭文档管家", {"doc-tracker", "paperwork-checklist"}),
}


def _run(script: Path, *args: str | Path) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=120,
    )


def _load(path: Path) -> ModuleType:
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location(f"_life2_{path.stem}", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
def test_expert_catalog_manifest_and_skills(expert_id: str) -> None:
    label_zh, skills = EXPERTS[expert_id]
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get(expert_id)
    assert expert is not None
    assert expert.summary.label_zh == label_zh
    assert expert.summary.icon_url == f"/experts/avatars/{expert_id}.svg"

    expert_dir = LIBRARY / expert_id
    manifest = json.loads((expert_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == expert_id
    for rel in manifest["prompt_files"]:
        assert (expert_dir / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (expert_dir / "skills").iterdir() if p.is_dir()} == skills
    soul = (expert_dir / "SOUL.md").read_text(encoding="utf-8")
    for name in skills:
        raw = (expert_dir / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (expert_dir / "skills" / name / rel).is_file(), f"{name}: {rel}"


@pytest.mark.parametrize("expert_id", sorted(EXPERTS))
async def test_expert_seeds_into_workspace(tmp_path: Path, expert_id: str) -> None:
    expert_dir = LIBRARY / expert_id
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(expert_dir)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=expert_dir, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (expert_dir / "manifest.json").read_bytes()
    )


# --------------------------------------------------------------------- legal


def test_severance_rules() -> None:
    out = _run(
        LEGAL,
        "severance",
        "--start",
        "2019-03-01",
        "--end",
        "2026-10-31",
        "--wage",
        "12000",
        "--city-avg",
        "12307",
    ).stdout
    assert "7 年 7 个月 30 天 → 经济补偿 N = 8" in out and "经济补偿（N）：96,000.00 元" in out

    out = _run(LEGAL, "severance", "--start", "2024-06-10", "--end", "2026-09-01", "--wage", "40000",
               "--city-avg", "12000", "--illegal").stdout  # fmt: skip
    assert "N = 2.5" in out and "计算基数：36,000.00" in out  # capped at 3× the city average
    assert "违法解除赔偿金（2N）：180,000.00 元" in out

    out = _run(LEGAL, "severance", "--start", "2010-01-01", "--end", "2026-03-01", "--wage", "50000",
               "--city-avg", "12000", "--no-notice").stdout  # fmt: skip
    assert "N = 12" in out and "最多按 12 年" in out and "代通知金（+1 个月）：50,000.00 元" in out
    assert (
        "N = 1\n"
        in _run(
            LEGAL, "severance", "--start", "2025-01-01", "--end", "2026-01-01", "--wage", "8000"
        ).stdout
    )
    assert (
        "N = 0.5"
        in _run(
            LEGAL, "severance", "--start", "2026-05-01", "--end", "2026-07-15", "--wage", "8000"
        ).stdout
    )
    assert (
        _run(
            LEGAL, "severance", "--start", "2026-05-01", "--end", "2026-01-01", "--wage", "8000"
        ).returncode
        == 1
    )


def test_overtime_and_deadlines() -> None:
    out = _run(
        LEGAL, "overtime", "--wage", "10000", "--weekday", "20", "--restday", "8", "--holiday", "8"
    ).stdout
    assert "日工资 459.77 元，小时工资 57.47 元" in out and "加班费合计：4,022.99 元" in out

    out = _run(LEGAL, "deadline", "--type", "appeal-judgment", "--from", "2026-10-18").stdout
    assert "最后一日：2026-11-02（一）" in out and "《民事诉讼法》第 171 条" in out
    out = _run(LEGAL, "deadline", "--type", "labor-arbitration", "--from", "2026-02-28").stdout
    assert "最后一日：2027-02-28（日）" in out and "顺延至 2027-03-01" in out
    assert "online-return" in _run(LEGAL, "deadline", "--list").stdout
    assert _run(LEGAL, "deadline", "--type", "nope", "--from", "2026-01-01").returncode == 1


# ------------------------------------------------------------------- finance


def test_mortgage_methods_and_prepayment(tmp_path: Path) -> None:
    out_xlsx = tmp_path / "m.xlsx"
    result = _run(FINANCE, "mortgage", "--principal", "1000000", "--rate", "3.1", "--years", "30",
                  "--prepay", "36:200000:shorten", "-o", out_xlsx)  # fmt: skip
    assert result.returncode == 0, result.stderr
    out = result.stdout
    assert "等额本息：首月 4,270.16 元（每月相同），总利息 537,259.04 元" in out
    assert "等额本金：首月 5,361.11 元" in out and "总利息 466,291.67 元" in out
    assert "共 265 期（少还 95 期）" in out and "节省利息 207,447.61 元" in out
    wb = load_workbook(out_xlsx)
    assert wb.sheetnames == ["等额本息", "等额本息（提前还款）", "等额本金", "等额本金（提前还款）"]
    ws = wb["等额本息（提前还款）"]
    assert ws.cell(37, 5).value == 200000 and ws.cell(ws.max_row, 1).value == "合计"

    reduce = _run(
        FINANCE,
        "mortgage",
        "--principal",
        "1000000",
        "--rate",
        "3.1",
        "--years",
        "30",
        "--prepay",
        "36:200000:reduce",
    ).stdout
    assert "共 360 期（少还 0 期）" in reduce and "月供约 3,358.17 元" in reduce
    assert (
        _run(
            FINANCE, "mortgage", "--principal", "1", "--rate", "3", "--years", "30", "--prepay", "x"
        ).returncode
        == 1
    )


def test_income_tax_and_bonus() -> None:
    result = _run(FINANCE, "tax", "--salary", "20000", "--base", "15000", "--children", "1", "--elderly", "3000",
                  "--loan", "--bonus", "36001")  # fmt: skip
    assert result.returncode == 0, result.stderr
    out = result.stdout
    assert "个人社保公积金 2,625.00 元/月，专项附加扣除 6,000.00 元/月" in out
    assert "| 1 月 | 191.25 |" in out and "| 6 月 | 348.75 |" in out and "| 12 月 | 637.50 |" in out
    assert "全年工资个税 5,130.00 元" in out
    assert "单独计税 3,390.10 元，并入综合所得增加 3,600.10 元 → 选择「单独计税」" in out
    assert "临界点" in out and "10%" in out

    mod = _load(FINANCE)
    assert mod.annual_tax(36000) == pytest.approx(1080) and mod.annual_tax(144000) == pytest.approx(
        11880
    )
    assert mod.bonus_tax(36000) == pytest.approx(1080) and mod.bonus_tax(144000) == pytest.approx(
        14190
    )
    bad = _run(FINANCE, "tax", "--salary", "20000", "--loan", "--rent", "1500")
    assert bad.returncode == 1 and "二选一" in bad.stderr


def test_social_insurance_caps_base() -> None:
    out = _run(FINANCE, "insurance", "--base", "40000", "--max", "36921", "--fund", "12").stdout
    assert "缴费基数 36,921.00 元（按当地基数上限调整" in out
    assert "| 养老保险 8% | 2,953.68 |" in out and "| 合计 | 8,307.23 |" in out
    assert "个人账户每月共入账 8,861.04 元" in out


# ---------------------------------------------------------------------- menu


def test_menu_plan_respects_avoid_and_merges(tmp_path: Path) -> None:
    out_xlsx = tmp_path / "menu.xlsx"
    result = _run(
        MENU,
        "plan",
        "--days",
        "7",
        "--people",
        "3",
        "--avoid",
        "辣,海鲜",
        "--seed",
        "1",
        "-o",
        out_xlsx,
    )
    assert result.returncode == 0, result.stderr
    out = result.stdout
    meals = re.findall(r"^- 第 \d+ 天 (?:午餐|晚餐)：(.+)$", out, re.M)
    assert len(meals) == 14
    dishes = [d for m in meals for d in m.split("、")]
    for banned in ("宫保鸡丁", "麻婆豆腐", "清蒸鲈鱼", "油焖大虾", "虾仁炒蛋", "酸辣土豆丝"):
        assert banned not in dishes
    recipes = _load(MENU).load_recipes(MENU.parents[1] / "references" / "recipes.csv")
    meat_count = sum(
        1 for r in recipes.values() if r.kind == "荤菜" and not (r.tags & {"辣", "海鲜"})
    )
    assert len(set(dishes)) >= min(len(dishes), meat_count)  # no repeats until the pool runs out
    assert "**肉禽**" in out and "**蛋奶豆制品**：鸡蛋" in out  # eggs aren't counted as poultry
    assert "盐" not in out.split("## 买菜清单")[1].split("常备调料")[0]
    wb = load_workbook(out_xlsx)
    assert wb.sheetnames == ["菜单", "买菜清单"] and wb["买菜清单"]["E1"].value == "已买"
    assert (
        _run(MENU, "plan", "--days", "7", "--seed", "1").stdout
        == _run(MENU, "plan", "--days", "7", "--seed", "1").stdout
    )


def test_menu_shopping_from_markdown(tmp_path: Path) -> None:
    md = tmp_path / "m.md"
    md.write_text(
        "- 周六：红烧肉、蒜蓉西兰花、紫菜蛋花汤\n- 周日：可乐鸡翅、番茄炒菜花、酸菜鱼\n",
        encoding="utf-8",
    )
    out = _run(MENU, "shopping", md, "--people", "2").stdout
    assert "共 5 道菜，2 人份" in out
    assert "鸡翅中 300 克" in out and "五花肉 240 克" in out and "西兰花 300 克" in out
    assert "**其他**：可乐 200 毫升" in out
    assert "不在菜谱库中的菜（食材需要手动补充）：酸菜鱼" in out
    assert _run(MENU, "plan", "--avoid", "肉,蛋,豆腐,鸡,虾,鱼,排骨,猪,牛,羊").returncode == 1


# ---------------------------------------------------------------------- docs


def test_doc_tracker_build_due_and_ics(tmp_path: Path) -> None:
    script = DOCS / "scripts" / "doc_tracker.py"
    src = tmp_path / "docs.md"
    src.write_text(
        (DOCS / "references" / "example-docs.md")
        .read_text(encoding="utf-8")
        .replace("| 妈妈 | 1234 |", "| 妈妈 | 110101199001011234 |"),
        encoding="utf-8",
    )
    book = tmp_path / "book.xlsx"
    result = _run(script, "build", src, "-o", book, "--today", "2026-10-18")
    assert result.returncode == 0, result.stderr
    assert "共 6 项：已过期 1，即将到期 3，长期有效 1" in result.stdout
    assert "- 冰箱保修（家庭）：2026-09-01，已过期 47 天，存放：发票夹" in result.stdout
    assert (
        "- 护照（爸爸）：2027-03-15，还剩 148 天" in result.stdout
    )  # 180-day default lead for passports
    ws = load_workbook(book)["台账"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["身份证"][3] == "****1234"  # full ID numbers are masked
    assert rows["房屋租赁合同"][5] == 60 and rows["房屋租赁合同"][9] == "正常"
    assert rows["户口本"][9] == "长期"
    assert [r[0] for r in ws.iter_rows(min_row=2, max_row=2, values_only=True)] == ["冰箱保修"]
    assert ws["A2"].fill.fgColor.rgb.endswith("FECACA")

    due = _run(script, "due", book, "--days", "300", "--today", "2026-10-18").stdout
    assert "房屋租赁合同" in due

    ics = tmp_path / "r.ics"
    result = _run(script, "ics", book, "-o", ics, "--today", "2026-10-18")
    assert "4 个提醒" in result.stdout
    text = ics.read_text(encoding="utf-8")
    assert ics.read_bytes().startswith(b"BEGIN:VCALENDAR\r\n") and text.count("BEGIN:VEVENT") == 4
    assert "冰箱保修" not in text and "户口本" not in text
    assert "DTSTART;VALUE=DATE:20270501" in text  # lease: 60 days before 2027-06-30
    assert "DTSTART;VALUE=DATE:20261018" in text  # already inside the lead window → today

    bad = tmp_path / "bad.md"
    bad.write_text("| 名称 | 到期日 |\n|---|---|\n| 护照 | 明年 |\n", encoding="utf-8")
    assert _run(script, "build", bad, "-o", tmp_path / "x.xlsx").returncode == 1
