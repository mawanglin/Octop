"""Contract tests for the bundled 版权登记与存证专家 (copyright-advisor) expert."""

from __future__ import annotations

import hashlib
import json
import os
import random
import re
import subprocess
import sys
from pathlib import Path

import pytest
from deepagents.backends.local_shell import LocalShellBackend
from octop_harness.backends.workspace import BackendWorkspace
from openpyxl import load_workbook

from octop.infra.agents.experts.catalog import (
    WORKSPACE_MANIFEST_PATH,
    ExpertCatalog,
    default_library_root,
    discover_seed_paths,
    seed_expert_directory,
)

LIBRARY = default_library_root()
EXPERT_DIR = LIBRARY / "copyright-advisor"
SKILLS = {
    "sample-inspect",
    "evidence-hash",
    "softcopy-doc",
    "copyright-registration",
    "copyright-policy",
}
SAMPLE = EXPERT_DIR / "skills" / "sample-inspect" / "scripts" / "sample_inspect.py"
EVIDENCE = EXPERT_DIR / "skills" / "evidence-hash" / "scripts" / "evidence_hash.py"
SOFTCOPY = EXPERT_DIR / "skills" / "softcopy-doc" / "scripts" / "softcopy_doc.py"
POLICY = EXPERT_DIR / "skills" / "copyright-policy"


def _run(
    script: Path, *args: str | Path, cwd: Path | None = None
) -> subprocess.CompletedProcess[str]:
    # No __pycache__ inside the template dir: seeding copies every file it finds.
    env = {
        **os.environ,
        "PYTHONIOENCODING": "utf-8",
        "PYTHONUTF8": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
    }
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        timeout=180,
        cwd=cwd,
    )


def test_expert_catalog_manifest_and_skills() -> None:
    catalog = ExpertCatalog(LIBRARY)
    catalog.refresh()
    expert = catalog.get("copyright-advisor")
    assert expert is not None
    assert expert.summary.label_zh == "版权登记与存证专家"
    assert expert.summary.icon_url == "/experts/avatars/copyright-advisor.svg"

    manifest = json.loads((EXPERT_DIR / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["id"] == "copyright-advisor"
    for rel in manifest["prompt_files"]:
        assert (EXPERT_DIR / rel).is_file()
    assert len(manifest["quick_prompts"]) == 6
    assert len(manifest["task_examples"]["zh"]) == len(manifest["task_examples"]["en"]) == 6

    assert {p.name for p in (EXPERT_DIR / "skills").iterdir() if p.is_dir()} == SKILLS
    soul = (EXPERT_DIR / "SOUL.md").read_text(encoding="utf-8")
    for name in SKILLS:
        raw = (EXPERT_DIR / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
        assert raw.startswith(f"---\nname: {name}\n")
        assert "\ndescription:" in raw.split("---", 2)[1]
        assert f"`{name}`" in soul
        for rel in re.findall(r"`((?:scripts|references)/[\w.-]+)`", raw):
            assert (EXPERT_DIR / "skills" / name / rel).is_file(), f"{name}: {rel}"


async def test_expert_seeds_into_workspace(tmp_path: Path) -> None:
    workspace = BackendWorkspace(
        LocalShellBackend(root_dir=str(tmp_path), virtual_mode=False), tmp_path
    )
    paths = discover_seed_paths(EXPERT_DIR)
    assert not [p for p in paths if "__pycache__" in p]
    copied = await seed_expert_directory(expert_dir=EXPERT_DIR, workspace=workspace)
    assert copied == len(paths) + 1
    assert (
        await workspace.adownload_bytes(WORKSPACE_MANIFEST_PATH)
        == (EXPERT_DIR / "manifest.json").read_bytes()
    )


# ------------------------------------------------------------------ samples


ESSAY = (
    "春风又绿江南岸，明月何时照我还。这是一篇关于故乡与成长的散文，作者回忆了童年在江南小镇度过的时光，描写了青石板路、乌篷船和外婆的桂花糕。"
    * 3
)


def _samples(root: Path) -> Path:
    from docx import Document
    from PIL import Image, ImageDraw, ImageOps, PngImagePlugin
    from pypdf import PdfWriter

    root.mkdir()
    rng = random.Random(1)
    art = Image.new("RGB", (640, 480), "white")
    draw = ImageDraw.Draw(art)
    for _ in range(30):
        x, y = rng.randint(0, 600), rng.randint(0, 440)
        color = (rng.randint(0, 255), rng.randint(0, 255), rng.randint(0, 255))
        draw.ellipse([x, y, x + rng.randint(20, 120), y + rng.randint(20, 120)], fill=color)
    exif = Image.Exif()
    exif[0x013B] = "张三".encode()  # UTF-8 in an ASCII tag, as many tools write it
    exif[0x8298] = "© 2024 张三".encode()
    exif[0x0131] = "Adobe Photoshop 25.0"
    exif[0x0132] = "2024:05:01 10:00:00"
    art.save(root / "原作.jpg", exif=exif.tobytes(), quality=95)

    copy = Image.eval(
        ImageOps.mirror(art).resize((500, 375)), lambda v: min(255, int(v * 0.9) + 15)
    )
    copy.save(root / "疑似盗用.jpg", quality=80)

    ai = Image.new("RGB", (512, 512), (30, 60, 90))
    info = PngImagePlugin.PngInfo()
    info.add_text(
        "parameters", "a cat in ink style, Negative prompt: blurry, Steps: 30, Model: sdxl_base"
    )
    ai.save(root / "AI插画.png", pnginfo=info)

    other = Image.new("RGB", (640, 480), "black")
    ImageDraw.Draw(other).line([0, 0, 640, 480], fill="yellow", width=40)
    other.save(root / "无关.png")

    doc = Document()
    doc.add_paragraph(ESSAY)
    doc.core_properties.author = "李四"
    doc.core_properties.last_modified_by = "张三"
    doc.save(str(root / "散文原稿.docx"))
    doc = Document()
    doc.add_paragraph("前言：本文转载自网络。" + ESSAY[:150] + "以上内容仅供学习。")
    doc.core_properties.author = "Administrator"
    doc.save(str(root / "网文转载.docx"))
    (root / "无关文章.txt").write_text(
        "区块链存证利用分布式账本记录电子数据的哈希值。" * 5, encoding="utf-8"
    )

    pdf = PdfWriter()
    pdf.add_blank_page(595, 842)
    pdf.add_metadata({"/Author": "张三", "/CreationDate": "D:20240301120000+08'00'"})
    with (root / "扫描件.pdf").open("wb") as fh:
        pdf.write(fh)
    return root


def test_sample_scan_reports_metadata_and_clues(tmp_path: Path) -> None:
    root = _samples(tmp_path / "samples")
    report = tmp_path / "out" / "样本报告.xlsx"
    proc = _run(SAMPLE, "scan", root, "--author", "张三", "--json", "-o", report)
    assert proc.returncode == 0, proc.stderr
    data = json.loads(proc.stdout.split("\n\n报告：")[0])
    rows = {Path(r["文件"]).name: r for r in data["samples"]}
    art = rows["原作.jpg"]
    assert art["作者"] == "张三" and art["版权"] == "© 2024 张三"
    assert art["软件"] == "Adobe Photoshop 25.0" and art["修改时间"] == "2024-05-01 10:00:00"
    assert rows["AI插画.png"]["AI痕迹"] == "是"
    assert "sdxl_base" in rows["AI插画.png"]["生成参数"]
    assert rows["散文原稿.docx"]["作者"] == "李四"
    assert (
        rows["扫描件.pdf"]["作者"] == "张三"
        and rows["扫描件.pdf"]["创建时间"] == "2024-03-01 12:00:00"
    )
    assert len(art["SHA-256"]) == 64

    notes = "\n".join(data["findings"])
    assert "AI插画.png：发现 AI 生成痕迹" in notes
    assert "散文原稿.docx：作者字段为「李四」" in notes
    assert "最后修改人字段为「张三」" not in notes
    assert "网文转载.docx：作者字段是软件默认值「Administrator」" in notes
    assert "扫描件.pdf：没有文字层" in notes
    assert "疑似盗用.jpg：图片没有任何元数据" in notes
    assert "原作.jpg" not in notes and "AI插画.png：图片没有任何元数据" not in notes

    wb = load_workbook(report)
    assert wb.sheetnames == ["样本清单", "线索提示"]
    assert wb["样本清单"].max_row == 9


def test_sample_compare_finds_copies(tmp_path: Path) -> None:
    root = _samples(tmp_path / "samples")
    proc = _run(SAMPLE, "compare", root, "--json")
    assert proc.returncode == 0, proc.stderr
    pairs = {(Path(r["作品A"]).name, Path(r["作品B"]).name): r for r in json.loads(proc.stdout)}
    image = pairs[("原作.jpg", "疑似盗用.jpg")]
    assert image["结论"] == "高度相似" and "镜像" in image["说明"]
    assert pairs[("原作.jpg", "无关.png")]["结论"] == "不相似"
    text = pairs[("散文原稿.docx", "网文转载.docx")]
    assert text["结论"] == "高度相似"
    assert any("青石板路" in p for p in text["相同段落"])
    assert pairs[("散文原稿.docx", "无关文章.txt")]["结论"] == "不相似"

    proc = _run(SAMPLE, "compare", root, "--ref", root / "原作.jpg", "-o", tmp_path / "比对.xlsx")
    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.splitlines()[0].startswith("- [高度相似] 原作.jpg ↔ 疑似盗用.jpg")
    assert "构成侵权" not in proc.stdout.replace("是否构成侵权", "")
    assert (tmp_path / "比对.xlsx").is_file()

    lone = tmp_path / "lone"
    lone.mkdir()
    (lone / "a.txt").write_text("只有一个文件", encoding="utf-8")
    assert _run(SAMPLE, "compare", lone).returncode == 1
    assert _run(SAMPLE, "scan", tmp_path / "missing").returncode == 1


# ----------------------------------------------------------------- evidence


def test_evidence_manifest_make_and_verify(tmp_path: Path) -> None:
    works = tmp_path / "works"
    (works / "系列").mkdir(parents=True)
    (works / "封面.png").write_bytes(b"\x89PNG fake image bytes")
    (works / "系列" / "第一章.txt").write_text("第一章 正文", encoding="utf-8")
    out = tmp_path / "存证" / "清单.csv"
    proc = _run(EVIDENCE, "make", works, "-o", out, "--owner", "张三", "--title", "江南", "--xlsx")
    assert proc.returncode == 0, proc.stderr

    sha_file = out.with_suffix(".sha256")
    lines = sha_file.read_text(encoding="utf-8").splitlines()
    assert lines == [
        f"{hashlib.sha256((works / '封面.png').read_bytes()).hexdigest()}  封面.png",
        f"{hashlib.sha256((works / '系列' / '第一章.txt').read_bytes()).hexdigest()}  系列/第一章.txt",
    ]
    total = hashlib.sha256(sha_file.read_bytes()).hexdigest()
    assert total in proc.stdout
    csv_text = out.read_text(encoding="utf-8-sig")
    assert f"# 清单总指纹（.sha256 文件的 SHA-256）,{total}" in csv_text
    assert "# 权利人,张三" in csv_text
    assert (
        load_workbook(out.with_suffix(".xlsx"))["存证清单"]["A1"].value == "作品存证文件清单：江南"
    )

    ok = _run(EVIDENCE, "verify", out)
    assert ok.returncode == 0, ok.stdout
    assert "一致 2，被修改 0，缺失 0" in ok.stdout

    (works / "系列" / "第一章.txt").write_text("第一章 改过", encoding="utf-8")
    (works / "封面.png").unlink()
    (works / "新增.txt").write_text("新", encoding="utf-8")
    bad = _run(EVIDENCE, "verify", sha_file, "--root", works)
    assert bad.returncode == 3
    assert "[被修改] 系列/第一章.txt" in bad.stdout
    assert "[缺失] 封面.png" in bad.stdout
    assert "[清单外] 新增.txt" in bad.stdout

    assert _run(EVIDENCE, "make", tmp_path / "nope", "-o", out).returncode == 1


# ----------------------------------------------------------------- softcopy


def _repo(root: Path, functions: int) -> None:
    (root / "app").mkdir(parents=True)
    (root / "node_modules" / "lib").mkdir(parents=True)
    body = "".join(
        f"def func_{j}(x):\n    # 计算第 {j} 项\n    return x * {j}\n\n" for j in range(functions)
    )
    (root / "app" / "main.py").write_text(body, encoding="utf-8")
    (root / "app" / "config.py").write_text(
        'DB_PASSWORD = "Super' + 'Secret123"\nTITLE = "' + "很长的标题" * 40 + '"\n',
        encoding="utf-8",
    )
    (root / "node_modules" / "lib" / "index.js").write_text(
        "console.log(1)\n" * 500, encoding="utf-8"
    )
    (root / "app" / "bundle.min.js").write_text("var a=1;" * 50, encoding="utf-8")


def _docx_pages(path: Path) -> tuple[list[str], int, str]:
    from docx import Document

    doc = Document(str(path))
    breaks = sum(1 for p in doc.paragraphs for r in p.runs if 'w:type="page"' in r._r.xml)
    return [p.text for p in doc.paragraphs], breaks, doc.sections[0].header.paragraphs[0].text


def test_softcopy_takes_first_and_last_30_pages(tmp_path: Path) -> None:
    _repo(tmp_path / "repo", functions=1200)  # 3600 source lines -> 73 pages
    out = tmp_path / "源程序.docx"
    proc = _run(
        SOFTCOPY, tmp_path / "repo", "--name", "演示管理系统", "--version", "V2.1", "-o", out
    )
    assert proc.returncode == 0, proc.stderr
    assert "收录 2 个文件" in proc.stdout
    assert "提交前 30 页和后 30 页，文档共 60 页" in proc.stdout
    assert "疑似密码或密钥" in proc.stdout and "config.py:1" in proc.stdout

    lines, breaks, header = _docx_pages(out)
    assert len(lines) == 60 * 50 and breaks == 59
    assert header.startswith("演示管理系统 V2.1\t第 ")
    assert lines[0].startswith("DB_PASSWORD")  # files in path order: config.py before main.py
    assert lines[-1] == "    return x * 1199"
    assert all(line.strip() for line in lines)
    assert not any("console.log" in line or "var a=1" in line for line in lines)
    widths = [sum(2 if ord(ch) > 0x2E80 else 1 for ch in line) for line in lines]
    assert max(widths) <= 90


def test_softcopy_small_project_and_errors(tmp_path: Path) -> None:
    _repo(tmp_path / "repo", functions=50)
    out = tmp_path / "小.docx"
    proc = _run(
        SOFTCOPY,
        tmp_path / "repo",
        "--name",
        "小工具",
        "-o",
        out,
        "--strip-comments",
        "--ext",
        "py",
    )
    assert proc.returncode == 0, proc.stderr
    assert "全部提交" in proc.stdout
    lines, breaks, header = _docx_pages(out)
    assert not any(line.strip().startswith("#") for line in lines)
    assert breaks == len(lines) // 50 - (len(lines) % 50 == 0)
    assert header.startswith("小工具 V1.0")

    assert _run(SOFTCOPY, tmp_path / "missing", "--name", "x", "-o", out).returncode == 1
    empty = tmp_path / "empty"
    empty.mkdir()
    assert _run(SOFTCOPY, empty, "--name", "x", "-o", out).returncode == 1


# ------------------------------------------------------------------- policy


@pytest.mark.parametrize("name", ["example-2010.md", "example-2020.md", "copyright-laws.md"])
def test_policy_references_exist(name: str) -> None:
    assert (POLICY / "references" / name).read_text(encoding="utf-8").strip()


def test_policy_diff_shows_2020_amendments() -> None:
    refs = POLICY / "references"
    proc = _run(
        POLICY / "scripts" / "text_diff.py", refs / "example-2010.md", refs / "example-2020.md"
    )
    assert proc.returncode == 0, proc.stderr
    assert "【+视听+】" in proc.stdout
    assert "【+百元以上五百+】" in proc.stdout
    assert "| 新增 |" in proc.stdout
