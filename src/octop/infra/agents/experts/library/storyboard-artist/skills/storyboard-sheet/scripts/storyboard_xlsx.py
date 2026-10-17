#!/usr/bin/env python3
"""分镜表 Markdown → 拍摄现场用的 Excel（分镜表 / 场景汇总 / 拍摄顺序）。

Usage:
    python scripts/storyboard_xlsx.py storyboard.md -o storyboard.xlsx [--prompts]

输入：可选 front matter（项目、版本、画幅、风格）+ 一张 Markdown 表格，表头使用：
    镜号 | 场景 | 景别 | 运镜 | 画面 | 台词/旁白 | 字幕 | 音效/音乐 | 时长 | 道具 | 备注
（列可以少，名称需一致；其他列原样保留。）时长写秒数，如 3、3s、0:03。
--prompts 追加一列「AI 绘图提示词」（由景别、画面、风格、画幅拼成，可直接粘贴到绘图/视频生成工具）。
Exit codes: 0 ok（可能有提示）, 1 输入错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

SHOT_SIZES = ("大远景", "远景", "全景", "中景", "中近景", "近景", "特写", "大特写")
MOVES = (
    "固定", "推", "拉", "摇", "移", "跟", "升", "降", "升降", "环绕", "甩", "变焦",
    "手持", "航拍", "滑轨", "延时", "升格", "降格", "希区柯克变焦", "主观镜头",
)  # fmt: skip
WIDTHS = {
    "镜号": 6, "场景": 12, "景别": 8, "运镜": 10, "画面": 40, "台词/旁白": 32, "字幕": 20,
    "音效/音乐": 16, "时长": 7, "累计": 8, "道具": 14, "备注": 16, "AI 绘图提示词": 48,
}  # fmt: skip
SIZE_EN = {
    "大远景": "extreme wide shot", "远景": "wide shot", "全景": "full shot", "中景": "medium shot",
    "中近景": "medium close-up", "近景": "close-up", "特写": "close-up detail", "大特写": "extreme close-up",
}  # fmt: skip
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
SCENE_FILLS = ("F3F4F6", "FFFFFF")
THIN = Side(style="thin", color="D1D5DB")


class InputError(ValueError):
    pass


def parse_duration(value: str) -> float | None:
    v = value.strip().lower().rstrip("s秒").strip()
    if not v:
        return None
    try:
        if ":" in v:
            m, s = v.split(":", 1)
            return int(m) * 60 + float(s)
        return float(v)
    except ValueError:
        return None


def parse(source: str) -> tuple[dict[str, str], list[str], list[list[str]]]:
    lines = source.replace("\r\n", "\n").split("\n")
    meta: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                for raw in lines[1:i]:
                    key, sep, value = raw.replace("：", ":", 1).partition(":")
                    if sep:
                        meta[key.strip()] = value.strip()
                lines = lines[i + 1 :]
                break
    rows = [
        [c.strip() for c in ln.strip().strip("|").split("|")]
        for ln in lines
        if ln.strip().startswith("|") and ln.strip().endswith("|")
    ]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    if len(rows) < 2:
        raise InputError("没有找到分镜表（需要一张带表头的 Markdown 表格）")
    header, body = rows[0], rows[1:]
    if "镜号" not in header or "画面" not in header:
        raise InputError("表头至少需要「镜号」和「画面」两列")
    for i, r in enumerate(body, start=1):
        if len(r) != len(header):
            raise InputError(f"第 {i} 行有 {len(r)} 列，表头有 {len(header)} 列")
    return meta, header, body


def prompt_for(row: dict[str, str], meta: dict[str, str]) -> str:
    parts = [row.get("画面", "")]
    size = row.get("景别", "")
    if size in SIZE_EN:
        parts.append(f"{size}（{SIZE_EN[size]}）")
    if row.get("场景"):
        parts.append(f"场景：{row['场景']}")
    if meta.get("风格"):
        parts.append(f"风格：{meta['风格']}")
    parts.append(f"画幅 {meta.get('画幅', '9:16')}")
    parts.append("电影感光影，无文字，无水印")
    return "，".join(p for p in parts if p)


def build(
    meta: dict[str, str], header: list[str], body: list[list[str]], out: Path, prompts: bool
) -> list[str]:
    warnings: list[str] = []
    records = [dict(zip(header, r, strict=True)) for r in body]
    durations: list[float | None] = []
    for rec in records:
        shot = rec.get("镜号", "?")
        size = rec.get("景别", "")
        if size and size not in SHOT_SIZES:
            warnings.append(f"镜号 {shot}：景别「{size}」不是常用写法（{'/'.join(SHOT_SIZES)}）")
        move = rec.get("运镜", "")
        if move and not any(m in move for m in MOVES):
            warnings.append(f"镜号 {shot}：运镜「{move}」不是常用术语，确认摄影师能看懂")
        d = parse_duration(rec.get("时长", "")) if "时长" in rec else None
        if "时长" in rec and d is None:
            warnings.append(f"镜号 {shot}：时长「{rec.get('时长')}」无法识别，请写秒数")
        if d is not None and d > 15:
            warnings.append(f"镜号 {shot}：单镜头 {d:g} 秒偏长，短视频单镜头一般 1–6 秒")
        durations.append(d)

    columns = list(header)
    if "时长" in columns:
        columns.insert(columns.index("时长") + 1, "累计")
    if prompts:
        columns.append("AI 绘图提示词")

    wb = Workbook()
    ws = wb.active
    ws.title = "分镜表"
    title = meta.get("项目", "分镜头脚本")
    info = "　".join(f"{k}：{meta[k]}" for k in ("版本", "画幅", "风格") if meta.get(k))
    ws.cell(1, 1, title).font = Font(size=16, bold=True)
    ws.cell(2, 1, info).font = Font(size=10, color="6B7280")
    head_row = 4
    for j, name in enumerate(columns, start=1):
        c = ws.cell(head_row, j, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(name, 14)
    scene_order: list[str] = []
    for i, rec in enumerate(records):
        r = head_row + 1 + i
        scene = rec.get("场景", "")
        if scene not in scene_order:
            scene_order.append(scene)
        fill = PatternFill("solid", fgColor=SCENE_FILLS[scene_order.index(scene) % 2])
        for j, name in enumerate(columns, start=1):
            if name == "累计":
                dur_col = get_column_letter(columns.index("时长") + 1)
                value: object = f"=SUM(${dur_col}${head_row + 1}:{dur_col}{r})"
            elif name == "时长":
                value = durations[i] if durations[i] is not None else rec.get("时长", "")
            elif name == "AI 绘图提示词":
                value = prompt_for(rec, meta)
            elif name == "镜号" and rec.get("镜号", "").isdigit():
                value = int(rec["镜号"])
            else:
                value = rec.get(name, "")
            c = ws.cell(r, j, value)
            c.alignment = Alignment(
                vertical="top",
                wrap_text=True,
                horizontal="center" if name in ("镜号", "景别", "时长", "累计") else "left",
            )
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            c.fill = fill
    last = head_row + len(records)
    if "时长" in columns:
        dur_col = get_column_letter(columns.index("时长") + 1)
        total_row = last + 1
        ws.cell(total_row, 1, "合计").font = Font(bold=True)
        ws.cell(
            total_row, columns.index("时长") + 1, f"=SUM({dur_col}{head_row + 1}:{dur_col}{last})"
        ).font = Font(bold=True)
    ws.freeze_panes = ws.cell(head_row + 1, 2)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(columns))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{head_row}:{head_row}"

    # 场景汇总
    sm = wb.create_sheet("场景汇总")
    for j, name in enumerate(("场景", "镜头数", "总时长（秒）", "镜号"), start=1):
        c = sm.cell(1, j, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for i, scene in enumerate(scene_order, start=2):
        shots = [rec for rec in records if rec.get("场景", "") == scene]
        total = sum(
            d or 0
            for rec, d in zip(records, durations, strict=True)
            if rec.get("场景", "") == scene
        )
        sm.cell(i, 1, scene or "（未填场景）")
        sm.cell(i, 2, len(shots))
        sm.cell(i, 3, round(total, 1))
        sm.cell(i, 4, "、".join(rec.get("镜号", "") for rec in shots))
    for col, width in zip("ABCD", (16, 10, 14, 40), strict=True):
        sm.column_dimensions[col].width = width

    # 拍摄顺序：同一场景集中拍，减少转场
    so = wb.create_sheet("拍摄顺序")
    keep = [c for c in ("场景", "镜号", "景别", "运镜", "画面", "道具", "时长") if c in header]
    for j, name in enumerate(["顺序", *keep], start=1):
        c = so.cell(1, j, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        so.column_dimensions[get_column_letter(j)].width = WIDTHS.get(name, 8)
    ordered = sorted(records, key=lambda rec: scene_order.index(rec.get("场景", "")))
    for i, rec in enumerate(ordered, start=2):
        so.cell(i, 1, i - 1)
        for j, name in enumerate(keep, start=2):
            c = so.cell(i, j, rec.get(name, ""))
            c.alignment = Alignment(vertical="top", wrap_text=True)
    so.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    total = sum(d or 0 for d in durations)
    print(f"已生成 {out}：{len(records)} 个镜头，{len(scene_order)} 个场景，总时长 {total:g} 秒")
    return warnings


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input")
    parser.add_argument("-o", "--output", required=True)
    parser.add_argument("--prompts", action="store_true", help="追加 AI 绘图提示词列")
    args = parser.parse_args(argv)
    try:
        meta, header, body = parse(Path(args.input).read_text(encoding="utf-8"))
        warnings = build(meta, header, body, Path(args.output), args.prompts)
    except (OSError, InputError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    for w in warnings:
        print(f"提示：{w}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
