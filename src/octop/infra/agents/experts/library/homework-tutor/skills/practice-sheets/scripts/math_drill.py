#!/usr/bin/env python3
"""数学练习卷：按题型批量出题（答案由程序计算，保证正确），导出可打印的 Word，答案另起一页。

Usage:
    python scripts/math_drill.py --type add --max 100 --carry require -n 40 -o 练习.docx
    python scripts/math_drill.py --type mixed -n 20 --cols 2 --title "三年级混合运算" -o 练习.docx
    python scripts/math_drill.py --list

题型（--type）：
    add / sub        加法 / 减法；--max 最大数，--carry require|avoid|any 是否要求进位（减法为退位）
    addsub           加减混合两步
    table            乘法口诀（1–9）
    mul              两位数×一位数（--max 控制第一个因数上限，如 999 为三位数×一位数）
    mul2             两位数×两位数
    div              整除（表内或多位数÷一位数，--max 控制被除数上限）
    divr             有余数的除法
    mixed            三步混合运算（含括号，结果为正整数）
    decimal          小数加减（一到两位小数）
    fraction         分数加减（同分母或异分母，结果化为最简）
    equation         一元一次方程 ax + b = c（整数解）
--seed 固定随机种子可复现同一张卷子。Exit codes: 0 ok, 1 参数错误, 2 缺少 python-docx。
"""

from __future__ import annotations

import argparse
import random
import re
import sys
from collections.abc import Callable
from fractions import Fraction
from pathlib import Path

Problem = tuple[str, str]  # (题目, 答案)
TYPES: dict[str, str] = {
    "add": "加法", "sub": "减法", "addsub": "加减混合", "table": "乘法口诀", "mul": "多位数×一位数",
    "mul2": "两位数×两位数", "div": "整除", "divr": "有余数除法", "mixed": "三步混合运算",
    "decimal": "小数加减", "fraction": "分数加减", "equation": "一元一次方程",
}  # fmt: skip


def has_carry(a: int, b: int) -> bool:
    while a or b:
        if a % 10 + b % 10 >= 10:
            return True
        a, b = a // 10, b // 10
    return False


def has_borrow(a: int, b: int) -> bool:
    while b:
        if a % 10 < b % 10:
            return True
        a, b = a // 10, b // 10
    return False


def gen_add(r: random.Random, mx: int, carry: str) -> Problem:
    while True:
        a, b = r.randint(1, mx - 1), r.randint(1, mx - 1)
        if a + b > mx:
            continue
        c = has_carry(a, b)
        if carry == "any" or (carry == "require") == c:
            return f"{a} + {b} =", str(a + b)


def gen_sub(r: random.Random, mx: int, carry: str) -> Problem:
    while True:
        a = r.randint(2, mx)
        b = r.randint(1, a - 1)
        c = has_borrow(a, b)
        if carry == "any" or (carry == "require") == c:
            return f"{a} − {b} =", str(a - b)


def gen_addsub(r: random.Random, mx: int, _: str) -> Problem:
    while True:
        a, b, c = r.randint(1, mx), r.randint(1, mx), r.randint(1, mx)
        op1, op2 = r.choice("+-"), r.choice("+-")
        first = a + b if op1 == "+" else a - b
        result = first + c if op2 == "+" else first - c
        if 0 <= first <= mx and 0 <= result <= mx:
            sym = {"+": "+", "-": "−"}
            return f"{a} {sym[op1]} {b} {sym[op2]} {c} =", str(result)


def gen_table(r: random.Random, _mx: int, _c: str) -> Problem:
    a, b = r.randint(1, 9), r.randint(1, 9)
    return f"{a} × {b} =", str(a * b)


def gen_mul(r: random.Random, mx: int, _: str) -> Problem:
    a, b = r.randint(10, max(11, min(mx, 9999))), r.randint(2, 9)
    return f"{a} × {b} =", str(a * b)


def gen_mul2(r: random.Random, _mx: int, _c: str) -> Problem:
    a, b = r.randint(11, 99), r.randint(11, 99)
    return f"{a} × {b} =", str(a * b)


def gen_div(r: random.Random, mx: int, _: str) -> Problem:
    while True:
        b = r.randint(2, 9)
        q = r.randint(2, max(3, mx // b))
        if b * q <= max(mx, 81):
            return f"{b * q} ÷ {b} =", str(q)


def gen_divr(r: random.Random, mx: int, _: str) -> Problem:
    while True:
        b = r.randint(2, 9)
        q, rem = r.randint(1, max(2, mx // b)), r.randint(1, b - 1)
        a = b * q + rem
        if a <= max(mx, 89):
            return f"{a} ÷ {b} =", f"{q}……{rem}"


def gen_mixed(r: random.Random, mx: int, _: str) -> Problem:
    """Three-step expressions with brackets; evaluated exactly, kept to positive integers."""
    limit = max(mx, 100)
    while True:
        a, b, c, d = (r.randint(2, 50) for _ in range(4))
        pattern = r.choice(("({a} + {b}) × {c} − {d}", "{a} × {b} − {c} × {d}", "({a} − {b}) × {c} + {d}",
                            "{a} + {b} × {c} − {d}", "({a} + {b}) ÷ {c} + {d}", "{a} × ({b} + {c}) ÷ {d}"))  # fmt: skip
        expr = pattern.format(a=a, b=b, c=c, d=d)
        try:
            value = _eval(expr)
        except ZeroDivisionError:
            continue
        groups_positive = all(_eval(g) > 0 for g in re.findall(r"\([^()]*\)", expr))
        if (
            value.denominator == 1
            and 0 < value <= limit
            and groups_positive
            and _divides_cleanly(expr)
        ):
            return expr + " =", str(value.numerator)


def _eval(expr: str) -> Fraction:
    """Exact value of an expression built only from digits, brackets and + − × ÷ above."""
    py = expr.replace("−", "-").replace("×", "*").replace("÷", "/")
    return Fraction(eval(re.sub(r"(\d+)", r"Fraction(\1)", py), {"Fraction": Fraction}))  # noqa: S307


def _divides_cleanly(expr: str) -> bool:
    """Every ÷ in *expr* must produce an integer at that step (primary-school convention)."""
    for m in re.finditer(r"(\([^()]*\)|\d+)\s*÷\s*(\d+)", expr):
        if (_eval(m.group(1)) / int(m.group(2))).denominator != 1:
            return False
    return True


def gen_decimal(r: random.Random, _mx: int, _c: str) -> Problem:
    places = r.choice((1, 2))
    scale = 10**places
    a, b = r.randint(1, 99 * scale) / scale, r.randint(1, 99 * scale) / scale
    if a.is_integer() or b.is_integer():
        return gen_decimal(r, _mx, _c)
    if r.random() < 0.5:
        return f"{a:g} + {b:g} =", f"{round(a + b, places):g}"
    a, b = max(a, b), min(a, b)
    return f"{a:g} − {b:g} =", f"{round(a - b, places):g}"


def frac_text(f: Fraction) -> str:
    if f.denominator == 1:
        return str(f.numerator)
    whole, rest = divmod(f.numerator, f.denominator)
    return (
        f"{f.numerator}/{f.denominator}"
        if whole == 0
        else f"{f.numerator}/{f.denominator}（{whole}又{rest}/{f.denominator}）"
    )


def gen_fraction(r: random.Random, _mx: int, _c: str) -> Problem:
    while True:
        d1, d2 = r.randint(2, 12), r.randint(2, 12)
        if r.random() < 0.4:
            d2 = d1
        a, b = Fraction(r.randint(1, d1 - 1), d1), Fraction(r.randint(1, d2 - 1), d2)
        if r.random() < 0.5:
            return f"{a.numerator}/{a.denominator} + {b.numerator}/{b.denominator} =", frac_text(
                a + b
            )
        if a > b:
            return f"{a.numerator}/{a.denominator} − {b.numerator}/{b.denominator} =", frac_text(
                a - b
            )


def gen_equation(r: random.Random, _mx: int, _c: str) -> Problem:
    x = r.randint(-12, 15)
    a = r.choice([n for n in range(-9, 10) if n != 0])
    b = r.randint(-30, 30)
    left = {1: "x", -1: "-x"}.get(a, f"{a}x")
    if b:
        left += f" {'+' if b > 0 else '-'} {abs(b)}"
    return f"{left} = {a * x + b}".replace("-", "−"), f"x = {x}".replace("-", "−")


GENERATORS: dict[str, Callable[[random.Random, int, str], Problem]] = {
    "add": gen_add, "sub": gen_sub, "addsub": gen_addsub, "table": gen_table, "mul": gen_mul, "mul2": gen_mul2,
    "div": gen_div, "divr": gen_divr, "mixed": gen_mixed, "decimal": gen_decimal, "fraction": gen_fraction,
    "equation": gen_equation,
}  # fmt: skip


def generate(kind: str, n: int, mx: int, carry: str, seed: int | None) -> list[Problem]:
    r = random.Random(seed)
    out: list[Problem] = []
    seen: set[str] = set()
    tries = 0
    while len(out) < n and tries < n * 200:
        tries += 1
        q, a = GENERATORS[kind](r, mx, carry)
        if q in seen:
            continue
        seen.add(q)
        out.append((q, a))
    return out


def write_docx(problems: list[Problem], out: Path, title: str, cols: int) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 python-docx：请运行 `pip install python-docx`", file=sys.stderr)
        sys.exit(2)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Mm(18))

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size, run.font.bold = Pt(18), True

    def grid(items: list[str], size: float) -> None:
        rows = -(-len(items) // cols)
        table = doc.add_table(rows=rows, cols=cols)
        for i, text in enumerate(items):
            row, col = i % rows, i // rows  # fill column by column, like printed drills
            cell = table.cell(row, col)
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(size)
            cell.paragraphs[0].paragraph_format.space_after = Pt(10)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        table._tbl.tblPr.append(borders)

    heading(title)
    info = doc.add_paragraph(
        "姓名：__________　　日期：__________　　用时：______ 分钟　　得分：______"
    )
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    grid([f"{i}. {q}" for i, (q, _) in enumerate(problems, start=1)], 14)
    doc.add_page_break()
    heading(f"{title}（答案）")
    grid([f"{i}. {q} {a}" for i, (q, a) in enumerate(problems, start=1)], 11)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--type", choices=sorted(TYPES), help="题型")
    parser.add_argument("-n", type=int, default=30, help="题数")
    parser.add_argument("--max", type=int, default=100, help="数的范围上限")
    parser.add_argument("--carry", choices=["require", "avoid", "any"], default="any")
    parser.add_argument("--cols", type=int, default=3, help="每行几列")
    parser.add_argument("--title")
    parser.add_argument("--seed", type=int)
    parser.add_argument("-o", "--output")
    parser.add_argument("--list", action="store_true", help="列出题型")
    args = parser.parse_args(argv)
    if args.list or not args.type:
        for key, label in TYPES.items():
            print(f"{key:10s} {label}")
        return 0 if args.list else 1
    if not 1 <= args.n <= 200 or not 1 <= args.cols <= 5 or args.max < 10:
        print("题数 1–200、列数 1–5、--max 至少 10", file=sys.stderr)
        return 1
    problems = generate(args.type, args.n, args.max, args.carry, args.seed)
    if len(problems) < args.n:
        print(f"提示：在当前条件下只能生成 {len(problems)} 道不重复的题，请放宽 --max 或 --carry")
    title = args.title or f"{TYPES[args.type]}练习（{len(problems)} 题）"
    if args.output:
        write_docx(problems, Path(args.output), title, args.cols)
        print(f"已生成 {args.output}：{len(problems)} 道题，答案在第 2 页")
    else:
        for i, (q, a) in enumerate(problems, start=1):
            print(f"{i}. {q} {a}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
