#!/usr/bin/env python3
"""Build a styled .pptx from a JSON deck spec (python-pptx only, no Node/LibreOffice).

Usage:
    python scripts/build_deck.py deck.json -o output.pptx [--theme business-blue]
    python scripts/build_deck.py --list-themes

The spec format is documented in ``references/deck-spec.md``.
Exit codes: 0 ok (warnings may be printed), 1 spec error, 2 missing dependency.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 python-pptx：请运行 `pip install python-pptx`", file=sys.stderr)
    sys.exit(2)

from textfit import fit_size, text_height_in

THEMES: dict[str, dict[str, Any]] = {
    "business-blue": {
        "label": "商务蓝 — 汇报、方案、路演",
        "bg": "FFFFFF",
        "surface": "EFF6FF",
        "text": "1F2937",
        "muted": "6B7280",
        "primary": "1E40AF",
        "secondary": "3B82F6",
        "accent": "F59E0B",
        "cover_bg": "1E3A8A",
        "cover_text": "FFFFFF",
        "cover_muted": "BFDBFE",
        "chart": ["1E40AF", "3B82F6", "F59E0B", "10B981", "8B5CF6", "EF4444"],
    },
    "tech-dark": {
        "label": "科技深色 — 产品发布、技术分享",
        "bg": "0F172A",
        "surface": "1E293B",
        "text": "E2E8F0",
        "muted": "94A3B8",
        "primary": "38BDF8",
        "secondary": "818CF8",
        "accent": "22D3EE",
        "cover_bg": "020617",
        "cover_text": "F8FAFC",
        "cover_muted": "94A3B8",
        "chart": ["38BDF8", "818CF8", "22D3EE", "F472B6", "FACC15", "4ADE80"],
    },
    "fresh-green": {
        "label": "清新绿 — 教育培训、公益、健康",
        "bg": "FFFFFF",
        "surface": "ECFDF5",
        "text": "1F2937",
        "muted": "6B7280",
        "primary": "047857",
        "secondary": "10B981",
        "accent": "F97316",
        "cover_bg": "064E3B",
        "cover_text": "FFFFFF",
        "cover_muted": "A7F3D0",
        "chart": ["047857", "10B981", "F97316", "0EA5E9", "A855F7", "EAB308"],
    },
    "warm-orange": {
        "label": "活力橙 — 营销、活动、品牌",
        "bg": "FFFFFF",
        "surface": "FFF7ED",
        "text": "292524",
        "muted": "78716C",
        "primary": "C2410C",
        "secondary": "FB923C",
        "accent": "0EA5E9",
        "cover_bg": "7C2D12",
        "cover_text": "FFFFFF",
        "cover_muted": "FED7AA",
        "chart": ["C2410C", "FB923C", "0EA5E9", "65A30D", "A21CAF", "FACC15"],
    },
    "minimal-gray": {
        "label": "极简黑白 — 咨询报告、学术答辩",
        "bg": "FFFFFF",
        "surface": "F3F4F6",
        "text": "111827",
        "muted": "6B7280",
        "primary": "111827",
        "secondary": "4B5563",
        "accent": "DC2626",
        "cover_bg": "111827",
        "cover_text": "FFFFFF",
        "cover_muted": "D1D5DB",
        "chart": ["111827", "6B7280", "DC2626", "9CA3AF", "2563EB", "D97706"],
    },
}

DEFAULT_FONTS = {"latin": "Arial", "ea": "Microsoft YaHei"}
MARGIN = 0.6
CONTENT_TOP = 1.5


class SpecError(ValueError):
    """Raised for invalid deck specs; message is shown to the model verbatim."""


@dataclass
class Ctx:
    prs: Any
    theme: dict[str, Any]
    fonts: dict[str, str]
    footer: str
    width: float
    height: float
    warnings: list[str] = field(default_factory=list)
    slide_no: int = 0


# ---------------------------------------------------------------- primitives


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _set_alpha(shape: Any, percent: int) -> None:
    """Make a solid-filled shape partially transparent (percent = opacity)."""
    fill = shape.fill._xPr.find(qn("a:solidFill"))
    if fill is None or len(fill) == 0:
        return
    alpha = fill[0].makeelement(qn("a:alpha"), {"val": str(percent * 1000)})
    fill[0].append(alpha)


def _style_run(ctx: Ctx, run: Any, size: float, color: str, bold: bool = False) -> None:
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = _rgb(color)
    font.name = ctx.fonts["latin"]
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", ctx.fonts["ea"])


def _set_bullet(paragraph: Any, level: int, color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    indent = 0.3 + 0.35 * level
    ppr.set("marL", str(int(Inches(indent))))
    ppr.set("indent", str(int(-Inches(0.3))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)
    bu_clr = ppr.makeelement(qn("a:buClr"), {})
    srgb = bu_clr.makeelement(qn("a:srgbClr"), {"val": color})
    bu_clr.append(srgb)
    ppr.append(bu_clr)
    ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"}))


def _rect(
    ctx: Ctx,
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    color: str,
    shape: Any = MSO_SHAPE.RECTANGLE,
) -> Any:
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.08
    return shp


def _text(
    ctx: Ctx,
    slide: Any,
    box: tuple[float, float, float, float],
    paragraphs: list[Any],
    size: float,
    color: str,
    *,
    bold: bool = False,
    align: Any = PP_ALIGN.LEFT,
    anchor: Any = MSO_ANCHOR.TOP,
    bullets: bool = False,
    bullet_color: str | None = None,
) -> Any:
    """Add a textbox. *paragraphs* items are ``str`` or ``(text, level)``."""
    x, y, w, h = box
    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.1))
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    for i, item in enumerate(paragraphs):
        text, level = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.15
        if i > 0:
            p.space_before = Pt(size * 0.35)
        if bullets:
            _set_bullet(p, level, bullet_color or ctx.theme["primary"])
        run = p.add_run()
        run.text = text
        _style_run(ctx, run, size - 2 * level, color, bold)
    return shp


def _fit(
    ctx: Ctx,
    paragraphs: list[str],
    width: float,
    height: float,
    max_pt: float,
    min_pt: float,
    where: str,
    indent: float = 0.0,
) -> float:
    usable_w = width - 0.2 - indent
    usable_h = height - 0.1
    size = fit_size(paragraphs, usable_w, usable_h, max_pt, min_pt)
    if text_height_in(paragraphs, size, usable_w) > usable_h:
        ctx.warnings.append(
            f"第 {ctx.slide_no} 页 {where}：内容过多，最小字号 {min_pt:g}pt 仍可能溢出，"
            "请精简文字或拆成两页"
        )
    return size


def _plain(items: list[Any]) -> list[str]:
    return [i if isinstance(i, str) else i[0] for i in items]


def _bullet_items(raw: Any, where: str) -> list[tuple[str, int]]:
    if not isinstance(raw, list) or not raw:
        raise SpecError(f"{where} 必须是非空数组")
    items: list[tuple[str, int]] = []
    for entry in raw:
        if isinstance(entry, str):
            items.append((entry, 0))
        elif isinstance(entry, dict) and isinstance(entry.get("text"), str):
            items.append((entry["text"], min(int(entry.get("level", 0)), 2)))
        else:
            raise SpecError(f"{where} 的每一项必须是字符串或 {{text, level}} 对象")
    return items


def _req(data: dict[str, Any], key: str, where: str) -> Any:
    value = data.get(key)
    if value in (None, "", []):
        raise SpecError(f"{where} 缺少必填字段 `{key}`")
    return value


def _background(ctx: Ctx, slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _title_bar(ctx: Ctx, slide: Any, title: str) -> float:
    t = ctx.theme
    _rect(ctx, slide, MARGIN, 0.55, 0.09, 0.62, t["primary"])
    w = ctx.width - 2 * MARGIN - 0.25
    size = _fit(ctx, [title], w, 0.85, 30, 20, "标题")
    _text(
        ctx,
        slide,
        (MARGIN + 0.2, 0.43, w, 0.85),
        [title],
        size,
        t["text"],
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    return CONTENT_TOP


def _footer(ctx: Ctx, slide: Any) -> None:
    t = ctx.theme
    y = ctx.height - 0.5
    if ctx.footer:
        _text(ctx, slide, (MARGIN, y, ctx.width / 2, 0.32), [ctx.footer], 10, t["muted"])
    _text(
        ctx,
        slide,
        (ctx.width - MARGIN - 1.0, y, 1.0, 0.32),
        [str(ctx.slide_no)],
        10,
        t["muted"],
        align=PP_ALIGN.RIGHT,
    )


def _takeaway(ctx: Ctx, slide: Any, text: str) -> float:
    """Bottom callout bar; returns its top y so content can stop above it."""
    t = ctx.theme
    h = 0.75
    y = ctx.height - 0.65 - h
    w = ctx.width - 2 * MARGIN
    _rect(ctx, slide, MARGIN, y, w, h, t["surface"])
    _rect(ctx, slide, MARGIN, y, 0.08, h, t["accent"])
    size = _fit(ctx, [text], w - 0.3, h, 18, 13, "takeaway")
    _text(
        ctx,
        slide,
        (MARGIN + 0.25, y, w - 0.3, h),
        [text],
        size,
        t["text"],
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    return y - 0.2


def _content_bottom(ctx: Ctx, slide: Any, data: dict[str, Any]) -> float:
    if data.get("takeaway"):
        return _takeaway(ctx, slide, str(data["takeaway"]))
    return ctx.height - 0.75


# ------------------------------------------------------------------- layouts


def layout_cover(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    _background(ctx, slide, t["cover_bg"])
    circle = _rect(ctx, slide, ctx.width - 4.2, -1.6, 5.6, 5.6, t["secondary"], MSO_SHAPE.OVAL)
    _set_alpha(circle, 25)
    circle2 = _rect(
        ctx, slide, ctx.width - 2.4, ctx.height - 2.2, 3.4, 3.4, t["accent"], MSO_SHAPE.OVAL
    )
    _set_alpha(circle2, 20)
    top = ctx.height * 0.3
    _rect(ctx, slide, MARGIN, top + 0.1, 0.12, 2.0, t["accent"])
    w = ctx.width - 2 * MARGIN - 3.0
    title = str(_req(d, "title", f"第 {ctx.slide_no} 页 cover"))
    size = _fit(ctx, [title], w, 1.5, 44, 28, "封面标题")
    _text(
        ctx,
        slide,
        (MARGIN + 0.4, top, w, 1.5),
        [title],
        size,
        t["cover_text"],
        bold=True,
        anchor=MSO_ANCHOR.BOTTOM,
    )
    if d.get("subtitle"):
        _text(
            ctx,
            slide,
            (MARGIN + 0.4, top + 1.6, w, 0.8),
            [str(d["subtitle"])],
            20,
            t["cover_muted"],
        )
    if d.get("meta"):
        _text(
            ctx,
            slide,
            (MARGIN + 0.4, ctx.height - 1.3, w, 0.5),
            [str(d["meta"])],
            14,
            t["cover_muted"],
        )


def layout_section(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    _background(ctx, slide, t["cover_bg"])
    top = ctx.height * 0.26
    w = ctx.width - 2 * MARGIN
    if d.get("index"):
        _text(ctx, slide, (MARGIN, top, 4, 1.4), [str(d["index"])], 72, t["accent"], bold=True)
    title = str(_req(d, "title", f"第 {ctx.slide_no} 页 section"))
    _text(
        ctx,
        slide,
        (MARGIN, top + 1.5, w, 1.0),
        [title],
        _fit(ctx, [title], w, 1.0, 40, 26, "章节标题"),
        t["cover_text"],
        bold=True,
    )
    _rect(ctx, slide, MARGIN + 0.1, top + 2.6, 1.2, 0.06, t["accent"])
    if d.get("subtitle"):
        _text(ctx, slide, (MARGIN, top + 2.8, w, 0.8), [str(d["subtitle"])], 18, t["cover_muted"])


def layout_agenda(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    top = _title_bar(ctx, slide, str(d.get("title") or "目录"))
    items = [str(i) for i in _req(d, "items", f"第 {ctx.slide_no} 页 agenda")]
    cols = 2 if len(items) > 5 else 1
    per_col = -(-len(items) // cols)
    col_w = (ctx.width - 2 * MARGIN - 0.6 * (cols - 1)) / cols
    avail = ctx.height - 0.9 - top
    row_h = min(1.1, avail / per_col)
    top += (avail - row_h * per_col) / 2
    for i, item in enumerate(items):
        col, row = divmod(i, per_col)
        x = MARGIN + col * (col_w + 0.6)
        y = top + 0.1 + row * row_h
        dot = min(0.62, row_h - 0.15)
        _rect(ctx, slide, x, y + (row_h - dot) / 2 - 0.05, dot, dot, t["primary"], MSO_SHAPE.OVAL)
        _text(
            ctx,
            slide,
            (x, y + (row_h - dot) / 2 - 0.05, dot, dot),
            [f"{i + 1:02d}"],
            14,
            "FFFFFF",
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        tw = col_w - dot - 0.25
        _text(
            ctx,
            slide,
            (x + dot + 0.2, y, tw, row_h - 0.1),
            [item],
            _fit(ctx, [item], tw, row_h - 0.1, 22, 14, f"目录第 {i + 1} 项"),
            t["text"],
            anchor=MSO_ANCHOR.MIDDLE,
        )
    _footer(ctx, slide)


def layout_bullets(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 bullets"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    items = _bullet_items(d.get("bullets"), f"{where}.bullets")
    bottom = _content_bottom(ctx, slide, d)
    w = ctx.width - 2 * MARGIN
    size = _fit(ctx, _plain(items), w, bottom - top, 28, 14, "要点", indent=0.4)
    _text(ctx, slide, (MARGIN, top, w, bottom - top), items, size, t["text"], bullets=True)
    _footer(ctx, slide)


def _panel(
    ctx: Ctx,
    slide: Any,
    box: tuple[float, float, float, float],
    panel: dict[str, Any],
    where: str,
    color: str,
) -> None:
    t = ctx.theme
    x, y, w, h = box
    _rect(ctx, slide, x, y, w, h, t["surface"], MSO_SHAPE.ROUNDED_RECTANGLE)
    heading = str(_req(panel, "heading", where))
    _text(ctx, slide, (x + 0.25, y + 0.2, w - 0.5, 0.6), [heading], 20, color, bold=True)
    items = _bullet_items(panel.get("bullets"), f"{where}.bullets")
    bh = h - 1.0
    size = _fit(ctx, _plain(items), w - 0.5, bh, 24, 12, f"{where} 要点", indent=0.4)
    _text(
        ctx,
        slide,
        (x + 0.25, y + 0.85, w - 0.5, bh),
        items,
        size,
        t["text"],
        bullets=True,
        bullet_color=color,
    )


def layout_two_column(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 two-column"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    bottom = _content_bottom(ctx, slide, d)
    gap = 0.4
    w = (ctx.width - 2 * MARGIN - gap) / 2
    left = _req(d, "left", where)
    right = _req(d, "right", where)
    _panel(ctx, slide, (MARGIN, top, w, bottom - top), left, f"{where}.left", t["primary"])
    _panel(
        ctx, slide, (MARGIN + w + gap, top, w, bottom - top), right, f"{where}.right", t["accent"]
    )
    _footer(ctx, slide)


def layout_cards(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 cards"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    cards = _req(d, "cards", where)
    if not 2 <= len(cards) <= 4:
        raise SpecError(f"{where}.cards 需要 2–4 张卡片，当前 {len(cards)} 张")
    bottom = _content_bottom(ctx, slide, d)
    gap = 0.3
    n = len(cards)
    w = (ctx.width - 2 * MARGIN - gap * (n - 1)) / n
    h = bottom - top
    for i, card in enumerate(cards):
        x = MARGIN + i * (w + gap)
        color = t["chart"][i % len(t["chart"])]
        _rect(ctx, slide, x, top, w, h, t["surface"], MSO_SHAPE.ROUNDED_RECTANGLE)
        _rect(ctx, slide, x + 0.25, top + 0.25, 0.6, 0.08, color)
        y = top + 0.45
        if card.get("icon"):
            _text(
                ctx, slide, (x + 0.15, y, w - 0.3, 0.7), [str(card["icon"])], 28, color, bold=True
            )
            y += 0.75
        title = str(_req(card, "title", f"{where}.cards[{i}]"))
        ts = _fit(ctx, [title], w - 0.4, 0.9, 20, 14, f"卡片 {i + 1} 标题")
        _text(ctx, slide, (x + 0.15, y, w - 0.3, 0.9), [title], ts, t["text"], bold=True)
        y += 0.95
        body = str(card.get("body", ""))
        if body:
            bh = top + h - y - 0.2
            _text(
                ctx,
                slide,
                (x + 0.15, y, w - 0.3, bh),
                [body],
                _fit(ctx, [body], w - 0.3, bh, 16, 11, f"卡片 {i + 1} 正文"),
                t["muted"],
            )
    _footer(ctx, slide)


def layout_kpi(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 kpi"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    metrics = _req(d, "metrics", where)
    if not 1 <= len(metrics) <= 4:
        raise SpecError(f"{where}.metrics 需要 1–4 个指标，当前 {len(metrics)} 个")
    bottom = _content_bottom(ctx, slide, d)
    n = len(metrics)
    gap = 0.3
    w = (ctx.width - 2 * MARGIN - gap * (n - 1)) / n
    h = min(3.6, bottom - top)
    y0 = top + (bottom - top - h) / 2
    for i, m in enumerate(metrics):
        x = MARGIN + i * (w + gap)
        _rect(ctx, slide, x, y0, w, h, t["surface"], MSO_SHAPE.ROUNDED_RECTANGLE)
        value = str(_req(m, "value", f"{where}.metrics[{i}]"))
        vs = _fit(ctx, [value], w - 0.3, 1.3, 54, 28, f"指标 {i + 1} 数值")
        _text(
            ctx,
            slide,
            (x + 0.15, y0 + 0.35, w - 0.3, 1.3),
            [value],
            vs,
            t["primary"],
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        label = str(m.get("label", ""))
        _text(
            ctx,
            slide,
            (x + 0.15, y0 + 1.75, w - 0.3, 0.7),
            [label],
            _fit(ctx, [label], w - 0.3, 0.7, 18, 12, f"指标 {i + 1} 标签"),
            t["text"],
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        if m.get("note"):
            note = str(m["note"])
            nh = h - 2.55
            _text(
                ctx,
                slide,
                (x + 0.15, y0 + 2.45, w - 0.3, nh),
                [note],
                _fit(ctx, [note], w - 0.3, nh, 13, 10, f"指标 {i + 1} 备注"),
                t["muted"],
                align=PP_ALIGN.CENTER,
            )
    _footer(ctx, slide)


_CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked-column": XL_CHART_TYPE.COLUMN_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def layout_chart(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 chart"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    spec = _req(d, "chart", where)
    kind = spec.get("type", "column")
    if kind not in _CHART_TYPES:
        raise SpecError(f"{where}.chart.type 只支持 {sorted(_CHART_TYPES)}，当前 {kind!r}")
    categories = [str(c) for c in _req(spec, "categories", f"{where}.chart")]
    series = _req(spec, "series", f"{where}.chart")
    data = CategoryChartData()
    data.categories = categories
    for s in series:
        values = s.get("values") or []
        if len(values) != len(categories):
            raise SpecError(
                f"{where}.chart 序列 {s.get('name')!r} 有 {len(values)} 个值，但 categories 有 {len(categories)} 个"
            )
        data.add_series(str(s.get("name", "")), [float(v) for v in values])
    bottom = _content_bottom(ctx, slide, d)
    if d.get("caption"):
        bottom -= 0.5
        _text(
            ctx,
            slide,
            (MARGIN, bottom + 0.05, ctx.width - 2 * MARGIN, 0.45),
            [str(d["caption"])],
            13,
            t["muted"],
        )
    gf = slide.shapes.add_chart(
        _CHART_TYPES[kind],
        Inches(MARGIN),
        Inches(top),
        Inches(ctx.width - 2 * MARGIN),
        Inches(bottom - top),
        data,
    )
    chart = gf.chart
    chart.font.size = Pt(12)
    chart.font.color.rgb = _rgb(t["text"])
    chart.font.name = ctx.fonts["latin"]
    pie_like = kind in ("pie", "doughnut")
    chart.has_legend = pie_like or len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    palette = t["chart"]
    if pie_like:
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.show_percentage = True
        labels.show_value = False
        labels.number_format = "0%"
        labels.number_format_is_linked = False
    else:
        for i, s in enumerate(plot.series):
            color = _rgb(palette[i % len(palette)])
            if kind == "line":
                s.format.line.color.rgb = color
                s.format.line.width = Pt(2.5)
                s.smooth = False
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = color
        if len(series) == 1:
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(11)
            if spec.get("unit"):
                plot.data_labels.number_format = f'0.##"{spec["unit"]}"'
                plot.data_labels.number_format_is_linked = False
        if hasattr(plot, "gap_width"):
            plot.gap_width = 80
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _rgb(
            t["surface"] if t["bg"] != "FFFFFF" else "E5E7EB"
        )
        va.format.line.fill.background()
    _footer(ctx, slide)


def layout_table(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 table"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    columns = [str(c) for c in _req(d, "columns", where)]
    rows = _req(d, "rows", where)
    for i, r in enumerate(rows):
        if len(r) != len(columns):
            raise SpecError(f"{where}.rows[{i}] 有 {len(r)} 列，表头有 {len(columns)} 列")
    if len(rows) > 10:
        ctx.warnings.append(
            f"第 {ctx.slide_no} 页表格有 {len(rows)} 行，超过 10 行建议拆页或只保留关键行"
        )
    bottom = _content_bottom(ctx, slide, d)
    if d.get("caption"):
        bottom -= 0.5
        _text(
            ctx,
            slide,
            (MARGIN, bottom + 0.05, ctx.width - 2 * MARGIN, 0.45),
            [str(d["caption"])],
            13,
            t["muted"],
        )
    n_rows = len(rows) + 1
    row_h = min(0.6, (bottom - top) / n_rows)
    size = 16 if row_h >= 0.5 else 13 if row_h >= 0.38 else 11
    w = ctx.width - 2 * MARGIN
    shape = slide.shapes.add_table(
        n_rows, len(columns), Inches(MARGIN), Inches(top), Inches(w), Inches(row_h * n_rows)
    )
    table = shape.table
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for r in range(n_rows):
        table.rows[r].height = Emu(int(Inches(row_h)))
        for c in range(len(columns)):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            text = columns[c] if r == 0 else str(rows[r - 1][c])
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = _rgb(t["primary"])
            else:
                cell.fill.fore_color.rgb = _rgb(t["surface"] if r % 2 == 0 else t["bg"])
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            run = p.add_run()
            run.text = text
            header_color = "FFFFFF" if t["bg"] == "FFFFFF" else t["bg"]
            _style_run(ctx, run, size, header_color if r == 0 else t["text"], bold=r == 0)
    _footer(ctx, slide)


def layout_timeline(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 timeline"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    steps = _req(d, "steps", where)
    if not 2 <= len(steps) <= 6:
        raise SpecError(f"{where}.steps 需要 2–6 个节点，当前 {len(steps)} 个")
    bottom = _content_bottom(ctx, slide, d)
    n = len(steps)
    w = (ctx.width - 2 * MARGIN) / n
    line_y = top + max(1.25, (bottom - top) / 2 - 0.8)
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(MARGIN),
        Inches(line_y),
        Inches(ctx.width - MARGIN),
        Inches(line_y),
    )
    line.line.color.rgb = _rgb(t["muted"])
    line.line.width = Pt(2)
    for i, step in enumerate(steps):
        cx = MARGIN + w * i + w / 2
        color = t["chart"][i % len(t["chart"])]
        if step.get("label"):
            _text(
                ctx,
                slide,
                (cx - w / 2, line_y - 1.0, w, 0.6),
                [str(step["label"])],
                20,
                color,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
        _rect(ctx, slide, cx - 0.17, line_y - 0.17, 0.34, 0.34, color, MSO_SHAPE.OVAL)
        title = str(_req(step, "title", f"{where}.steps[{i}]"))
        _text(
            ctx,
            slide,
            (cx - w / 2 + 0.05, line_y + 0.35, w - 0.1, 0.8),
            [title],
            _fit(ctx, [title], w - 0.1, 0.8, 18, 12, f"节点 {i + 1} 标题"),
            t["text"],
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        if step.get("body"):
            body = str(step["body"])
            bh = bottom - (line_y + 1.2)
            _text(
                ctx,
                slide,
                (cx - w / 2 + 0.05, line_y + 1.2, w - 0.1, bh),
                [body],
                _fit(ctx, [body], w - 0.1, bh, 14, 10, f"节点 {i + 1} 说明"),
                t["muted"],
                align=PP_ALIGN.CENTER,
            )
    _footer(ctx, slide)


def layout_image(ctx: Ctx, slide: Any, d: dict[str, Any], base_dir: Path) -> None:
    t = ctx.theme
    where = f"第 {ctx.slide_no} 页 image"
    top = _title_bar(ctx, slide, str(_req(d, "title", where)))
    path = Path(str(_req(d, "image", where)))
    if not path.is_absolute():
        path = base_dir / path
    if not path.is_file():
        raise SpecError(f"{where}.image 文件不存在：{path}")
    bottom = _content_bottom(ctx, slide, d)
    if d.get("caption"):
        bottom -= 0.5
        _text(
            ctx,
            slide,
            (MARGIN, bottom + 0.05, ctx.width - 2 * MARGIN, 0.45),
            [str(d["caption"])],
            13,
            t["muted"],
            align=PP_ALIGN.CENTER,
        )
    box_w, box_h = ctx.width - 2 * MARGIN, bottom - top
    try:
        from PIL import Image

        with Image.open(path) as img:
            iw, ih = img.size
    except Exception:  # noqa: BLE001 - fall back to box aspect
        iw, ih = box_w, box_h
    scale = min(box_w / iw, box_h / ih)
    w, h = iw * scale, ih * scale
    slide.shapes.add_picture(
        str(path),
        Inches(MARGIN + (box_w - w) / 2),
        Inches(top + (box_h - h) / 2),
        Inches(w),
        Inches(h),
    )
    _footer(ctx, slide)


def layout_quote(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    _background(ctx, slide, t["surface"])
    _text(ctx, slide, (MARGIN, 0.3, 2.2, 2.6), ["“"], 140, t["primary"], bold=True)
    quote = str(_req(d, "quote", f"第 {ctx.slide_no} 页 quote"))
    w = ctx.width - 2 * MARGIN - 1.0
    h = ctx.height * 0.45
    _text(
        ctx,
        slide,
        (MARGIN + 0.5, ctx.height * 0.25, w, h),
        [quote],
        _fit(ctx, [quote], w, h, 34, 20, "引言"),
        t["text"],
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    if d.get("author"):
        _text(
            ctx,
            slide,
            (MARGIN + 0.5, ctx.height * 0.25 + h + 0.1, w, 0.5),
            [f"—— {d['author']}"],
            18,
            t["muted"],
            align=PP_ALIGN.RIGHT,
        )
    _footer(ctx, slide)


def layout_closing(ctx: Ctx, slide: Any, d: dict[str, Any]) -> None:
    t = ctx.theme
    _background(ctx, slide, t["cover_bg"])
    circle = _rect(ctx, slide, -1.5, ctx.height - 3.5, 5.0, 5.0, t["secondary"], MSO_SHAPE.OVAL)
    _set_alpha(circle, 25)
    w = ctx.width - 2 * MARGIN
    title = str(d.get("title") or "谢谢")
    _text(
        ctx,
        slide,
        (MARGIN, ctx.height * 0.3, w, 1.4),
        [title],
        _fit(ctx, [title], w, 1.4, 48, 28, "结束页标题"),
        t["cover_text"],
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    _rect(ctx, slide, ctx.width / 2 - 0.6, ctx.height * 0.3 + 1.55, 1.2, 0.06, t["accent"])
    if d.get("subtitle"):
        _text(
            ctx,
            slide,
            (MARGIN, ctx.height * 0.3 + 1.75, w, 0.7),
            [str(d["subtitle"])],
            20,
            t["cover_muted"],
            align=PP_ALIGN.CENTER,
        )
    if d.get("contact"):
        _text(
            ctx,
            slide,
            (MARGIN, ctx.height - 1.3, w, 0.5),
            [str(d["contact"])],
            14,
            t["cover_muted"],
            align=PP_ALIGN.CENTER,
        )


LAYOUTS = {
    "cover": layout_cover,
    "section": layout_section,
    "agenda": layout_agenda,
    "bullets": layout_bullets,
    "two-column": layout_two_column,
    "cards": layout_cards,
    "kpi": layout_kpi,
    "chart": layout_chart,
    "table": layout_table,
    "timeline": layout_timeline,
    "image": layout_image,
    "quote": layout_quote,
    "closing": layout_closing,
}


# ---------------------------------------------------------------------- main


def build(
    spec: dict[str, Any], output: Path, base_dir: Path, theme_name: str | None = None
) -> list[str]:
    """Render *spec* to *output*; returns non-fatal warnings."""
    name = theme_name or spec.get("theme") or "business-blue"
    if name not in THEMES:
        raise SpecError(f"未知主题 {name!r}，可选：{', '.join(THEMES)}")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise SpecError("spec.slides 必须是非空数组")
    prs = Presentation()
    size = spec.get("size", "16:9")
    if size not in ("16:9", "4:3"):
        raise SpecError("spec.size 只支持 '16:9' 或 '4:3'")
    width = 13.333 if size == "16:9" else 10.0
    prs.slide_width = Inches(width)
    prs.slide_height = Inches(7.5)
    if spec.get("title"):
        prs.core_properties.title = str(spec["title"])
    fonts = {**DEFAULT_FONTS, **(spec.get("font") or {})}
    overrides = spec.get("colors") or {}
    unknown = set(overrides) - set(THEMES[name])
    if unknown:
        raise SpecError(
            f"spec.colors 含未知键 {sorted(unknown)}，可用：{', '.join(k for k in THEMES[name] if k != 'label')}"
        )
    theme = dict(THEMES[name])
    for key, value in overrides.items():
        if key == "chart":
            theme[key] = [str(c).lstrip("#") for c in value]
        else:
            theme[key] = str(value).lstrip("#")
    ctx = Ctx(prs, theme, fonts, str(spec.get("footer", "")), width, 7.5)
    blank = prs.slide_layouts[6]
    for i, data in enumerate(slides, start=1):
        if not isinstance(data, dict):
            raise SpecError(f"第 {i} 页必须是对象")
        layout = data.get("layout")
        if layout not in LAYOUTS:
            raise SpecError(f"第 {i} 页 layout {layout!r} 不支持，可选：{', '.join(LAYOUTS)}")
        ctx.slide_no = i
        slide = prs.slides.add_slide(blank)
        _background(ctx, slide, ctx.theme["bg"])
        if layout == "image":
            layout_image(ctx, slide, data, base_dir)
        else:
            LAYOUTS[layout](ctx, slide, data)
        if data.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(data["notes"])
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return ctx.warnings


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("spec", nargs="?", help="deck spec JSON path")
    parser.add_argument("-o", "--output", help="output .pptx path")
    parser.add_argument("--theme", help="override spec.theme")
    parser.add_argument("--list-themes", action="store_true", help="list built-in themes")
    args = parser.parse_args(argv)
    if args.list_themes:
        for key, theme in THEMES.items():
            print(f"{key:14s} {theme['label']}")
        return 0
    if not args.spec or not args.output:
        parser.error("需要 spec 路径和 -o 输出路径")
    spec_path = Path(args.spec)
    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        print(f"无法读取 spec：{exc}", file=sys.stderr)
        return 1
    output = Path(args.output)
    try:
        warnings = build(spec, output, spec_path.resolve().parent, args.theme)
    except SpecError as exc:
        print(f"spec 错误：{exc}", file=sys.stderr)
        return 1
    print(f"已生成 {output}（{len(spec['slides'])} 页）")
    for w in warnings:
        print(f"警告：{w}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
