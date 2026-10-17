#!/usr/bin/env python3
"""Read any .pptx: print a per-slide outline and flag likely layout problems.

Usage:
    python scripts/inspect_deck.py deck.pptx            # outline + issues (Markdown)
    python scripts/inspect_deck.py deck.pptx --json     # machine-readable
    python scripts/inspect_deck.py deck.pptx --issues-only

Checks (estimates, not a renderer): text likely overflowing its box, shapes
outside the slide, empty text boxes, slides with too much text.
Exit code: 0 no issues, 3 issues found, 1 error.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 python-pptx：请运行 `pip install python-pptx`", file=sys.stderr)
    sys.exit(2)

from textfit import text_height_in

EMU_PER_IN = 914400
MAX_CHARS_PER_SLIDE = 300
DEFAULT_PT = 18.0


def _inches(value: int | None) -> float:
    return (value or 0) / EMU_PER_IN


def _shape_font_pt(shape: Any) -> float:
    sizes = [
        run.font.size.pt
        for p in shape.text_frame.paragraphs
        for run in p.runs
        if run.font.size is not None
    ]
    return max(sizes) if sizes else DEFAULT_PT


def _iter_shapes(shapes: Any) -> Any:
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def inspect(path: Path) -> dict[str, Any]:
    prs = Presentation(str(path))
    sw, sh = _inches(prs.slide_width), _inches(prs.slide_height)
    slides: list[dict[str, Any]] = []
    issues: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        candidates: list[tuple[bool, float, str]] = []
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        for shape in _iter_shapes(slide.shapes):
            x, y = _inches(shape.left), _inches(shape.top)
            w, h = _inches(shape.width), _inches(shape.height)
            name = shape.name
            outside = x < -0.01 or y < -0.01 or x + w > sw + 0.01 or y + h > sh + 0.01
            # Decorative shapes may bleed on purpose; only flag ones with text.
            if outside and shape.has_text_frame and shape.text_frame.text.strip():
                issues.append(f"第 {idx} 页「{name}」含文字但超出页面边界")
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text_frame.text.strip() for c in row.cells))
                continue
            if not shape.has_text_frame:
                continue
            paras = [p.text for p in shape.text_frame.paragraphs]
            body = "\n".join(paras).strip()
            if not body:
                if shape.shape_type == 17:  # TEXT_BOX
                    issues.append(f"第 {idx} 页「{name}」是空文本框")
                continue
            texts.append(body)
            tf = shape.text_frame
            inset_w = _inches(tf.margin_left) + _inches(tf.margin_right)
            inset_h = _inches(tf.margin_top) + _inches(tf.margin_bottom)
            pt = _shape_font_pt(shape)
            # Title guess: prefer large text near the top, else the largest text.
            first_line = body.split("\n", 1)[0]
            if any(ch.isalpha() for ch in first_line):  # skip "01", page numbers, quote marks
                candidates.append((y < sh * 0.2, pt, first_line))
            need = text_height_in([p for p in paras if p.strip()], pt, w - inset_w)
            if w > 0 and h > 0 and need > (h - inset_h) * 1.15:
                issues.append(
                    f"第 {idx} 页「{name}」文字可能溢出：约需 {need:.1f}in 高，框高 {h:.1f}in"
                    f"（{pt:g}pt）。精简文字、缩小字号或拆页"
                )
        if not title and candidates:
            title = max(candidates)[2]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        chars = sum(len(t) for t in texts)
        if chars > MAX_CHARS_PER_SLIDE:
            issues.append(
                f"第 {idx} 页文字约 {chars} 字，超过 {MAX_CHARS_PER_SLIDE} 字，建议拆页或转成图表"
            )
        slides.append({"index": idx, "title": title, "texts": texts, "notes": notes})
    return {
        "file": str(path),
        "size_in": [round(sw, 3), round(sh, 3)],
        "slide_count": len(slides),
        "slides": slides,
        "issues": issues,
    }


def to_markdown(report: dict[str, Any], issues_only: bool) -> str:
    lines: list[str] = []
    if not issues_only:
        w, h = report["size_in"]
        lines.append(f"# {Path(report['file']).name}（{report['slide_count']} 页，{w}×{h} in）\n")
        for s in report["slides"]:
            lines.append(f"## {s['index']}. {s['title'] or '（无标题）'}")
            for t in s["texts"]:
                if t != s["title"]:
                    lines.extend(f"- {line}" for line in t.splitlines() if line.strip())
            if s["notes"]:
                lines.append(f"> 备注：{s['notes']}")
            lines.append("")
    lines.append("## 检查结果")
    if report["issues"]:
        lines.extend(f"- {i}" for i in report["issues"])
    else:
        lines.append("- 未发现溢出、越界或空文本框问题（估算结果，仍建议视觉预览）")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("pptx")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--issues-only", action="store_true")
    args = parser.parse_args(argv)
    path = Path(args.pptx)
    try:
        report = inspect(path)
    except Exception as exc:  # noqa: BLE001 - report any parse error to the caller
        print(f"无法读取 {path}：{exc}", file=sys.stderr)
        return 1
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(to_markdown(report, args.issues_only))
    return 3 if report["issues"] else 0


if __name__ == "__main__":
    sys.exit(main())
