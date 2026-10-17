#!/usr/bin/env python3
"""Small, safe edits on an existing .pptx while keeping its original styling.

Usage (always write to a new file unless --in-place is given):
    python scripts/edit_deck.py in.pptx -o out.pptx replace --find 旧 --with 新 [--slide 3]
    python scripts/edit_deck.py in.pptx -o out.pptx delete --slides 3,5
    python scripts/edit_deck.py in.pptx -o out.pptx move --slide 5 --to 2
    python scripts/edit_deck.py in.pptx -o out.pptx duplicate --slide 3
    python scripts/edit_deck.py in.pptx -o out.pptx notes --slide 2 --text "讲稿"

Slide numbers are 1-based, as shown by inspect_deck.py.
"""

from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 python-pptx：请运行 `pip install python-pptx`", file=sys.stderr)
    sys.exit(2)


class EditError(ValueError):
    pass


def _slide_ids(prs: Any) -> Any:
    return prs.slides._sldIdLst


def _check_index(prs: Any, number: int) -> int:
    count = len(prs.slides)
    if not 1 <= number <= count:
        raise EditError(f"页码 {number} 超出范围（共 {count} 页）")
    return number - 1


def _paragraphs(slide: Any) -> Any:
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield from shape.text_frame.paragraphs
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs


def replace_text(prs: Any, find: str, new: str, slide_no: int | None) -> int:
    targets = [prs.slides[_check_index(prs, slide_no)]] if slide_no else list(prs.slides)
    count = 0
    for slide in targets:
        for p in _paragraphs(slide):
            if find not in p.text:
                continue
            hit = False
            for run in p.runs:
                if find in run.text:
                    count += run.text.count(find)
                    run.text = run.text.replace(find, new)
                    hit = True
            if not hit and p.runs:
                # Match spans several runs: collapse into the first run's formatting.
                count += p.text.count(find)
                merged = p.text.replace(find, new)
                p.runs[0].text = merged
                for run in p.runs[1:]:
                    run.text = ""
    return count


def delete_slides(prs: Any, numbers: list[int]) -> None:
    ids = _slide_ids(prs)
    for idx in sorted({_check_index(prs, n) for n in numbers}, reverse=True):
        sld = ids[idx]
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def move_slide(prs: Any, number: int, to: int) -> None:
    ids = _slide_ids(prs)
    src = _check_index(prs, number)
    dst = _check_index(prs, to)
    sld = ids[src]
    ids.remove(sld)
    ids.insert(dst, sld)


def duplicate_slide(prs: Any, number: int) -> int:
    """Copy slide *number* right after itself; returns the new 1-based number."""
    src_idx = _check_index(prs, number)
    src = prs.slides[src_idx]
    new = prs.slides.add_slide(src.slide_layout)
    for shape in list(new.shapes):
        shape._element.getparent().remove(shape._element)
    rid_map: dict[str, str] = {}
    for rel in src.part.rels.values():
        if "notesSlide" in rel.reltype:
            continue
        if rel.is_external:
            rid_map[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        clone = copy.deepcopy(el)
        for node in clone.iter():
            for attr, value in list(node.attrib.items()):
                if attr.endswith(("}embed", "}link", "}id")) and value in rid_map:
                    node.set(attr, rid_map[value])
        new.shapes._spTree.append(clone)
    bg = src._element.cSld.bg
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    if src.has_notes_slide:
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    move_slide(prs, len(prs.slides), number + 1)
    return number + 1


def set_notes(prs: Any, number: int, text: str) -> None:
    slide = prs.slides[_check_index(prs, number)]
    slide.notes_slide.notes_text_frame.text = text


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("pptx")
    out = parser.add_mutually_exclusive_group(required=True)
    out.add_argument("-o", "--output")
    out.add_argument("--in-place", action="store_true")
    sub = parser.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("replace")
    r.add_argument("--find", required=True)
    r.add_argument("--with", dest="new", required=True)
    r.add_argument("--slide", type=int)
    d = sub.add_parser("delete")
    d.add_argument("--slides", required=True, help="comma-separated, e.g. 3,5")
    m = sub.add_parser("move")
    m.add_argument("--slide", type=int, required=True)
    m.add_argument("--to", type=int, required=True)
    c = sub.add_parser("duplicate")
    c.add_argument("--slide", type=int, required=True)
    n = sub.add_parser("notes")
    n.add_argument("--slide", type=int, required=True)
    n.add_argument("--text", required=True)
    args = parser.parse_args(argv)

    src = Path(args.pptx)
    if not src.is_file():
        print(f"文件不存在：{src}", file=sys.stderr)
        return 1
    prs = Presentation(str(src))
    try:
        if args.cmd == "replace":
            hits = replace_text(prs, args.find, args.new, args.slide)
            if hits == 0:
                raise EditError(f"未找到「{args.find}」，未做修改")
            message = f"替换 {hits} 处"
        elif args.cmd == "delete":
            numbers = [int(x) for x in args.slides.split(",") if x.strip()]
            delete_slides(prs, numbers)
            message = f"删除第 {', '.join(map(str, numbers))} 页"
        elif args.cmd == "move":
            move_slide(prs, args.slide, args.to)
            message = f"第 {args.slide} 页移到第 {args.to} 页"
        elif args.cmd == "duplicate":
            new_no = duplicate_slide(prs, args.slide)
            message = f"复制第 {args.slide} 页为第 {new_no} 页"
        else:
            set_notes(prs, args.slide, args.text)
            message = f"更新第 {args.slide} 页备注"
    except (EditError, ValueError) as exc:
        print(f"编辑失败：{exc}", file=sys.stderr)
        return 1
    target = src if args.in_place else Path(args.output)
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    print(f"{message}，已保存 {target}（共 {len(prs.slides)} 页）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
