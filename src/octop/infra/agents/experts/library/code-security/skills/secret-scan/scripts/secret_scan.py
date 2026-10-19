#!/usr/bin/env python3
"""代码安全扫描：找出泄露的密钥 / Token / 密码、不该提交的敏感文件和常见危险写法。只读，结果中的密钥会打码。

Usage:
    python scripts/secret_scan.py 项目目录 [--json] [-o 扫描报告.xlsx] [--exclude "tests/*,*.min.js"] [--min-level 中]

检查项（严重程度 高 / 中 / 低）：
  高：云厂商和平台密钥（AWS、腾讯云、阿里云、GitHub、GitLab、Slack、Google、OpenAI、Anthropic 等）、私钥文件内容、
      带密码的数据库连接串、代码里写死的密码 / Token / Secret；
  中：提交了 .env / 私钥 / 证书 / 凭证文件且没有被 .gitignore 忽略；拼接 SQL、eval / exec、shell=True、os.system、
      不安全的反序列化（pickle.loads、yaml.load 不带 SafeLoader）、innerHTML 赋值、child_process.exec、关闭证书校验；
  低：MD5 / SHA1 用于口令、DEBUG = True、chmod 777、JWT 字符串。
在行尾加注释 nosecret 可以忽略该行（如 # nosecret）。是 git 仓库时会额外标出已被 git 跟踪的敏感文件。
本工具基于规则匹配，会有误报和漏报，不能替代专业的安全审计。Exit codes: 0 无高 / 中风险, 3 发现高 / 中风险, 1 输入错误。
"""

from __future__ import annotations

import argparse
import fnmatch
import json
import os
import re
import shutil
import subprocess
import sys
from dataclasses import asdict, dataclass
from pathlib import Path

SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", "dist", "build", ".next", ".idea", ".vscode", "coverage",
             ".pytest_cache", ".mypy_cache", ".ruff_cache", "vendor", "target", ".tox", ".terraform"}  # fmt: skip
BINARY_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".ico", ".pdf", ".zip", ".gz", ".tar", ".7z", ".rar", ".exe", ".dll",
              ".so", ".dylib", ".class", ".jar", ".woff", ".woff2", ".ttf", ".otf", ".mp3", ".mp4", ".mov", ".db", ".sqlite",
              ".pyc", ".bin", ".iso", ".xlsx", ".docx", ".pptx"}  # fmt: skip
LEVELS = {"高": 3, "中": 2, "低": 1}
PLACEHOLDER = re.compile(r"^(?:x+|\*+|\.+|<[^>]*>|\$\{[^}]*\}|\{\{[^}]*\}\}|%\(?\w*\)?s?|changeme|change_me|your[_-]?\w*|example\w*|"
                         r"dummy\w*|test\w*|sample\w*|placeholder|password|secret|token|none|null|todo|redacted|\w*_here)$", re.I)  # fmt: skip

SECRET_RULES: list[tuple[str, str, re.Pattern[str]]] = [
    ("AWS Access Key", "高", re.compile(r"\b(AKIA|ASIA)[0-9A-Z]{16}\b")),
    (
        "AWS Secret Key",
        "高",
        re.compile(r"aws_secret_access_key\s*[:=]\s*['\"]?(?P<v>[A-Za-z0-9/+=]{40})"),
    ),
    ("腾讯云 SecretId", "高", re.compile(r"\bAKID[0-9A-Za-z]{32}\b")),
    ("阿里云 AccessKey", "高", re.compile(r"\bLTAI[0-9A-Za-z]{12,20}\b")),
    (
        "GitHub Token",
        "高",
        re.compile(r"\b(ghp|gho|ghu|ghs|ghr)_[0-9A-Za-z]{36}\b|\bgithub_pat_[0-9A-Za-z_]{60,}"),
    ),
    ("GitLab Token", "高", re.compile(r"\bglpat-[0-9A-Za-z_\-]{20}\b")),
    ("Slack Token", "高", re.compile(r"\bxox[abprs]-[0-9A-Za-z\-]{10,}")),
    ("Google API Key", "高", re.compile(r"\bAIza[0-9A-Za-z_\-]{35}\b")),
    ("Anthropic API Key", "高", re.compile(r"\bsk-ant-[0-9A-Za-z_\-]{20,}")),
    ("OpenAI / 兼容接口 API Key", "高", re.compile(r"\bsk-(?!ant-)(?:proj-)?[0-9A-Za-z_\-]{20,}")),
    (
        "私钥",
        "高",
        re.compile(
            r"-----BEGIN (?:RSA |EC |DSA |OPENSSH |PGP |ENCRYPTED )?PRIVATE KEY(?: BLOCK)?-----"
        ),
    ),
    (
        "带密码的连接串",
        "高",
        re.compile(r"\b[a-z][a-z0-9+]*://[^\s:/@'\"]+:(?P<v>[^\s@/'\"]{3,})@[\w.\-]+", re.I),
    ),
    (
        "写死的密码 / 密钥",
        "高",
        re.compile(
            r"""(?i)\b[\w.\-]*(?:password|passwd|pwd|secret|token|api[_-]?key|access[_-]?key|private[_-]?key|client[_-]?secret)\b["']?\s*[:=]\s*["'](?P<v>[^"'\s]{6,})["']"""
        ),
    ),
    (
        "JWT",
        "低",
        re.compile(r"\beyJ[0-9A-Za-z_\-]{10,}\.eyJ[0-9A-Za-z_\-]{10,}\.[0-9A-Za-z_\-]{10,}"),
    ),
]
CODE_RULES: list[tuple[str, str, tuple[str, ...], re.Pattern[str]]] = [
    (
        "拼接 SQL",
        "中",
        (".py",),
        re.compile(r"\.execute\(\s*(?:f[\"']|[\"'][^\"']*[\"']\s*(?:%|\+|\.format\())", re.I),
    ),
    (
        "拼接 SQL",
        "中",
        (".js", ".ts", ".java", ".php", ".go", ".cs"),
        re.compile(
            r"(?:query|execute|exec|prepare)\w*\(\s*(?:`[^`]*\$\{|[\"'][^\"']*(?:SELECT|INSERT|UPDATE|DELETE)[^\"']*[\"']\s*\+)",
            re.I,
        ),
    ),
    ("eval / exec", "中", (".py",), re.compile(r"(?<![\w.])(?:eval|exec)\s*\(")),
    ("eval", "中", (".js", ".ts", ".php"), re.compile(r"(?<![\w.])eval\s*\(|new Function\s*\(")),
    ("shell=True", "中", (".py",), re.compile(r"subprocess\.\w+\([^)]*shell\s*=\s*True")),
    ("os.system", "中", (".py",), re.compile(r"\bos\.(?:system|popen)\s*\(")),
    (
        "不安全的反序列化",
        "中",
        (".py",),
        re.compile(
            r"\bpickle\.loads?\s*\(|\byaml\.load\s*\((?![^)]*Loader\s*=\s*(?:yaml\.)?(?:Safe|CSafe)Loader)"
        ),
    ),
    (
        "innerHTML 赋值",
        "中",
        (".js", ".ts", ".jsx", ".tsx", ".vue", ".html"),
        re.compile(r"\.(?:innerHTML|outerHTML)\s*=|dangerouslySetInnerHTML|document\.write\s*\("),
    ),
    (
        "child_process.exec",
        "中",
        (".js", ".ts"),
        re.compile(r"\bchild_process\b.*\bexec\s*\(|\bexec(?:Sync)?\s*\(\s*`"),
    ),
    (
        "关闭证书校验",
        "中",
        (".py", ".js", ".ts", ".go", ".java"),
        re.compile(
            r"verify\s*=\s*False|rejectUnauthorized\s*:\s*false|InsecureSkipVerify\s*:\s*true|NODE_TLS_REJECT_UNAUTHORIZED\s*=\s*['\"]?0"
        ),
    ),
    (
        "弱哈希",
        "低",
        (".py", ".js", ".ts", ".java", ".php", ".go"),
        re.compile(
            r"\b(?:hashlib\.(?:md5|sha1)|createHash\(\s*['\"](?:md5|sha1)['\"]|MessageDigest\.getInstance\(\s*\"(?:MD5|SHA-1)\"|\bmd5\s*\()",
            re.I,
        ),
    ),
    (
        "DEBUG 开启",
        "低",
        (".py", ".env", ".ini", ".cfg", ".toml", ".yaml", ".yml"),
        re.compile(r"^\s*DEBUG\s*[:=]\s*(?:True|true|1)\b"),
    ),
    (
        "chmod 777",
        "低",
        (".sh", ".py", ".yml", ".yaml", ".js", ".ts", ""),
        re.compile(r"chmod\s+(?:-R\s+)?0?777|os\.chmod\([^)]*0o777"),
    ),
]
SENSITIVE_FILES = [
    (re.compile(r"^\.env(\..+)?$"), "环境变量文件"),
    (re.compile(r"^(id_rsa|id_dsa|id_ecdsa|id_ed25519)$"), "SSH 私钥"),
    (re.compile(r".*\.(pem|key|p12|pfx|jks|keystore)$", re.I), "私钥 / 证书文件"),
    (
        re.compile(r"^(credentials|service[-_]account|client_secret)[\w\-]*\.json$", re.I),
        "云服务凭证文件",
    ),
    (re.compile(r"^\.(npmrc|pypirc|netrc|git-credentials)$"), "包管理 / Git 凭证配置"),
]
SAFE_ENV_SUFFIX = (".example", ".sample", ".template", ".dist", ".defaults")


@dataclass
class Finding:
    level: str
    rule: str
    file: str
    line: int
    snippet: str


class ScanError(ValueError):
    pass


def mask(value: str) -> str:
    v = value.strip()
    if len(v) <= 8:
        return v[:2] + "****"
    return f"{v[:4]}****{v[-2:]}"


def iter_files(root: Path, excludes: list[str]) -> list[Path]:
    out = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(d for d in dirnames if d not in SKIP_DIRS)
        for name in sorted(filenames):
            p = Path(dirpath) / name
            rel = p.relative_to(root).as_posix()
            if any(fnmatch.fnmatch(rel, pat) or fnmatch.fnmatch(name, pat) for pat in excludes):
                continue
            if p.suffix.lower() in BINARY_EXT or p.is_symlink():
                continue
            try:
                if p.stat().st_size > 2 * 1024 * 1024:
                    continue
            except OSError:
                continue
            out.append(p)
    return out


def tracked_files(root: Path) -> set[str] | None:
    exe = shutil.which("git")
    if not exe or not (root / ".git").exists():
        return None
    r = subprocess.run(
        [exe, "-C", str(root), "ls-files"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
    )
    return set(r.stdout.splitlines()) if r.returncode == 0 else None


def gitignore_patterns(root: Path) -> list[str]:
    gi = root / ".gitignore"
    if not gi.is_file():
        return []
    return [
        ln.strip()
        for ln in gi.read_text(encoding="utf-8", errors="ignore").splitlines()
        if ln.strip() and not ln.startswith("#")
    ]


def ignored(rel: str, patterns: list[str]) -> bool:
    name = rel.rsplit("/", 1)[-1]
    hit = False
    for pat in patterns:
        neg = pat.startswith("!")
        p = pat[1:] if neg else pat
        p = p.rstrip("/").lstrip("/")
        if (
            fnmatch.fnmatch(name, p)
            or fnmatch.fnmatch(rel, p)
            or fnmatch.fnmatch(rel, f"**/{p}")
            or rel.startswith(p + "/")
        ):
            hit = not neg
    return hit


def scan(root: Path, excludes: list[str]) -> list[Finding]:
    findings: list[Finding] = []
    tracked = tracked_files(root)
    gi = gitignore_patterns(root)
    for p in iter_files(root, excludes):
        rel = p.relative_to(root).as_posix()
        for pat, label in SENSITIVE_FILES:
            if pat.match(p.name) and not p.name.endswith(SAFE_ENV_SUFFIX):
                if tracked is not None and rel in tracked:
                    findings.append(
                        Finding(
                            "高",
                            f"{label}已被 git 跟踪",
                            rel,
                            0,
                            "会随仓库一起推送，需从仓库和历史中移除并更换其中的密钥",
                        )
                    )
                elif not ignored(rel, gi):
                    findings.append(
                        Finding(
                            "中",
                            f"{label}没有被 .gitignore 忽略",
                            rel,
                            0,
                            "提交时容易被带上，建议加入 .gitignore",
                        )
                    )
                break
        try:
            text = p.read_text(encoding="utf-8")
        except (UnicodeDecodeError, OSError):
            continue
        suffix = p.suffix.lower() if p.suffix else ("" if not p.name.startswith(".env") else ".env")
        for no, line in enumerate(text.splitlines(), start=1):
            if "nosecret" in line or len(line) > 2000:
                continue
            for rule, level, rx in SECRET_RULES:
                m = rx.search(line)
                if not m:
                    continue
                value = m.group("v") if "v" in rx.groupindex else m.group(0)
                if rule in ("写死的密码 / 密钥", "带密码的连接串") and PLACEHOLDER.match(
                    value.strip()
                ):
                    continue
                if rule == "写死的密码 / 密钥" and re.search(
                    r"(?:os\.environ|getenv|process\.env|config\.|settings\.)", line
                ):
                    continue
                snippet = (
                    line.strip() if rule == "私钥" else line.strip().replace(value, mask(value))
                )[:160]
                findings.append(Finding(level, rule, rel, no, snippet))
                break
            for rule, level, exts, rx in CODE_RULES:
                if exts and suffix not in exts and not ("" in exts and not p.suffix):
                    continue
                if rx.search(line):
                    findings.append(Finding(level, rule, rel, no, line.strip()[:160]))
    order = {"高": 0, "中": 1, "低": 2}
    return sorted(findings, key=lambda f: (order[f.level], f.file, f.line))


ADVICE = {
    "写死的密码 / 密钥": "改为从环境变量或密钥管理服务读取；已经提交过的要立即更换",
    "拼接 SQL": "改用参数化查询（占位符 + 参数），不要把用户输入拼进 SQL",
    "eval / exec": "避免执行动态代码；解析数据用 json / ast.literal_eval",
    "shell=True": "改用参数列表调用 subprocess.run([...])，不要拼接命令字符串",
    "os.system": "改用 subprocess.run([...]) 并检查返回值",
    "不安全的反序列化": "不要反序列化不可信数据；YAML 用 yaml.safe_load",
    "innerHTML 赋值": "用 textContent 或框架的转义渲染，必须插 HTML 时先做净化",
    "关闭证书校验": "保持证书校验开启，自签名证书请配置 CA 文件",
    "弱哈希": "口令哈希改用 bcrypt / scrypt / Argon2",
}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("root")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("-o", "--output")
    parser.add_argument("--exclude", default="")
    parser.add_argument("--min-level", choices=list(LEVELS), default="低")
    args = parser.parse_args(argv)
    root = Path(args.root)
    if not root.is_dir():
        print(f"失败：{root} 不是目录", file=sys.stderr)
        return 1
    excludes = [e.strip() for e in args.exclude.split(",") if e.strip()]
    findings = [f for f in scan(root, excludes) if LEVELS[f.level] >= LEVELS[args.min_level]]
    counts = {lv: sum(1 for f in findings if f.level == lv) for lv in LEVELS}
    if args.json:
        print(json.dumps([asdict(f) for f in findings], ensure_ascii=False, indent=2))
    else:
        print(f"扫描 {root}：高 {counts['高']}，中 {counts['中']}，低 {counts['低']}")
        for f in findings:
            loc = f"{f.file}:{f.line}" if f.line else f.file
            print(f"- [{f.level}] {f.rule}  {loc}\n    {f.snippet}")
        rules = {f.rule for f in findings}
        tips = [f"{r}：{ADVICE[r]}" for r in ADVICE if r in rules]
        if any(f.level == "高" and f.line for f in findings):
            tips.insert(
                0,
                "发现疑似真实密钥：先到对应平台作废并重新生成，再从代码中移除；只删代码不换密钥没有用，git 历史里仍然能找到",
            )
        if tips:
            print("\n处理建议：")
            for t in tips:
                print(f"- {t}")
        print(
            "\n本工具基于规则匹配，可能有误报和漏报；确认是误报的行可在行尾加 nosecret 注释忽略。"
        )
    if args.output:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:  # pragma: no cover - environment check
            print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
            return 2
        wb = Workbook()
        ws = wb.active
        ws.title = "扫描结果"
        ws.append(["严重程度", "规则", "文件", "行号", "内容（已打码）", "处理状态"])
        fills = {"高": "FECACA", "中": "FED7AA", "低": "FEF9C3"}
        for f in findings:
            ws.append([f.level, f.rule, f.file, f.line or "", f.snippet, ""])
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=fills[f.level])
        for c in ws[1]:
            c.font = Font(bold=True)
        for col, w in zip("ABCDEF", (8, 22, 40, 6, 70, 12), strict=True):
            ws.column_dimensions[col].width = w
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.output)
        print(f"已生成 {args.output}")
    return 3 if counts["高"] or counts["中"] else 0


if __name__ == "__main__":
    sys.exit(main())
