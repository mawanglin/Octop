#!/usr/bin/env python3
"""预算明细 Markdown → 带公式的 Excel 预算表（分类小计、管理费、预备费、税费、总计）。

Usage:
    python scripts/budget_xlsx.py budget.md -o budget.xlsx

输入：可选 front matter + 一张 Markdown 表格，表头使用：
    类别 | 项目 | 单位 | 单价 | 数量 | 天数 | 备注
（"天数"可省略；单价、数量、天数填数字，可带"元""天"等单位文字。）
front matter：
    项目: 2026 城市宣传片
    币种: 元
    管理费率: 10%      # 可选，按小计合计计算
    预备费率: 5%       # 可选
    税率: 6%           # 可选，按（小计 + 管理费 + 预备费）计算
所有金额都是 Excel 公式，用户改单价或数量后自动重算。
Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
GROUP_FILL = PatternFill("solid", fgColor="E5E7EB")
TOTAL_FILL = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0.00"


class InputError(ValueError):
    pass


def number(value: str, where: str) -> float:
    text = re.sub(r"[,，\s元天次人个套项台小时¥￥]", "", value)
    if text == "":
        return 1.0 if "天数" in where else 0.0
    try:
        return float(text)
    except ValueError as exc:
        raise InputError(f"{where}「{value}」不是数字") from exc


def rate(value: str | None, name: str) -> float:
    if not value:
        return 0.0
    text = value.strip().rstrip("%％")
    try:
        r = float(text)
    except ValueError as exc:
        raise InputError(f"{name}「{value}」应写成百分比，如 6%") from exc
    return r / 100 if value.strip().endswith(("%", "％")) or r > 1 else r


def parse(source: str) -> tuple[dict[str, str], list[dict[str, str]]]:
    lines = source.replace("\r\n", "\n").split("\n")
    meta: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                for raw in lines[1:i]:
                    key, sep, value = raw.split("#", 1)[0].replace("：", ":", 1).partition(":")
                    if sep:
                        meta[key.strip()] = value.strip()
                lines = lines[i + 1 :]
                break
    rows = [
        [c.strip() for c in ln.strip().strip("|").split("|")]
        for ln in lines
        if ln.strip().startswith("|") and ln.strip().endswith("|")
    ]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    if len(rows) < 2:
        raise InputError("没有找到预算表（需要一张带表头的 Markdown 表格）")
    header = rows[0]
    for need in ("项目", "单价", "数量"):
        if need not in header:
            raise InputError(f"表头缺少「{need}」列")
    items = []
    for i, r in enumerate(rows[1:], start=1):
        if len(r) != len(header):
            raise InputError(f"第 {i} 行有 {len(r)} 列，表头有 {len(header)} 列")
        items.append(dict(zip(header, r, strict=True)))
    return meta, items


def build(meta: dict[str, str], items: list[dict[str, str]], out: Path) -> float:
    currency = meta.get("币种", "元")
    mgmt, reserve, tax = (rate(meta.get(k), k) for k in ("管理费率", "预备费率", "税率"))
    wb = Workbook()
    ws = wb.active
    ws.title = "预算表"
    ws.cell(1, 1, meta.get("项目", "预算表")).font = Font(size=16, bold=True)
    ws.cell(2, 1, f"单位：{currency}　（金额均为公式，修改单价/数量/天数后自动重算）").font = Font(
        size=10, color="6B7280"
    )
    cols = ["序号", "项目", "单位", "单价", "数量", "天数", "金额", "备注"]
    widths = [6, 28, 8, 12, 8, 8, 14, 28]
    head = 4
    for j, (name, w) in enumerate(zip(cols, widths, strict=True), start=1):
        c = ws.cell(head, j, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[chr(64 + j)].width = w

    groups: dict[str, list[dict[str, str]]] = {}
    for it in items:
        groups.setdefault(it.get("类别", "") or "其他", []).append(it)
    r = head + 1
    subtotal_cells: list[str] = []
    estimate = 0.0
    for g_index, (group, rows) in enumerate(groups.items(), start=1):
        ws.cell(r, 1, f"{'一二三四五六七八九十'[g_index - 1] if g_index <= 10 else g_index}、")
        ws.cell(r, 2, group).font = Font(bold=True)
        for j in range(1, 9):
            ws.cell(r, j).fill = GROUP_FILL
            ws.cell(r, j).border = BORDER
        group_row = r
        r += 1
        first = r
        for n, it in enumerate(rows, start=1):
            price = number(it.get("单价", ""), f"「{it.get('项目')}」的单价")
            qty = number(it.get("数量", ""), f"「{it.get('项目')}」的数量")
            days = number(it.get("天数", ""), f"「{it.get('项目')}」的天数")
            estimate += price * qty * days
            values = [
                n,
                it.get("项目", ""),
                it.get("单位", ""),
                price,
                qty,
                days,
                f"=D{r}*E{r}*F{r}",
                it.get("备注", ""),
            ]
            for j, v in enumerate(values, start=1):
                c = ws.cell(r, j, v)
                c.border = BORDER
                if j in (4, 7):
                    c.number_format = MONEY
                c.alignment = Alignment(wrap_text=j in (2, 8), vertical="top")
            r += 1
        sub = ws.cell(group_row, 7, f"=SUM(G{first}:G{r - 1})")
        sub.number_format = MONEY
        sub.font = Font(bold=True)
        subtotal_cells.append(f"G{group_row}")

    def total_row(
        label: str, formula: str, bold: bool = False, fill: PatternFill | None = None
    ) -> str:
        nonlocal r
        ws.cell(r, 2, label).font = Font(bold=bold)
        c = ws.cell(r, 7, formula)
        c.number_format = MONEY
        c.font = Font(bold=bold)
        for j in range(1, 9):
            ws.cell(r, j).border = BORDER
            if fill:
                ws.cell(r, j).fill = fill
        ref = f"G{r}"
        r += 1
        return ref

    base = total_row("小计合计", "=" + "+".join(subtotal_cells), bold=True)
    parts = [base]
    if mgmt:
        parts.append(total_row(f"管理费（{mgmt:.0%}）", f"={base}*{mgmt}"))
    if reserve:
        parts.append(total_row(f"预备费（{reserve:.0%}）", f"={base}*{reserve}"))
    pre_tax = "+".join(parts)
    if tax:
        parts.append(total_row(f"税费（{tax:.0%}）", f"=({pre_tax})*{tax}"))
    total_row("总计", "=" + "+".join(parts), bold=True, fill=TOTAL_FILL)
    ws.freeze_panes = ws.cell(head + 1, 1)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    grand = estimate * (1 + mgmt + reserve) * (1 + tax)
    return grand


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input")
    parser.add_argument("-o", "--output", required=True)
    args = parser.parse_args(argv)
    try:
        meta, items = parse(Path(args.input).read_text(encoding="utf-8"))
        grand = build(meta, items, Path(args.output))
    except (OSError, InputError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    print(f"已生成 {args.output}：{len(items)} 项，估算总计 {grand:,.2f} {meta.get('币种', '元')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
