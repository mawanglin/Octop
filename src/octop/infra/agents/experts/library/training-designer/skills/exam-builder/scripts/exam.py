#!/usr/bin/env python3
"""题库与组卷：检查题库，按题型数量随机组卷（选项顺序打乱、答案同步调整），导出 Word 试卷和答案页。

Usage:
    python scripts/exam.py check 题库.csv
    python scripts/exam.py paper 题库.csv -o 试卷.docx --single 10 --multi 5 --judge 5 [--blank 0] [--short 0]
                           [--title "安全生产培训考试"] [--topic 消防,用电] [--seed 1] [--no-shuffle]

题库 CSV（UTF-8 或 GBK）列：题型,题目,选项,答案,分值,知识点,难度,解析
  - 题型：单选 / 多选 / 判断 / 填空 / 简答；
  - 选项：单选多选用 | 分隔（A、B 前缀可写可不写），判断题留空；
  - 答案：单选写 B，多选写 ACD，判断写 对/错（√/×），填空多个空用 | 分隔，简答写参考要点；
  - 分值留空时按 单选 2、多选 3、判断 1、填空 2、简答 10。
check：缺字段、答案不在选项内、多选只有一个答案、重复题目、各题型和知识点数量。
paper：每个题型从题库中随机抽取，--topic 只从指定知识点抽题；同一份试卷不重复；--seed 固定后可以重新生成同一份。
Exit codes: 0 ok, 1 输入错误或 check 发现问题, 2 缺少 python-docx。
"""

from __future__ import annotations

import argparse
import csv
import io
import random
import re
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

TYPES = ["单选", "多选", "判断", "填空", "简答"]
DEFAULT_POINTS = {"单选": 2.0, "多选": 3.0, "判断": 1.0, "填空": 2.0, "简答": 10.0}
LETTERS = "ABCDEFGH"
CN_NUM = "一二三四五"


class ExamError(ValueError):
    pass


@dataclass
class Question:
    row: int
    kind: str
    text: str
    options: list[str]
    answer: str
    points: float
    topic: str
    level: str
    note: str


def read_bank(path: Path) -> list[Question]:
    raw = path.read_bytes()
    text = next((raw.decode(e) for e in ("utf-8-sig", "gb18030") if _ok(raw, e)), None)
    if text is None:
        raise ExamError("题库编码无法识别")
    rows = [
        {(k or "").strip(): (v or "").strip() for k, v in r.items()}
        for r in csv.DictReader(io.StringIO(text))
    ]
    if not rows or not {"题型", "题目", "答案"} <= set(rows[0]):
        raise ExamError("题库表头至少需要：题型,题目,答案（可选：选项,分值,知识点,难度,解析）")
    out = []
    for i, r in enumerate(rows, start=2):
        kind = r["题型"].replace("题", "").replace("单项选择", "单选").replace("多项选择", "多选")
        options = [
            re.sub(r"^[A-Ha-h][.、．:：)）]\s*", "", o.strip())
            for o in r.get("选项", "").split("|")
            if o.strip()
        ]
        pts = r.get("分值", "")
        out.append(Question(i, kind, r["题目"], options, r["答案"].replace(" ", ""), float(pts) if pts else DEFAULT_POINTS.get(kind, 0),
                            r.get("知识点", ""), r.get("难度", ""), r.get("解析", "")))  # fmt: skip
    return out


def _ok(raw: bytes, enc: str) -> bool:
    try:
        raw.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def judge_value(ans: str) -> bool | None:
    a = ans.strip().upper()
    if a in ("对", "√", "正确", "T", "TRUE", "Y", "是"):
        return True
    if a in ("错", "×", "X", "错误", "F", "FALSE", "N", "否"):
        return False
    return None


def validate(bank: list[Question]) -> list[str]:
    issues = []
    seen: dict[str, int] = {}
    for q in bank:
        where = f"第 {q.row} 行"
        if q.kind not in TYPES:
            issues.append(f"{where}：题型「{q.kind}」应为 {'/'.join(TYPES)}")
            continue
        if not q.text or not q.answer:
            issues.append(f"{where}：题目或答案为空")
            continue
        if q.kind in ("单选", "多选"):
            letters = set(q.answer.upper())
            valid = set(LETTERS[: len(q.options)])
            if len(q.options) < 2:
                issues.append(f"{where}：选择题至少需要 2 个选项")
            elif not letters <= valid:
                issues.append(f"{where}：答案「{q.answer}」不在选项 {''.join(sorted(valid))} 内")
            elif q.kind == "单选" and len(letters) != 1:
                issues.append(f"{where}：单选题答案应只有一个字母")
            elif q.kind == "多选" and len(letters) < 2:
                issues.append(f"{where}：多选题答案只有一个，是否应为单选")
        if q.kind == "判断" and judge_value(q.answer) is None:
            issues.append(f"{where}：判断题答案应为 对/错")
        if (
            q.kind == "填空"
            and q.text.count("___") + q.text.count("（ ）") + q.text.count("()") == 0
        ):
            issues.append(f"{where}：填空题题目里没有空（用 ___ 表示）")
        key = re.sub(r"\W", "", q.text)
        if key in seen:
            issues.append(f"{where}：与第 {seen[key]} 行题目重复")
        seen[key] = q.row
    return issues


def shuffled(q: Question, rng: random.Random, shuffle: bool) -> tuple[list[str], str]:
    if q.kind not in ("单选", "多选") or not shuffle:
        return q.options, q.answer.upper()
    idx = list(range(len(q.options)))
    if not any(
        re.search(r"以上|都(对|不)", o) for o in q.options
    ):  # keep "以上都对" style options in place
        rng.shuffle(idx)
    opts = [q.options[i] for i in idx]
    correct = {LETTERS.index(c) for c in q.answer.upper()}
    return opts, "".join(sorted(LETTERS[new] for new, old in enumerate(idx) if old in correct))


def cmd_check(bank: list[Question]) -> int:
    issues = validate(bank)
    kinds = Counter(q.kind for q in bank)
    topics = Counter(q.topic or "未分类" for q in bank)
    print(f"题库共 {len(bank)} 题：" + "、".join(f"{k} {kinds[k]}" for k in TYPES if kinds[k]))
    print("知识点：" + "、".join(f"{k} {v}" for k, v in topics.most_common()))
    print(f"发现 {len(issues)} 个问题" + ("：" if issues else ""))
    for i in issues:
        print(f"- {i}")
    return 1 if issues else 0


def cmd_paper(bank: list[Question], a: argparse.Namespace) -> int:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 python-docx：请运行 `pip install python-docx`", file=sys.stderr)
        return 2
    issues = validate(bank)
    if issues:
        raise ExamError(f"题库有 {len(issues)} 个问题，请先运行 check 修正")
    topics = {t.strip() for t in re.split(r"[,，、]", a.topic or "") if t.strip()}
    pool = [q for q in bank if not topics or q.topic in topics]
    want = {"单选": a.single, "多选": a.multi, "判断": a.judge, "填空": a.blank, "简答": a.short}
    rng = random.Random(a.seed)
    picked: dict[str, list[Question]] = {}
    for kind, n in want.items():
        cands = [q for q in pool if q.kind == kind]
        if n > len(cands):
            raise ExamError(
                f"{kind}题需要 {n} 道，题库{'（指定知识点）' if topics else ''}只有 {len(cands)} 道"
            )
        if n:
            picked[kind] = rng.sample(cands, n)
    if not picked:
        raise ExamError("请至少指定一种题型的数量，如 --single 10")
    total = sum(q.points for qs in picked.values() for q in qs)
    doc = Document()

    def center(text: str, size: float, bold: bool = True) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)

    center(a.title, 18)
    center(f"满分 {total:g} 分　　考试时间 ______ 分钟", 11, bold=False)
    doc.add_paragraph("姓名：__________　部门：__________　日期：__________　得分：______")
    answers: list[str] = []
    n = 0
    for sec, kind in enumerate(k for k in TYPES if k in picked):
        qs = picked[kind]
        pts = {q.points for q in qs}
        each = f"每题 {qs[0].points:g} 分，" if len(pts) == 1 else ""
        head = doc.add_paragraph()
        head.add_run(
            f"{CN_NUM[sec]}、{kind}题（{each}共 {sum(q.points for q in qs):g} 分）"
        ).bold = True
        answers.append(f"{CN_NUM[sec]}、{kind}题")
        for q in qs:
            n += 1
            opts, ans = shuffled(q, rng, not a.no_shuffle)
            stem = q.text if q.kind != "判断" else f"{q.text}（　）"
            if q.kind in ("单选", "多选"):
                stem += "（　）"
            doc.add_paragraph(f"{n}. {stem}" + ("" if len(pts) == 1 else f"（{q.points:g} 分）"))
            if opts:
                line = "　　".join(f"{LETTERS[i]}. {o}" for i, o in enumerate(opts))
                doc.add_paragraph(line).paragraph_format.left_indent = Pt(18)
            if q.kind == "简答":
                for _ in range(4):
                    doc.add_paragraph("_" * 60)
            shown = ("对" if judge_value(ans) else "错") if q.kind == "判断" else ans
            answers.append(f"{n}. {shown}" + (f"（{q.note}）" if q.note else ""))
    doc.add_page_break()
    center(f"{a.title}（参考答案）", 16)
    for line in answers:
        p = doc.add_paragraph(line)
        if not re.match(r"\d", line):
            p.runs[0].bold = True
    out = Path(a.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(
        f"已生成 {out}：{n} 题，满分 {total:g} 分（"
        + "、".join(f"{k} {len(v)}" for k, v in picked.items())
        + "），答案在最后一页"
    )
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check")
    c.add_argument("bank")
    p = sub.add_parser("paper")
    p.add_argument("bank")
    p.add_argument("-o", "--output", required=True)
    for opt in ("single", "multi", "judge", "blank", "short"):
        p.add_argument(f"--{opt}", type=int, default=0)
    p.add_argument("--title", default="培训考试")
    p.add_argument("--topic")
    p.add_argument("--seed", type=int)
    p.add_argument("--no-shuffle", action="store_true")
    args = parser.parse_args(argv)
    try:
        bank = read_bank(Path(args.bank))
        return cmd_check(bank) if args.cmd == "check" else cmd_paper(bank, args)
    except (OSError, ExamError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
