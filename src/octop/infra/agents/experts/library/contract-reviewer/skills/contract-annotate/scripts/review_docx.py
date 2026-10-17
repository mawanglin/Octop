#!/usr/bin/env python3
"""合同审阅辅助：导出带段落编号的合同全文；把审阅意见作为 Word 批注写回合同副本。

Usage:
    python scripts/review_docx.py extract 合同.docx                 # 带 [P12] 段落编号的全文
    python scripts/review_docx.py annotate 合同.docx review.md -o 合同-审阅.docx [--author 审阅人] [--report 风险清单.xlsx]

review.md 是一张 Markdown 表格：
    | 段落 | 原文片段 | 风险等级 | 问题 | 修改建议 |
- 段落：extract 输出的编号（如 P12），可留空；
- 原文片段：从合同中原样复制的一段文字（10–40 字为宜），批注会精确标在这段文字上；
  找不到时退回到整段批注，并在输出中提示；
- 风险等级：高 / 中 / 低。
原合同不会被修改；批注写在新文件里。Exit codes: 0 ok, 3 有片段未找到, 1 输入错误, 2 缺少依赖。
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from docx import Document
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 python-docx：请运行 `pip install python-docx`", file=sys.stderr)
    sys.exit(2)

LEVELS = ("高", "中", "低")


class ReviewError(ValueError):
    pass


@dataclass
class Finding:
    para: str
    snippet: str
    level: str
    issue: str
    advice: str


def iter_paragraphs(doc: Any) -> list[Any]:
    """Body paragraphs followed by table-cell paragraphs, in document order per table."""
    out: list[Any] = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                out.extend(p for p in cell.paragraphs if p not in out)
    return out


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", text)


def cmd_extract(path: Path) -> int:
    doc = Document(str(path))
    for i, p in enumerate(iter_paragraphs(doc), start=1):
        if p.text.strip():
            print(f"[P{i}] {p.text.strip()}")
    return 0


def parse_review(path: Path) -> list[Finding]:
    rows = [
        [c.strip() for c in ln.strip().strip("|").split("|")]
        for ln in path.read_text(encoding="utf-8").splitlines()
        if ln.strip().startswith("|") and ln.strip().endswith("|")
    ]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    if len(rows) < 2:
        raise ReviewError("审阅意见需要一张带表头的 Markdown 表格")
    header = rows[0]
    for need in ("原文片段", "问题", "修改建议"):
        if need not in header:
            raise ReviewError(f"审阅表格缺少「{need}」列")
    findings = []
    for i, r in enumerate(rows[1:], start=1):
        rec = dict(zip(header, r + [""] * (len(header) - len(r)), strict=False))
        level = rec.get("风险等级", "").strip() or "中"
        if level not in LEVELS:
            raise ReviewError(f"第 {i} 条风险等级「{level}」应为 高 / 中 / 低")
        if not rec.get("原文片段") and not rec.get("段落"):
            raise ReviewError(f"第 {i} 条需要「原文片段」或「段落」")
        findings.append(
            Finding(
                rec.get("段落", ""), rec.get("原文片段", ""), level, rec["问题"], rec["修改建议"]
            )
        )
    return findings


def split_run(run: Any, offset: int) -> Any:
    """Split *run* at *offset*; returns the new run holding the tail (same formatting)."""
    text = run.text
    tail = copy.deepcopy(run._r)
    run._r.addnext(tail)
    run.text = text[:offset]
    from docx.text.run import Run

    new = Run(tail, run._parent)
    new.text = text[offset:]
    return new


def isolate(paragraph: Any, start: int, end: int) -> list[Any]:
    """Return runs covering exactly paragraph.text[start:end], splitting runs as needed."""
    pos = 0
    covered: list[Any] = []
    for run in list(paragraph.runs):
        length = len(run.text)
        r_start, r_end = pos, pos + length
        pos = r_end
        if r_end <= start or r_start >= end or length == 0:
            continue
        current = run
        if start > r_start:
            current = split_run(current, start - r_start)
            r_start = start
        if end < r_end:
            split_run(current, end - r_start)
        covered.append(current)
    return covered


def locate(paragraphs: list[Any], f: Finding) -> tuple[Any, list[Any], bool]:
    """Find the paragraph and runs for a finding. Returns (paragraph, runs, exact)."""
    candidates = paragraphs
    if f.para:
        m = re.fullmatch(r"[Pp]?(\d+)", f.para.strip())
        if m and 1 <= int(m.group(1)) <= len(paragraphs):
            candidates = [paragraphs[int(m.group(1)) - 1], *paragraphs]
    snippet = f.snippet.strip()
    if snippet:
        for p in candidates:
            idx = p.text.find(snippet)
            if idx >= 0:
                return p, isolate(p, idx, idx + len(snippet)), True
        target = normalize(snippet)
        for p in candidates:
            if target and target in normalize(p.text) and p.runs:
                return p, list(p.runs), False
    if f.para and candidates is not paragraphs and candidates[0].runs:
        return candidates[0], list(candidates[0].runs), False
    raise ReviewError(f"找不到「{snippet[:20]}」，请从 extract 输出中原样复制片段或填写段落编号")


def cmd_annotate(args: argparse.Namespace) -> int:
    src, out = Path(args.contract), Path(args.output)
    if src.resolve() == out.resolve():
        raise ReviewError("输出文件不能覆盖原合同")
    if out.exists() and not args.overwrite:
        raise ReviewError(f"{out} 已存在；确认要覆盖时加 --overwrite")
    findings = parse_review(Path(args.review))
    doc = Document(str(src))
    paragraphs = iter_paragraphs(doc)
    fuzzy: list[str] = []
    missing: list[str] = []
    placed = 0
    for i, f in enumerate(findings, start=1):
        try:
            p, runs, exact = locate(paragraphs, f)
        except ReviewError as exc:
            missing.append(f"第 {i} 条：{exc}")
            continue
        text = f"【{f.level}风险】{f.issue}\n建议：{f.advice}"
        doc.add_comment(runs, text=text, author=args.author, initials=args.author[:2])
        placed += 1
        if not exact:
            fuzzy.append(f"第 {i} 条：片段与原文空格或标点不完全一致，已标注整段")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    counts = {lv: sum(1 for f in findings if f.level == lv) for lv in LEVELS}
    print(
        f"已生成 {out}：{placed}/{len(findings)} 条批注（高 {counts['高']}、中 {counts['中']}、低 {counts['低']}），原合同未修改"
    )
    if args.report:
        write_report(findings, Path(args.report))
        print(f"风险清单：{args.report}")
    for msg in fuzzy:
        print(f"提示：{msg}")
    for msg in missing:
        print(f"未标注：{msg}")
    return 3 if missing else 0


def write_report(findings: list[Finding], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:  # pragma: no cover - environment check
        raise ReviewError("生成风险清单需要 openpyxl") from exc
    fills = {"高": "FEE2E2", "中": "FEF3C7", "低": "F3F4F6"}
    order = {lv: i for i, lv in enumerate(LEVELS)}
    wb = Workbook()
    ws = wb.active
    ws.title = "风险清单"
    ws.append(["序号", "风险等级", "段落", "原文片段", "问题", "修改建议"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    for i, f in enumerate(sorted(findings, key=lambda f: order[f.level]), start=1):
        ws.append([i, f.level, f.para, f.snippet, f.issue, f.advice])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fills[f.level])
    for col, width in zip("ABCDEF", (6, 10, 8, 36, 40, 44), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extract", help="导出带段落编号的全文")
    e.add_argument("contract")
    a = sub.add_parser("annotate", help="把审阅意见写成批注")
    a.add_argument("contract")
    a.add_argument("review", help="审阅意见 Markdown 表格")
    a.add_argument("-o", "--output", required=True)
    a.add_argument("--author", default="Octop 合同审阅")
    a.add_argument("--report", help="同时输出风险清单 .xlsx")
    a.add_argument("--overwrite", action="store_true")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "extract":
            return cmd_extract(Path(args.contract))
        return cmd_annotate(args)
    except (OSError, ReviewError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
