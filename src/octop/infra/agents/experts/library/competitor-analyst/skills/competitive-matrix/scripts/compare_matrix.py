#!/usr/bin/env python3
"""竞品对比矩阵：按产品类型生成带权重的评分模板；打分后计算加权总分、排名、各维度领先者和我方差距，
检查没有证据支撑的评分，输出带雷达图和柱状图的 Excel。

Usage:
    python scripts/compare_matrix.py template --type internet|saas|ai|hardware|consumer \
        --competitors 我方,竞品A,竞品B [--self 我方] -o 竞品矩阵.xlsx
    python scripts/compare_matrix.py score 竞品矩阵.xlsx [--self 我方] [-o 评分结果.xlsx] [--json]

模板有两张表：
  评分   每行一个维度：维度、说明、权重，之后每个竞品一列，填 1–5 分（可留空表示未评估）
  证据   每条评分的依据：竞品、维度、结论、来源链接、获取日期、可信度（高/中/低）
score 读取模板（也接受同样表头的 CSV），权重自动归一化；没有证据或证据缺少来源链接的评分会被列出。
Exit codes: 0 ok, 3 有需要补充的地方（缺分、缺证据）, 1 输入错误。
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

PRESETS: dict[str, list[tuple[str, str, float]]] = {
    "internet": [
        ("用户与场景", "目标人群、核心场景、需求强度是否清晰", 10),
        ("核心功能", "主流程功能完整度、差异化功能", 20),
        ("体验与交互", "上手成本、关键路径步数、稳定性和性能", 15),
        ("内容与生态", "内容供给、社区、第三方生态和网络效应", 10),
        ("增长与渠道", "获客渠道、用户规模和增速、留存", 15),
        ("商业化与定价", "收费模式、价格、付费转化", 15),
        ("技术与数据", "技术壁垒、数据积累、迭代速度", 5),
        ("品牌与口碑", "知名度、应用商店评分、社媒口碑", 10),
    ],
    "saas": [
        ("功能覆盖", "核心模块是否齐全、深度", 20),
        ("易用与上手", "实施周期、学习成本、模板和引导", 15),
        ("集成与开放", "API、Webhook、常用系统集成", 10),
        ("安全与合规", "权限、审计、等保/ISO 等认证、数据驻留", 10),
        ("部署方式", "公有云、私有化、混合部署", 5),
        ("定价与套餐", "定价透明度、单价、免费版和试用", 15),
        ("服务与支持", "客户成功、响应时效、文档和社区", 10),
        ("客户与案例", "标杆客户、行业覆盖、续约口碑", 15),
    ],
    "ai": [
        ("核心效果", "关键任务的输出质量和准确性（附测试样例）", 25),
        ("任务完成度", "能否自主多步完成任务、工具调用和自动化程度", 15),
        ("模型与技术", "自研或接入的模型、上下文长度、多模态", 10),
        ("数据安全", "数据用途说明、私有化、隐私合规", 10),
        ("易用性", "上手成本、交互方式、速度", 10),
        ("集成与生态", "插件、API、与常用工具打通", 10),
        ("定价与额度", "免费额度、付费价格、限制条件", 10),
        ("口碑与增长", "用户规模、评分、社媒讨论度", 10),
    ],
    "hardware": [
        ("性能参数", "核心规格和实测性能", 25),
        ("做工与设计", "材质、工艺、外观、重量", 10),
        ("功能特性", "差异化功能和使用场景覆盖", 15),
        ("软件与生态", "配套 App、系统更新、生态联动", 10),
        ("价格与性价比", "定价、促销价、同价位对比", 20),
        ("渠道与服务", "购买渠道、售后网点、保修政策", 10),
        ("品牌与口碑", "品牌力、电商评价、媒体测评", 10),
    ],
    "consumer": [
        ("产品力", "配方/成分/功能、品质和差异化卖点", 25),
        ("价格带", "主销价位、规格单价、促销力度", 15),
        ("渠道覆盖", "线上平台、线下网点、区域覆盖", 15),
        ("品牌与营销", "品牌认知、内容营销、达人投放", 15),
        ("包装与体验", "包装设计、开箱和使用体验", 10),
        ("供应链与品控", "产能、品控、交付时效", 10),
        ("用户口碑", "电商评分、复购、社媒评价", 10),
    ],
}
FIXED = ("维度", "说明", "权重")
EVIDENCE_COLS = ("竞品", "维度", "结论", "来源链接", "获取日期", "可信度")


class MatrixError(ValueError):
    pass


def split_names(raw: str) -> list[str]:
    names = [x.strip() for x in raw.replace("，", ",").split(",") if x.strip()]
    if len(names) < 2:
        raise MatrixError("--competitors 至少两个（包括我方）")
    if len(set(names)) != len(names):
        raise MatrixError("竞品名称重复")
    return names


def template(kind: str, names: list[str], out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "评分"
    ws.append([*FIXED, *names])
    for dim, note, weight in PRESETS[kind]:
        ws.append([dim, note, weight, *[None] * len(names)])
    dv = DataValidation(
        type="whole", operator="between", formula1="1", formula2="5", allow_blank=True
    )
    dv.error, dv.prompt = "请填 1–5 的整数", "1 很差 · 2 较差 · 3 一般 · 4 较好 · 5 很好"
    ws.add_data_validation(dv)
    last_col = ws.cell(1, len(FIXED) + len(names)).column_letter
    dv.add(f"D2:{last_col}{ws.max_row}")
    ev = wb.create_sheet("证据")
    ev.append(list(EVIDENCE_COLS))
    for dim, _, _ in PRESETS[kind]:
        for name in names:
            ev.append([name, dim, "", "", "", ""])
    guide = wb.create_sheet("说明")
    for line in (
        "评分标准：1 很差 · 2 较差 · 3 一般 · 4 较好 · 5 很好；未评估留空，不要猜。",
        "每个分数都要在“证据”表里写结论和来源链接（官网、应用商店、评测、实测截图等）和获取日期。",
        "可信度：高＝官方或实测；中＝权威媒体或第三方数据；低＝用户评论、转述。",
        "权重可以按业务重点调整，脚本会自动归一化。可以增删维度行和竞品列。",
    ):
        guide.append([line])
    for sheet in (ws, ev):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
        sheet.freeze_panes = "A2"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 14, 40
    for col in ("A", "B", "C", "D", "E", "F"):
        ev.column_dimensions[col].width = {"C": 40, "D": 36}.get(col, 14)
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 90
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def _rows_from_sheet(path: Path, sheet: str) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        return []
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def read_inputs(path: Path) -> tuple[list[list[Any]], list[list[Any]]]:
    if not path.is_file():
        raise MatrixError(f"找不到 {path}")
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as fh:
            return [list(r) for r in csv.reader(fh)], []
    return _rows_from_sheet(path, "评分") or _rows_from_sheet(path, path.stem), _rows_from_sheet(
        path, "证据"
    )


def to_score(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        v = float(value)
    except (TypeError, ValueError):
        raise MatrixError(f"分数“{value}”不是数字") from None
    if not 1 <= v <= 5:
        raise MatrixError(f"分数 {v:g} 不在 1–5 之间")
    return v


def score(path: Path, self_name: str | None) -> dict[str, Any]:
    rows, evidence_rows = read_inputs(path)
    if not rows or [str(c or "").strip() for c in rows[0][:3]] != list(FIXED):
        raise MatrixError("评分表前三列必须是：维度、说明、权重（用 template 生成）")
    names = [str(c).strip() for c in rows[0][3:] if str(c or "").strip()]
    if len(names) < 2:
        raise MatrixError("评分表至少要有两个竞品列")
    dims = []
    for r in rows[1:]:
        if not r or not str(r[0] or "").strip():
            continue
        try:
            weight = float(r[2] or 0)
        except (TypeError, ValueError):
            raise MatrixError(f"维度“{r[0]}”的权重不是数字") from None
        scores = {n: to_score(r[3 + i] if 3 + i < len(r) else None) for i, n in enumerate(names)}
        dims.append({"维度": str(r[0]).strip(), "权重": weight, "分数": scores})
    total_w = sum(d["权重"] for d in dims)
    if total_w <= 0:
        raise MatrixError("权重之和必须大于 0")
    for d in dims:
        d["权重占比"] = round(d["权重"] / total_w, 4)

    results = []
    for n in names:
        rated = [d for d in dims if d["分数"][n] is not None]
        covered = sum(d["权重占比"] for d in rated)
        weighted = sum(d["分数"][n] * d["权重占比"] for d in rated)
        results.append(
            {
                "竞品": n,
                "加权得分": round(weighted / covered, 2) if covered else None,
                "评估覆盖": round(covered, 2),
                "已评维度": f"{len(rated)}/{len(dims)}",
            }
        )
    results.sort(key=lambda x: (x["加权得分"] is None, -(x["加权得分"] or 0)))
    for i, r in enumerate(results, 1):
        r["排名"] = i if r["加权得分"] is not None else "-"

    leaders = []
    for d in dims:
        rated = {n: s for n, s in d["分数"].items() if s is not None}
        if rated:
            best = max(rated.values())
            leaders.append(
                {
                    "维度": d["维度"],
                    "领先": "、".join(n for n, s in rated.items() if s == best),
                    "最高分": best,
                }
            )

    gaps = []
    if self_name:
        if self_name not in names:
            raise MatrixError(f"评分表里没有“{self_name}”这一列")
        for d in dims:
            mine = d["分数"][self_name]
            others = {n: s for n, s in d["分数"].items() if n != self_name and s is not None}
            if mine is None or not others:
                continue
            best_name, best = max(others.items(), key=lambda kv: kv[1])
            gaps.append(
                {
                    "维度": d["维度"],
                    "我方": mine,
                    "最强竞品": best_name,
                    "竞品分": best,
                    "差距": round(mine - best, 1),
                    "加权影响": round((mine - best) * d["权重占比"], 3),
                }
            )
        gaps.sort(key=lambda g: g["加权影响"])

    missing_scores = [f"{n}·{d['维度']}" for d in dims for n in names if d["分数"][n] is None]
    ev_index: dict[tuple[str, str], dict[str, str]] = {}
    if evidence_rows:
        header = [str(c or "").strip() for c in evidence_rows[0]]
        idx = {c: header.index(c) for c in EVIDENCE_COLS if c in header}
        for r in evidence_rows[1:]:

            def get(c: str, r: list[Any] = r) -> str:
                return str(r[idx[c]] or "").strip() if c in idx and idx[c] < len(r) else ""

            key = (get("竞品"), get("维度"))
            if key[0] and key[1] and (get("结论") or get("来源链接")):
                ev_index[key] = {c: get(c) for c in EVIDENCE_COLS}
    no_evidence, no_source = [], []
    for d in dims:
        for n in names:
            if d["分数"][n] is None:
                continue
            ev = ev_index.get((n, d["维度"]))
            if not ev:
                no_evidence.append(f"{n}·{d['维度']}")
            elif not ev["来源链接"]:
                no_source.append(f"{n}·{d['维度']}")
    return {
        "竞品": names,
        "维度": dims,
        "排名": results,
        "维度领先": leaders,
        "我方差距": gaps,
        "缺分": missing_scores,
        "缺证据": no_evidence if evidence_rows else [],
        "缺来源": no_source,
        "有证据表": bool(evidence_rows),
    }


def write_result(path: Path, res: dict[str, Any]) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, RadarChart, Reference
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "排名"
    ws.append(["排名", "竞品", "加权得分", "评估覆盖", "已评维度"])
    for r in res["排名"]:
        ws.append([r["排名"], r["竞品"], r["加权得分"], r["评估覆盖"], r["已评维度"]])
    bar = BarChart()
    bar.title, bar.y_axis.title = "加权得分", "分（1–5）"
    bar.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=2, max_row=ws.max_row))
    bar.y_axis.scaling.min, bar.y_axis.scaling.max = 0, 5
    bar.height, bar.width = 8, 14
    ws.add_chart(bar, "G2")

    m = wb.create_sheet("维度对比")
    m.append(["维度", "权重占比", *res["竞品"]])
    for d in res["维度"]:
        m.append([d["维度"], d["权重占比"], *[d["分数"][n] for n in res["竞品"]]])
    radar = RadarChart()
    radar.type, radar.style, radar.title = "marker", 26, "竞品能力雷达图"
    data = Reference(m, min_col=3, max_col=2 + len(res["竞品"]), min_row=1, max_row=m.max_row)
    radar.add_data(data, titles_from_data=True)
    radar.set_categories(Reference(m, min_col=1, min_row=2, max_row=m.max_row))
    radar.y_axis.scaling.min, radar.y_axis.scaling.max = 0, 5
    radar.height, radar.width = 12, 16
    m.add_chart(radar, f"{m.cell(1, len(res['竞品']) + 4).column_letter}2")

    for title, rows in (("维度领先", res["维度领先"]), ("我方差距", res["我方差距"])):
        sheet = wb.create_sheet(title)
        if rows:
            sheet.append(list(rows[0]))
            for r in rows:
                sheet.append(list(r.values()))
    todo = wb.create_sheet("待补充")
    todo.append(["类型", "竞品·维度"])
    for kind in ("缺分", "缺证据", "缺来源"):
        for item in res[kind]:
            todo.append([kind, item])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
        sheet.column_dimensions["A"].width = 16
        sheet.column_dimensions["B"].width = 18
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("template")
    t.add_argument("--type", choices=sorted(PRESETS), default="internet")
    t.add_argument("--competitors", required=True, help="逗号分隔，包括我方")
    t.add_argument("-o", "--output", required=True)
    s = sub.add_parser("score")
    s.add_argument("input")
    s.add_argument("-o", "--output")
    s.add_argument("--json", action="store_true")
    for p in (t, s):
        p.add_argument("--self", help="我方产品名（列名），用于计算差距")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "template":
            names = split_names(args.competitors)
            if args.self and args.self not in names:
                raise MatrixError(f"--self “{args.self}” 不在 --competitors 里")
            template(args.type, names, Path(args.output))
            print(
                f"已生成 {args.type} 模板：{args.output}（{len(PRESETS[args.type])} 个维度 × {len(names)} 个竞品）"
            )
            print("下一步：在“评分”表填 1–5 分，在“证据”表写结论和来源链接，然后运行 score。")
            return 0
        res = score(Path(args.input), args.self)
    except MatrixError as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    if args.output:
        write_result(Path(args.output), res)
    if args.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    else:
        print("加权得分排名（只按已评维度计算，覆盖不足时注意）：")
        for r in res["排名"]:
            score_text = "-" if r["加权得分"] is None else f"{r['加权得分']:.2f}"
            print(
                f"  {r['排名']}. {r['竞品']}：{score_text} 分，已评 {r['已评维度']}，覆盖 {r['评估覆盖']:.0%} 权重"
            )
        print(
            "各维度领先："
            + "；".join(f"{x['维度']}→{x['领先']}（{x['最高分']:g}）" for x in res["维度领先"])
        )
        if res["我方差距"]:
            print("我方差距（按加权影响从大到小）：")
            for g in res["我方差距"]:
                sign = "落后" if g["差距"] < 0 else ("持平" if g["差距"] == 0 else "领先")
                print(
                    f"  - {g['维度']}：我方 {g['我方']:g}，{g['最强竞品']} {g['竞品分']:g}，{sign} {abs(g['差距']):g} 分"
                )
        for kind, hint in (
            ("缺分", "未评估"),
            ("缺证据", "有分数但证据表里没有对应结论"),
            ("缺来源", "证据缺少来源链接"),
        ):
            if res[kind]:
                print(
                    f"{kind} {len(res[kind])} 项（{hint}）：{'、'.join(res[kind][:12])}{' …' if len(res[kind]) > 12 else ''}"
                )
        if not res["有证据表"]:
            print("提示：没有证据表，评分无法追溯来源；正式报告建议用 template 生成的 Excel。")
        if args.output:
            print(f"已保存：{args.output}（含雷达图、柱状图、待补充清单）")
    return 3 if res["缺分"] or res["缺证据"] or res["缺来源"] else 0


if __name__ == "__main__":
    sys.exit(main())
