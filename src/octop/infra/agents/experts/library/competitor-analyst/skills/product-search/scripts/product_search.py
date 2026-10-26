#!/usr/bin/env python3
"""产品搜索：从公开数据源取结构化的产品信息，每条结果都带来源链接和获取时间。

Usage:
    python scripts/product_search.py appstore 关键词 [--country cn] [--limit 10] [-o 结果.xlsx] [--json]
    python scripts/product_search.py app 应用ID... [--country cn] [--json]
    python scripts/product_search.py reviews 应用ID [--country cn] [--pages 3] [-o 评论.xlsx] [--json]
    python scripts/product_search.py github 关键词 [--limit 10] [--json]
    python scripts/product_search.py site 网址... [--history] [--json]

appstore  App Store 搜索：名称、开发者、评分和评分人数、价格、分类、版本、最近更新、上架时间、大小
app       按应用 ID 查详情：以上字段加更新说明、简介、最低系统、支持语言
reviews   App Store 最新用户评论（每页约 50 条，最多 10 页）：评分分布、各版本平均分、
          差评和好评里的高频词、典型差评原文
github    GitHub 开源项目：星标、Fork、最近提交、许可证、主要语言、话题（设置 GITHUB_TOKEN 可提高频率上限）
site      官网信息：标题和描述、定价/下载/文档/更新日志/招聘等关键页面链接、ICP 备案号、页脚主体公司、
          技术和统计工具线索、App Store 链接；--history 用网页时光机查最早收录时间（推断上线时间）
只读取公开接口和页面；数据以来源页面为准，引用时注明获取时间。
Exit codes: 0 ok, 1 输入错误, 2 网络或接口错误。
"""

from __future__ import annotations

import argparse
import html
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import UTC, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

ITUNES = os.environ.get("PRODUCT_SEARCH_ITUNES", "https://itunes.apple.com")
GITHUB = os.environ.get("PRODUCT_SEARCH_GITHUB", "https://api.github.com")
WAYBACK = os.environ.get("PRODUCT_SEARCH_WAYBACK", "https://web.archive.org")
UA = "Mozilla/5.0 (compatible; OctopProductSearch/1.0)"
TIMEOUT = 20

KEY_PAGES = {
    "产品介绍": r"products?\b|features?|solutions?|产品|功能|方案|能力|介绍|tour",
    "定价": r"pricing|price|plans?|定价|价格|套餐|购买|会员|vip",
    "下载": r"download|下载|客户端|app\b",
    "文档": r"docs?|documentation|help|guide|文档|帮助|教程|手册",
    "更新日志": r"changelog|release[-_ ]?notes|what'?s[-_ ]?new|更新日志|版本|动态",
    "博客": r"blog|news|博客|资讯|新闻",
    "招聘": r"careers?|jobs?|join|招聘|加入我们",
    "关于": r"about|company|关于|公司",
    "API": r"\bapi\b|developer|开放平台|开发者",
}
TECH_MARKERS = {
    "Next.js": r"__NEXT_DATA__|/_next/",
    "Nuxt": r"__NUXT__|/_nuxt/",
    "React": r"data-reactroot|react-dom",
    "Vue": r"data-v-[0-9a-f]{8}|vue(?:\.min)?\.js",
    "Angular": r"ng-version=",
    "WordPress": r"wp-content|wp-includes",
    "Shopify": r"cdn\.shopify\.com",
    "百度统计": r"hm\.baidu\.com",
    "Google Analytics": r"googletagmanager\.com|google-analytics\.com|gtag\(",
    "神策数据": r"sensorsdata|sensors\.js",
    "GrowingIO": r"growingio",
    "友盟/CNZZ": r"umeng\.com|cnzz\.com",
    "Intercom": r"intercom",
    "Stripe": r"js\.stripe\.com",
}
ICP = re.compile(
    r"[京津沪渝冀豫云辽黑湘皖鲁新苏浙赣鄂桂甘晋蒙陕吉闽贵粤青藏川宁琼]ICP[备证]\d{6,}号(?:-\d+)?"
)
COMPANY = re.compile(
    r"(?:©|&copy;|版权所有|Copyright)[^<\n]{0,40}?([\u4e00-\u9fa5（）()]{2,40}?(?:股份有限公司|有限责任公司|有限公司))"
)
# Function characters that never start or end a meaningful phrase.
STOP = set("的了是我也就都很还在吧啊呢吗么这那它他她你")
STOP_PHRASES = {
    "自己",
    "你们",
    "我们",
    "他们",
    "知道",
    "时候",
    "发现",
    "已经",
    "之前",
    "以后",
    "现在",
    "怎么",
    "为什么",
    "一样",
    "东西",
    "还是",
    "一直",
    "直接",
    "其实",
    "而且",
    "不过",
    "或者",
    "而已",
    "希望",
    "就是",
    "没有",
    "一个",
    "什么",
    "可以",
    "但是",
    "因为",
    "所以",
    "如果",
    "然后",
    "觉得",
    "感觉",
    "真的",
    "非常",
    "一下",
    "有点",
    "不是",
    "这个",
    "那个",
    "出来",
    "完全",
    "每次",
    "一次",
    "一点",
    "能不能",
}


class SearchError(Exception):
    def __init__(self, message: str, code: int = 2) -> None:
        super().__init__(message)
        self.code = code


def now() -> str:
    return datetime.now(UTC).astimezone().isoformat(timespec="seconds")


def fetch(
    url: str, *, accept: str = "*/*", limit: int = 3_000_000
) -> tuple[bytes, dict[str, str], str]:
    headers = {"User-Agent": UA, "Accept": accept, "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8"}
    if url.startswith(GITHUB) and os.environ.get("GITHUB_TOKEN"):
        headers["Authorization"] = f"Bearer {os.environ['GITHUB_TOKEN']}"
    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=TIMEOUT) as resp:  # noqa: S310 - user supplied public URL
            return resp.read(limit), {k.lower(): v for k, v in resp.headers.items()}, resp.geturl()
    except urllib.error.HTTPError as exc:
        hint = (
            "（GitHub 未登录时每分钟限 10 次，可设置 GITHUB_TOKEN）"
            if exc.code == 403 and "github" in url
            else ""
        )
        raise SearchError(f"{url} 返回 HTTP {exc.code}{hint}") from None
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        raise SearchError(f"无法访问 {url}：{getattr(exc, 'reason', exc)}") from None


def fetch_json(url: str) -> Any:
    body, _, _ = fetch(url, accept="application/json")
    try:
        return json.loads(body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        raise SearchError(f"{url} 返回的不是 JSON") from None


# ------------------------------------------------------------------ App Store


def _app_row(a: dict[str, Any], detail: bool = False) -> dict[str, Any]:
    row = {
        "名称": a.get("trackName", ""),
        "开发者": a.get("sellerName") or a.get("artistName", ""),
        "评分": round(float(a.get("averageUserRating") or 0), 2),
        "评分人数": int(a.get("userRatingCount") or 0),
        "价格": a.get("formattedPrice") or ("免费" if not a.get("price") else a.get("price")),
        "分类": a.get("primaryGenreName", ""),
        "版本": a.get("version", ""),
        "最近更新": str(a.get("currentVersionReleaseDate", ""))[:10],
        "上架时间": str(a.get("releaseDate", ""))[:10],
        "大小MB": round(int(a.get("fileSizeBytes") or 0) / 1048576, 1),
        "应用ID": a.get("trackId", ""),
        "链接": a.get("trackViewUrl", "").split("?")[0],
    }
    if detail:
        row.update(
            {
                "更新说明": (a.get("releaseNotes") or "").strip(),
                "简介": (a.get("description") or "").strip()[:800],
                "最低系统": a.get("minimumOsVersion", ""),
                "年龄分级": a.get("contentAdvisoryRating", ""),
                "支持语言": "、".join(a.get("languageCodesISO2A") or []),
            }
        )
    return row


def appstore_search(term: str, country: str, limit: int) -> dict[str, Any]:
    q = urllib.parse.urlencode(
        {"term": term, "country": country, "entity": "software", "limit": limit}
    )
    url = f"{ITUNES}/search?{q}"
    data = fetch_json(url)
    rows = [_app_row(a) for a in data.get("results", []) if a.get("trackId")]
    return {"来源": url, "获取时间": now(), "结果": rows}


def appstore_lookup(ids: list[str], country: str) -> dict[str, Any]:
    if not all(i.isdigit() for i in ids):
        raise SearchError(
            "应用 ID 是 App Store 链接里 id 后面的数字，如 id414478124 → 414478124", 1
        )
    url = f"{ITUNES}/lookup?{urllib.parse.urlencode({'id': ','.join(ids), 'country': country})}"
    data = fetch_json(url)
    rows = [_app_row(a, detail=True) for a in data.get("results", []) if a.get("trackId")]
    if not rows:
        raise SearchError(f"没有查到应用 {','.join(ids)}（检查 ID 和 --country）", 1)
    return {"来源": url, "获取时间": now(), "结果": rows}


def _label(entry: dict[str, Any], key: str) -> str:
    value = entry.get(key)
    if isinstance(value, dict):
        return str(value.get("label", ""))
    return ""


def phrases(texts: list[str], top: int = 15) -> list[tuple[str, int]]:
    counter: Counter[str] = Counter()
    for t in texts:
        seen: set[str] = set()
        for run in re.findall(r"[\u4e00-\u9fff]{2,}", t):
            for n in (2, 3, 4, 5, 6):
                for i in range(len(run) - n + 1):
                    gram = run[i : i + n]
                    if gram[0] in STOP or gram[-1] in STOP or gram in STOP_PHRASES:
                        continue
                    seen.add(gram)
        for w in re.findall(r"[A-Za-z]{3,}", t.lower()):
            seen.add(w)
        counter.update(seen)  # count each phrase once per review
    picked: list[tuple[str, int]] = []
    for gram, n in counter.most_common(top * 4):
        if n < 2 or any(gram in g and n <= m for g, m in picked):
            continue
        picked = [(g, m) for g, m in picked if not (g in gram and m <= n)]
        picked.append((gram, n))
        if len(picked) >= top:
            break
    return picked


def reviews(app_id: str, country: str, pages: int) -> dict[str, Any]:
    if not app_id.isdigit():
        raise SearchError("应用 ID 必须是数字", 1)
    items: list[dict[str, Any]] = []
    sources = []
    for page in range(1, min(max(pages, 1), 10) + 1):
        url = (
            f"{ITUNES}/{country}/rss/customerreviews/page={page}/id={app_id}/sortby=mostrecent/json"
        )
        sources.append(url)
        entries = fetch_json(url).get("feed", {}).get("entry", [])
        if isinstance(entries, dict):
            entries = [entries]
        batch = []
        for e in entries:
            rating = _label(e, "im:rating")
            if not rating.isdigit():
                continue  # the first entry of old feeds describes the app itself
            batch.append(
                {
                    "评分": int(rating),
                    "标题": _label(e, "title"),
                    "内容": _label(e, "content"),
                    "版本": _label(e, "im:version"),
                    "时间": _label(e, "updated")[:10],
                    "作者": (e.get("author") or {}).get("name", {}).get("label", ""),
                }
            )
        if not batch:
            break
        items.extend(batch)
    if not items:
        raise SearchError(f"应用 {app_id} 在 {country} 区没有取到评论", 1)
    dist = Counter(r["评分"] for r in items)
    by_version: dict[str, list[int]] = {}
    for r in items:
        by_version.setdefault(r["版本"] or "未知", []).append(r["评分"])
    neg = [f"{r['标题']} {r['内容']}" for r in items if r["评分"] <= 2]
    pos = [f"{r['标题']} {r['内容']}" for r in items if r["评分"] >= 4]
    return {
        "来源": sources,
        "获取时间": now(),
        "评论数": len(items),
        "平均分": round(sum(r["评分"] for r in items) / len(items), 2),
        "评分分布": {f"{k}星": dist.get(k, 0) for k in range(5, 0, -1)},
        "各版本": {
            v: {"评论数": len(s), "平均分": round(sum(s) / len(s), 2)}
            for v, s in list(by_version.items())[:8]
        },
        "差评高频词": phrases(neg),
        "好评高频词": phrases(pos),
        "典型差评": [r for r in items if r["评分"] <= 2][:8],
        "评论": items,
    }


# --------------------------------------------------------------------- GitHub


def github_search(term: str, limit: int) -> dict[str, Any]:
    q = urllib.parse.urlencode(
        {"q": term, "sort": "stars", "order": "desc", "per_page": min(limit, 50)}
    )
    url = f"{GITHUB}/search/repositories?{q}"
    data = fetch_json(url)
    rows = []
    for r in data.get("items", [])[:limit]:
        rows.append(
            {
                "项目": r.get("full_name", ""),
                "简介": (r.get("description") or "")[:120],
                "星标": r.get("stargazers_count", 0),
                "Fork": r.get("forks_count", 0),
                "未关闭Issue": r.get("open_issues_count", 0),
                "最近提交": str(r.get("pushed_at", ""))[:10],
                "创建时间": str(r.get("created_at", ""))[:10],
                "语言": r.get("language") or "",
                "许可证": ((r.get("license") or {}).get("spdx_id") or ""),
                "话题": "、".join(r.get("topics") or []),
                "链接": r.get("html_url", ""),
            }
        )
    return {
        "来源": url,
        "获取时间": now(),
        "总数": data.get("total_count", len(rows)),
        "结果": rows,
    }


# ----------------------------------------------------------------------- site


class _PageParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.meta: dict[str, str] = {}
        self.links: list[tuple[str, str]] = []
        self._in_title = False
        self._in_a: str | None = None
        self._a_text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        a = {k.lower(): (v or "") for k, v in attrs}
        if tag == "title":
            self._in_title = True
        elif tag == "meta":
            key = (a.get("name") or a.get("property") or "").lower()
            if key and a.get("content"):
                self.meta.setdefault(key, a["content"].strip())
        elif tag == "a" and a.get("href"):
            self._in_a = a["href"]
            self._a_text = []

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False
        elif tag == "a" and self._in_a is not None:
            self.links.append((self._in_a, " ".join("".join(self._a_text).split())))
            self._in_a = None

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        if self._in_a is not None:
            self._a_text.append(data)


def _decode(body: bytes, headers: dict[str, str]) -> str:
    m = re.search(r"charset=([\w-]+)", headers.get("content-type", ""), re.I) or re.search(
        rb'<meta[^>]+charset=["\']?([\w-]+)', body[:4000], re.I
    )
    enc = (m.group(1) if isinstance(m.group(1), str) else m.group(1).decode()) if m else "utf-8"
    try:
        return body.decode(enc, errors="replace")
    except LookupError:
        return body.decode("utf-8", errors="replace")


def first_capture(url: str) -> dict[str, str] | None:
    host = urllib.parse.urlparse(url).netloc or url
    api = f"{WAYBACK}/cdx/search/cdx?{urllib.parse.urlencode({'url': host, 'limit': 1, 'output': 'json', 'fl': 'timestamp,original'})}"
    try:
        rows = fetch_json(api)
    except SearchError:
        return None
    if isinstance(rows, list) and len(rows) >= 2:
        ts = rows[1][0]
        return {
            "最早收录": f"{ts[:4]}-{ts[4:6]}-{ts[6:8]}",
            "说明": "域名最早被收录的时间，不等于产品上线时间（域名可能更早属于他人），需看快照内容确认",
            "快照": f"{WAYBACK}/web/{ts}/{rows[1][1]}",
            "来源": api,
        }
    return None


def site_info(url: str, history: bool) -> dict[str, Any]:
    if not re.match(r"^https?://", url):
        url = "https://" + url
    body, headers, final = fetch(url, accept="text/html,application/xhtml+xml")
    text = _decode(body, headers)
    p = _PageParser()
    p.feed(text)
    found: dict[str, list[str]] = {}
    base_host = urllib.parse.urlparse(final).netloc
    for href, label in p.links:
        if href.startswith(("javascript:", "mailto:", "#", "tel:")):
            continue
        absolute = urllib.parse.urljoin(final, href)
        hay = f"{label} {urllib.parse.urlparse(absolute).path}".lower()
        for name, pattern in KEY_PAGES.items():
            if re.search(pattern, hay, re.I) and absolute not in found.get(name, []):
                found.setdefault(name, []).append(absolute)
    tech = [name for name, pattern in TECH_MARKERS.items() if re.search(pattern, text, re.I)]
    for h in ("server", "x-powered-by"):
        if headers.get(h):
            tech.append(f"{h}: {headers[h]}")
    company = sorted({html.unescape(m.group(1)) for m in COMPANY.finditer(text)})
    apps = sorted(set(re.findall(r"apps\.apple\.com/[\w/-]*?id(\d{6,})", text)))
    gplay = sorted(set(re.findall(r"play\.google\.com/store/apps/details\?id=([\w.]+)", text)))
    info: dict[str, Any] = {
        "网址": final,
        "获取时间": now(),
        "标题": " ".join(p.title.split()),
        "描述": p.meta.get("description") or p.meta.get("og:description", ""),
        "关键词": p.meta.get("keywords", ""),
        "站点名": p.meta.get("og:site_name", ""),
        "关键页面": {k: v[:5] for k, v in found.items()},
        "ICP备案": sorted(set(ICP.findall(text))),
        "页脚主体": company,
        "技术线索": tech,
        "AppStore应用ID": apps,
        "GooglePlay包名": gplay,
        "外链域名数": len(
            {urllib.parse.urlparse(urllib.parse.urljoin(final, h)).netloc for h, _ in p.links}
            - {base_host}
        ),
    }
    if history:
        info["网页时光机"] = first_capture(final) or "未查到收录"
    return info


# --------------------------------------------------------------------- output


def write_xlsx(path: Path, sheets: dict[str, list[dict[str, Any]]], meta: dict[str, str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        if not rows:
            ws.append(["（无数据）"])
            continue
        cols = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(cols)
        for r in rows:
            ws.append(
                [
                    ("、".join(map(str, v)) if isinstance(v, list) else v)
                    for v in (r.get(c, "") for c in cols)
                ]
            )
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
        for j, c in enumerate(cols, 1):
            width = max([len(str(c))] + [min(len(str(r.get(c, ""))), 50) for r in rows])
            ws.column_dimensions[ws.cell(1, j).column_letter].width = max(8, width * 1.15)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    src = wb.create_sheet("来源")
    for k, v in meta.items():
        src.append([k, v])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def print_table(rows: list[dict[str, Any]], keys: list[str]) -> None:
    for i, r in enumerate(rows, 1):
        print(f"{i}. " + "，".join(f"{k} {r[k]}" for k in keys if r.get(k) not in ("", None)))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("appstore")
    s.add_argument("term")
    s.add_argument("--limit", type=int, default=10)
    a = sub.add_parser("app")
    a.add_argument("ids", nargs="+")
    r = sub.add_parser("reviews")
    r.add_argument("id")
    r.add_argument("--pages", type=int, default=3)
    g = sub.add_parser("github")
    g.add_argument("term")
    g.add_argument("--limit", type=int, default=10)
    w = sub.add_parser("site")
    w.add_argument("urls", nargs="+")
    w.add_argument("--history", action="store_true")
    for p in (s, a, r):
        p.add_argument("--country", default="cn", help="App Store 区域，如 cn、us、jp")
    for p in (s, a, r, g, w):
        p.add_argument("-o", "--output", help="保存为 Excel")
        p.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    try:
        if args.cmd == "appstore":
            res = appstore_search(args.term, args.country, max(1, min(args.limit, 50)))
            rows, sheet = res["结果"], "App Store 搜索"
            keys = [
                "名称",
                "开发者",
                "评分",
                "评分人数",
                "价格",
                "分类",
                "版本",
                "最近更新",
                "应用ID",
            ]
        elif args.cmd == "app":
            res = appstore_lookup(args.ids, args.country)
            rows, sheet = res["结果"], "应用详情"
            keys = [
                "名称",
                "开发者",
                "评分",
                "评分人数",
                "价格",
                "版本",
                "最近更新",
                "上架时间",
                "大小MB",
                "最低系统",
            ]
        elif args.cmd == "reviews":
            res = reviews(args.id, args.country, args.pages)
            rows, sheet, keys = res["评论"], "评论", []
        elif args.cmd == "github":
            res = github_search(args.term, max(1, args.limit))
            rows, sheet = res["结果"], "GitHub"
            keys = ["项目", "星标", "Fork", "最近提交", "语言", "许可证", "简介"]
        else:
            rows = [site_info(u, args.history) for u in args.urls]
            res, sheet, keys = {"获取时间": now(), "结果": rows}, "官网", []
    except SearchError as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return exc.code

    if args.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    elif args.cmd == "reviews":
        print(
            f"应用 {args.id}（{args.country} 区）最新 {res['评论数']} 条评论，平均 {res['平均分']} 分"
        )
        print("评分分布：" + "，".join(f"{k} {v}" for k, v in res["评分分布"].items()))
        print(
            "各版本："
            + "；".join(
                f"{v} 共 {d['评论数']} 条、{d['平均分']} 分" for v, d in res["各版本"].items()
            )
        )
        print(
            "差评高频词："
            + ("、".join(f"{p}({n})" for p, n in res["差评高频词"]) or "（差评太少）")
        )
        print(
            "好评高频词："
            + ("、".join(f"{p}({n})" for p, n in res["好评高频词"]) or "（好评太少）")
        )
        for item in res["典型差评"]:
            print(
                f"- [{item['评分']}星 {item['版本']} {item['时间']}] {item['标题']}：{item['内容'][:80]}"
            )
        print("说明：App Store 接口只返回最新的几百条评论，代表近期口碑，不是全部用户。")
    elif args.cmd == "site":
        for info in rows:
            print(f"{info['网址']}：{info['标题']}")
            for k in (
                "描述",
                "站点名",
                "ICP备案",
                "页脚主体",
                "技术线索",
                "AppStore应用ID",
                "GooglePlay包名",
                "网页时光机",
            ):
                v = info.get(k)
                if isinstance(v, dict):
                    v = f"最早收录 {v['最早收录']}（{v['说明']}）快照 {v['快照']}"
                if v:
                    print(f"  {k}：{'、'.join(v) if isinstance(v, list) else v}")
            for name, links in info["关键页面"].items():
                print(f"  {name}页：{' '.join(links[:3])}")
    else:
        print(f"共 {len(rows)} 条（来源：{res['来源']}，获取时间 {res['获取时间']}）")
        print_table(rows, keys)
    if args.output:
        sheets = {sheet: rows}
        if args.cmd == "reviews":
            sheets["高频词"] = [
                {"类型": t, "词": p, "出现条数": n}
                for t, key in (("差评", "差评高频词"), ("好评", "好评高频词"))
                for p, n in res[key]
            ]
        if args.cmd == "site":
            sheets = {
                sheet: [
                    {
                        k: (json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v)
                        for k, v in i.items()
                    }
                    for i in rows
                ]
            }
        src = res.get("来源", "")
        write_xlsx(
            Path(args.output),
            sheets,
            {
                "来源": "\n".join(src) if isinstance(src, list) else str(src),
                "获取时间": res["获取时间"],
            },
        )
        if not args.json:
            print(f"已保存：{args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
