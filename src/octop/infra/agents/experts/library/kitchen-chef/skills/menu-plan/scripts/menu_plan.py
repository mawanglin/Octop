#!/usr/bin/env python3
"""一周菜单与买菜清单：从菜谱库排出荤素搭配、不重复的菜单，按人数合并食材，生成分类买菜清单。

Usage:
    python scripts/menu_plan.py plan --days 7 --people 3 [--meals lunch,dinner] [--avoid 辣,海鲜] [--soup] [--seed 1] [-o 菜单.xlsx]
    python scripts/menu_plan.py shopping 菜单.md --people 3 [-o 买菜清单.xlsx]
    python scripts/menu_plan.py list [--tag 快手]

菜谱库默认读取 references/recipes.csv（列：菜名,类型,标签,食材（每人份）,做法要点；食材写 名称:数量:单位，用 ; 分隔），
可用 --recipes 指定自己的菜谱。每餐 1 荤 1 素，--soup 或 4 人以上加 1 个汤；一周内同一道菜不重复（菜谱不够时才重复），
相邻两餐不用同一种主料。--avoid 按标签或食材排除（如 辣、海鲜、花生、羊肉）。
shopping：从 Markdown（列表或表格里出现的菜名）汇总食材；不在菜谱库里的菜会单独列出。
盐、糖、油、酱油等常备调料不列入买菜清单。Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl（仅 -o）。
"""

from __future__ import annotations

import argparse
import csv
import random
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

DEFAULT_RECIPES = Path(__file__).resolve().parent.parent / "references" / "recipes.csv"
MEALS = {"breakfast": "早餐", "lunch": "午餐", "dinner": "晚餐"}
PANTRY = {"盐", "糖", "冰糖", "油", "食用油", "酱油", "生抽", "老抽", "醋", "香醋", "料酒", "淀粉", "蚝油", "八角", "花椒",
          "干辣椒", "郫县豆瓣", "胡椒粉", "鸡精", "香油"}  # fmt: skip
CATEGORY_RULES = [
    (
        "蛋奶豆制品",
        ("蛋", "奶", "豆腐", "豆干", "腐竹"),
    ),  # before 肉禽 so 鸡蛋 isn't counted as meat
    ("水产", ("鱼", "虾", "蟹", "贝", "鱿", "海带", "紫菜")),
    ("肉禽", ("肉", "鸡", "鸭", "排骨", "牛", "羊", "猪", "五花", "里脊", "翅")),
    ("主食干货", ("米", "面", "粉", "花生", "木耳", "香菇", "玉米")),
    ("其他", ("可乐", "酒", "饮料")),
]


class MenuError(ValueError):
    pass


@dataclass
class Recipe:
    name: str
    kind: str
    tags: set[str]
    items: list[tuple[str, float, str]]
    note: str

    @property
    def main(self) -> str:
        return self.items[0][0] if self.items else ""


def category(name: str) -> str:
    for cat, keys in CATEGORY_RULES:
        if any(k in name for k in keys):
            return cat
    return "蔬菜水果"


def load_recipes(path: Path) -> dict[str, Recipe]:
    if not path.is_file():
        raise MenuError(f"找不到菜谱库 {path}")
    out: dict[str, Recipe] = {}
    with path.open(encoding="utf-8-sig", newline="") as f:
        for i, row in enumerate(csv.DictReader(f), start=2):
            name = (row.get("菜名") or "").strip()
            if not name:
                continue
            items = []
            for part in (row.get("食材（每人份）") or "").split(";"):
                bits = [b.strip() for b in part.split(":")]
                if len(bits) != 3:
                    raise MenuError(f"菜谱第 {i} 行「{name}」的食材「{part}」应写成 名称:数量:单位")
                items.append((bits[0], float(bits[1]), bits[2]))
            tags = {t.strip() for t in re.split(r"[;；,，]", row.get("标签") or "") if t.strip()}
            out[name] = Recipe(
                name,
                (row.get("类型") or "").strip(),
                tags,
                items,
                (row.get("做法要点") or "").strip(),
            )
    if not out:
        raise MenuError(f"{path} 里没有菜谱")
    return out


def allowed(r: Recipe, avoid: set[str]) -> bool:
    return not (r.tags & avoid) and not any(a in name for a in avoid for name, _q, _u in r.items)


def pick(pool: list[Recipe], used: set[str], last_mains: set[str], rng: random.Random) -> Recipe:
    fresh = [r for r in pool if r.name not in used and r.main not in last_mains]
    if not fresh:
        fresh = (
            [r for r in pool if r.name not in used]
            or [r for r in pool if r.main not in last_mains]
            or pool
        )
    choice = rng.choice(fresh)
    used.add(choice.name)
    return choice


def plan(
    recipes: dict[str, Recipe],
    days: int,
    meals: list[str],
    avoid: set[str],
    soup: bool,
    seed: int | None,
) -> list[tuple[int, str, list[Recipe]]]:
    rng = random.Random(seed)
    pools = {
        k: [r for r in recipes.values() if r.kind == k and allowed(r, avoid)]
        for k in ("荤菜", "素菜", "汤")
    }
    if not pools["荤菜"] or not pools["素菜"]:
        raise MenuError("排除忌口后荤菜或素菜不够，请减少 --avoid 或补充菜谱")
    used: set[str] = set()
    out = []
    last: set[str] = set()
    for d in range(days):
        for meal in meals:
            dishes = [pick(pools["荤菜"], used, last, rng), pick(pools["素菜"], used, last, rng)]
            if soup and pools["汤"]:
                dishes.append(pick(pools["汤"], used, last, rng))
            last = {r.main for r in dishes}
            out.append((d, meal, dishes))
    return out


def merge(dishes: list[Recipe], people: int) -> dict[str, dict[tuple[str, str], float]]:
    total: dict[str, dict[tuple[str, str], float]] = defaultdict(lambda: defaultdict(float))
    for r in dishes:
        for name, qty, unit in r.items:
            if name in PANTRY:
                continue
            total[category(name)][(name, unit)] += qty * people
    return total


def fmt_qty(q: float, unit: str) -> str:
    if unit == "克":
        return f"{q / 500:.1f} 斤（{q:.0f} 克）" if q >= 500 else f"{q:.0f} 克"
    if unit in ("个", "根"):
        return f"{-(-q // 1):.0f} {unit}"
    return f"{q:g} {unit}"


def print_shopping(total: dict[str, dict[tuple[str, str], float]]) -> None:
    order = ["蔬菜水果", "肉禽", "水产", "蛋奶豆制品", "主食干货", "其他"]
    print("\n## 买菜清单")
    for cat in order:
        if cat in total:
            print(
                f"**{cat}**："
                + "、".join(
                    f"{n} {fmt_qty(q, u)}"
                    for (n, u), q in sorted(total[cat].items(), key=lambda kv: -kv[1])
                )
            )


def write_xlsx(
    out: Path, menu: list[tuple[str, str, str]], total: dict[str, dict[tuple[str, str], float]]
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
        sys.exit(2)
    wb = Workbook()
    ws = wb.active
    ws.title = "菜单"
    ws.append(["日期", "餐次", "菜品"])
    for row in menu:
        ws.append(list(row))
    sh = wb.create_sheet("买菜清单")
    sh.append(["类别", "食材", "数量", "单位", "已买"])
    for cat, items in total.items():
        for (name, unit), q in sorted(items.items()):
            sh.append([cat, name, round(q, 1), unit, "☐"])
    for sheet in (ws, sh):
        for c in sheet[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="C2410C")
        sheet.column_dimensions["A"].width = 12
        sheet.column_dimensions["B"].width = 14
        sheet.column_dimensions["C"].width = 36 if sheet is ws else 10
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def cmd_plan(a: argparse.Namespace) -> int:
    recipes = load_recipes(Path(a.recipes))
    meals = [m.strip() for m in a.meals.split(",") if m.strip()]
    bad = [m for m in meals if m not in MEALS]
    if bad or not 1 <= a.days <= 14 or not 1 <= a.people <= 20:
        raise MenuError("--meals 可选 breakfast,lunch,dinner；--days 1–14；--people 1–20")
    avoid = {x.strip() for x in re.split(r"[,，、]", a.avoid or "") if x.strip()}
    soup = a.soup or a.people >= 4
    menu = plan(recipes, a.days, meals, avoid, soup, a.seed)
    rows = []
    print(f"## {a.days} 天菜单（{a.people} 人，每餐 1 荤 1 素{' 1 汤' if soup else ''}）")
    for d, meal, dishes in menu:
        names = "、".join(r.name for r in dishes)
        print(f"- 第 {d + 1} 天 {MEALS[meal]}：{names}")
        rows.append((f"第 {d + 1} 天", MEALS[meal], names))
    total = merge([r for _d, _m, ds in menu for r in ds], a.people)
    print_shopping(total)
    print(
        "\n常备调料（盐、油、酱油、糖、醋、料酒、淀粉等）请自行检查是否够用；生鲜建议分两次购买，前三天和后四天。"
    )
    if a.output:
        write_xlsx(Path(a.output), rows, total)
        print(f"已生成 {a.output}")
    return 0


def cmd_shopping(a: argparse.Namespace) -> int:
    recipes = load_recipes(Path(a.recipes))
    text = Path(a.input).read_text(encoding="utf-8")
    found: list[Recipe] = []
    for name in sorted(
        recipes, key=len, reverse=True
    ):  # longer names first: 番茄炒菜花 before 炒菜花
        count = text.count(name)
        if count:
            found.extend([recipes[name]] * count)
            text = text.replace(name, " ")
    unknown = sorted(
        {
            w
            for w in re.findall(r"[\u4e00-\u9fff]{2,8}", text)
            if w.endswith(("汤", "炒", "肉", "鱼", "虾", "菜", "蛋", "豆腐"))
        }
    )
    if not found:
        raise MenuError("菜单里没有找到菜谱库中的菜名，用 list 查看可用的菜")
    total = merge(found, a.people)
    print(f"共 {len(found)} 道菜，{a.people} 人份")
    print_shopping(total)
    if unknown:
        print(f"\n不在菜谱库中的菜（食材需要手动补充）：{'、'.join(unknown)}")
    if a.output:
        write_xlsx(Path(a.output), [("", "", r.name) for r in found], total)
        print(f"已生成 {a.output}")
    return 0


def cmd_list(a: argparse.Namespace) -> int:
    recipes = load_recipes(Path(a.recipes))
    for kind in ("荤菜", "素菜", "汤", "主食"):
        items = [r for r in recipes.values() if r.kind == kind and (not a.tag or a.tag in r.tags)]
        if items:
            print(f"{kind}（{len(items)}）：" + "、".join(r.name for r in items))
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--recipes", default=str(DEFAULT_RECIPES))
    sub = parser.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("plan")
    p.add_argument("--days", type=int, default=7)
    p.add_argument("--people", type=int, default=3)
    p.add_argument("--meals", default="lunch,dinner")
    p.add_argument("--avoid", default="")
    p.add_argument("--soup", action="store_true")
    p.add_argument("--seed", type=int)
    p.add_argument("-o", "--output")
    s = sub.add_parser("shopping")
    s.add_argument("input")
    s.add_argument("--people", type=int, default=3)
    s.add_argument("-o", "--output")
    lst = sub.add_parser("list")
    lst.add_argument("--tag")
    args = parser.parse_args(argv)
    try:
        return {"plan": cmd_plan, "shopping": cmd_shopping, "list": cmd_list}[args.cmd](args)
    except (OSError, MenuError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
