#!/usr/bin/env python3
"""宠物疫苗和驱虫日程：按出生日期（幼宠）或上次接种 / 驱虫日期（成年宠物）排出日程表，列出近期要做的事，导出日历提醒。

Usage:
    python scripts/pet_care.py plan --species dog --name 旺财 --birth 2026-07-01 [-o 日程.xlsx] [--ics 提醒.ics]
    python scripts/pet_care.py plan --species cat --name 咪咪 --birth 2023-05-01 --last-core 2026-03-10 --last-rabies 2026-03-10 \
                                    --last-deworm-in 2026-09-01 --last-deworm-out 2026-10-01 [--months 12]

日程依据常见的兽医推荐（各地区、各医院和疫苗品牌会不同，以兽医建议为准）：
  - 核心疫苗：犬多联苗（犬瘟热、细小、腺病毒、副流感等）/ 猫三联（猫瘟、鼻气管炎、杯状病毒），
    8、12、16 周龄各一针，首免完成 1 年后加强，之后每 1–3 年一次（本工具按每年提醒）；
  - 狂犬疫苗：16 周龄（约 4 月龄）首针，之后每年一次（多数地区犬只登记要求每年接种）；
  - 体内驱虫：2–12 周龄每 2 周，3–6 月龄每月，之后每 3 个月；
  - 体外驱虫：8 周龄起每月（按所用药品说明调整）。
成年宠物没有提供上次日期的项目，会提示尽快咨询兽医补种或驱虫。
Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl（仅 -o）。
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
import uuid
from pathlib import Path

SPECIES = {"dog": ("狗", "犬多联疫苗（核心苗）"), "cat": ("猫", "猫三联疫苗（核心苗）")}


class PetError(ValueError):
    pass


def weeks(birth: dt.date, w: float) -> dt.date:
    return birth + dt.timedelta(days=round(w * 7))


def periodic(start: dt.date, every_days: int, until: dt.date) -> list[dt.date]:
    out = []
    d = start + dt.timedelta(days=every_days)
    while d <= until:
        out.append(d)
        d += dt.timedelta(days=every_days)
    return out


def schedule(
    species: str,
    birth: dt.date | None,
    last: dict[str, dt.date | None],
    today: dt.date,
    months: int,
) -> list[tuple[dt.date, str, str]]:
    core_name = SPECIES[species][1]
    until = today + dt.timedelta(days=round(months * 30.4))
    young = birth is not None and (today - birth).days < 365
    ev: list[tuple[dt.date, str, str]] = []

    # core vaccines
    if last["core"]:
        ev += [
            (d, core_name, "年度加强（部分疫苗可 3 年一次，听兽医建议）")
            for d in periodic(last["core"], 365, until)
        ]
    elif young and birth:
        shots = [weeks(birth, w) for w in (8, 12, 16)]
        ev += [
            (d, core_name, f"首免第 {i} 针（{w} 周龄）")
            for i, (d, w) in enumerate(zip(shots, (8, 12, 16), strict=True), 1)
        ]
        booster = shots[-1] + dt.timedelta(days=365)
        ev.append((booster, core_name, "首免完成 1 年后加强"))
        ev += [(d, core_name, "年度加强") for d in periodic(booster, 365, until)]
    else:
        ev.append((today, core_name, "没有上次接种记录：请尽快咨询兽医是否需要补种"))

    # rabies
    if last["rabies"]:
        ev += [(d, "狂犬疫苗", "每年一次") for d in periodic(last["rabies"], 365, until)]
    elif young and birth:
        first = weeks(birth, 16)
        ev.append((first, "狂犬疫苗", "首针（约 4 月龄）"))
        ev += [(d, "狂犬疫苗", "每年一次") for d in periodic(first, 365, until)]
    else:
        ev.append((today, "狂犬疫苗", "没有上次接种记录：请尽快咨询兽医补种"))

    # internal deworming
    if last["in"]:
        ev += [(d, "体内驱虫", "每 3 个月") for d in periodic(last["in"], 91, until)]
    elif young and birth:
        pts = [weeks(birth, w) for w in (2, 4, 6, 8, 10, 12)] + [
            weeks(birth, w) for w in (16, 20, 24)
        ]
        ev += [(d, "体内驱虫", "幼宠期（12 周龄前每 2 周，6 月龄前每月）") for d in pts]
        ev += [(d, "体内驱虫", "每 3 个月") for d in periodic(pts[-1], 91, until)]
    else:
        ev.append((today, "体内驱虫", "没有上次记录：建议尽快驱虫，之后每 3 个月"))

    # external deworming
    if last["out"]:
        ev += [(d, "体外驱虫", "每月（按药品说明）") for d in periodic(last["out"], 30, until)]
    elif young and birth:
        first = weeks(birth, 8)
        ev.append((first, "体外驱虫", "8 周龄起"))
        ev += [(d, "体外驱虫", "每月（按药品说明）") for d in periodic(first, 30, until)]
    else:
        ev.append((today, "体外驱虫", "没有上次记录：建议尽快驱虫，之后每月"))
    return sorted((e for e in ev if e[0] <= until), key=lambda e: (e[0], e[1]))


def write_xlsx(rows: list[tuple[dt.date, str, str]], name: str, today: dt.date, out: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
        sys.exit(2)
    wb = Workbook()
    ws = wb.active
    ws.title = "日程"
    ws.append(["日期", "宠物", "项目", "说明", "状态", "实际完成日期", "医院 / 药品"])
    past = PatternFill("solid", fgColor="F3F4F6")
    for d, item, note in rows:
        ws.append([d, name, item, note, "已过（请确认是否完成）" if d < today else "待办", "", ""])
        ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"
        if d < today:
            for c in ws[ws.max_row]:
                c.fill = past
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="A16207")
    for col, w in zip("ABCDEFG", (12, 8, 20, 40, 18, 14, 18), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def write_ics(rows: list[tuple[dt.date, str, str]], name: str, today: dt.date, out: Path) -> int:
    stamp = dt.datetime.now(dt.UTC).strftime("%Y%m%dT%H%M%SZ")
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//Octop//pet-care//CN",
        "CALSCALE:GREGORIAN",
    ]
    n = 0
    for d, item, note in rows:
        if d < today:
            continue
        uid = uuid.uuid5(uuid.NAMESPACE_URL, f"octop-pet-care/{name}/{item}/{d}")
        lines += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{stamp}", f"DTSTART;VALUE=DATE:{d:%Y%m%d}",
                  f"DTEND;VALUE=DATE:{d + dt.timedelta(days=1):%Y%m%d}", f"SUMMARY:{name} {item}", f"DESCRIPTION:{note}",
                  "BEGIN:VALARM", "ACTION:DISPLAY", f"DESCRIPTION:{name} {item}", "TRIGGER:-P1D", "END:VALARM", "END:VEVENT"]  # fmt: skip
        n += 1
    lines.append("END:VCALENDAR")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")
    return n


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("plan")
    p.add_argument("--species", choices=sorted(SPECIES), required=True)
    p.add_argument("--name", default="")
    p.add_argument("--birth")
    for key in ("core", "rabies", "deworm-in", "deworm-out"):
        p.add_argument(f"--last-{key}")
    p.add_argument("--months", type=int, default=12)
    p.add_argument("--days", type=int, default=30, help="终端列出多少天内的待办")
    p.add_argument("--today")
    p.add_argument("-o", "--output")
    p.add_argument("--ics")
    args = parser.parse_args(argv)
    try:
        today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
        birth = dt.date.fromisoformat(args.birth) if args.birth else None
        if birth and birth > today:
            raise PetError("出生日期不能晚于今天")
        last = {k: dt.date.fromisoformat(v) if v else None for k, v in
                (("core", args.last_core), ("rabies", args.last_rabies), ("in", args.last_deworm_in), ("out", args.last_deworm_out))}  # fmt: skip
        if not birth and not any(last.values()):
            raise PetError("请提供 --birth（幼宠）或上次接种 / 驱虫日期（成年宠物）")
        if not 1 <= args.months <= 36:
            raise PetError("--months 为 1–36")
    except (PetError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    label = args.name or SPECIES[args.species][0]
    rows = schedule(args.species, birth, last, today, args.months)
    age = (
        f"，{(today - birth).days // 7} 周龄"
        if birth and (today - birth).days < 365
        else (f"，约 {(today - birth).days / 365:.1f} 岁" if birth else "")
    )
    print(
        f"{label}（{SPECIES[args.species][0]}{age}）未来 {args.months} 个月共 {sum(1 for r in rows if r[0] >= today)} 项"
    )
    soon = [r for r in rows if today <= r[0] <= today + dt.timedelta(days=args.days)]
    print(f"## {args.days} 天内")
    for d, item, note in soon or [(today, "暂无", "")]:
        print(f"- {d}：{item}" + (f"（{note}）" if note else ""))
    missed = [r for r in rows if r[0] < today]
    if missed:
        print(
            f"另有 {len(missed)} 项日期已过（如 {missed[-1][0]} {missed[-1][1]}），请确认是否已经完成；未完成的请咨询兽医如何补上。"
        )
    print("提示：疫苗种类和间隔以兽医建议为准；接种前后一周避免洗澡和剧烈运动，留观 30 分钟。")
    if args.output:
        write_xlsx(rows, label, today, Path(args.output))
        print(f"已生成 {args.output}")
    if args.ics:
        n = write_ics(rows, label, today, Path(args.ics))
        print(f"已生成 {args.ics}：{n} 个提醒（提前 1 天）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
