#!/usr/bin/env python3
"""家庭财税计算：房贷（等额本息 / 等额本金 / 提前还款）、工资个税（累计预扣法 + 专项附加扣除 + 年终奖）、五险一金。

Usage:
    python scripts/finance_calc.py mortgage --principal 1000000 --rate 3.1 --years 30 [--prepay 36:200000:shorten] [-o 还款计划.xlsx]
    python scripts/finance_calc.py tax --salary 20000 [--base 15000 | --insurance 2200] [--children 1] [--elderly 3000]
                                   [--loan | --rent 1500] [--edu] [--cert] [--infants 0] [--bonus 60000] [-o 个税.xlsx]
    python scripts/finance_calc.py insurance --base 15000 [--fund 7] [--min 7384 --max 36921]

mortgage：同时计算等额本息和等额本金；--prepay 第几期后:金额:shorten（缩短年限）/ reduce（减少月供），可多次指定。
tax：综合所得按月累计预扣（基本减除 5000/月），专项附加扣除按 2023 年起标准：子女教育、3 岁以下婴幼儿照护每孩 2000/月，
     赡养老人独生子女 3000/月（非独生填分摊额，最多 1500），住房贷款利息 1000/月，住房租金 1500/1100/800（与房贷利息二选一），
     继续教育学历 400/月、职业资格当年 3600；年终奖比较“单独计税”（政策执行至 2027 年底）和“并入综合所得”两种方式。
insurance：个人缴费默认养老 8%、医疗 2%、失业 0.5%、公积金 --fund%；缴费基数按 --min/--max 封顶保底（各地每年公布）。
所有结果为估算，以当地社保、公积金和税务部门的规定为准。Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl（仅 -o）。
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path

ANNUAL_BRACKETS = [(36000, 0.03, 0), (144000, 0.10, 2520), (300000, 0.20, 16920), (420000, 0.25, 31920),
                   (660000, 0.30, 52920), (960000, 0.35, 85920), (float("inf"), 0.45, 181920)]  # fmt: skip
MONTHLY_BRACKETS = [(3000, 0.03, 0), (12000, 0.10, 210), (25000, 0.20, 1410), (35000, 0.25, 2660),
                    (55000, 0.30, 4410), (80000, 0.35, 7160), (float("inf"), 0.45, 15160)]  # fmt: skip


class CalcError(ValueError):
    pass


def yuan(x: float) -> str:
    return f"{x:,.2f}"


# --------------------------------------------------------------- mortgage


@dataclass
class Row:
    period: int
    payment: float
    principal: float
    interest: float
    balance: float
    extra: float = 0.0


def schedule(principal: float, annual_rate: float, months: int, method: str,
             prepay: dict[int, tuple[float, str]] | None = None) -> list[Row]:  # fmt: skip
    r = annual_rate / 100 / 12
    rows: list[Row] = []
    balance = principal
    remaining = months
    payment = _annuity(balance, r, remaining) if method == "equal-payment" else 0.0
    base_principal = balance / remaining
    period = 0
    while balance > 0.005 and period < months * 2:
        period += 1
        interest = balance * r
        if method == "equal-payment":
            pay = min(payment, balance + interest)
            prin = pay - interest
        else:
            prin = min(base_principal, balance)
            pay = prin + interest
        balance -= prin
        remaining -= 1
        extra = 0.0
        if prepay and period in prepay and balance > 0:
            amount, mode = prepay[period]
            extra = min(amount, balance)
            balance -= extra
            if balance > 0.005:
                if mode == "reduce":
                    if method == "equal-payment":
                        payment = _annuity(balance, r, remaining)
                    else:
                        base_principal = balance / remaining
                elif method == "equal-payment":
                    remaining = _months_for(balance, r, payment)
                else:
                    remaining = max(1, round(balance / base_principal + 0.4999))
        rows.append(Row(period, pay, prin, interest, max(balance, 0.0), extra))
        if remaining <= 0 and balance > 0.005:  # rounding leftovers go into the last payment
            rows[-1].principal += balance
            rows[-1].payment += balance
            rows[-1].balance = 0.0
            balance = 0
    return rows


def _annuity(p: float, r: float, n: int) -> float:
    if n <= 0:
        return p
    if r == 0:
        return p / n
    f = (1 + r) ** n
    return p * r * f / (f - 1)


def _months_for(balance: float, r: float, payment: float) -> int:
    n = 0
    b = balance
    while b > 0.005 and n < 1200:
        b = b * (1 + r) - payment
        n += 1
    return n


def parse_prepay(specs: list[str]) -> dict[int, tuple[float, str]]:
    out = {}
    for s in specs:
        parts = s.split(":")
        if len(parts) != 3 or parts[2] not in ("shorten", "reduce"):
            raise CalcError(
                f"--prepay「{s}」格式应为 期数:金额:shorten|reduce，如 36:200000:shorten"
            )
        out[int(parts[0])] = (float(parts[1]), parts[2])
    return out


def cmd_mortgage(
    principal: float, rate: float, years: int, prepay_specs: list[str], out: Path | None
) -> int:
    if principal <= 0 or not 0 <= rate < 30 or not 1 <= years <= 40:
        raise CalcError("贷款金额需大于 0，利率 0–30%，年限 1–40 年")
    months = years * 12
    prepay = parse_prepay(prepay_specs)
    results = {}
    print(f"贷款 {yuan(principal)} 元，年利率 {rate}%，{years} 年（{months} 期）")
    for method, label in (("equal-payment", "等额本息"), ("equal-principal", "等额本金")):
        base = schedule(principal, rate, months, method)
        interest = sum(r.interest for r in base)
        first = base[0].payment
        line = f"- {label}：首月 {yuan(first)} 元"
        line += (
            "（每月相同）"
            if method == "equal-payment"
            else f"，末月 {yuan(base[-1].payment)} 元，每月递减约 {yuan(base[0].payment - base[1].payment)} 元"
        )
        print(line + f"，总利息 {yuan(interest)} 元，还款总额 {yuan(interest + principal)} 元")
        results[label] = base
        if prepay:
            pre = schedule(principal, rate, months, method, prepay)
            pi = sum(r.interest for r in pre)
            last_pay = next((r.payment for r in pre if r.period > max(prepay)), pre[-1].payment)
            print(f"  提前还款后：共 {len(pre)} 期（少还 {len(base) - len(pre)} 期），总利息 {yuan(pi)} 元，节省利息 {yuan(interest - pi)} 元，"
                  f"提前还款后月供约 {yuan(last_pay)} 元")  # fmt: skip
            results[label + "（提前还款）"] = pre
    diff = sum(r.interest for r in results["等额本息"]) - sum(
        r.interest for r in results["等额本金"]
    )
    print(
        f"等额本金比等额本息少付利息 {yuan(diff)} 元，但前期月供更高；选择时先看每月现金流能否承受。"
    )
    if prepay:
        print(
            "提示：提前还款前确认银行是否收违约金、需要提前多久预约；缩短年限比减少月供省利息更多。"
        )
    if out:
        write_schedules(results, out)
        print(f"已生成 {out}")
    return 0


def write_schedules(results: dict[str, list[Row]], out: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
        sys.exit(2)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in results.items():
        ws = wb.create_sheet(name[:31])
        ws.append(["期数", "月供", "本金", "利息", "提前还款", "剩余本金"])
        for r in rows:
            ws.append([r.period, round(r.payment, 2), round(r.principal, 2), round(r.interest, 2), round(r.extra, 2) or None,
                       round(r.balance, 2)])  # fmt: skip
        n = ws.max_row
        ws.append(
            ["合计", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", ""]
        )
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F766E")
        for col in "BCDEF":
            ws.column_dimensions[col].width = 14
            for cell in ws[col][1:]:
                cell.number_format = "#,##0.00"
        ws.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


# -------------------------------------------------------------------- tax


def annual_tax(taxable: float) -> float:
    if taxable <= 0:
        return 0.0
    for limit, rate, quick in ANNUAL_BRACKETS:
        if taxable <= limit:
            return taxable * rate - quick
    return 0.0  # pragma: no cover


def bonus_tax(bonus: float) -> float:
    if bonus <= 0:
        return 0.0
    monthly = bonus / 12
    for limit, rate, quick in MONTHLY_BRACKETS:
        if monthly <= limit:
            return bonus * rate - quick
    return 0.0  # pragma: no cover


def insurance_amounts(
    base: float, lo: float | None, hi: float | None, fund: float
) -> dict[str, float]:
    b = base
    if lo is not None:
        b = max(b, lo)
    if hi is not None:
        b = min(b, hi)
    return {"缴费基数": b, "养老保险 8%": b * 0.08, "医疗保险 2%": b * 0.02, "失业保险 0.5%": b * 0.005,
            f"住房公积金 {fund:g}%": b * fund / 100}  # fmt: skip


def special_deductions(a: argparse.Namespace) -> dict[str, float]:
    d: dict[str, float] = {}
    if a.children:
        d[f"子女教育（{a.children} 孩）"] = 2000 * a.children
    if a.infants:
        d[f"3 岁以下婴幼儿照护（{a.infants} 孩）"] = 2000 * a.infants
    if a.elderly:
        if a.elderly > 3000:
            raise CalcError("赡养老人每月最多 3000 元（非独生子女每人最多分摊 1500 元）")
        d["赡养老人"] = a.elderly
    if a.loan and a.rent:
        raise CalcError("住房贷款利息和住房租金只能二选一")
    if a.loan:
        d["住房贷款利息"] = 1000
    if a.rent:
        if a.rent not in (1500, 1100, 800):
            raise CalcError(
                "住房租金扣除标准为 1500（直辖市、省会等）/ 1100（市辖区户籍人口超过 100 万）/ 800"
            )
        d["住房租金"] = a.rent
    if a.edu:
        d["继续教育（学历）"] = 400
    return d


def cmd_tax(a: argparse.Namespace) -> int:
    months = a.months
    if a.salary <= 0 or not 1 <= months <= 12:
        raise CalcError("--salary 需大于 0，--months 为 1–12")
    if a.base is not None:
        ins = insurance_amounts(a.base, a.min, a.max, a.fund)
        insurance = sum(v for k, v in ins.items() if k != "缴费基数")
    else:
        insurance = a.insurance or 0.0
    deductions = special_deductions(a)
    monthly_special = sum(deductions.values())
    cert = 3600.0 if a.cert else 0.0
    rows = []
    cum_tax = 0.0
    for m in range(1, months + 1):
        cum_income = a.salary * m
        cum_deduct = (
            5000 * m
            + insurance * m
            + monthly_special * m
            + (cert if m == months else 0)
            + a.other * m
        )
        taxable = max(0.0, cum_income - cum_deduct)
        due = max(0.0, annual_tax(taxable) - cum_tax)
        cum_tax += due
        rows.append(
            (m, a.salary, insurance, monthly_special, taxable, due, a.salary - insurance - due)
        )
    print(f"月薪 {yuan(a.salary)} 元，个人社保公积金 {yuan(insurance)} 元/月，专项附加扣除 {yuan(monthly_special)} 元/月"
          + (f"（{'、'.join(f'{k} {v:g}' for k, v in deductions.items())}）" if deductions else ""))  # fmt: skip
    print("| 月份 | 当月预扣个税 | 税后到手 |\n|------|------|------|")
    for m, _s, _i, _d, _t, due, net in rows:
        print(f"| {m} 月 | {yuan(due)} | {yuan(net)} |")
    print(f"全年工资个税 {yuan(cum_tax)} 元，税后工资合计 {yuan(sum(r[6] for r in rows))} 元"
          + ("（职业资格继续教育 3600 元在最后一个月扣除，实际可在年度汇算时扣除）" if cert else ""))  # fmt: skip
    if rows[0][5] < rows[-1][5]:
        print(
            "说明：累计预扣法下，收入累计越多适用税率越高，所以下半年每月扣税可能比上半年多，这是正常的。"
        )
    if a.bonus:
        sep = bonus_tax(a.bonus)
        annual_taxable = max(
            0.0, a.salary * months - (5000 + insurance + monthly_special + a.other) * months - cert
        )
        combined = annual_tax(annual_taxable + a.bonus) - annual_tax(annual_taxable)
        better = "单独计税" if sep <= combined else "并入综合所得"
        print(
            f"年终奖 {yuan(a.bonus)} 元：单独计税 {yuan(sep)} 元，并入综合所得增加 {yuan(combined)} 元 → 选择「{better}」少缴 {yuan(abs(sep - combined))} 元"
        )
        for (limit, _r, _q), (_l, rate, _q2) in zip(
            MONTHLY_BRACKETS, MONTHLY_BRACKETS[1:], strict=False
        ):
            edge = limit * 12
            if edge < a.bonus and bonus_tax(a.bonus) - bonus_tax(edge) > a.bonus - edge:
                print(
                    f"提示：年终奖略高于 {yuan(edge)} 元的临界点，整笔奖金适用税率跳到 {rate:.0%}，多发的部分全部被税吃掉还不止，可与公司沟通调整"
                )
    print("提示：以上为估算；实际以个税 App 的预扣和年度汇算为准，多退少补。")
    if a.output:
        write_tax(rows, Path(a.output))
        print(f"已生成 {a.output}")
    return 0


def write_tax(rows: list[tuple[int, float, float, float, float, float, float]], out: Path) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
        sys.exit(2)
    wb = Workbook()
    ws = wb.active
    ws.title = "个税预扣"
    ws.append(
        [
            "月份",
            "应发工资",
            "社保公积金",
            "专项附加扣除",
            "累计应纳税所得额",
            "当月预扣个税",
            "税后到手",
        ]
    )
    for r in rows:
        ws.append([r[0], *[round(x, 2) for x in r[1:]]])
    n = ws.max_row
    ws.append(
        [
            "合计",
            f"=SUM(B2:B{n})",
            f"=SUM(C2:C{n})",
            f"=SUM(D2:D{n})",
            "",
            f"=SUM(F2:F{n})",
            f"=SUM(G2:G{n})",
        ]
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def cmd_insurance(base: float, lo: float | None, hi: float | None, fund: float) -> int:
    if base <= 0 or not 5 <= fund <= 12:
        raise CalcError("缴费基数需大于 0；公积金比例一般为 5%–12%")
    ins = insurance_amounts(base, lo, hi, fund)
    b = ins.pop("缴费基数")
    note = (
        ""
        if b == base
        else f"（按当地基数{'下限' if b > base else '上限'}调整，原工资 {yuan(base)}）"
    )
    print(f"缴费基数 {yuan(b)} 元{note}")
    print("| 项目 | 个人每月 |\n|------|------|")
    for k, v in ins.items():
        print(f"| {k} | {yuan(v)} |")
    total = sum(ins.values())
    print(f"| 合计 | {yuan(total)} |")
    print(
        f"公积金单位同比例缴纳 {yuan(b * fund / 100)} 元，个人账户每月共入账 {yuan(b * fund / 50)} 元。"
    )
    print(
        "提示：医疗保险部分城市另有大病统筹固定金额；单位缴费比例（养老 16% 等）和基数上下限各地每年调整，以当地公布为准。"
    )
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("mortgage")
    m.add_argument("--principal", type=float, required=True)
    m.add_argument("--rate", type=float, required=True, help="年利率 %，如 3.1")
    m.add_argument("--years", type=int, required=True)
    m.add_argument("--prepay", action="append", default=[])
    m.add_argument("-o", "--output")
    t = sub.add_parser("tax")
    t.add_argument("--salary", type=float, required=True)
    t.add_argument("--months", type=int, default=12)
    t.add_argument("--insurance", type=float)
    t.add_argument("--base", type=float)
    t.add_argument("--fund", type=float, default=7)
    t.add_argument("--min", type=float)
    t.add_argument("--max", type=float)
    t.add_argument("--children", type=int, default=0)
    t.add_argument("--infants", type=int, default=0)
    t.add_argument("--elderly", type=float, default=0)
    t.add_argument("--loan", action="store_true")
    t.add_argument("--rent", type=int, default=0)
    t.add_argument("--edu", action="store_true")
    t.add_argument("--cert", action="store_true")
    t.add_argument("--other", type=float, default=0, help="其他扣除（如个人养老金 12000/年 ÷12）")
    t.add_argument("--bonus", type=float, default=0)
    t.add_argument("-o", "--output")
    i = sub.add_parser("insurance")
    i.add_argument("--base", type=float, required=True)
    i.add_argument("--fund", type=float, default=7)
    i.add_argument("--min", type=float)
    i.add_argument("--max", type=float)
    args = parser.parse_args(argv)
    try:
        if args.cmd == "mortgage":
            return cmd_mortgage(
                args.principal,
                args.rate,
                args.years,
                args.prepay,
                Path(args.output) if args.output else None,
            )
        if args.cmd == "tax":
            return cmd_tax(args)
        return cmd_insurance(args.base, args.min, args.max, args.fund)
    except (CalcError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
