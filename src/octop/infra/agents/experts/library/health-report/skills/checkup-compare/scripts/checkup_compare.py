#!/usr/bin/env python3
"""体检报告对比：历年体检指标整理成一张表，标出超出参考范围的指标、连续异常和变化趋势。只做整理和解释，不做诊断。

Usage:
    python scripts/checkup_compare.py 体检数据.csv [-o 体检对比.xlsx] [--json]

体检数据 CSV（UTF-8 或 GBK）两种写法都支持：
  宽表：指标,单位,参考范围,2024,2025,2026          （每年一列）
  长表：年份,指标,结果,单位,参考范围              （每条一行，可以从多份报告逐条录入）
参考范围写法：3.5-5.5、3.5~5.5、<5.2、≤1.7、>1.04、≥40、阴性；结果可写数字、阴性/阳性、+/-。
输出：最近一次异常指标（↑偏高 / ↓偏低 / 阳性）、连续两年及以上异常、持续变差（连续远离参考范围）的指标，
并对常见指标给出通俗解释和一般建议就诊科室。所有内容不能代替医生诊断，异常结果请咨询医生。
Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl（仅 -o）。
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

# 名称别名 → (标准名, 通俗解释, 一般就诊科室)
KNOWLEDGE: dict[str, tuple[str, str, str]] = {
    "总胆固醇": ("总胆固醇", "血液中胆固醇总量，偏高与动脉粥样硬化风险相关", "心内科 / 内分泌科"),
    "TC": ("总胆固醇", "血液中胆固醇总量，偏高与动脉粥样硬化风险相关", "心内科 / 内分泌科"),
    "甘油三酯": (
        "甘油三酯",
        "血脂的一种，受饮食、饮酒影响大，明显偏高可能诱发胰腺炎",
        "内分泌科 / 心内科",
    ),
    "TG": (
        "甘油三酯",
        "血脂的一种，受饮食、饮酒影响大，明显偏高可能诱发胰腺炎",
        "内分泌科 / 心内科",
    ),
    "低密度脂蛋白": ("低密度脂蛋白胆固醇", "俗称“坏胆固醇”，是评估心血管风险的重要指标", "心内科"),
    "LDL-C": ("低密度脂蛋白胆固醇", "俗称“坏胆固醇”，是评估心血管风险的重要指标", "心内科"),
    "高密度脂蛋白": ("高密度脂蛋白胆固醇", "俗称“好胆固醇”，偏低时心血管风险增加", "心内科"),
    "HDL-C": ("高密度脂蛋白胆固醇", "俗称“好胆固醇”，偏低时心血管风险增加", "心内科"),
    "空腹血糖": ("空腹血糖", "反映血糖水平，偏高需复查并结合糖化血红蛋白判断", "内分泌科"),
    "GLU": ("空腹血糖", "反映血糖水平，偏高需复查并结合糖化血红蛋白判断", "内分泌科"),
    "糖化血红蛋白": ("糖化血红蛋白", "反映近 2–3 个月的平均血糖", "内分泌科"),
    "HbA1c": ("糖化血红蛋白", "反映近 2–3 个月的平均血糖", "内分泌科"),
    "尿酸": (
        "尿酸",
        "嘌呤代谢产物，偏高与痛风、肾脏问题相关，受饮食和饮水影响",
        "内分泌科 / 风湿免疫科",
    ),
    "UA": (
        "尿酸",
        "嘌呤代谢产物，偏高与痛风、肾脏问题相关，受饮食和饮水影响",
        "内分泌科 / 风湿免疫科",
    ),
    "谷丙转氨酶": (
        "谷丙转氨酶",
        "肝细胞损伤的敏感指标，熬夜、饮酒、脂肪肝、药物都可能使其升高",
        "消化内科 / 肝病科",
    ),
    "ALT": (
        "谷丙转氨酶",
        "肝细胞损伤的敏感指标，熬夜、饮酒、脂肪肝、药物都可能使其升高",
        "消化内科 / 肝病科",
    ),
    "谷草转氨酶": ("谷草转氨酶", "存在于肝脏和心肌等，常与谷丙转氨酶一起看", "消化内科"),
    "AST": ("谷草转氨酶", "存在于肝脏和心肌等，常与谷丙转氨酶一起看", "消化内科"),
    "肌酐": ("肌酐", "评估肾功能的常用指标，受肌肉量影响", "肾内科"),
    "Cr": ("肌酐", "评估肾功能的常用指标，受肌肉量影响", "肾内科"),
    "血红蛋白": ("血红蛋白", "偏低提示可能贫血", "血液科 / 消化内科"),
    "HGB": ("血红蛋白", "偏低提示可能贫血", "血液科 / 消化内科"),
    "白细胞": ("白细胞计数", "与感染、炎症等有关，轻微波动常见", "血液科 / 感染科"),
    "WBC": ("白细胞计数", "与感染、炎症等有关，轻微波动常见", "血液科 / 感染科"),
    "血小板": ("血小板计数", "与凝血功能相关", "血液科"),
    "PLT": ("血小板计数", "与凝血功能相关", "血液科"),
    "收缩压": ("收缩压", "血压高的一端，多次测量偏高才考虑高血压", "心内科"),
    "舒张压": ("舒张压", "血压低的一端", "心内科"),
    "BMI": ("体重指数", "体重(kg) ÷ 身高(m)²，中国成人 18.5–23.9 为正常", "营养科"),
    "促甲状腺激素": ("促甲状腺激素", "反映甲状腺功能，偏高或偏低需结合 FT3、FT4", "内分泌科"),
    "TSH": ("促甲状腺激素", "反映甲状腺功能，偏高或偏低需结合 FT3、FT4", "内分泌科"),
    "癌胚抗原": (
        "癌胚抗原",
        "肿瘤标志物之一，吸烟、炎症也可能轻度升高，单次轻度升高不能说明问题",
        "消化内科 / 肿瘤科",
    ),
    "CEA": (
        "癌胚抗原",
        "肿瘤标志物之一，吸烟、炎症也可能轻度升高，单次轻度升高不能说明问题",
        "消化内科 / 肿瘤科",
    ),
    "幽门螺杆菌": ("幽门螺杆菌", "与胃炎、胃溃疡相关，阳性可咨询是否需要根除治疗", "消化内科"),
    "尿蛋白": (
        "尿蛋白",
        "阳性可能与肾脏问题有关，也可能是剧烈运动后的一过性表现，需复查",
        "肾内科",
    ),
    "尿潜血": ("尿潜血", "女性经期前后易出现假阳性，需复查", "肾内科 / 泌尿外科"),
}
NEGATIVE = {"阴性", "-", "neg", "negative", "正常", "未见异常", "未检出"}
POSITIVE = {"阳性", "+", "++", "+++", "pos", "positive", "弱阳性", "±"}


class ReportError(ValueError):
    pass


@dataclass
class Range:
    low: float | None = None
    high: float | None = None
    qualitative: bool = False
    text: str = ""

    def judge(self, value: str) -> str:
        """'' normal, '↑' high, '↓' low, '阳性' positive, '?' unparseable."""
        v = value.strip()
        if not v:
            return ""
        if self.qualitative or v.lower() in NEGATIVE | POSITIVE:
            if v.lower() in POSITIVE:
                return "阳性"
            return "" if v.lower() in NEGATIVE else "?"
        num = to_float(v)
        if num is None:
            return "?"
        if self.high is not None and num > self.high:
            return "↑"
        if self.low is not None and num < self.low:
            return "↓"
        return ""

    def distance(self, value: str) -> float:
        num = to_float(value)
        if num is None:
            return 0.0
        if self.high is not None and num > self.high:
            return (num - self.high) / (abs(self.high) or 1)
        if self.low is not None and num < self.low:
            return (self.low - num) / (abs(self.low) or 1)
        return 0.0


def to_float(s: str) -> float | None:
    m = re.search(r"-?\d+(?:\.\d+)?", s.replace(",", ""))
    return float(m.group(0)) if m else None


def parse_range(text: str) -> Range:
    t = text.strip().replace("～", "~").replace("—", "-").replace("–", "-").replace(" ", "")
    if not t:
        return Range(text=text)
    if t in NEGATIVE or t.lower() in NEGATIVE:
        return Range(qualitative=True, text=text)
    m = re.fullmatch(r"(-?\d+(?:\.\d+)?)[-~至](-?\d+(?:\.\d+)?)", t)
    if m:
        return Range(float(m.group(1)), float(m.group(2)), text=text)
    m = re.fullmatch(r"(<|≤|<=|＜)(\d+(?:\.\d+)?)", t)
    if m:
        return Range(None, float(m.group(2)), text=text)
    m = re.fullmatch(r"(>|≥|>=|＞)(\d+(?:\.\d+)?)", t)
    if m:
        return Range(float(m.group(2)), None, text=text)
    raise ReportError(f"参考范围「{text}」无法识别，请写成 3.5-5.5、<5.2、>1.04 或 阴性")


@dataclass
class Item:
    name: str
    unit: str
    ref: Range
    values: dict[str, str] = field(default_factory=dict)


def read_rows(path: Path) -> list[dict[str, str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ReportError("文件编码无法识别")
    return [
        {(k or "").strip(): (v or "").strip() for k, v in r.items()}
        for r in csv.DictReader(io.StringIO(text))
    ]


def load(path: Path) -> tuple[list[Item], list[str]]:
    rows = read_rows(path)
    if not rows:
        raise ReportError("没有数据")
    cols = list(rows[0])
    items: dict[str, Item] = {}
    years: list[str] = []
    if {"年份", "指标", "结果"} <= set(cols):
        for r in rows:
            name = r["指标"]
            item = items.setdefault(
                name, Item(name, r.get("单位", ""), parse_range(r.get("参考范围", "")))
            )
            if r.get("参考范围") and not item.ref.text:
                item.ref = parse_range(r["参考范围"])
            item.values[r["年份"]] = r["结果"]
            if r["年份"] not in years:
                years.append(r["年份"])
    elif "指标" in cols and "参考范围" in cols:
        years = [c for c in cols if re.fullmatch(r"(19|20)\d{2}(-\d{1,2})?", c)]
        if not years:
            raise ReportError("宽表需要年份列（如 2025）")
        for r in rows:
            if r["指标"]:
                items[r["指标"]] = Item(
                    r["指标"],
                    r.get("单位", ""),
                    parse_range(r["参考范围"]),
                    {y: r.get(y, "") for y in years},
                )
    else:
        raise ReportError("表头应为“指标,单位,参考范围,2024,2025…”或“年份,指标,结果,单位,参考范围”")
    return list(items.values()), sorted(years)


def knowledge(name: str) -> tuple[str, str, str] | None:
    key = re.sub(r"[（(].*?[)）]", "", name).strip()
    for alias, info in KNOWLEDGE.items():
        if alias == key or alias in name or info[0] == key:
            return info
    return None


def analyze(items: list[Item], years: list[str]) -> dict[str, list[dict[str, str]]]:
    latest_abn, persistent, worsening = [], [], []
    for it in items:
        flags = [
            (y, it.values.get(y, ""), it.ref.judge(it.values.get(y, "")))
            for y in years
            if it.values.get(y, "")
        ]
        if not flags:
            continue
        info = knowledge(it.name)
        base = {
            "指标": it.name,
            "单位": it.unit,
            "参考范围": it.ref.text,
            "解释": info[1] if info else "",
            "科室": info[2] if info else "",
        }
        last_y, last_v, last_f = flags[-1]
        if last_f and last_f != "?":
            latest_abn.append({**base, "年份": last_y, "结果": last_v, "判断": last_f})
        run = 0
        for _y, _v, f in reversed(flags):
            if f and f != "?":
                run += 1
            else:
                break
        if run >= 2:
            persistent.append(
                {
                    **base,
                    "连续异常": f"{run} 次",
                    "结果": " → ".join(v for _y, v, _f in flags[-run:]),
                }
            )
        dists = [it.ref.distance(v) for _y, v, _f in flags]
        if len(dists) >= 3 and dists[-1] > dists[-2] > dists[-3] and dists[-1] > 0:
            worsening.append({**base, "趋势": " → ".join(v for _y, v, _f in flags[-3:])})
    return {"最近异常": latest_abn, "连续异常": persistent, "持续变差": worsening}


def write_xlsx(
    items: list[Item], years: list[str], res: dict[str, list[dict[str, str]]], out: Path
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
        sys.exit(2)
    wb = Workbook()
    ws = wb.active
    ws.title = "历年对比"
    ws.append(["指标", "单位", "参考范围", *years, "说明"])
    hi, lo = PatternFill("solid", fgColor="FECACA"), PatternFill("solid", fgColor="BFDBFE")
    for it in items:
        info = knowledge(it.name)
        row = [it.name, it.unit, it.ref.text]
        for y in years:
            v = it.values.get(y, "")
            f = it.ref.judge(v)
            row.append(f"{v} {f}".strip() if f and f != "?" else v)
        ws.append([*row, info[1] if info else ""])
        for j, y in enumerate(years, start=4):
            f = it.ref.judge(it.values.get(y, ""))
            if f in ("↑", "阳性"):
                ws.cell(ws.max_row, j).fill = hi
            elif f == "↓":
                ws.cell(ws.max_row, j).fill = lo
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="BE123C")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions[ws.cell(1, len(years) + 4).column_letter].width = 50
    ws.freeze_panes = "D2"
    for title, rows in res.items():
        if rows:
            sh = wb.create_sheet(title)
            keys = list(rows[0])
            sh.append(keys)
            for r in rows:
                sh.append([r[k] for k in keys])
            for c in sh[1]:
                c.font = Font(bold=True)
    note = wb.create_sheet("说明")
    note.append(
        [
            "本表只整理体检数据并标出超出参考范围的指标，不能代替医生诊断；异常结果请携带完整报告咨询医生。"
        ]
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input")
    parser.add_argument("-o", "--output")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    try:
        items, years = load(Path(args.input))
    except (OSError, ReportError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    res = analyze(items, years)
    if args.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    else:
        print(f"{len(items)} 项指标，{len(years)} 次体检（{'、'.join(years)}）")
        labels = {
            "最近异常": "最近一次异常",
            "连续异常": "连续两次及以上异常",
            "持续变差": "连续三次越来越偏离参考范围",
        }
        for key, label in labels.items():
            rows = res[key]
            print(f"\n## {label}（{len(rows)}）")
            for r in rows:
                detail = r.get("结果") if key != "持续变差" else r["趋势"]
                extra = f"{r['判断']} " if key == "最近异常" else ""
                print(
                    f"- {r['指标']}：{extra}{detail} {r['单位']}（参考 {r['参考范围']}）"
                    + (f" —— {r['解释']}；可咨询{r['科室']}" if r["解释"] else "")
                )
        unknown = [
            it.name for it in items if any(it.ref.judge(v) == "?" for v in it.values.values())
        ]
        if unknown:
            print(f"\n无法判断（结果不是数字或阴阳性）：{'、'.join(unknown)}")
        print("\n以上只是对报告数据的整理，不能代替医生诊断。")
    if args.output:
        write_xlsx(items, years, res, Path(args.output))
        print(f"已生成 {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
