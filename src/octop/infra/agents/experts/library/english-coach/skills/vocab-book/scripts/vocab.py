#!/usr/bin/env python3
"""单词本：Markdown / CSV 生词 → Excel 单词本（按遗忘曲线排复习日期），列出今天要复习的词，生成可打印的默写测验卷。

Usage:
    python scripts/vocab.py build 生词.md -o 单词本.xlsx [--merge 单词本.xlsx] [--today 2026-10-18]
    python scripts/vocab.py due 单词本.xlsx [--today 2026-10-18]
    python scripts/vocab.py quiz 单词本.xlsx -o 默写.docx [-n 20] [--mode en2zh|zh2en|mixed] [--due-only] [--seed 1]

生词表头：| 单词 | 音标 | 释义 | 例句 | 来源 |（单词、释义必填；也可以是同样列名的 CSV）。
复习安排在加入后第 1、2、4、7、15、30 天；在 Excel“已掌握”列填“是”后不再复习，测验也不再抽到。
合并时同一单词（不区分大小写）只保留一条，保留原来的加入日期和已掌握标记。
Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl / python-docx。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import random
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

FIELDS = ["单词", "音标", "释义", "例句", "来源"]
INTERVALS = (1, 2, 4, 7, 15, 30)
COLUMNS = [*FIELDS, "加入日期", *[f"复习{i + 1}" for i in range(len(INTERVALS))], "已掌握"]
DONE = ("是", "y", "Y", "✓", "√")


class VocabError(ValueError):
    pass


def read_input(path: Path) -> list[dict[str, str]]:
    text = path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() == ".csv":
        rows = [
            {(k or "").strip(): (v or "").strip() for k, v in r.items()}
            for r in csv.DictReader(io.StringIO(text))
        ]
    else:
        table = [
            [c.strip() for c in ln.strip().strip("|").split("|")]
            for ln in text.splitlines()
            if ln.strip().startswith("|")
        ]
        table = [r for r in table if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
        if len(table) < 2:
            raise VocabError("没有找到生词表格")
        header = table[0]
        rows = []
        for r in table[1:]:
            if len(r) != len(header):
                raise VocabError(f"「{' | '.join(r)}」的列数和表头不一致")
            rows.append(dict(zip(header, r, strict=True)))
    if not rows or "单词" not in rows[0] or "释义" not in rows[0]:
        raise VocabError("表头至少需要「单词」「释义」")
    return [{f: r.get(f, "") for f in FIELDS} for r in rows if r.get("单词")]


def read_book(path: Path) -> list[dict[str, object]]:
    ws = load_workbook(path)["单词本"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r and r[0]:
            rec = {k: ("" if v is None else v) for k, v in zip(header, r, strict=False)}
            d = rec.get("加入日期")
            rec["加入日期"] = (
                d.date()
                if isinstance(d, dt.datetime)
                else d
                if isinstance(d, dt.date)
                else dt.date.fromisoformat(str(d))
            )
            out.append(rec)
    return out


def write_book(items: list[dict[str, object]], out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "单词本"
    ws.append(COLUMNS)
    for rec in sorted(items, key=lambda r: (r["加入日期"], str(r["单词"]).lower())):  # type: ignore[arg-type]
        added = rec["加入日期"]
        assert isinstance(added, dt.date)
        ws.append(
            [
                *(rec.get(f, "") for f in FIELDS),
                added,
                *(added + dt.timedelta(days=n) for n in INTERVALS),
                rec.get("已掌握", ""),
            ]
        )
        for j in range(len(FIELDS) + 1, len(COLUMNS)):
            ws.cell(ws.max_row, j).number_format = "yyyy-mm-dd" if j == len(FIELDS) + 1 else "mm-dd"
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F46E5")
    for col, w in zip("ABCDE", (16, 14, 28, 40, 12), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{ws.max_row}"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def due_items(
    items: list[dict[str, object]], today: dt.date
) -> list[tuple[dict[str, object], int]]:
    out = []
    for r in items:
        if str(r.get("已掌握", "")).strip() in DONE:
            continue
        added = r["加入日期"]
        assert isinstance(added, dt.date)
        stages = [i for i, n in enumerate(INTERVALS) if added + dt.timedelta(days=n) <= today]
        if stages and added + dt.timedelta(days=INTERVALS[stages[-1]]) >= today - dt.timedelta(
            days=2
        ):
            out.append((r, stages[-1] + 1))
    return out


def cmd_build(src: Path, out: Path, merge: Path | None, today: dt.date) -> int:
    items = read_book(merge) if merge else []
    seen = {str(r["单词"]).lower() for r in items}
    added = 0
    for r in read_input(src):
        key = r["单词"].lower()
        if key in seen:
            continue
        items.append({**r, "加入日期": today, "已掌握": ""})
        seen.add(key)
        added += 1
    write_book(items, out)
    print(
        f"已生成 {out}：共 {len(items)} 个单词（本次新增 {added} 个），复习安排在第 {'/'.join(map(str, INTERVALS))} 天"
    )
    return 0


def cmd_due(book: Path, today: dt.date) -> int:
    items = read_book(book)
    due = due_items(items, today)
    mastered = sum(1 for r in items if str(r.get("已掌握", "")).strip() in DONE)
    print(f"{today} 待复习 {len(due)} 个（单词本共 {len(items)} 个，已掌握 {mastered} 个）")
    for r, stage in due:
        print(f"- {r['单词']} {r.get('音标', '')} {r['释义']}（第 {stage} 轮）".replace("  ", " "))
    return 0


def cmd_quiz(
    book: Path, out: Path, n: int, mode: str, due_only: bool, seed: int | None, today: dt.date
) -> int:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 python-docx：请运行 `pip install python-docx`", file=sys.stderr)
        return 2
    items = read_book(book)
    pool = (
        [r for r, _s in due_items(items, today)]
        if due_only
        else [r for r in items if str(r.get("已掌握", "")).strip() not in DONE]
    )
    if not pool:
        raise VocabError("没有可以出题的单词（全部已掌握，或今天没有要复习的）")
    rng = random.Random(seed)
    picks = rng.sample(pool, min(n, len(pool)))
    qs = []
    for r in picks:
        m = rng.choice(("en2zh", "zh2en")) if mode == "mixed" else mode
        if m == "en2zh":
            qs.append((f"{r['单词']}  ____________", f"{r['单词']}：{r['释义']}"))
        else:
            qs.append((f"{r['释义']}  ____________", f"{r['释义']}：{r['单词']}"))
    doc = Document()

    def title(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold, run.font.size = True, Pt(16)

    title(f"单词默写（{len(qs)} 题）")
    doc.add_paragraph(f"日期：{today}　　姓名：________　　得分：________")
    for i, (q, _a) in enumerate(qs, start=1):
        doc.add_paragraph(f"{i}. {q}").paragraph_format.space_after = Pt(8)
    doc.add_page_break()
    title("答案")
    for i, (_q, a) in enumerate(qs, start=1):
        doc.add_paragraph(f"{i}. {a}")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(
        f"已生成 {out}：{len(qs)} 题（{ {'en2zh': '看英文写中文', 'zh2en': '看中文写英文', 'mixed': '混合'}[mode] }），答案在第 2 页"
    )
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("input")
    b.add_argument("-o", "--output", required=True)
    b.add_argument("--merge")
    b.add_argument("--today")
    d = sub.add_parser("due")
    d.add_argument("book")
    d.add_argument("--today")
    q = sub.add_parser("quiz")
    q.add_argument("book")
    q.add_argument("-o", "--output", required=True)
    q.add_argument("-n", type=int, default=20)
    q.add_argument("--mode", choices=["en2zh", "zh2en", "mixed"], default="mixed")
    q.add_argument("--due-only", action="store_true")
    q.add_argument("--seed", type=int)
    q.add_argument("--today")
    args = parser.parse_args(argv)
    try:
        today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
        if args.cmd == "build":
            return cmd_build(
                Path(args.input), Path(args.output), Path(args.merge) if args.merge else None, today
            )
        if args.cmd == "due":
            return cmd_due(Path(args.book), today)
        return cmd_quiz(
            Path(args.book), Path(args.output), args.n, args.mode, args.due_only, args.seed, today
        )
    except (OSError, VocabError, ValueError, KeyError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
