#!/usr/bin/env python3
"""家庭账本：导入支付宝 / 微信 / 银行流水，统一格式、去重、自动分类，并生成收支报告。

Usage:
    python scripts/ledger.py import 账单1.csv 账单2.xlsx ... -o 账本.xlsx [--rules 分类规则.csv] [--merge 旧账本.xlsx]
    python scripts/ledger.py report 账本.xlsx -o 收支报告.xlsx [--budget 预算.csv] [--month 2026-09]

import：
  - 自动识别表头（支付宝“交易时间/收/支/金额”、微信“交易时间/收/支/金额(元)”、银行“交易日期/收入/支出/余额”等），
    CSV 自动识别 UTF-8 / GBK 编码；
  - 跳过“不计收支”、交易关闭、退款中、余额宝收益转入等不影响收支的记录；转账、红包标为“转账”类，不计入消费；
  - 按“来源 + 订单号”去重（同一笔支付宝交易导出两次只算一次）；无订单号时按 时间+金额+对方 去重；
  - 按关键词自动分类，可用 --rules 追加自定义规则（两列：关键词,类别；先匹配自定义规则）。
report：月度收支汇总、分类 × 月份透视、收入/支出/结余/储蓄率、Top 商户、疑似订阅扣费、大额异常支出、预算对比，带图表。
只读原始账单，不修改。Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import statistics
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

LEDGER_COLUMNS = [
    "日期",
    "时间",
    "来源",
    "收支",
    "金额",
    "类别",
    "对方",
    "说明",
    "支付方式",
    "订单号",
]
HEADER_FILL = PatternFill("solid", fgColor="1F2937")

SYNONYMS = {
    "time": ("交易时间", "交易日期", "记账日期", "交易日", "日期", "时间", "入账时间", "记账时间"),
    "direction": ("收/支", "收支", "收支类型", "借贷标志", "借贷"),
    "amount": ("金额(元)", "金额（元）", "金额", "交易金额", "发生额"),
    "income": ("收入金额", "收入", "贷方金额", "存入"),
    "expense": ("支出金额", "支出", "借方金额", "取出"),
    "party": ("交易对方", "对方户名", "对方名称", "对方账户名", "商户名称", "收款方"),
    "desc": ("商品说明", "商品", "摘要", "交易摘要", "用途", "附言", "备注", "交易类型"),
    "category": ("交易分类",),
    "type": ("交易类型",),
    "method": ("收/付款方式", "支付方式", "交易渠道"),
    "status": ("交易状态", "当前状态", "状态"),
    "order": ("交易订单号", "交易单号", "流水号", "交易流水号", "订单号"),
}

DEFAULT_RULES: list[tuple[str, str]] = [
    # 先排除转账、红包、理财等不属于消费的记录
    # 具体的在前：工资、报销常以“转账”方式到账，要先于“转账”匹配
    ("工资", "工资"), ("薪", "工资"), ("报销", "报销"), ("退款", "退款"),
    ("转账", "转账"), ("红包", "转账"), ("还款", "转账"), ("信用卡还款", "转账"), ("余额宝", "理财"), ("理财", "理财"),
    ("基金", "理财"), ("零钱通", "理财"),
    ("美团", "餐饮"), ("饿了么", "餐饮"), ("肯德基", "餐饮"), ("麦当劳", "餐饮"), ("星巴克", "餐饮"), ("瑞幸", "餐饮"),
    ("奶茶", "餐饮"), ("咖啡", "餐饮"), ("餐", "餐饮"), ("饭", "餐饮"), ("面馆", "餐饮"), ("烧烤", "餐饮"),
    ("盒马", "买菜"), ("叮咚", "买菜"), ("朴朴", "买菜"), ("菜市场", "买菜"), ("超市", "买菜"), ("便利店", "买菜"),
    ("水果", "买菜"), ("生鲜", "买菜"),
    ("滴滴", "交通"), ("高德打车", "交通"), ("地铁", "交通"), ("公交", "交通"), ("12306", "交通"), ("铁路", "交通"),
    ("航空", "交通"), ("加油", "交通"), ("中石化", "交通"), ("中石油", "交通"), ("停车", "交通"), ("高速", "交通"),
    ("ETC", "交通"), ("哈啰", "交通"), ("充电", "交通"),
    ("淘宝", "购物"), ("天猫", "购物"), ("京东", "购物"), ("拼多多", "购物"), ("抖音电商", "购物"), ("唯品会", "购物"),
    ("房租", "居住"), ("物业", "居住"), ("电费", "居住"), ("水费", "居住"), ("燃气", "居住"), ("国家电网", "居住"),
    ("话费", "通讯"), ("宽带", "通讯"), ("中国移动", "通讯"), ("中国联通", "通讯"), ("中国电信", "通讯"),
    ("会员", "娱乐"), ("爱奇艺", "娱乐"), ("腾讯视频", "娱乐"), ("优酷", "娱乐"), ("哔哩哔哩", "娱乐"), ("网易云", "娱乐"),
    ("QQ音乐", "娱乐"), ("电影", "娱乐"), ("游戏", "娱乐"), ("Steam", "娱乐"), ("KTV", "娱乐"),
    ("医院", "医疗"), ("药房", "医疗"), ("药店", "医疗"), ("诊所", "医疗"), ("体检", "医疗"),
    ("学费", "教育"), ("培训", "教育"), ("课程", "教育"), ("书店", "教育"), ("得到", "教育"),
    ("保险", "保险"), ("人寿", "保险"), ("理发", "个人护理"), ("美容", "个人护理"), ("健身", "运动"),
]  # fmt: skip
NON_CONSUMPTION = {"转账", "理财", "退款", "报销", "工资"}


class LedgerError(ValueError):
    pass


@dataclass
class Entry:
    when: dt.datetime
    source: str
    direction: str  # 收入 / 支出
    amount: float
    category: str
    party: str
    desc: str
    method: str
    order: str


# ------------------------------------------------------------------ reading


def read_rows(path: Path) -> list[list[str]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = [["" if c is None else (c.strftime("%Y-%m-%d %H:%M:%S") if isinstance(c, dt.datetime) else str(c)) for c in r]
                for r in wb.worksheets[0].iter_rows(values_only=True)]  # fmt: skip
        wb.close()
        return rows
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise LedgerError(f"{path.name}：无法识别编码")
    return [[c.strip().strip("\t") for c in row] for row in csv.reader(text.splitlines())]


def detect(rows: list[list[str]]) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(rows[:60]):
        cells = [c.strip() for c in row]
        cols: dict[str, int] = {}
        for key, names in SYNONYMS.items():
            for name in names:
                if name in cells:
                    cols[key] = cells.index(name)
                    break
        has_amount = "amount" in cols or ("income" in cols and "expense" in cols)
        if "time" in cols and has_amount:
            return i, cols
    raise LedgerError("没有找到账单表头（需要包含“交易时间/交易日期”和“金额/收入/支出”列）")


def parse_time(text: str) -> dt.datetime | None:
    text = text.strip().replace("/", "-").replace(".", "-")
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y%m%d %H:%M:%S",
        "%Y%m%d",
        "%Y年%m月%d日 %H:%M:%S",
        "%Y年%m月%d日",
    ):
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def parse_amount(text: str) -> float | None:
    cleaned = re.sub(r"[¥￥,，\s元+]", "", str(text))
    if cleaned in ("", "-", "/"):
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def source_name(path: Path, rows: list[list[str]]) -> str:
    head = " ".join(" ".join(r) for r in rows[:25])
    if "支付宝" in head or "alipay" in path.name.lower():
        return "支付宝"
    if "微信" in head or "wechat" in path.name.lower() or "微信" in path.name:
        return "微信"
    return path.stem[:20]


def categorize(text: str, rules: list[tuple[str, str]], hint: str = "") -> str:
    for keyword, category in rules:
        if keyword and keyword.lower() in text.lower():
            return category
    hint_map = {"餐饮美食": "餐饮", "交通出行": "交通", "日用百货": "购物", "服饰装扮": "购物", "数码电器": "购物",
                "充值缴费": "通讯", "医疗健康": "医疗", "教育培训": "教育", "文化休闲": "娱乐", "住房物业": "居住",
                "转账红包": "转账", "投资理财": "理财", "退款": "退款", "运动户外": "运动", "美容美发": "个人护理"}  # fmt: skip
    return hint_map.get(hint, "其他")


def load_rules(path: Path | None) -> list[tuple[str, str]]:
    if path is None:
        return DEFAULT_RULES
    custom = []
    for row in read_rows(path):
        if len(row) >= 2 and row[0] and row[0] not in ("关键词", "keyword"):
            custom.append((row[0].strip(), row[1].strip()))
    return custom + DEFAULT_RULES


SKIP_STATUS = re.compile(r"关闭|失败|撤销|退款中|已退款|全额退款|对方已退还|已退还")


def parse_file(path: Path, rules: list[tuple[str, str]]) -> tuple[list[Entry], list[str]]:
    rows = read_rows(path)
    start, cols = detect(rows)
    source = source_name(path, rows)
    entries: list[Entry] = []
    skipped: Counter[str] = Counter()

    def cell(row: list[str], key: str) -> str:
        i = cols.get(key)
        return row[i].strip() if i is not None and i < len(row) else ""

    for row in rows[start + 1 :]:
        if not any(row) or row[0].startswith(("-", "共", "合计")):
            continue
        when = parse_time(cell(row, "time"))
        if when is None:
            continue
        status = cell(row, "status")
        direction_text = cell(row, "direction")
        if "不计收支" in direction_text or direction_text == "/":
            skipped["不计收支"] += 1
            continue
        if status and SKIP_STATUS.search(status):
            skipped[f"状态：{status}"] += 1
            continue
        if "income" in cols and "expense" in cols and "amount" not in cols:
            inc, exp = parse_amount(cell(row, "income")), parse_amount(cell(row, "expense"))
            if inc:
                direction, amount = "收入", inc
            elif exp:
                direction, amount = "支出", exp
            else:
                continue
        else:
            amount_value = parse_amount(cell(row, "amount"))
            if amount_value is None or amount_value == 0:
                continue
            if "收入" in direction_text or direction_text in ("贷", "C"):
                direction = "收入"
            elif "支出" in direction_text or direction_text in ("借", "D"):
                direction = "支出"
            elif direction_text == "":
                direction = "支出" if amount_value < 0 else "收入"
            else:
                skipped[f"收支：{direction_text}"] += 1
                continue
            amount = abs(amount_value)
        party, desc = cell(row, "party"), cell(row, "desc")
        text = f"{party} {desc} {cell(row, 'type')} {cell(row, 'category')}"
        category = categorize(text, rules, cell(row, "category"))
        if (
            direction == "收入"
            and category not in NON_CONSUMPTION
            and category in ("餐饮", "购物", "交通", "其他")
        ):
            category = "其他收入"
        entries.append(Entry(when, source, direction, round(amount, 2), category, party, desc, cell(row, "method"),
                             cell(row, "order").strip("\t ")))  # fmt: skip
    notes = [f"{path.name}：识别为{source}账单，读取 {len(entries)} 笔"]
    notes += [f"{path.name}：跳过 {n} 笔（{why}）" for why, n in skipped.most_common()]
    return entries, notes


def dedupe(entries: list[Entry]) -> tuple[list[Entry], int]:
    seen: set[tuple[Any, ...]] = set()
    out = []
    for e in sorted(entries, key=lambda x: x.when):
        key: tuple[Any, ...] = (
            (e.source, e.order) if e.order else (e.when, e.amount, e.party, e.direction)
        )
        if key in seen:
            continue
        seen.add(key)
        out.append(e)
    return out, len(entries) - len(out)


def read_ledger(path: Path) -> list[Entry]:
    wb = load_workbook(path, data_only=True)
    ws = wb["账本"] if "账本" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    if header[: len(LEDGER_COLUMNS)] != LEDGER_COLUMNS:
        raise LedgerError(
            f"{path.name} 不是本工具生成的账本（表头应为：{','.join(LEDGER_COLUMNS)}）"
        )
    out = []
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        d = r[0] if isinstance(r[0], dt.date) else dt.date.fromisoformat(str(r[0])[:10])
        t = str(r[1] or "00:00:00")
        hh, mm, ss = (int(x) for x in (t.split(":") + ["0", "0"])[:3])
        out.append(Entry(dt.datetime(d.year, d.month, d.day, hh, mm, ss), str(r[2] or ""), str(r[3]), float(r[4]),
                         str(r[5] or "其他"), str(r[6] or ""), str(r[7] or ""), str(r[8] or ""), str(r[9] or "")))  # fmt: skip
    return out


def write_ledger(entries: list[Entry], out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "账本"
    ws.append(LEDGER_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for e in entries:
        ws.append([e.when.date(), e.when.strftime("%H:%M:%S"), e.source, e.direction, e.amount, e.category,
                   e.party, e.desc, e.method, e.order])  # fmt: skip
        ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"
        ws.cell(ws.max_row, 5).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJ", (12, 10, 8, 6, 12, 10, 24, 32, 14, 30), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def cmd_import(files: list[Path], out: Path, rules_path: Path | None, merge: Path | None) -> int:
    rules = load_rules(rules_path)
    entries: list[Entry] = read_ledger(merge) if merge else []
    notes = []
    for f in files:
        got, file_notes = parse_file(f, rules)
        entries.extend(got)
        notes.extend(file_notes)
    entries, dupes = dedupe(entries)
    if not entries:
        raise LedgerError("没有读取到任何交易记录")
    write_ledger(entries, out)
    for n in notes:
        print(n)
    span = f"{entries[0].when:%Y-%m-%d} ~ {entries[-1].when:%Y-%m-%d}"
    others = sum(1 for e in entries if e.category == "其他")
    print(
        f"已生成 {out}：共 {len(entries)} 笔（{span}），去重 {dupes} 笔；“其他”类 {others} 笔，可用 --rules 补充分类规则"
    )
    return 0


# ------------------------------------------------------------------- report


def cmd_report(ledger: Path, out: Path, budget_path: Path | None, month: str | None) -> int:
    entries = read_ledger(ledger)
    if month:
        entries = [e for e in entries if e.when.strftime("%Y-%m") == month]
        if not entries:
            raise LedgerError(f"{month} 没有记录")
    spend = [e for e in entries if e.direction == "支出" and e.category not in NON_CONSUMPTION]
    income = [
        e for e in entries if e.direction == "收入" and e.category not in ("转账", "退款", "理财")
    ]
    months = sorted({e.when.strftime("%Y-%m") for e in entries})
    by_mc: dict[tuple[str, str], float] = defaultdict(float)
    by_cat: Counter[str] = Counter()
    for e in spend:
        by_mc[(e.when.strftime("%Y-%m"), e.category)] += e.amount
        by_cat[e.category] += e.amount
    cats = [c for c, _ in by_cat.most_common()]  # largest spend first
    inc_m: dict[str, float] = defaultdict(float)
    exp_m: dict[str, float] = defaultdict(float)
    for e in income:
        inc_m[e.when.strftime("%Y-%m")] += e.amount
    for e in spend:
        exp_m[e.when.strftime("%Y-%m")] += e.amount

    wb = Workbook()
    ws = wb.active
    ws.title = "月度汇总"
    ws.append(["月份", "收入", "支出", "结余", "储蓄率"])
    for m in months:
        inc, exp = round(inc_m[m], 2), round(exp_m[m], 2)
        ws.append([m, inc, exp, round(inc - exp, 2), (inc - exp) / inc if inc else None])
        ws.cell(ws.max_row, 5).number_format = "0.0%"
    _style_header(ws)
    chart = BarChart()
    chart.title = "每月收入与支出"
    chart.add_data(
        Reference(ws, min_col=2, max_col=3, min_row=1, max_row=len(months) + 1),
        titles_from_data=True,
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(months) + 1))
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, f"A{len(months) + 4}")

    pv = wb.create_sheet("分类×月份")
    pv.append(["类别", *months, "合计", "占比"])
    total_spend = sum(e.amount for e in spend) or 1
    for c in cats:
        vals = [round(by_mc[(m, c)], 2) for m in months]
        pv.append([c, *vals, round(sum(vals), 2), sum(vals) / total_spend])
        pv.cell(pv.max_row, len(months) + 3).number_format = "0.0%"
    _style_header(pv)
    if cats:
        pie = PieChart()
        pie.title = "支出构成"
        pie.add_data(
            Reference(pv, min_col=len(months) + 2, min_row=1, max_row=len(cats) + 1),
            titles_from_data=True,
        )
        pie.set_categories(Reference(pv, min_col=1, min_row=2, max_row=len(cats) + 1))
        pv.add_chart(pie, f"A{len(cats) + 4}")

    top = wb.create_sheet("Top商户")
    top.append(["对方", "笔数", "金额"])
    merchants: dict[str, list[float]] = defaultdict(list)
    for e in spend:
        merchants[e.party or e.desc or "（未知）"].append(e.amount)
    ranked = sorted(merchants.items(), key=lambda kv: -sum(kv[1]))[:20]
    for name, amounts in ranked:
        top.append([name, len(amounts), round(sum(amounts), 2)])
    _style_header(top)

    # 疑似订阅：同一对方、金额相同、出现在 ≥3 个不同月份，且像订阅（会员/包月/续费，或娱乐、通讯等类别）
    subs: list[tuple[str, float, int]] = []
    groups: dict[tuple[str, float], set[str]] = defaultdict(set)
    sub_words = re.compile(r"会员|包月|包年|续费|订阅|月卡|年卡|VIP|vip|自动扣款")
    for e in spend:
        if sub_words.search(f"{e.party} {e.desc}") or e.category in (
            "娱乐",
            "通讯",
            "运动",
            "教育",
        ):
            groups[(e.party or e.desc, e.amount)].add(e.when.strftime("%Y-%m"))
    for (name, amount), ms in groups.items():
        if len(ms) >= 3:
            subs.append((name, amount, len(ms)))
    sub_ws = wb.create_sheet("疑似订阅")
    sub_ws.append(["对方", "每次金额", "出现月数", "按月估算年费"])
    for name, amount, n in sorted(subs, key=lambda s: -s[1]):
        sub_ws.append([name, amount, n, round(amount * 12, 2)])
    _style_header(sub_ws)

    # 异常大额：超过该类别中位数 5 倍且 ≥ 200 元
    big: list[Entry] = []
    per_cat: dict[str, list[float]] = defaultdict(list)
    for e in spend:
        per_cat[e.category].append(e.amount)
    for e in spend:
        values = per_cat[e.category]
        if len(values) >= 5 and e.amount >= 200 and e.amount > 5 * statistics.median(values):
            big.append(e)
    big_ws = wb.create_sheet("大额支出")
    big_ws.append(["日期", "类别", "对方", "说明", "金额"])
    for e in sorted(big, key=lambda x: -x.amount)[:30]:
        big_ws.append([e.when.date(), e.category, e.party, e.desc, e.amount])
        big_ws.cell(big_ws.max_row, 1).number_format = "yyyy-mm-dd"
    _style_header(big_ws)

    budget_lines: list[str] = []
    if budget_path:
        budget = {
            r[0].strip(): parse_amount(r[1])
            for r in read_rows(budget_path)
            if len(r) >= 2 and parse_amount(r[1])
        }
        bd = wb.create_sheet("预算对比")
        bd.append(["类别", "月预算", "月均实际", "差额", "执行率"])
        for c, b in budget.items():
            actual = sum(by_mc[(m, c)] for m in months) / max(len(months), 1)
            bd.append(
                [c, b, round(actual, 2), round((b or 0) - actual, 2), actual / b if b else None]
            )
            bd.cell(bd.max_row, 5).number_format = "0.0%"
            if b and actual > b:
                budget_lines.append(f"{c} 月均 {actual:,.0f} 元，超出预算 {b:,.0f} 元")
        _style_header(bd)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    total_inc, total_exp = sum(inc_m.values()), sum(exp_m.values())
    print(f"已生成 {out}（{months[0]} ~ {months[-1]}，{len(months)} 个月）")
    print(f"- 收入 {total_inc:,.2f} 元，消费支出 {total_exp:,.2f} 元，结余 {total_inc - total_exp:,.2f} 元"
          + (f"，储蓄率 {(total_inc - total_exp) / total_inc:.1%}" if total_inc else ""))  # fmt: skip
    print(
        "- 支出前三："
        + "、".join(f"{c} {sum(by_mc[(m, c)] for m in months):,.0f}" for c in cats[:3])
    )
    if subs:
        print(
            f"- 疑似订阅扣费 {len(subs)} 项，按月估算年费合计 {sum(a * 12 for _, a, _ in subs):,.0f} 元"
        )
    if big:
        print(f"- 大额异常支出 {len(big)} 笔，见“大额支出”表")
    for line in budget_lines:
        print(f"- 超预算：{line}")
    return 0


def _style_header(ws: Any) -> None:
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    ws.column_dimensions["A"].width = 16
    ws.freeze_panes = "B2"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    i = sub.add_parser("import")
    i.add_argument("files", nargs="+")
    i.add_argument("-o", "--output", required=True)
    i.add_argument("--rules", help="自定义分类规则 CSV：关键词,类别")
    i.add_argument("--merge", help="合并到已有账本（去重）")
    r = sub.add_parser("report")
    r.add_argument("ledger")
    r.add_argument("-o", "--output", required=True)
    r.add_argument("--budget", help="预算 CSV：类别,月预算")
    r.add_argument("--month", help="只看某月，如 2026-09")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "import":
            return cmd_import([Path(f) for f in args.files], Path(args.output),
                              Path(args.rules) if args.rules else None, Path(args.merge) if args.merge else None)  # fmt: skip
        return cmd_report(
            Path(args.ledger),
            Path(args.output),
            Path(args.budget) if args.budget else None,
            args.month,
        )
    except (OSError, LedgerError, ValueError, KeyError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
