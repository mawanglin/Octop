#!/usr/bin/env python3
"""投流素材测试：生成"钩子 × 卖点 × 行动引导"的测试矩阵和统一的素材命名；导入投放后台导出的数据，按素材和各维度复盘，给出放量、关停、改稿建议。

Usage:
    python scripts/ad_test.py matrix --product 小智 --hooks 痛点,数字,反常识 --points 省时,降本 \
        --ctas 免费试用,预约演示 [--durations 30] [--max 24] [--date 1022] -o 测试矩阵.xlsx
    python scripts/ad_test.py review 投放数据.csv|xlsx [--fields 日期,产品,钩子,卖点,引导,时长,序号]
        [--min-spend 100] [-o 复盘.xlsx] [--json]

matrix  素材命名为 "日期_产品_钩子_卖点_引导_时长_序号"，例如 1022_小智_痛点_省时_免费试用_30s_01；
        组合数超过 --max 时均匀抽取，保证每个钩子、卖点、引导都被测到。
review  自动识别列名（素材/创意名称、展示/曝光、点击、消耗/花费、转化/表单/注册/激活/线索），
        计算点击率、转化率、单次转化成本、千次展示成本；按素材名里的各段（钩子、卖点……）汇总，
        看哪类开头、哪类卖点更好。消耗低于 --min-spend 的素材标为"数据不足"，不下结论。
Exit codes: 0 ok, 1 输入错误。
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import sys
from collections import Counter, defaultdict
from datetime import date
from itertools import product
from pathlib import Path
from typing import Any

ALIASES = {
    "name": [
        "素材名称",
        "创意名称",
        "视频名称",
        "广告名称",
        "素材",
        "创意",
        "广告",
        "name",
        "creative",
    ],
    "impressions": ["展示数", "展现数", "曝光量", "曝光数", "曝光", "展示", "展现", "impressions"],
    "clicks": ["点击数", "点击量", "点击", "clicks"],
    "cost": [
        "消耗",
        "花费",
        "总花费",
        "消耗(元)",
        "消耗（元）",
        "花费(元)",
        "花费（元）",
        "cost",
        "spend",
    ],
    "conversions": [
        "转化数",
        "转化量",
        "转化",
        "表单数",
        "表单提交",
        "注册数",
        "激活数",
        "线索数",
        "下载数",
        "conversions",
    ],
}


class AdError(ValueError):
    pass


# ------------------------------------------------------------------- matrix


def split_list(raw: str) -> list[str]:
    return [x.strip() for x in raw.replace("，", ",").split(",") if x.strip()]


def balanced_pick(combos: list[tuple[str, ...]], limit: int) -> list[tuple[str, ...]]:
    """Greedy sample: prefer combos that test unseen levels, then the least-used ones."""
    used: list[Counter[str]] = [Counter() for _ in combos[0]]
    remaining, picked = list(combos), []
    while remaining and len(picked) < limit:
        best = max(
            remaining,
            key=lambda c: (
                sum(used[d][v] == 0 for d, v in enumerate(c)),
                -sum(used[d][v] for d, v in enumerate(c)),
            ),
        )
        remaining.remove(best)
        picked.append(best)
        for d, v in enumerate(best):
            used[d][v] += 1
    return picked


def matrix(
    product_name: str,
    hooks: list[str],
    points: list[str],
    ctas: list[str],
    durations: list[str],
    limit: int,
    day: str,
) -> list[dict[str, Any]]:
    if not hooks or not points or not ctas:
        raise AdError("--hooks、--points、--ctas 都至少要有一项")
    combos = list(product(hooks, points, ctas, durations or ["30"]))
    if len(combos) > limit:
        combos = balanced_pick(combos, limit)
    rows = []
    for i, (hook, point, cta, dur) in enumerate(combos, 1):
        dur_s = dur if dur.endswith("s") else f"{dur}s"
        rows.append(
            {
                "序号": i,
                "素材命名": f"{day}_{product_name}_{hook}_{point}_{cta}_{dur_s}_{i:02d}",
                "钩子类型": hook,
                "核心卖点": point,
                "行动引导": cta,
                "时长": dur_s,
                "口播稿": "",
                "制作状态": "待写",
                "上线日期": "",
                "备注": "",
            }
        )
    return rows


# ------------------------------------------------------------------- review


def read_table(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise AdError(f"找不到 {path}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        wb.close()
        header_idx = next(
            (i for i, r in enumerate(rows[:10]) if sum(1 for c in r if c not in (None, "")) >= 3), 0
        )
        header = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
        return [
            {h: ("" if v is None else str(v)) for h, v in zip(header, r, strict=False)}
            for r in rows[header_idx + 1 :]
        ]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise AdError("无法识别文件编码（支持 UTF-8 和 GBK）")
    return [
        {(k or "").strip(): (v or "").strip() for k, v in row.items()}
        for row in csv.DictReader(io.StringIO(text))
    ]


def find_columns(header: list[str]) -> dict[str, str]:
    found: dict[str, str] = {}
    lowered = {h.lower(): h for h in header}
    for key, names in ALIASES.items():
        for alias in names:
            if alias.lower() in lowered:
                found[key] = lowered[alias.lower()]
                break
        else:
            for h in header:
                if any(alias in h for alias in names[:3]):
                    found[key] = h
                    break
    missing = [
        k for k in ("name", "impressions", "clicks", "cost", "conversions") if k not in found
    ]
    if missing:
        label = {
            "name": "素材名称",
            "impressions": "展示",
            "clicks": "点击",
            "cost": "消耗",
            "conversions": "转化",
        }
        raise AdError(
            f"没有识别到这些列：{'、'.join(label[m] for m in missing)}；表头为：{'、'.join(header)}"
        )
    return found


def num(value: str) -> float:
    text = str(value).replace(",", "").replace("¥", "").replace("￥", "").replace("%", "").strip()
    try:
        return float(text) if text not in ("", "-", "--") else 0.0
    except ValueError:
        return 0.0


def metrics(imp: float, clk: float, cost: float, conv: float) -> dict[str, float | None]:
    return {
        "展示": imp,
        "点击": clk,
        "消耗": round(cost, 2),
        "转化": conv,
        "点击率": round(clk / imp, 4) if imp else None,
        "转化率": round(conv / clk, 4) if clk else None,
        "转化成本": round(cost / conv, 2) if conv else None,
        "千次展示成本": round(cost / imp * 1000, 2) if imp else None,
    }


def verdict(m: dict[str, Any], avg: dict[str, Any], min_spend: float) -> tuple[str, str]:
    if m["消耗"] < min_spend:
        return "数据不足", f"消耗不到 {min_spend:g} 元，继续观察"
    ctr, cvr, cpa = m["点击率"] or 0, m["转化率"] or 0, m["转化成本"]
    a_ctr, a_cvr, a_cpa = avg["点击率"] or 0, avg["转化率"] or 0, avg["转化成本"]
    if a_cpa and m["消耗"] >= 2 * a_cpa:
        if m["转化"] == 0:
            return "关停", f"已消耗 {m['消耗']:g} 元（超过平均转化成本 2 倍）仍无转化"
        if cpa is not None and cpa >= 2.5 * a_cpa:
            return "关停", f"转化成本 {cpa:g} 元，是平均的 {cpa / a_cpa:.1f} 倍"
    if cpa is not None and a_cpa and cpa <= 0.8 * a_cpa:
        return "放量", "转化成本低于平均 20% 以上：加预算，并围绕同一钩子和卖点衍生新版本"
    if a_ctr and ctr < 0.5 * a_ctr:
        return "改开头", "点击率不到平均一半：开头没抓住人，换钩子或前 3 秒画面"
    if a_ctr and a_cvr and ctr >= a_ctr and cvr < 0.6 * a_cvr:
        return "改承接", "点得多、转化少：口播承诺和落地页不一致，或人群不准，检查行动引导和落地页"
    if a_ctr and a_cvr and ctr < a_ctr and cvr >= 1.2 * a_cvr:
        return "改开头", "转化率高但点击率偏低：人群精准，开头不够吸引，保留中后段只换开头"
    return "观察", "表现接近平均，继续积累数据"


def review(rows: list[dict[str, str]], fields: list[str], min_spend: float) -> dict[str, Any]:
    if not rows:
        raise AdError("表格里没有数据行")
    cols = find_columns(list(rows[0]))
    agg: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
    for r in rows:
        name = r.get(cols["name"], "").strip()
        if not name or name in {"合计", "总计", "汇总"}:
            continue
        a = agg[name]
        for i, key in enumerate(("impressions", "clicks", "cost", "conversions")):
            a[i] += num(r.get(cols[key], ""))
    if not agg:
        raise AdError("没有有效的素材行")
    total = [sum(v[i] for v in agg.values()) for i in range(4)]
    avg = metrics(*total)
    items = []
    for name, (imp, clk, cost, conv) in agg.items():
        m = {"素材": name, **metrics(imp, clk, cost, conv)}
        m["结论"], m["建议"] = verdict(m, avg, min_spend)
        items.append(m)
    order = {"放量": 0, "观察": 1, "改开头": 2, "改承接": 3, "关停": 4, "数据不足": 5}
    items.sort(key=lambda m: (order[m["结论"]], m["转化成本"] is None, m["转化成本"] or 0))

    dims: dict[str, list[dict[str, Any]]] = {}
    parts = {name: name.split("_") for name in agg}
    width = max(len(p) for p in parts.values())
    labels = fields + [f"第{i + 1}段" for i in range(len(fields), width)]
    for idx, label in enumerate(labels[:width]):
        groups: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
        for name, segs in parts.items():
            if idx < len(segs):
                g = groups[segs[idx]]
                for i in range(4):
                    g[i] += agg[name][i]
        if 1 < len(groups) < len(agg):  # skip constant fields and per-material ids
            dims[label] = sorted(
                ({"取值": k, **metrics(*v)} for k, v in groups.items()),
                key=lambda m: (m["转化成本"] is None, m["转化成本"] or 0),
            )
    return {"columns": cols, "overall": avg, "materials": items, "dimensions": dims}


# ------------------------------------------------------------------- output


def write_xlsx(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    pct = {"点击率", "转化率"}
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        if not rows:
            continue
        cols = list(rows[0])
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="BE185D")
        for j, c in enumerate(cols, 1):
            letter = ws.cell(1, j).column_letter
            ws.column_dimensions[letter].width = max(
                10, min(48, max(len(str(r.get(c, ""))) for r in rows) * 1.3, len(c) * 2.2)
            )
            if c in pct:
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    row[0].number_format = "0.00%"
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def fmt_pct(v: float | None) -> str:
    return "-" if v is None else f"{v:.2%}"


def fmt_money(v: float | None) -> str:
    return "-" if v is None else f"{v:,.2f}"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    mx = sub.add_parser("matrix")
    mx.add_argument("--product", required=True, help="产品简称，用于素材命名")
    mx.add_argument("--hooks", required=True, help="钩子类型，逗号分隔")
    mx.add_argument("--points", required=True, help="核心卖点，逗号分隔")
    mx.add_argument("--ctas", required=True, help="行动引导，逗号分隔")
    mx.add_argument("--durations", default="30", help="时长（秒），逗号分隔")
    mx.add_argument("--max", type=int, default=24, help="最多生成多少条")
    mx.add_argument("--date", default=date.today().strftime("%m%d"))
    mx.add_argument("-o", "--output")
    rv = sub.add_parser("review")
    rv.add_argument("input")
    rv.add_argument(
        "--fields", default="日期,产品,钩子,卖点,引导,时长,序号", help="素材名按 _ 分段后各段的含义"
    )
    rv.add_argument("--min-spend", type=float, default=100, help="消耗低于此值不下结论")
    rv.add_argument("-o", "--output")
    rv.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "matrix":
            if args.max < 1:
                raise AdError("--max 必须大于 0")
            if "_" in args.product:
                raise AdError("产品简称里不能有下划线 _（命名用它分段）")
            rows = matrix(
                args.product,
                split_list(args.hooks),
                split_list(args.points),
                split_list(args.ctas),
                split_list(args.durations),
                args.max,
                args.date,
            )
            print(f"测试矩阵 {len(rows)} 条（命名规则：日期_产品_钩子_卖点_引导_时长_序号）")
            for r in rows:
                print(f"- {r['素材命名']}")
            if args.output:
                write_xlsx(Path(args.output), {"测试矩阵": rows})
                print(f"已保存：{args.output}")
            return 0
        result = review(read_table(Path(args.input)), split_list(args.fields), args.min_spend)
        if args.json:
            print(json.dumps(result, ensure_ascii=False, indent=2))
        else:
            o = result["overall"]
            print(
                f"整体：消耗 {fmt_money(o['消耗'])} 元，展示 {o['展示']:,.0f}，点击率 {fmt_pct(o['点击率'])}，转化率 {fmt_pct(o['转化率'])}，转化成本 {fmt_money(o['转化成本'])} 元"
            )
            print(f"\n素材排行（按结论和转化成本，{len(result['materials'])} 条）：")
            for m in result["materials"]:
                print(
                    f"- [{m['结论']}] {m['素材']}：消耗 {fmt_money(m['消耗'])}，点击率 {fmt_pct(m['点击率'])}，转化 {m['转化']:g}，转化成本 {fmt_money(m['转化成本'])} —— {m['建议']}"
                )
            for dim, groups in result["dimensions"].items():
                print(f"\n按「{dim}」汇总：")
                for g in groups:
                    print(
                        f"- {g['取值']}：消耗 {fmt_money(g['消耗'])}，点击率 {fmt_pct(g['点击率'])}，转化率 {fmt_pct(g['转化率'])}，转化成本 {fmt_money(g['转化成本'])}"
                    )
            print(
                "\n说明：样本量小时差异可能只是随机波动，单条素材至少积累 2–3 倍目标转化成本的消耗再下结论。"
            )
        if args.output:
            dim_rows = [{"维度": d, **g} for d, gs in result["dimensions"].items() for g in gs]
            write_xlsx(
                Path(args.output),
                {
                    "素材排行": result["materials"],
                    "维度汇总": dim_rows,
                    "整体": [result["overall"]],
                },
            )
            if not args.json:
                print(f"已保存：{args.output}")
    except AdError as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
