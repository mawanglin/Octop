#!/usr/bin/env python3
"""作品样本分析：批量读取文档和图片的元数据与正文，找出创作时间、作者、编辑软件、AI 生成痕迹等权属线索；比对样本之间的相似度。

Usage:
    python scripts/sample_inspect.py scan 文件或目录... [-o 样本报告.xlsx] [--author 张三] [--json]
    python scripts/sample_inspect.py compare 文件或目录... [--ref 参照作品] [-o 比对.xlsx] [--json]

scan 读取：
  图片  尺寸、格式、EXIF（拍摄时间、修改时间、设备、软件、作者 Artist、版权 Copyright、说明），
        PNG 文本块（Stable Diffusion / ComfyUI 生成参数等），XMP 作者和版权，C2PA 内容凭证、
        IPTC "AI 生成"来源类型标记
  文档  docx / pptx / xlsx 的作者、最后修改人、创建和修改时间、修订次数、字数与正文摘录；
        pdf 的 Author / Creator / Producer / 创建和修改时间、页数、正文摘录（无文字层时提示是扫描件）
  文本  txt / md 的字数和摘录；音频的标签（艺术家、标题、日期）
  全部  文件大小、文件修改时间、SHA-256
  线索  AI 生成痕迹、作者字段与 --author 不一致、修改时间早于创建时间、元数据被清空、扫描件
compare：图片用感知哈希（dHash，含镜像）找改色、缩放、轻微裁切后的近似图；
        文字作品用 8 字片段重合度比较，并列出最长的相同段落。只做初筛，是否构成侵权需专业判断。
只读，不修改任何文件。
Exit codes: 0 ok, 1 输入错误。
"""

from __future__ import annotations

import argparse
import contextlib
import difflib
import hashlib
import json
import re
import sys
import zipfile
from datetime import datetime
from itertools import combinations
from pathlib import Path
from typing import Any

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp", ".gif"}
DOC_EXT = {".docx", ".pptx", ".xlsx", ".pdf"}
TEXT_EXT = {".txt", ".md"}
AUDIO_EXT = {".mp3", ".flac", ".m4a", ".ogg", ".wav", ".aac"}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv"}

EXIF_TAGS = {
    0x9003: "拍摄时间",
    0x0132: "修改时间",
    0x010F: "设备厂商",
    0x0110: "设备型号",
    0x0131: "软件",
    0x013B: "作者",
    0x8298: "版权",
    0x010E: "说明",
}
AI_PATTERNS = re.compile(
    r"stable[\s-]?diffusion|comfyui|automatic1111|midjourney|dall[\s·-]?e|novelai|"
    r"firefly|leonardo|\bflux\b|sdxl|negative prompt|trainedAlgorithmicMedia|"
    r"compositeWithTrainedAlgorithmicMedia|文心一格|通义万相|即梦|可灵|豆包",
    re.I,
)
XP_TAGS = {0x9C9D: "作者", 0x9C9B: "标题", 0x9C9C: "说明"}  # Windows UTF-16 fields
DEFAULT_AUTHORS = {
    "python-docx",
    "administrator",
    "admin",
    "user",
    "microsoft office user",
    "microsoft",
    "unknown",
    "author",
    "微软用户",
    "用户",
    "lenovo",
    "dell",
    "hp",
}
PNG_AI_KEYS = {"parameters", "prompt", "workflow", "negative_prompt", "sd-metadata", "dream"}


class InspectError(ValueError):
    pass


def collect(paths: list[str]) -> list[Path]:
    out: list[Path] = []
    for raw in paths:
        p = Path(raw)
        if p.is_file():
            out.append(p)
        elif p.is_dir():
            for f in sorted(p.rglob("*")):
                supported = f.suffix.lower() in IMAGE_EXT | DOC_EXT | TEXT_EXT | AUDIO_EXT
                if supported and f.is_file() and not set(f.relative_to(p).parts[:-1]) & SKIP_DIRS:
                    out.append(f)
        else:
            raise InspectError(f"找不到 {raw}")
    if not out:
        raise InspectError("没有找到可分析的文件（支持图片、docx/pptx/xlsx/pdf、txt/md、常见音频）")
    return out


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bytes):
        value = value.decode("utf-8", "replace")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return re.sub(r"\s+", " ", str(value).replace("\x00", "")).strip()


def _exif_text(value: Any) -> str:
    """EXIF ASCII fields often carry UTF-8 that Pillow decodes as Latin-1."""
    if isinstance(value, str):
        with contextlib.suppress(UnicodeEncodeError, UnicodeDecodeError):
            value = value.encode("latin-1").decode("utf-8")
    return _clean(value)


def _excerpt(text: str, n: int = 120) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    return text[:n] + ("…" if len(text) > n else "")


def _pdf_date(raw: str) -> str:
    m = re.match(r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?", raw or "")
    if not m:
        return _clean(raw)
    parts = [p or "00" for p in m.groups()]
    parts[1] = parts[1] if parts[1] != "00" else "01"
    parts[2] = parts[2] if parts[2] != "00" else "01"
    return f"{parts[0]}-{parts[1]}-{parts[2]} {parts[3]}:{parts[4]}:{parts[5]}"


def inspect_image(path: Path, info: dict[str, Any]) -> None:
    from PIL import ExifTags, Image

    raw = path.read_bytes()
    with Image.open(path) as im:
        info["类型"] = "图片"
        info["尺寸"] = f"{im.width}×{im.height}"
        info["格式"] = im.format or ""
        exif = im.getexif()
        merged = dict(exif)
        with contextlib.suppress(KeyError, AttributeError):
            merged.update(exif.get_ifd(ExifTags.IFD.Exif))
        for tag, name in EXIF_TAGS.items():
            if merged.get(tag):
                value = _exif_text(merged[tag])
                info[name] = value.replace(":", "-", 2) if "时间" in name else value
        for tag, name in XP_TAGS.items():
            if isinstance(merged.get(tag), bytes) and not info.get(name):
                info[name] = _clean(merged[tag].decode("utf-16le", "replace"))
        texts = {
            k: _clean(v)
            for k, v in (getattr(im, "text", None) or im.info).items()
            if isinstance(v, (str, bytes)) and k not in ("exif", "icc_profile")
        }
    ai_keys = sorted(k for k in texts if k.lower() in PNG_AI_KEYS)
    if ai_keys:
        info["生成参数"] = _excerpt(" | ".join(f"{k}: {texts[k]}" for k in ai_keys), 300)
    for key in ("Author", "Copyright", "Software", "Description"):
        if texts.get(key) and not info.get(key):
            info.setdefault(
                {"Author": "作者", "Copyright": "版权", "Software": "软件", "Description": "说明"}[
                    key
                ],
                texts[key],
            )
    xmp = re.search(rb"<x:xmpmeta.*?</x:xmpmeta>", raw, re.S)
    if xmp:
        x = xmp.group(0).decode("utf-8", "replace")
        creator = re.search(r"<dc:creator>.*?<rdf:li[^>]*>(.*?)</rdf:li>", x, re.S)
        rights = re.search(r"<dc:rights>.*?<rdf:li[^>]*>(.*?)</rdf:li>", x, re.S)
        tool = re.search(r'xmp:CreatorTool(?:="|>)([^"<]+)', x)
        if creator and not info.get("作者"):
            info["作者"] = _clean(creator.group(1))
        if rights and not info.get("版权"):
            info["版权"] = _clean(rights.group(1))
        if tool and not info.get("软件"):
            info["软件"] = _clean(tool.group(1))
        if "DigitalSourceType" in x:
            m = re.search(r"digitalsourcetype/(\w+)", x, re.I)
            if m:
                info["来源类型"] = m.group(1)
    if b"c2pa" in raw:
        info["内容凭证"] = "含 C2PA 内容凭证"
    haystack = " ".join(str(info.get(k, "")) for k in ("软件", "生成参数", "来源类型", "说明"))
    ai_text = " ".join(texts.values())
    if ai_keys or AI_PATTERNS.search(haystack) or AI_PATTERNS.search(ai_text):
        info["AI痕迹"] = "是"


def _ooxml_core(path: Path) -> dict[str, str]:
    out: dict[str, str] = {}
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("docProps/core.xml").decode("utf-8", "replace")
            app = (
                zf.read("docProps/app.xml").decode("utf-8", "replace")
                if "docProps/app.xml" in zf.namelist()
                else ""
            )
    except (KeyError, zipfile.BadZipFile):
        return out
    fields = {
        "dc:creator": "作者",
        "cp:lastModifiedBy": "最后修改人",
        "dcterms:created": "创建时间",
        "dcterms:modified": "修改时间",
        "cp:revision": "修订次数",
        "dc:title": "标题",
    }
    for tag, name in fields.items():
        m = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", xml, re.S)
        if m and m.group(1).strip():
            value = m.group(1).strip()
            if "时间" in name:
                value = value.replace("T", " ").rstrip("Z")
            out[name] = value
    m = re.search(r"<Application>(.*?)</Application>", app)
    if m:
        out["软件"] = m.group(1)
    return out


def _doc_text(path: Path) -> tuple[str, int | None]:
    """Plain text and page/slide count (None when unknown)."""
    ext = path.suffix.lower()
    if ext == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts), None
    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n".join(parts), len(prs.slides)
    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=200, values_only=True):
                parts.extend(str(v) for v in row if v is not None)
        wb.close()
        return "\n".join(parts), None
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "") for page in reader.pages[:50]]
        return "\n".join(pages), len(reader.pages)
    return path.read_text(encoding="utf-8", errors="replace"), None


def inspect_doc(path: Path, info: dict[str, Any]) -> None:
    ext = path.suffix.lower()
    info["类型"] = {".docx": "Word", ".pptx": "PPT", ".xlsx": "Excel", ".pdf": "PDF"}.get(
        ext, "文本"
    )
    if ext in {".docx", ".pptx", ".xlsx"}:
        info.update(_ooxml_core(path))
    elif ext == ".pdf":
        from pypdf import PdfReader

        meta = PdfReader(str(path)).metadata or {}
        for key, name in (
            ("/Author", "作者"),
            ("/Creator", "软件"),
            ("/Producer", "生成工具"),
            ("/Title", "标题"),
        ):
            if meta.get(key):
                info[name] = _clean(meta.get(key))
        for key, name in (("/CreationDate", "创建时间"), ("/ModDate", "修改时间")):
            if meta.get(key):
                info[name] = _pdf_date(str(meta.get(key)))
    text, pages = _doc_text(path)
    if pages is not None:
        info["页数"] = pages
    info["字数"] = len(re.sub(r"\s", "", text))
    info["摘录"] = _excerpt(text)
    if ext == ".pdf" and info["字数"] < 20:
        info["提示"] = "没有文字层，可能是扫描件或图片 PDF：请用支持看图的模型直接阅读，或先做 OCR"


def inspect_audio(path: Path, info: dict[str, Any]) -> None:
    info["类型"] = "音频"
    try:
        import mutagen
    except ImportError:
        return
    try:
        audio = mutagen.File(str(path), easy=True)
    except Exception:  # noqa: BLE001 - damaged files are reported, not fatal
        info["提示"] = "无法读取音频标签"
        return
    if audio is None:
        return
    for key, name in (
        ("artist", "作者"),
        ("title", "标题"),
        ("date", "创建时间"),
        ("copyright", "版权"),
    ):
        if audio.get(key):
            info[name] = _clean(audio[key][0])
    if getattr(audio, "info", None) and getattr(audio.info, "length", None):
        info["时长"] = f"{audio.info.length:.0f} 秒"


def inspect_file(path: Path) -> dict[str, Any]:
    stat = path.stat()
    info: dict[str, Any] = {
        "文件": str(path),
        "大小": stat.st_size,
        "文件修改时间": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "SHA-256": sha256(path),
    }
    ext = path.suffix.lower()
    try:
        if ext in IMAGE_EXT:
            inspect_image(path, info)
        elif ext in DOC_EXT | TEXT_EXT:
            inspect_doc(path, info)
        elif ext in AUDIO_EXT:
            inspect_audio(path, info)
    except Exception as exc:  # noqa: BLE001 - one broken sample must not stop the batch
        info["提示"] = f"读取失败：{type(exc).__name__}: {exc}"
    return info


def findings(rows: list[dict[str, Any]], author: str | None) -> list[str]:
    out = []
    for r in rows:
        name = Path(r["文件"]).name
        if r.get("AI痕迹"):
            detail = (
                r.get("生成参数") or r.get("软件") or r.get("来源类型") or r.get("内容凭证") or ""
            )
            out.append(
                f"{name}：发现 AI 生成痕迹（{_excerpt(str(detail), 80)}）。登记和维权时需说明人的独创性投入"
            )
        for key in ("作者", "最后修改人"):
            value = str(r.get(key) or "")
            if value.lower() in DEFAULT_AUTHORS:
                out.append(f"{name}：{key}字段是软件默认值「{value}」，不能作为权属依据")
            elif author and value and author not in value:
                out.append(
                    f"{name}：{key}字段为「{value}」，与声称的作者「{author}」不一致，需核实或准备说明"
                )
        created, modified = str(r.get("创建时间") or ""), str(r.get("修改时间") or "")
        if created and modified and modified[:19] < created[:19]:
            out.append(
                f"{name}：修改时间（{modified}）早于创建时间（{created}），元数据可能被改动或转存过"
            )
        if r.get("类型") == "图片" and not any(
            r.get(k) for k in ("拍摄时间", "修改时间", "设备型号", "软件", "作者", "生成参数")
        ):
            out.append(
                f"{name}：图片没有任何元数据（可能经过平台压缩或被清除），创作时间需靠底稿、原图等其他证据"
            )
        if r.get("提示"):
            out.append(f"{name}：{r['提示']}")
    return out


# ------------------------------------------------------------------ compare


def dhash_bits(path: Path) -> tuple[int, int]:
    """64-bit difference hash of the image and of its mirror."""
    from PIL import Image, ImageOps

    with Image.open(path) as im:
        gray = ImageOps.exif_transpose(im).convert("L")
        results = []
        for img in (gray, ImageOps.mirror(gray)):
            small = img.resize((9, 8), Image.Resampling.LANCZOS)
            px = small.tobytes()
            bits = 0
            for row in range(8):
                for col in range(8):
                    bits = (bits << 1) | (px[row * 9 + col] > px[row * 9 + col + 1])
            results.append(bits)
    return results[0], results[1]


def norm_text(text: str) -> str:
    return re.sub(r"[\W_]+", "", text.lower())


def shingles(text: str, k: int = 8) -> set[str]:
    return {text[i : i + k] for i in range(max(len(text) - k + 1, 0))}


def common_passages(a: str, b: str, min_len: int = 20, limit: int = 5) -> list[str]:
    sm = difflib.SequenceMatcher(None, a[:20000], b[:20000], autojunk=False)
    blocks = sorted(
        (m for m in sm.get_matching_blocks() if m.size >= min_len), key=lambda m: -m.size
    )
    return [a[m.a : m.a + m.size] for m in blocks[:limit]]


def image_level(distance: int) -> str:
    return "高度相似" if distance <= 10 else "可能相似" if distance <= 18 else "不相似"


def text_level(score: float) -> str:
    return "高度相似" if score >= 0.5 else "部分相似" if score >= 0.15 else "不相似"


def compare(files: list[Path], ref: Path | None) -> list[dict[str, Any]]:
    images = [f for f in files if f.suffix.lower() in IMAGE_EXT]
    texts = [f for f in files if f.suffix.lower() in DOC_EXT | TEXT_EXT]
    if len(images) + len(texts) > 300:
        raise InspectError("一次最多比对 300 个文件，请缩小范围或用 --ref 指定参照作品")
    hashes = {
        f: dhash_bits(f)
        for f in images + ([ref] if ref and ref.suffix.lower() in IMAGE_EXT else [])
    }
    norm: dict[Path, str] = {}
    for f in texts + ([ref] if ref and ref.suffix.lower() in DOC_EXT | TEXT_EXT else []):
        norm[f] = norm_text(_doc_text(f)[0])

    def pairs(group: list[Path]) -> list[tuple[Path, Path]]:
        if ref:
            same = (
                ref.suffix.lower() in IMAGE_EXT
                if group is images
                else ref.suffix.lower() in DOC_EXT | TEXT_EXT
            )
            return [(ref, f) for f in group if f.resolve() != ref.resolve()] if same else []
        return list(combinations(group, 2))

    rows: list[dict[str, Any]] = []
    for a, b in pairs(images):
        (ha, _), (hb, hb_mirror) = hashes[a], hashes[b]
        direct, mirrored = bin(ha ^ hb).count("1"), bin(ha ^ hb_mirror).count("1")
        dist = min(direct, mirrored)
        rows.append(
            {
                "类型": "图片",
                "作品A": str(a),
                "作品B": str(b),
                "相似度": round(1 - dist / 64, 3),
                "结论": image_level(dist),
                "说明": f"哈希距离 {dist}/64" + ("（镜像后）" if mirrored < direct else ""),
            }
        )
    for a, b in pairs(texts):
        sa, sb = shingles(norm[a]), shingles(norm[b])
        if not sa or not sb:
            continue
        overlap = len(sa & sb)
        jaccard = overlap / len(sa | sb)
        contain = overlap / min(len(sa), len(sb))
        passages = common_passages(norm[a], norm[b]) if contain >= 0.05 else []
        rows.append(
            {
                "类型": "文字",
                "作品A": str(a),
                "作品B": str(b),
                "相似度": round(contain, 3),
                "结论": text_level(contain),
                "说明": f"片段重合 {contain:.0%}（较短一方被覆盖的比例），整体重合 {jaccard:.0%}"
                + (f"；最长相同段落：{_excerpt(passages[0], 60)}" if passages else ""),
                "相同段落": passages,
            }
        )
    rows.sort(key=lambda r: -r["相似度"])
    return rows


# ------------------------------------------------------------------- output


def write_xlsx(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        columns: list[str] = []
        for r in rows:
            columns += [k for k in r if k not in columns]
        ws.append(columns)
        for r in rows:
            ws.append(
                [
                    ("\n".join(v) if isinstance(v, list) else v)
                    for v in (r.get(c, "") for c in columns)
                ]
            )
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="7C2D12")
        for i, col in enumerate(columns, 1):
            width = max([len(str(col))] + [min(len(str(r.get(col, ""))), 60) for r in rows])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(8, width * 1.2)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    scan = sub.add_parser("scan")
    scan.add_argument("paths", nargs="+")
    scan.add_argument("--author", help="声称的作者或单位名称，用于核对元数据")
    cmp_p = sub.add_parser("compare")
    cmp_p.add_argument("paths", nargs="+")
    cmp_p.add_argument("--ref", help="参照作品：只比较它和其他文件")
    for p in (scan, cmp_p):
        p.add_argument("-o", "--output", help="输出 Excel 报告")
        p.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    try:
        files = collect(args.paths)
        if args.cmd == "scan":
            rows = [inspect_file(f) for f in files]
            notes = findings(rows, args.author)
            if args.json:
                print(
                    json.dumps({"samples": rows, "findings": notes}, ensure_ascii=False, indent=2)
                )
            else:
                print(f"共 {len(rows)} 个样本")
                for r in rows:
                    keys = [
                        k
                        for k in (
                            "类型",
                            "尺寸",
                            "页数",
                            "字数",
                            "作者",
                            "最后修改人",
                            "拍摄时间",
                            "创建时间",
                            "修改时间",
                            "软件",
                            "版权",
                            "AI痕迹",
                        )
                        if r.get(k) not in (None, "")
                    ]
                    print(f"- {Path(r['文件']).name}：" + "，".join(f"{k} {r[k]}" for k in keys))
                print(f"\n线索 {len(notes)} 条：" if notes else "\n没有发现异常线索。")
                for n in notes:
                    print(f"- {n}")
            if args.output:
                write_xlsx(
                    Path(args.output), {"样本清单": rows, "线索提示": [{"线索": n} for n in notes]}
                )
                print(f"\n报告：{args.output}")
        else:
            ref = Path(args.ref) if args.ref else None
            if ref and not ref.is_file():
                raise InspectError(f"找不到参照作品 {args.ref}")
            rows = compare(files, ref)
            if not rows:
                raise InspectError("没有可以比对的同类文件（图片和图片比、文字和文字比，至少两个）")
            if args.json:
                print(json.dumps(rows, ensure_ascii=False, indent=2))
            else:
                for r in rows:
                    print(
                        f"- [{r['结论']}] {Path(r['作品A']).name} ↔ {Path(r['作品B']).name}：{r['说明']}"
                    )
                print(
                    "\n说明：相似度只用于初筛，是否构成侵权需结合独创性、接触可能性等由专业人士判断。"
                )
            if args.output:
                write_xlsx(Path(args.output), {"比对结果": rows})
                print(f"报告：{args.output}")
    except InspectError as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
