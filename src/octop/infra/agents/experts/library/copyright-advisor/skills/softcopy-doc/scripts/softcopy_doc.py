#!/usr/bin/env python3
"""软著源代码材料：从代码目录生成计算机软件著作权登记用的源程序鉴别材料 Word 文档。

Usage:
    python scripts/softcopy_doc.py 源码目录 --name 软件全称 [--version V1.0] -o 源程序.docx
                                   [--ext .py,.js] [--exclude "tests/*"] [--order 文件顺序.txt]
                                   [--lines-per-page 50] [--pages 30] [--keep-blank] [--strip-comments]

规则（《计算机软件著作权登记办法》的常见要求，提交前以中国版权保护中心最新说明为准）：
  - 源程序共不超过 60 页时全部提交；超过时提交前 30 页和后 30 页（连续），可用 --pages 调整；
  - 每页 50 行（--lines-per-page），每页页眉为"软件名称 + 版本号"，右侧为页码；
  - 默认去掉空行；超长行按页面宽度折行并计入行数，保证每页正好 50 行；
  - 自动跳过依赖目录、构建产物、压缩过的代码（*.min.js）和锁文件。
生成前会检查代码里疑似密码、密钥的内容并提示：提交材料会被存档，务必先去除。
Exit codes: 0 ok, 1 输入错误。
"""

from __future__ import annotations

import argparse
import fnmatch
import re
import sys
import unicodedata
from pathlib import Path

DEFAULT_EXT = [
    ".py",
    ".java",
    ".kt",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".vue",
    ".go",
    ".rs",
    ".c",
    ".h",
    ".cc",
    ".cpp",
    ".hpp",
    ".cs",
    ".php",
    ".rb",
    ".swift",
    ".m",
    ".scala",
    ".dart",
    ".lua",
    ".sql",
    ".sh",
    ".html",
    ".css",
    ".scss",
    ".less",
    ".xml",
    ".wxml",
    ".wxss",
]
SKIP_DIRS = {
    ".git",
    "node_modules",
    ".venv",
    "venv",
    "__pycache__",
    "dist",
    "build",
    "target",
    "out",
    ".idea",
    ".vscode",
    ".next",
    "coverage",
    "vendor",
    "Pods",
    ".gradle",
    "bin",
    "obj",
}
SKIP_FILES = (
    "*.min.js",
    "*.min.css",
    "*.lock",
    "package-lock.json",
    "*.map",
    "*.pb.go",
    "*_pb2.py",
)
HASH_COMMENT = {".py", ".sh", ".rb"}
SLASH_COMMENT = {
    ".java",
    ".kt",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".go",
    ".rs",
    ".c",
    ".h",
    ".cc",
    ".cpp",
    ".hpp",
    ".cs",
    ".php",
    ".swift",
    ".m",
    ".scala",
    ".dart",
}
SECRET = re.compile(
    r"""(?i)(password|passwd|pwd|secret|api[_-]?key|access[_-]?key|token|private[_-]?key)\s*[:=]\s*['"][^'"\s]{6,}['"]"""
    r"|-----BEGIN [A-Z ]*PRIVATE KEY-----|AKID[0-9A-Za-z]{16,}|sk-[A-Za-z0-9]{20,}"
)
COLUMNS = 90


class SoftcopyError(ValueError):
    pass


def collect(root: Path, exts: set[str], excludes: list[str], order: Path | None) -> list[Path]:
    if not root.is_dir():
        raise SoftcopyError(f"找不到目录 {root}")
    if order:
        names = [
            ln.strip()
            for ln in order.read_text(encoding="utf-8").splitlines()
            if ln.strip() and not ln.startswith("#")
        ]
        files = [root / n for n in names]
        missing = [n for n, f in zip(names, files, strict=True) if not f.is_file()]
        if missing:
            raise SoftcopyError(f"--order 中的文件不存在：{'、'.join(missing[:5])}")
        return files
    out = []
    for f in sorted(root.rglob("*")):
        rel = f.relative_to(root)
        if not f.is_file() or f.suffix.lower() not in exts or set(rel.parts[:-1]) & SKIP_DIRS:
            continue
        if any(fnmatch.fnmatch(f.name, p) for p in SKIP_FILES):
            continue
        if any(fnmatch.fnmatch(rel.as_posix(), p) for p in excludes):
            continue
        out.append(f)
    if not out:
        raise SoftcopyError(f"{root} 下没有找到源代码文件（扩展名：{' '.join(sorted(exts))}）")
    return out


def width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def wrap(line: str, columns: int = COLUMNS) -> list[str]:
    if width(line) <= columns:
        return [line]
    parts, cur, cur_w = [], "", 0
    for ch in line:
        w = 2 if unicodedata.east_asian_width(ch) in "WF" else 1
        if cur_w + w > columns:
            parts.append(cur)
            cur, cur_w = "", 0
        cur += ch
        cur_w += w
    if cur:
        parts.append(cur)
    return parts


def source_lines(
    files: list[Path], keep_blank: bool, strip_comments: bool
) -> tuple[list[str], list[str]]:
    lines: list[str] = []
    warnings: list[str] = []
    for f in files:
        try:
            text = f.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = f.read_bytes().decode("gb18030", errors="replace")
        ext = f.suffix.lower()
        in_block = False
        for no, raw in enumerate(text.splitlines(), 1):
            line = raw.expandtabs(4).rstrip()
            if SECRET.search(line):
                warnings.append(f"{f.name}:{no}: {line.strip()[:60]}")
            stripped = line.strip()
            if strip_comments:
                if ext in SLASH_COMMENT:
                    if in_block:
                        in_block = "*/" not in stripped
                        continue
                    if stripped.startswith("/*"):
                        in_block = "*/" not in stripped
                        continue
                    if stripped.startswith("//"):
                        continue
                if (
                    ext in HASH_COMMENT
                    and stripped.startswith("#")
                    and not stripped.startswith("#!")
                ):
                    continue
            if not stripped and not keep_blank:
                continue
            if len(line) > 2000:
                continue  # minified or generated content
            lines.extend(wrap(line))
    return lines, warnings


def select_pages(lines: list[str], per_page: int, pages_each: int) -> tuple[list[list[str]], str]:
    pages = [lines[i : i + per_page] for i in range(0, len(lines), per_page)]
    total = len(pages)
    if total <= pages_each * 2:
        return pages, f"共 {total} 页，全部提交"
    # Back part: the last N full pages, counted from the end so every page holds per_page lines.
    tail_lines = lines[-pages_each * per_page :]
    tail = [tail_lines[i : i + per_page] for i in range(0, len(tail_lines), per_page)]
    return pages[:pages_each] + tail, f"共 {total} 页，提交前 {pages_each} 页和后 {pages_each} 页"


def build_docx(pages: list[list[str]], name: str, version: str, output: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.6)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(1.5)
    sec.header_distance = Cm(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(9)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = style.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(13)

    header = sec.header.paragraphs[0]
    header.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_TAB_ALIGNMENT.RIGHT)
    header.add_run(f"{name} {version}\t第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    header._p.append(fld)
    header.add_run(" 页")

    body = doc.paragraphs[0] if doc.paragraphs else None
    first = True
    for index, page in enumerate(pages):
        for i, line in enumerate(page):
            para = body if first and body is not None else doc.add_paragraph()
            first = False
            run = para.add_run(line)
            if i == len(page) - 1 and index < len(pages) - 1:
                run.add_break(WD_BREAK.PAGE)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("source")
    parser.add_argument("--name", required=True, help="软件全称，须与申请表一致")
    parser.add_argument("--version", default="V1.0")
    parser.add_argument("-o", "--output", required=True)
    parser.add_argument("--ext", help="只收录这些扩展名，逗号分隔")
    parser.add_argument(
        "--exclude", action="append", default=[], help="排除的相对路径通配符，可多次使用"
    )
    parser.add_argument("--order", help="文件顺序清单（每行一个相对路径），入口文件放最前")
    parser.add_argument("--lines-per-page", type=int, default=50)
    parser.add_argument("--pages", type=int, default=30, help="超过 2×N 页时前后各取 N 页")
    parser.add_argument("--keep-blank", action="store_true", help="保留空行")
    parser.add_argument("--strip-comments", action="store_true", help="去掉整行注释")
    args = parser.parse_args(argv)
    try:
        if args.lines_per_page < 1 or args.pages < 1:
            raise SoftcopyError("--lines-per-page 和 --pages 必须大于 0")
        exts = (
            {
                e.strip().lower() if e.strip().startswith(".") else "." + e.strip().lower()
                for e in args.ext.split(",")
            }
            if args.ext
            else set(DEFAULT_EXT)
        )
        root = Path(args.source)
        files = collect(root, exts, args.exclude, Path(args.order) if args.order else None)
        lines, warnings = source_lines(files, args.keep_blank, args.strip_comments)
        if not lines:
            raise SoftcopyError("源代码去掉空行后没有内容")
        pages, note = select_pages(lines, args.lines_per_page, args.pages)
        build_docx(pages, args.name, args.version, Path(args.output))
    except SoftcopyError as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    print(f"已生成：{args.output}")
    print(
        f"收录 {len(files)} 个文件，{len(lines)} 行（每页 {args.lines_per_page} 行），{note}，文档共 {len(pages)} 页"
    )
    print(f"页眉：{args.name} {args.version}　（须与申请表中的软件全称和版本号完全一致）")
    print(
        f"首个文件：{files[0].relative_to(root).as_posix()}（入口文件建议放在最前，可用 --order 指定顺序）"
    )
    if warnings:
        print(f"\n[注意] 发现 {len(warnings)} 处疑似密码或密钥，提交前请删除或替换：")
        for w in warnings[:20]:
            print(f"  - {w}")
    print("\n提示：用 Word 打开后核对页数和每页行数（页面视图），确认无误再转 PDF 提交。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
