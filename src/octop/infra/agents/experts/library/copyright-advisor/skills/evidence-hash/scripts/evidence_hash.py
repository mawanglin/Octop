#!/usr/bin/env python3
"""存证清单：给作品文件逐个计算 SHA-256 指纹，生成存证清单和清单总指纹；之后随时核验文件是否被改动。

Usage:
    python scripts/evidence_hash.py make 文件或目录... -o 存证清单.csv [--owner 权利人] [--title 作品名称] [--xlsx]
    python scripts/evidence_hash.py verify 存证清单.csv|存证清单.sha256 [--root 目录]

make   输出三个文件（同名不同后缀）：
         .csv     清单（序号、相对路径、大小、SHA-256、文件修改时间），Excel 可直接打开
         .sha256  标准格式（与 sha256sum -c / certutil 兼容），可脱离本工具核验
         .xlsx    （--xlsx）带标题、权利人和清单总指纹的表格，便于打印附在存证材料后
       "清单总指纹"是 .sha256 文件本身的 SHA-256：只需把这一个值提交到可信时间戳、区块链存证
       平台或公证，即可锁定清单里所有文件在该时间点的状态。
verify 重新计算并逐个比对，列出一致、被修改、缺失的文件；--root 下有清单外的新文件也会列出。
指纹只能证明"文件没有被改过"，证明"在某个时间已经存在"还需要时间戳、存证平台或公证。
不修改任何作品文件。
Exit codes: 0 全部一致 / 生成成功, 3 有文件被修改或缺失, 1 输入错误。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import sys
from datetime import datetime
from pathlib import Path

SKIP_NAMES = {".DS_Store", "Thumbs.db", "desktop.ini"}
SKIP_DIRS = {".git", "__pycache__", "node_modules"}


class EvidenceError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def collect(paths: list[str], exclude: set[Path]) -> tuple[Path, list[Path]]:
    """Common root and the files under it (sorted, manifest outputs excluded)."""
    files: list[Path] = []
    roots: list[Path] = []
    for raw in paths:
        p = Path(raw).resolve()
        if p.is_file():
            files.append(p)
            roots.append(p.parent)
        elif p.is_dir():
            roots.append(p)
            for f in sorted(p.rglob("*")):
                rel_parts = set(f.relative_to(p).parts[:-1])
                if f.is_file() and f.name not in SKIP_NAMES and not rel_parts & SKIP_DIRS:
                    files.append(f)
        else:
            raise EvidenceError(f"找不到 {raw}")
    files = sorted({f for f in files if f not in exclude})
    if not files:
        raise EvidenceError("没有找到要存证的文件")
    return Path(os.path.commonpath([str(r) for r in roots])), files


def make(paths: list[str], output: Path, owner: str, title: str, xlsx: bool) -> str:
    out = output.resolve()
    sha_path = out.with_suffix(".sha256")
    xlsx_path = out.with_suffix(".xlsx")
    root, files = collect(paths, {out, sha_path, xlsx_path})
    rows = []
    for i, f in enumerate(files, 1):
        stat = f.stat()
        rows.append(
            {
                "序号": i,
                "文件": f.relative_to(root).as_posix(),
                "大小（字节）": stat.st_size,
                "SHA-256": sha256_file(f),
                "文件修改时间": datetime.fromtimestamp(stat.st_mtime)
                .astimezone()
                .isoformat(timespec="seconds"),
            }
        )
    out.parent.mkdir(parents=True, exist_ok=True)
    sha_text = "".join(f"{r['SHA-256']}  {r['文件']}\n" for r in rows)
    sha_path.write_text(sha_text, encoding="utf-8", newline="\n")
    total = hashlib.sha256(sha_path.read_bytes()).hexdigest()
    made_at = datetime.now().astimezone().isoformat(timespec="seconds")
    with out.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
        fh.write(
            f"\n# 根目录,{root}\n# 生成时间,{made_at}\n# 清单总指纹（.sha256 文件的 SHA-256）,{total}\n"
        )
        if owner:
            fh.write(f"# 权利人,{owner}\n")
        if title:
            fh.write(f"# 作品名称,{title}\n")
    if xlsx:
        write_xlsx(xlsx_path, rows, root, made_at, total, owner, title)
    return total


def write_xlsx(
    path: Path,
    rows: list[dict[str, object]],
    root: Path,
    made_at: str,
    total: str,
    owner: str,
    title: str,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "存证清单"
    ws.append([f"作品存证文件清单{('：' + title) if title else ''}"])
    ws["A1"].font = Font(bold=True, size=14)
    meta = [
        ("权利人", owner),
        ("生成时间", made_at),
        ("文件数量", len(rows)),
        ("清单总指纹（SHA-256）", total),
        ("根目录", str(root)),
    ]
    for k, v in meta:
        if v not in ("", None):
            ws.append([k, v])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(list(rows[0]))
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(list(r.values()))
    for col, width in zip("ABCDE", (8, 48, 14, 68, 28), strict=True):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def read_manifest(path: Path) -> tuple[dict[str, str], Path | None]:
    entries: dict[str, str] = {}
    root: Path | None = None
    if path.suffix.lower() == ".sha256":
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                digest, _, name = line.partition("  ")
                entries[name.lstrip("*")] = digest.strip().lower()
        return entries, None
    with path.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.reader(fh):
            if not row:
                continue
            if row[0] == "# 根目录" and len(row) > 1:
                root = Path(row[1])
            elif row[0].isdigit() and len(row) >= 4:
                entries[row[1]] = row[3].strip().lower()
    if not entries:
        raise EvidenceError(f"{path} 里没有读到清单条目")
    return entries, root


def verify(manifest: Path, root_arg: str | None) -> int:
    if not manifest.is_file():
        raise EvidenceError(f"找不到清单 {manifest}")
    entries, recorded_root = read_manifest(manifest)
    root = (
        Path(root_arg)
        if root_arg
        else (recorded_root if recorded_root and recorded_root.is_dir() else manifest.parent)
    )
    ok, changed, missing = [], [], []
    for name, digest in entries.items():
        f = root / name
        if not f.is_file():
            missing.append(name)
        elif sha256_file(f) != digest:
            changed.append(name)
        else:
            ok.append(name)
    extra: list[str] = []
    if root_arg:
        skip = {
            manifest.resolve(),
            manifest.with_suffix(".sha256").resolve(),
            manifest.with_suffix(".csv").resolve(),
            manifest.with_suffix(".xlsx").resolve(),
        }
        for f in sorted(Path(root_arg).rglob("*")):
            rel = f.relative_to(root_arg)
            if (
                f.is_file()
                and f.resolve() not in skip
                and f.name not in SKIP_NAMES
                and not set(rel.parts[:-1]) & SKIP_DIRS
            ) and rel.as_posix() not in entries:
                extra.append(rel.as_posix())
    print(f"核验目录：{root}")
    print(
        f"一致 {len(ok)}，被修改 {len(changed)}，缺失 {len(missing)}"
        + (f"，清单外文件 {len(extra)}" if root_arg else "")
    )
    for name in changed:
        print(f"- [被修改] {name}")
    for name in missing:
        print(f"- [缺失] {name}")
    for name in extra:
        print(f"- [清单外] {name}")
    if not changed and not missing:
        print("全部文件与清单一致。")
    return 3 if changed or missing else 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    mk = sub.add_parser("make")
    mk.add_argument("paths", nargs="+")
    mk.add_argument("-o", "--output", required=True, help="清单文件名（.csv）")
    mk.add_argument("--owner", default="", help="权利人")
    mk.add_argument("--title", default="", help="作品名称")
    mk.add_argument("--xlsx", action="store_true", help="同时输出可打印的 Excel 清单")
    vf = sub.add_parser("verify")
    vf.add_argument("manifest")
    vf.add_argument("--root", help="作品所在目录，默认用清单里记录的根目录")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "make":
            out = Path(args.output)
            if out.suffix.lower() != ".csv":
                out = out.with_suffix(".csv")
            total = make(args.paths, out, args.owner, args.title, args.xlsx)
            print(
                f"已生成：{out}、{out.with_suffix('.sha256')}"
                + (f"、{out.with_suffix('.xlsx')}" if args.xlsx else "")
            )
            print(f"清单总指纹（提交到时间戳 / 存证平台 / 公证）：{total}")
            print("提示：生成后不要再修改作品文件和 .sha256 文件，否则指纹会对不上。")
            return 0
        return verify(Path(args.manifest), args.root)
    except EvidenceError as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
