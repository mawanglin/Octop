#!/usr/bin/env python3
"""新旧文件条文对比：按段落（条款）对齐，列出新增、删除、修改，修改处标出字词变化。

Usage:
    python scripts/text_diff.py 旧版.docx 新版.docx [-o 对比.md] [--xlsx 对比.xlsx]

支持 .docx / .md / .txt。输出 Markdown 对比表（修改处用【+新增+】【-删除-】标注），
可选同时输出 Excel 对比表。只做文字比对，条款含义的变化需要人工解读。
Exit codes: 0 ok, 1 读取失败。
"""

from __future__ import annotations

import argparse
import difflib
import re
import sys
from pathlib import Path


def read_paragraphs(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - environment check
            raise OSError("读取 .docx 需要 python-docx") from exc
        doc = Document(str(path))
        paras = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paras.append(" | ".join(c.text.strip() for c in row.cells))
    else:
        paras = path.read_text(encoding="utf-8").splitlines()
    return [re.sub(r"\s+", " ", p).strip() for p in paras if p.strip()]


def label(text: str) -> str:
    """Article number such as 第五条 / 一、 / 1. used to show where a change is."""
    m = re.match(
        r"^(第[一二三四五六七八九十百零〇\d]+[条章节款]|[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）|\d+[.．])",
        text,
    )
    return m.group(1) if m else ""


TOKEN = re.compile(r"\d+(?:[.,]\d+)*%?|[A-Za-z]+|\s+|.")


def inline_diff(old: str, new: str) -> str:
    """Token-level diff: numbers and words change as a whole (5000 → 8000, not 5 → 8)."""
    a, b = TOKEN.findall(old), TOKEN.findall(new)
    out: list[str] = []
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    for op, a1, a2, b1, b2 in sm.get_opcodes():
        removed, added = "".join(a[a1:a2]), "".join(b[b1:b2])
        if op == "equal":
            out.append(added)
        elif op == "insert":
            out.append(f"【+{added}+】")
        elif op == "delete":
            out.append(f"【-{removed}-】")
        else:
            out.append(f"【-{removed}-】【+{added}+】")
    return "".join(out)


def compare(old: list[str], new: list[str]) -> list[tuple[str, str, str, str]]:
    """Return (kind, location, old_text, new_text_or_marked) rows."""
    rows: list[tuple[str, str, str, str]] = []
    sm = difflib.SequenceMatcher(None, old, new, autojunk=False)
    for op, a1, a2, b1, b2 in sm.get_opcodes():
        if op == "equal":
            continue
        olds, news = old[a1:a2], new[b1:b2]
        if op == "replace":
            # Pair paragraphs that are still similar; the rest are pure adds/deletes.
            used: set[int] = set()
            for o in olds:
                best, score = None, 0.0
                for j, n in enumerate(news):
                    if j in used:
                        continue
                    ratio = difflib.SequenceMatcher(None, o, n, autojunk=False).ratio()
                    if ratio > score:
                        best, score = j, ratio
                if best is not None and score >= 0.45:
                    used.add(best)
                    n = news[best]
                    rows.append(("修改", label(n) or label(o), o, inline_diff(o, n)))
                else:
                    rows.append(("删除", label(o), o, ""))
            for j, n in enumerate(news):
                if j not in used:
                    rows.append(("新增", label(n), "", n))
        elif op == "delete":
            rows.extend(("删除", label(o), o, "") for o in olds)
        else:
            rows.extend(("新增", label(n), "", n) for n in news)
    return rows


def to_markdown(rows: list[tuple[str, str, str, str]], old_name: str, new_name: str) -> str:
    counts = {k: sum(1 for r in rows if r[0] == k) for k in ("新增", "删除", "修改")}
    lines = [
        f"# 条文对比：{old_name} → {new_name}",
        "",
        f"新增 {counts['新增']} 处，删除 {counts['删除']} 处，修改 {counts['修改']} 处。"
        "修改处【+…+】为新增文字，【-…-】为删除文字。",
        "",
        "| 序号 | 类型 | 位置 | 原文 | 新文 / 修改标注 |",
        "|------|------|------|------|-----------------|",
    ]
    for i, (kind, loc, old, new) in enumerate(rows, start=1):
        cell = lambda s: s.replace("|", "｜")  # noqa: E731 - tiny local helper
        lines.append(f"| {i} | {kind} | {loc} | {cell(old)} | {cell(new)} |")
    if not rows:
        lines.append("| — | 无差异 | | | |")
    return "\n".join(lines) + "\n"


def to_xlsx(rows: list[tuple[str, str, str, str]], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:  # pragma: no cover - environment check
        raise OSError("输出 Excel 需要 openpyxl") from exc
    fills = {"新增": "DCFCE7", "删除": "FEE2E2", "修改": "FEF3C7"}
    wb = Workbook()
    ws = wb.active
    ws.title = "条文对比"
    ws.append(["序号", "类型", "位置", "原文", "新文 / 修改标注", "解读（人工填写）"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    for i, (kind, loc, old, new) in enumerate(rows, start=1):
        ws.append([i, kind, loc, old, new, ""])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fills[kind])
    for col, width in zip("ABCDEF", (6, 8, 10, 50, 50, 40), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("old")
    parser.add_argument("new")
    parser.add_argument("-o", "--output", help="输出 Markdown（省略时打印）")
    parser.add_argument("--xlsx", help="同时输出 Excel 对比表")
    args = parser.parse_args(argv)
    try:
        old, new = read_paragraphs(Path(args.old)), read_paragraphs(Path(args.new))
        rows = compare(old, new)
        md = to_markdown(rows, Path(args.old).name, Path(args.new).name)
        if args.xlsx:
            to_xlsx(rows, Path(args.xlsx))
    except OSError as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    if args.output:
        out = Path(args.output)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(md, encoding="utf-8")
        print(md.split("\n")[2])
        print(f"已生成 {out}" + (f"，{args.xlsx}" if args.xlsx else ""))
    else:
        print(md)
    return 0


if __name__ == "__main__":
    sys.exit(main())
