#!/usr/bin/env python3
"""会议待办追踪表：Markdown 待办清单（或上次生成的 xlsx）→ 标出逾期和临期的 Excel，并打印跟进摘要。

Usage:
    python scripts/actions_xlsx.py actions.md -o 待办追踪.xlsx [--today 2026-10-18]
    python scripts/actions_xlsx.py 待办追踪.xlsx -o 待办追踪.xlsx --overwrite   # 用户改完状态后重新生成

Markdown 表头：
    待办 | 负责人 | 截止 | 状态 | 来源 | 备注
（"待办""负责人""截止"必填；状态写 未开始 / 进行中 / 已完成 / 已取消，空着视为未开始；
来源写会议名称和日期。）截止日期支持 2026-10-20、10/20、10月20日（无年份时按 --today 的年份）。
Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

COLUMNS = ["序号", "待办", "负责人", "截止", "状态", "剩余天数", "来源", "备注"]
DONE = {"已完成", "完成", "done", "已取消", "取消"}
FILLS = {
    "逾期": PatternFill("solid", fgColor="FEE2E2"),
    "临期": PatternFill("solid", fgColor="FEF3C7"),
    "完成": PatternFill("solid", fgColor="DCFCE7"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")


class InputError(ValueError):
    pass


def parse_date(value: object, today: dt.date) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip()
    if not text or text in ("—", "-", "待定"):
        return None
    m = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?", text)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.fullmatch(r"(\d{1,2})[-/.月](\d{1,2})日?", text)
    if m:
        return dt.date(today.year, int(m.group(1)), int(m.group(2)))
    raise InputError(f"截止日期「{text}」无法识别，请写成 2026-10-20")


def read_items(path: Path) -> list[dict[str, object]]:
    if path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, data_only=True)
        ws = wb["待办"] if "待办" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        start = next((i for i, r in enumerate(rows) if r and "待办" in [str(c) for c in r]), None)
        if start is None:
            raise InputError("xlsx 中没有找到「待办」表头")
        header = [str(c) if c is not None else "" for c in rows[start]]
        body = [r for r in rows[start + 1 :] if r and any(c not in (None, "") for c in r)]
    else:
        lines = path.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
        rows = [
            [c.strip() for c in ln.strip().strip("|").split("|")]
            for ln in lines
            if ln.strip().startswith("|") and ln.strip().endswith("|")
        ]
        rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
        if len(rows) < 2:
            raise InputError("没有找到待办表格（需要一张带表头的 Markdown 表格）")
        header, body = rows[0], rows[1:]
    for need in ("待办", "负责人", "截止"):
        if need not in header:
            raise InputError(f"表头缺少「{need}」列")
    items = []
    for r in body:
        rec = {h: (r[i] if i < len(r) else "") for i, h in enumerate(header)}
        if str(rec.get("待办") or "").strip() and str(rec.get("待办")) != "合计":
            items.append(rec)
    return items


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input")
    parser.add_argument("-o", "--output", required=True)
    parser.add_argument("--today", help="基准日期，默认今天")
    parser.add_argument("--soon", type=int, default=3, help="几天内到期算临期，默认 3")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args(argv)
    try:
        today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
        items = read_items(Path(args.input))
        out = Path(args.output)
        if out.exists() and not args.overwrite:
            raise InputError(f"{out} 已存在；确认要覆盖时加 --overwrite")
        rows = []
        for rec in items:
            due = parse_date(rec.get("截止"), today)
            status = str(rec.get("状态") or "").strip() or "未开始"
            done = status.lower() in DONE
            left = (due - today).days if due else None
            if done:
                flag = "完成"
            elif left is not None and left < 0:
                flag = "逾期"
            elif left is not None and left <= args.soon:
                flag = "临期"
            else:
                flag = ""
            rows.append((rec, due, status, left, flag))
    except (OSError, InputError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1

    order = {"逾期": 0, "临期": 1, "": 2, "完成": 3}
    rows.sort(key=lambda x: (order[x[4]], x[1] or dt.date.max))
    wb = Workbook()
    ws = wb.active
    ws.title = "待办"
    ws.cell(1, 1, f"会议待办追踪（统计日期 {today.isoformat()}）").font = Font(size=14, bold=True)
    for j, name in enumerate(COLUMNS, start=1):
        c = ws.cell(3, j, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    widths = [6, 40, 10, 12, 10, 10, 24, 24]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + j)].width = w
    for i, (rec, due, status, left, flag) in enumerate(rows, start=1):
        r = 3 + i
        values = [i, rec.get("待办"), rec.get("负责人"), due or "", status,
                  "" if left is None or flag == "完成" else left, rec.get("来源", ""), rec.get("备注", "")]  # fmt: skip
        for j, v in enumerate(values, start=1):
            c = ws.cell(r, j, v)
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 7, 8))
            if j == 4 and due:
                c.number_format = "yyyy-mm-dd"
            if flag:
                c.fill = FILLS[flag]
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:H{3 + len(rows)}"

    by_owner: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0])
    for rec, _, _, _, flag in rows:
        owner = str(rec.get("负责人") or "（未指定）")
        if flag != "完成":
            by_owner[owner][0] += 1
        if flag == "逾期":
            by_owner[owner][1] += 1
        if flag == "临期":
            by_owner[owner][2] += 1
    sm = wb.create_sheet("按负责人")
    sm.append(["负责人", "未完成", "逾期", "临期"])
    for owner, (open_, late, soon) in sorted(
        by_owner.items(), key=lambda kv: (-kv[1][1], -kv[1][0])
    ):
        sm.append([owner, open_, late, soon])
    sm.column_dimensions["A"].width = 16
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))

    late = [x for x in rows if x[4] == "逾期"]
    soon = [x for x in rows if x[4] == "临期"]
    done = sum(1 for x in rows if x[4] == "完成")
    print(
        f"已生成 {out}：共 {len(rows)} 项，已完成 {done}，逾期 {len(late)}，{args.soon} 天内到期 {len(soon)}"
    )
    for label, group in (("逾期", late), ("临期", soon)):
        if group:
            print(f"\n## {label}")
            for rec, due, status, left, _ in group:
                when = f"已逾期 {-left} 天" if left is not None and left < 0 else f"还剩 {left} 天"
                print(f"- {rec.get('待办')}｜{rec.get('负责人')}｜截止 {due}（{when}）｜{status}")
    unowned = [x for x in rows if not str(x[0].get("负责人") or "").strip() and x[4] != "完成"]
    undated = [x for x in rows if x[1] is None and x[4] != "完成"]
    if unowned:
        print(f"\n提示：{len(unowned)} 项没有负责人")
    if undated:
        print(f"提示：{len(undated)} 项没有截止日期")
    return 0


if __name__ == "__main__":
    sys.exit(main())
