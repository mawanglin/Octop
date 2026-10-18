#!/usr/bin/env python3
"""投标文件自查：响应矩阵是否填完、★ 条款有无负偏离；投标文件里的占位符、他司名称、编号、日期和金额大小写。

Usage:
    python scripts/bid_check.py matrix 响应矩阵.xlsx
    python scripts/bid_check.py doc 投标文件.docx --bidder "某某科技有限公司" [--project-no ZB-2026-0815]
                                [--other "某某公司,另一公司"] [--deadline 2026-11-20]

matrix：「实质性条款」表中 ★/否决/资格 条款未填写响应情况或页码、出现负偏离；评分标准自评合计。
doc：
  - 残留占位符：XX、××、***、【…】、___、待定、TBD、此处填写、（公司名称）等；
  - 出现其他公司名称（--other，如从旧标书复制时残留的名称）；投标人名称出现次数；
  - 项目编号与 --project-no 不一致的写法；
  - 落款日期晚于投标截止日（--deadline）或早 60 天以上（旧标书残留）、不存在的日期；
  - 金额大写与小写不一致（如“壹佰捌拾万元整（¥1,800,000.00）”）。
只读原文件。Exit codes: 0 无问题, 1 发现问题或输入错误, 2 缺少依赖。
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from decimal import Decimal
from pathlib import Path

PLACEHOLDER = re.compile(r"X{2,}|×{2,}|\*{3,}|＊{2,}|【[^】]{0,20}】|_{3,}|＿{2,}|待定|TBD|TODO|此处填写|在此填写|"
                         r"[（(](?:公司名称|投标人名称|项目名称|日期|金额)[)）]|XXXX")  # fmt: skip
DIGITS = "零壹贰叁肆伍陆柒捌玖"
UNITS = ["", "拾", "佰", "仟"]
BIG = ["", "万", "亿", "万亿"]
AMOUNT = re.compile(
    r"([零壹贰叁肆伍陆柒捌玖拾佰仟万亿元圆角分整正]{3,})\s*[（(]?\s*(?:人民币)?\s*[¥￥]\s*([\d,]+(?:\.\d{1,2})?)"
)


class CheckError(ValueError):
    pass


def to_upper(amount: Decimal) -> str:
    """Decimal yuan → 中文大写金额（财务写法），如 1800000 → 壹佰捌拾万元整，100500.05 → 壹拾万零伍佰元零伍分。"""
    amount = amount.quantize(Decimal("0.01"))
    yuan = int(amount)
    cents = int((amount - yuan) * 100)
    jiao, fen = cents // 10, cents % 10
    out = ""
    digits = str(yuan)
    zero = False
    for i, ch in enumerate(digits):
        d, pos = int(ch), len(digits) - 1 - i
        if d == 0:
            zero = True
        else:
            if zero:
                out += "零"
                zero = False
            out += DIGITS[d] + UNITS[pos % 4]
        if pos % 4 == 0 and pos > 0 and int(digits[max(0, i - 3) : i + 1]):
            out += BIG[pos // 4]
    if yuan:
        out += "元"
    if not cents:
        return (out or "零元") + "整"
    if jiao:
        out += DIGITS[jiao] + "角"
    elif yuan:
        out += "零"
    if fen:
        out += DIGITS[fen] + "分"
    return out


def _norm_upper(s: str) -> str:
    s = s.replace("圆", "元").replace("正", "整").replace("人民币", "")
    return s[:-1] if s.endswith("角整") else s  # "伍角整" and "伍角" are both accepted


def read_text(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError:  # pragma: no cover - environment check
            print("需要 python-docx：pip install python-docx", file=sys.stderr)
            sys.exit(2)
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for r in t.rows:
                lines.append(" ".join(dict.fromkeys(c.text for c in r.cells)))
        for sec in doc.sections:
            lines += [p.text for p in sec.header.paragraphs] + [
                p.text for p in sec.footer.paragraphs
            ]
        return [ln.strip() for ln in lines if ln.strip()]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return [ln.strip() for ln in raw.decode(enc).splitlines() if ln.strip()]
        except UnicodeDecodeError:
            continue
    raise CheckError(f"{path} 编码无法识别")


def check_doc(lines: list[str], bidder: str, project_no: str | None, others: list[str],
              deadline: dt.date | None) -> list[str]:  # fmt: skip
    issues: list[str] = []
    for i, ln in enumerate(lines, start=1):
        for m in PLACEHOLDER.finditer(ln):
            issues.append(f"第 {i} 段残留占位符「{m.group(0)}」：{ln[:40]}")
        for o in others:
            if o and o in ln:
                issues.append(f"第 {i} 段出现其他公司名称「{o}」：{ln[:40]}")
        for m in AMOUNT.finditer(ln):
            upper, num = _norm_upper(m.group(1)), Decimal(m.group(2).replace(",", ""))
            expect = to_upper(num)
            if upper != expect:
                issues.append(
                    f"第 {i} 段金额大小写不一致：大写「{m.group(1)}」，小写 ¥{m.group(2)} 应写作「{expect}」"
                )
    text = "\n".join(lines)
    if bidder:
        n = text.count(bidder)
        if n == 0:
            issues.append(f"全文没有出现投标人名称「{bidder}」，请检查名称是否与营业执照一致")
        else:
            short = re.sub(r"(股份)?有限(责任)?公司$", "", bidder)
            variants = {
                v for v in re.findall(re.escape(short) + r"[一-鿿]{0,8}公司", text) if v != bidder
            }
            if variants:
                issues.append(
                    f"投标人名称有不同写法：{'、'.join(sorted(variants))}（应统一为「{bidder}」）"
                )
    if project_no:
        for m in re.finditer(
            r"(?:项目编号|招标编号|采购编号)\s*[:：]?\s*([A-Za-z0-9\-—_（）()]+)", text
        ):
            if m.group(1) != project_no:
                issues.append(f"项目编号写成了「{m.group(1)}」，应为「{project_no}」")
    for y, mo, d in re.findall(
        r"((?:19|20)\d{2})\s*[年.\-/]\s*(\d{1,2})\s*[月.\-/]\s*(\d{1,2})", text
    ):
        try:
            dt.date(int(y), int(mo), int(d))
        except ValueError:
            issues.append(f"日期「{y}年{mo}月{d}日」不存在")
    if deadline:
        for ln in lines:
            sign_line = re.search(r"日期|落款|签署", ln) or re.fullmatch(
                r"(?:19|20)\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日", ln
            )
            if not sign_line:
                continue
            for y, mo, d in re.findall(
                r"((?:19|20)\d{2})\s*[年.\-/]\s*(\d{1,2})\s*[月.\-/]\s*(\d{1,2})", ln
            ):
                try:
                    when = dt.date(int(y), int(mo), int(d))
                except ValueError:
                    continue
                if when > deadline:
                    issues.append(
                        f"落款日期 {when.isoformat()} 晚于投标截止日 {deadline.isoformat()}：{ln[:30]}"
                    )
                elif (deadline - when).days > 60:
                    issues.append(
                        f"落款日期 {when.isoformat()} 比投标截止日早 {(deadline - when).days} 天，可能是旧标书残留：{ln[:30]}"
                    )
    return issues


def check_matrix(path: Path) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - environment check
        print("需要 openpyxl：pip install openpyxl", file=sys.stderr)
        sys.exit(2)
    wb = load_workbook(path, data_only=True)
    if "实质性条款" not in wb.sheetnames:
        raise CheckError("没有「实质性条款」工作表：请使用 tender_scan.py 生成的响应矩阵")
    issues: list[str] = []
    ws = wb["实质性条款"]
    header = [str(c.value or "") for c in ws[1]]
    col = {
        name: next((i for i, h in enumerate(header) if h.startswith(name)), None)
        for name in ("序号", "类型", "招标文件原文", "响应情况", "偏离", "投标文件页码")
    }
    if None in col.values():
        raise CheckError(
            "「实质性条款」表头被修改过，需要包含：序号、类型、招标文件原文、响应情况、偏离、投标文件页码"
        )
    total = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col["招标文件原文"]] in (None, ""):  # type: ignore[index]
            continue
        total += 1
        no, kind = row[col["序号"]], str(row[col["类型"]] or "")  # type: ignore[index]
        resp, dev, page = (
            str(row[col[k]] or "").strip() for k in ("响应情况", "偏离", "投标文件页码")
        )  # type: ignore[index]
        key = any(k in kind for k in ("★", "否决", "资格"))
        if "负偏离" in dev:
            issues.append(
                f"第 {no} 条（{kind}）为负偏离"
                + ("：★/否决/资格条款负偏离将导致投标无效" if key else "")
            )
        if key and not resp:
            issues.append(f"第 {no} 条（{kind}）未填写响应情况")
        if key and not page:
            issues.append(f"第 {no} 条（{kind}）未填写投标文件页码")
    if "评分标准" in wb.sheetnames:
        sc = wb["评分标准"]
        selfs = [
            r[4] for r in sc.iter_rows(min_row=2, values_only=True) if r[1] not in (None, "合计")
        ]
        filled = [v for v in selfs if isinstance(v, (int, float))]
        if selfs and len(filled) < len(selfs):
            issues.append(f"评分标准中 {len(selfs) - len(filled)} 项还没有自评得分")
        elif filled:
            print(f"评分自评合计 {sum(filled):g} 分")
    print(f"响应矩阵共 {total} 条")
    return issues


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("matrix")
    m.add_argument("input")
    d = sub.add_parser("doc")
    d.add_argument("input")
    d.add_argument("--bidder", default="")
    d.add_argument("--project-no")
    d.add_argument("--other", default="")
    d.add_argument("--deadline")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "matrix":
            issues = check_matrix(Path(args.input))
        else:
            deadline = dt.date.fromisoformat(args.deadline) if args.deadline else None
            others = [o.strip() for o in re.split(r"[,，]", args.other) if o.strip()]
            issues = check_doc(
                read_text(Path(args.input)), args.bidder, args.project_no, others, deadline
            )
    except (OSError, CheckError, ValueError, KeyError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    print(
        f"发现 {len(issues)} 个问题"
        + ("：" if issues else "，请仍按招标文件要求人工复核签字、盖章和装订")
    )
    for i in issues:
        print(f"- {i}")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
