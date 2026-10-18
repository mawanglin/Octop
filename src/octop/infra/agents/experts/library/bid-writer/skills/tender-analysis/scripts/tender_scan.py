#!/usr/bin/env python3
"""招标文件解析：提取关键信息、★/▲ 实质性条款、否决投标条款、资格要求和评分标准，生成响应矩阵 Excel。

Usage:
    python scripts/tender_scan.py 招标文件.docx -o 响应矩阵.xlsx
    python scripts/tender_scan.py 招标文件.pdf -o 响应矩阵.xlsx     （需要 pdftotext；扫描件 PDF 读不到文字）

输出：
  - 终端：关键信息（项目名称、编号、预算/最高限价、投标截止和开标时间、保证金、工期、质保期、有效期、联合体）和条款数量；
  - Excel「关键信息」「实质性条款」（★、▲、否决/无效投标、资格要求，带响应情况/偏离/投标文件页码空列）、
    「评分标准」（来自含“分值/评分”的表格或“（N 分）”文字，合计分值自动计算）。
提取基于关键词和版式规律，可能漏掉或多抓，必须对照招标文件原文逐条核对。Exit codes: 0 ok, 1 输入错误, 2 缺少依赖。
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

KEY_FIELDS: list[tuple[str, str]] = [
    ("项目名称", r"(?:项目名称|采购项目名称|招标项目名称)"),
    ("项目编号", r"(?:项目编号|招标编号|采购编号|项目代码)"),
    ("采购人/招标人", r"(?:采购人(?:名称)?|招标人(?:名称)?)"),
    ("代理机构", r"(?:采购代理机构|招标代理机构|代理机构)(?:名称)?"),
    ("预算/最高限价", r"(?:预算金额|采购预算|项目预算|最高限价|招标控制价|最高投标限价)"),
    (
        "投标截止时间",
        r"(?:投标截止时间|递交投标文件截止时间|响应文件提交截止时间|投标文件递交截止时间)",
    ),
    ("开标时间", r"(?:开标时间)"),
    ("投标保证金", r"(?:投标保证金(?:金额)?)"),
    ("工期/交付期", r"(?:工期|交货期|交付期|服务期限?|供货期|合同履行期限)"),
    ("质保期", r"(?:质保期|质量保证期|保修期)"),
    ("投标有效期", r"(?:投标有效期|响应有效期)"),
    ("联合体", r"(?:是否接受联合体投标|联合体投标)"),
]
REJECT = re.compile(
    r"否决(?:其)?投标|废标|无效投标|投标无效|不予受理|作无效(?:投标|响应)处理|视为无效|拒绝其投标"
)
QUAL = re.compile(r"营业执照|资质证书|资质要求|业绩要求|类似业绩|信用中国|失信被执行人|重大税收违法|社保|社会保险|财务(?:审计)?报告|"
                  r"财务报表|资格要求|资格条件|特定资格|ISO\s*\d+|安全生产许可证|项目经理.{0,6}(?:证书|资格)")  # fmt: skip
SCORE_INLINE = re.compile(r"^(.{2,40}?)[（(]\s*(\d+(?:\.\d+)?)\s*分\s*[)）]\s*[:：]?\s*(.*)$")
FILL = PatternFill("solid", fgColor="1E3A8A")
STAR_FILL = PatternFill("solid", fgColor="FEE2E2")
TRI_FILL = PatternFill("solid", fgColor="FEF3C7")


class TenderError(ValueError):
    pass


def read_doc(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    """Return (paragraph lines, tables as rows of cell text)."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError:  # pragma: no cover - environment check
            print("读取 .docx 需要 python-docx：pip install python-docx", file=sys.stderr)
            sys.exit(2)
        doc = Document(str(path))
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables = []
        for t in doc.tables:
            rows = []
            for r in t.rows:
                cells: list[str] = []
                for c in r.cells:
                    txt = c.text.strip()
                    if not cells or txt != cells[-1]:  # merged cells repeat their text
                        cells.append(txt)
                rows.append(cells)
            tables.append(rows)
        return lines, tables
    if suffix == ".pdf":
        exe = shutil.which("pdftotext")
        if not exe:
            raise TenderError("读取 PDF 需要 pdftotext（poppler-utils），或请提供 Word 版招标文件")
        text = subprocess.run(
            [exe, "-layout", str(path), "-"], capture_output=True, text=True, timeout=120
        ).stdout
        if len(text.strip()) < 100:
            raise TenderError("PDF 中读不到文字（可能是扫描件），请提供 Word 版或先做文字识别")
    else:
        raw = path.read_bytes()
        text = next(
            (raw.decode(e) for e in ("utf-8-sig", "gb18030") if _ok(raw, e)),
            raw.decode("utf-8", "replace"),
        )
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    tables = [[[c.strip() for c in ln.strip("|").split("|")] for ln in block]
              for block in _md_tables(lines)]  # fmt: skip
    return [ln for ln in lines if not ln.startswith("|")], tables


def _ok(raw: bytes, enc: str) -> bool:
    try:
        raw.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def _md_tables(lines: list[str]) -> list[list[str]]:
    blocks: list[list[str]] = []
    cur: list[str] = []
    for ln in [*lines, ""]:
        if ln.startswith("|"):
            if not re.fullmatch(r"\|?[\s:|-]+\|?", ln):
                cur.append(ln)
        elif cur:
            blocks.append(cur)
            cur = []
    return blocks


def key_info(lines: list[str], tables: list[list[list[str]]]) -> dict[str, str]:
    found: dict[str, str] = {}
    cells = list(lines)
    for t in tables:  # "项目名称 | XXX" style tables (投标人须知前附表)
        for row in t:
            cells.extend(f"{row[i]}：{row[i + 1]}" for i in range(len(row) - 1))
    for name, pat in KEY_FIELDS:
        rx = re.compile(pat + r"\s*(?:为|是)?\s*[:：]\s*(.{2,80})")
        for ln in cells:
            m = rx.search(ln)
            if m:
                val = re.split(r"\s{2,}|；(?=\S{2,8}[:：])", m.group(1).strip())[0].rstrip("。；;")
                if val:
                    found[name] = val
                    break
    return found


HEADING = re.compile(
    r"^(?:#{1,6}\s*|第[一二三四五六七八九十\d]+[章节部分]|[一二三四五六七八九十]+[、.])"
)
ITEM = re.compile(r"^\s*(?:[（(]?\d+[)）.、]|[①-⑳])")
MARK = re.compile(r"^\s*(?:[（(]?\d+[)）.、]\s*)?([★▲])")


def _is_score_table(t: list[list[str]]) -> bool:
    return bool(t) and any(re.search(r"分值|分数|满分", h) for h in t[0])


def clauses(lines: list[str], tables: list[list[list[str]]]) -> list[tuple[str, str, str]]:
    """Flag clauses; numbered items under a 资格 / 无效投标 heading or lead-in sentence inherit its type."""
    out: list[tuple[str, str, str]] = []
    seen: set[str] = set()
    sources = [(f"第{i}段", ln, False) for i, ln in enumerate(lines, start=1)]
    for ti, t in enumerate(tables, start=1):
        if _is_score_table(t):
            continue
        for ri, row in enumerate(t, start=1):
            sources.append((f"表{ti}第{ri}行", " ".join(row), True))
    ctx: str | None = None
    for where, text, in_table in sources:
        if not in_table and HEADING.match(text) and len(text) <= 40:
            ctx = (
                "资格"
                if "资格" in text
                else "否决/无效"
                if re.search(r"无效|废标|否决", text)
                else None
            )
            continue
        is_item = bool(ITEM.match(text))
        if not in_table and not is_item and not MARK.match(text):
            lead = REJECT.search(text) and text.rstrip().endswith(("：", ":"))
            ctx = "否决/无效" if lead else (ctx if ctx == "资格" else None)
            if lead:
                continue  # the lead-in sentence itself; its numbered items follow
        kinds = []
        m = MARK.match(text)
        if m:
            kinds.append("★实质性" if m.group(1) == "★" else "▲重要")
        if REJECT.search(text) or (ctx == "否决/无效" and is_item):
            kinds.append("否决/无效")
        if (QUAL.search(text) and len(text) < 300) or (ctx == "资格" and is_item):
            kinds.append("资格")
        if kinds and text not in seen:
            seen.add(text)
            out.append(("、".join(dict.fromkeys(kinds)), text, where))
    return out


def scoring(lines: list[str], tables: list[list[list[str]]]) -> list[tuple[str, float | None, str]]:
    rows: list[tuple[str, float | None, str]] = []
    for t in tables:
        if not t:
            continue
        header = t[0]
        hidx = next((i for i, h in enumerate(header) if re.search(r"分值|分数|权重|满分", h)), None)
        if hidx is None:
            continue
        name_idx = next(
            (
                i
                for i, h in enumerate(header)
                if re.search(r"评分(?:因素|项目|内容|项)|评审(?:因素|项目)|评分因素|项目", h)
            ),
            0,
        )
        std_idx = next(
            (i for i, h in enumerate(header) if re.search(r"标准|细则|说明|要求", h)),
            len(header) - 1,
        )
        for r in t[1:]:
            if len(r) <= max(hidx, name_idx):
                continue
            m = re.search(r"\d+(?:\.\d+)?", r[hidx])
            rows.append(
                (
                    r[name_idx],
                    float(m.group(0)) if m else None,
                    r[std_idx] if std_idx < len(r) else "",
                )
            )
    if not rows:
        for ln in lines:
            m = SCORE_INLINE.match(ln)
            if m:
                rows.append((m.group(1).strip(" 、.0123456789"), float(m.group(2)), m.group(3)))
    return rows


def write_xlsx(out: Path, info: dict[str, str], items: list[tuple[str, str, str]],
               scores: list[tuple[str, float | None, str]]) -> None:  # fmt: skip
    wb = Workbook()
    ws = wb.active
    ws.title = "关键信息"
    ws.append(["项目", "内容", "核对（是否与原文一致）"])
    for name, _ in KEY_FIELDS:
        ws.append([name, info.get(name, "未找到，请人工查阅"), ""])
    cl = wb.create_sheet("实质性条款")
    cl.append(
        [
            "序号",
            "类型",
            "招标文件原文",
            "位置",
            "响应情况",
            "偏离（无/正偏离/负偏离）",
            "投标文件页码",
            "负责人",
        ]
    )
    for i, (kind, text, where) in enumerate(items, start=1):
        cl.append([i, kind, text, where, "", "", "", ""])
        if "★" in kind or "否决" in kind:
            for c in cl[cl.max_row]:
                c.fill = STAR_FILL
        elif "▲" in kind:
            for c in cl[cl.max_row]:
                c.fill = TRI_FILL
    sc = wb.create_sheet("评分标准")
    sc.append(["序号", "评分项", "分值", "评分标准", "自评得分", "得分依据 / 投标文件页码"])
    for i, (name, pts, std) in enumerate(scores, start=1):
        sc.append([i, name, pts, std, "", ""])
    if scores:
        n = sc.max_row
        sc.append(["", "合计", f"=SUM(C2:C{n})", "", f"=SUM(E2:E{n})", ""])
        sc.cell(sc.max_row, 2).font = Font(bold=True)
    for sheet, widths in (
        (ws, [16, 60, 20]),
        (cl, [6, 14, 70, 12, 30, 16, 12, 10]),
        (sc, [6, 22, 8, 60, 10, 26]),
    ):
        for c in sheet[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), FILL
        for j, w in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(1, j).column_letter].width = w
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input")
    parser.add_argument("-o", "--output", required=True)
    args = parser.parse_args(argv)
    try:
        lines, tables = read_doc(Path(args.input))
    except (OSError, TenderError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    info = key_info(lines, tables)
    items = clauses(lines, tables)
    scores = scoring(lines, tables)
    write_xlsx(Path(args.output), info, items, scores)
    print(f"已生成 {args.output}")
    print("## 关键信息")
    for name, _ in KEY_FIELDS:
        print(f"- {name}：{info.get(name, '未找到')}")
    count = {k: sum(1 for kind, _t, _w in items if k in kind) for k in ("★", "▲", "否决", "资格")}
    print(
        f"## 条款：★ 实质性 {count['★']} 条，▲ 重要 {count['▲']} 条，否决/无效投标 {count['否决']} 条，资格要求 {count['资格']} 条"
    )
    if scores:
        total = sum(p for _n, p, _s in scores if p is not None)
        print(
            f"## 评分标准：{len(scores)} 项，合计 {total:g} 分"
            + ("" if abs(total - 100) < 0.01 else "（不是 100 分，请核对是否漏项或分值写在子项中）")
        )
    else:
        print("## 评分标准：没有识别到，请人工查阅“评标办法”章节")
    print("提示：以上为自动提取，请对照招标文件原文逐条核对。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
