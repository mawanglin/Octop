#!/usr/bin/env python3
"""Markdown → .docx with Chinese office layouts.

Usage:
    python scripts/docx_build.py input.md -o output.docx --style gov|report|speech

Styles:
    gov     党政机关公文格式（GB/T 9704-2012）：版头、红色分隔线、标题、主送机关、
            正文层次（一、（一）1.（1））、附件说明、署名日期、附注、版记、"— 1 —" 页码。
            不填版头字段时就是政府汇报材料的版式。
    report  工作汇报：黑体标题，宋体小四正文，1.5 倍行距，表格和引用块。
    speech  讲话稿：仿宋小二大字号，1.8 倍行距，方便照读。

Markdown 约定（中文文稿习惯：每个非空行就是一个段落）：
    开头可选的 front matter（--- 包围，"键: 值" 每行一条），键名见 SKILL.md；
    "# 标题"（第一个）、"## / ### / ####" 三级标题、"- " 列表、"1. " 编号、
    "| 表格 |"、"> 引用"、"**加粗**"、单独一行 "<!-- pagebreak -->" 分页；
    单独一行 "![说明](图片路径)" 插入图片（相对路径以 Markdown 文件所在目录为准，说明作为图注），
    "[文字](网址)" 输出为 "文字（网址）"。

Exit codes: 0 ok, 1 input error, 2 missing python-docx.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 python-docx：请运行 `pip install python-docx`", file=sys.stderr)
    sys.exit(2)

# --------------------------------------------------------------------- styles

GOV_LINE = 28.95  # 225mm 版心 / 22 行
STYLES: dict[str, dict[str, Any]] = {
    "gov": {
        "margins_mm": (37, 35, 28, 26),  # top, bottom, left, right → 版心 156×225mm
        "fonts": {
            "title": "方正小标宋简体",
            "h1": "黑体",
            "h2": "楷体_GB2312",
            "h3": "仿宋_GB2312",
            "body": "仿宋_GB2312",
            "header": "方正小标宋简体",
            "page": "宋体",
        },
        "sizes": {"title": 22, "h1": 16, "h2": 16, "h3": 16, "body": 16, "table": 14, "note": 14},
        "line_exact": GOV_LINE,
        "indent_chars": 2,
        "page_number": "gov",
    },
    "report": {
        "margins_mm": (25.4, 25.4, 31.8, 31.8),
        "fonts": {
            "title": "黑体",
            "h1": "黑体",
            "h2": "黑体",
            "h3": "宋体",
            "body": "宋体",
            "header": "黑体",
            "page": "宋体",
        },
        "sizes": {
            "title": 22,
            "h1": 16,
            "h2": 14,
            "h3": 12,
            "body": 12,
            "table": 10.5,
            "note": 10.5,
        },
        "line_multiple": 1.5,
        "indent_chars": 2,
        "page_number": "center",
    },
    "speech": {
        "margins_mm": (25, 25, 25, 25),
        "fonts": {
            "title": "方正小标宋简体",
            "h1": "黑体",
            "h2": "楷体_GB2312",
            "h3": "仿宋_GB2312",
            "body": "仿宋_GB2312",
            "header": "黑体",
            "page": "宋体",
        },
        "sizes": {"title": 26, "h1": 18, "h2": 18, "h3": 18, "body": 18, "table": 14, "note": 14},
        "line_multiple": 1.8,
        "indent_chars": 2,
        "page_number": "center",
    },
}

# 公文层次序数：一、 （一） 1. （1）
_CN_NUM = "一二三四五六七八九十百零〇"
LEVEL_PATTERNS = (
    (1, re.compile(rf"^[{_CN_NUM}]+、")),
    (2, re.compile(rf"^（[{_CN_NUM}]+）")),
    (3, re.compile(r"^\d+[．.]\s?")),
    (4, re.compile(r"^（\d+）")),
)
IMAGE = re.compile(r'^!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+"[^"]*")?\s*\)$')
PAGEBREAK = re.compile(r"^\s*(<!--\s*pagebreak\s*-->|\\pagebreak)\s*$", re.I)
RED = RGBColor(0xFF, 0x00, 0x00)


class InputError(ValueError):
    """Bad markdown / front matter; message is shown to the caller."""


# ------------------------------------------------------------------- parsing


@dataclass
class Block:
    kind: str  # title, h1, h2, h3, para, bullet, number, quote, table, pagebreak, image
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


def parse_markdown(source: str) -> tuple[dict[str, str], list[Block]]:
    lines = source.replace("\r\n", "\n").split("\n")
    meta: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        try:
            end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        except StopIteration as exc:
            raise InputError("front matter 缺少结束的 ---") from exc
        for raw in lines[1:end]:
            if not raw.strip() or raw.lstrip().startswith("#"):
                continue
            key, sep, value = raw.partition(":")
            if not sep:
                key, sep, value = raw.partition("：")
            if not sep:
                raise InputError(f"front matter 行缺少冒号：{raw!r}")
            meta[key.strip()] = value.strip()
        lines = lines[end + 1 :]

    blocks: list[Block] = []
    table: list[list[str]] = []

    def flush_table() -> None:
        if table:
            blocks.append(Block("table", rows=[r[:] for r in table]))
            table.clear()

    for raw in lines:
        line = raw.strip()
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # separator row
            table.append(cells)
            continue
        flush_table()
        if not line:
            continue
        if PAGEBREAK.match(line):
            blocks.append(Block("pagebreak"))
        elif m := IMAGE.match(line):
            blocks.append(Block("image", m.group(1).strip(), rows=[[m.group(2)]]))
        elif line.startswith("#### "):
            blocks.append(Block("h3", line[5:].strip()))
        elif line.startswith("### "):
            blocks.append(Block("h2", line[4:].strip()))
        elif line.startswith("## "):
            blocks.append(Block("h1", line[3:].strip()))
        elif line.startswith("# "):
            kind = "title" if not any(b.kind == "title" for b in blocks) else "h1"
            blocks.append(Block(kind, line[2:].strip()))
        elif line.startswith(("- ", "* ", "+ ")):
            blocks.append(Block("bullet", line[2:].strip()))
        elif re.match(r"^\d+\.\s", line):
            blocks.append(Block("number", line))
        elif line.startswith(">"):
            blocks.append(Block("quote", line.lstrip("> ").strip()))
        else:
            blocks.append(Block("para", line))
    flush_table()
    if "标题" in meta and not any(b.kind == "title" for b in blocks):
        blocks.insert(0, Block("title", meta["标题"]))
    return meta, blocks


def _inline(text: str) -> list[tuple[str, bool]]:
    """Split **bold** segments; links become "text（url）"; drop inline code ticks."""

    def link(m: re.Match[str]) -> str:
        label, url = m.group(1), m.group(2).strip()
        return label if label == url or not re.match(r"https?://", url) else f"{label}（{url}）"

    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", link, text).replace("`", "")
    parts: list[tuple[str, bool]] = []
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            parts.append((chunk, i % 2 == 1))
    return parts or [("", False)]


def _em_width(text: str) -> float:
    """Width in em: CJK / full-width = 1, ASCII = 0.5."""
    return sum(1.0 if ord(ch) > 0x2E80 else 0.5 for ch in text)


# ----------------------------------------------------------------- rendering


class Builder:
    def __init__(
        self, style: str, fonts: dict[str, str] | None = None, base_dir: Path | None = None
    ) -> None:
        self.base_dir = base_dir or Path(".")
        self.warnings: list[str] = []
        if style not in STYLES:
            raise InputError(f"未知版式 {style!r}，可选：{', '.join(STYLES)}")
        self.name = style
        self.s = STYLES[style]
        self.fonts = {**self.s["fonts"], **(fonts or {})}
        self.doc = Document()
        top, bottom, left, right = self.s["margins_mm"]
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin, sec.bottom_margin = Mm(top), Mm(bottom)
        sec.left_margin, sec.right_margin = Mm(left), Mm(right)
        sec.footer_distance = Mm(24 if style == "gov" else 15)
        self.text_width_mm = 210 - left - right
        body = self.size("body")
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(body)
        self._set_font(normal.element.get_or_add_rPr(), self.fonts["body"])
        if style == "gov":
            # 每页 22 行的文档网格，配合固定行距。
            grid = sec._sectPr.find(qn("w:docGrid"))
            if grid is None:
                grid = OxmlElement("w:docGrid")
                sec._sectPr.append(grid)
            grid.set(qn("w:type"), "lines")
            grid.set(qn("w:linePitch"), str(round(GOV_LINE * 20)))

    # ---- helpers

    def size(self, key: str) -> float:
        return float(self.s["sizes"][key])

    @staticmethod
    def _set_font(rpr: Any, name: str) -> None:
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), name)

    def _spacing(self, p: Any, before: float = 0, after: float = 0) -> None:
        pf = p.paragraph_format
        if "line_exact" in self.s:
            pf.line_spacing = Pt(self.s["line_exact"])
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        else:
            pf.line_spacing = self.s["line_multiple"]
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.widow_control = True

    def _indent_chars(self, p: Any, first: float = 0, left: float = 0, right: float = 0) -> None:
        """Indent in character units (Word's *Chars attributes), with pt fallback."""
        pt = self.size("body")
        pf = p.paragraph_format
        pf.first_line_indent = Pt(first * pt) if first else None
        pf.left_indent = Pt(left * pt) if left else None
        pf.right_indent = Pt(right * pt) if right else None
        ind = p._p.get_or_add_pPr().get_or_add_ind()
        if first:
            ind.set(qn("w:firstLineChars"), str(int(first * 100)))
        if left:
            ind.set(qn("w:leftChars"), str(int(left * 100)))
        if right:
            ind.set(qn("w:rightChars"), str(int(right * 100)))

    def run(
        self,
        p: Any,
        text: str,
        font: str,
        size: float,
        *,
        bold: bool = False,
        color: RGBColor | None = None,
    ) -> Any:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color
        self._set_font(r._r.get_or_add_rPr(), font)
        return r

    def para(
        self,
        text: str = "",
        *,
        font: str | None = None,
        size: float | None = None,
        align: Any = WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_indent: float = 0,
        bold: bool = False,
        before: float = 0,
        after: float = 0,
        color: RGBColor | None = None,
    ) -> Any:
        p = self.doc.add_paragraph()
        p.alignment = align
        self._spacing(p, before, after)
        if first_indent:
            self._indent_chars(p, first=first_indent)
        if text:
            for chunk, strong in _inline(text):
                self.run(
                    p,
                    chunk,
                    font or self.fonts["body"],
                    size or self.size("body"),
                    bold=bold or strong,
                    color=color,
                )
        return p

    def image(self, src: str, caption: str) -> None:
        """Embed a local picture (max text width) with its caption; placeholders otherwise."""
        if re.match(r"https?://", src):
            self.warnings.append(f"网络图片没有嵌入，请先下载到本地：{src}")
            self.para(f"〔图片：{caption or src}（{src}）〕", align=WD_ALIGN_PARAGRAPH.CENTER)
            return
        path = (self.base_dir / src).resolve() if not Path(src).is_absolute() else Path(src)
        if not path.is_file():
            self.warnings.append(f"找不到图片：{src}")
            self.para(f"〔图片缺失：{src}〕", align=WD_ALIGN_PARAGRAPH.CENTER)
            return
        sec = self.doc.sections[0]
        max_width = sec.page_width - sec.left_margin - sec.right_margin
        try:
            from PIL import Image

            with Image.open(path) as im:
                width_px, fmt = im.width, (im.format or "").upper()
                if fmt not in {"PNG", "JPEG", "GIF", "BMP", "TIFF"}:
                    import io

                    buf = io.BytesIO()
                    im.convert("RGBA" if "A" in im.getbands() else "RGB").save(buf, "PNG")
                    buf.seek(0)
                    stream: Any = buf
                else:
                    stream = str(path)
        except Exception as exc:  # noqa: BLE001 - unreadable image becomes a placeholder
            self.warnings.append(f"无法读取图片 {src}：{exc}")
            self.para(f"〔图片无法读取：{src}〕", align=WD_ALIGN_PARAGRAPH.CENTER)
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._spacing(p, 6, 2)
        width = min(max_width, Mm(width_px * 25.4 / 96))
        p.add_run().add_picture(stream, width=width)
        if caption:
            self.para(
                caption, size=self.size("body") - 1.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=6
            )

    def blank(self, n: int = 1) -> None:
        for _ in range(n):
            self.para()

    @staticmethod
    def _border(p: Any, side: str, *, color: str, size_eighths: int, space_pt: int = 1) -> None:
        ppr = p._p.get_or_add_pPr()
        bdr = ppr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = OxmlElement("w:pBdr")
            ppr.insert_element_before(
                bdr, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr"
            )
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size_eighths))
        edge.set(qn("w:space"), str(space_pt))
        edge.set(qn("w:color"), color)
        bdr.append(edge)

    def _right_tab(self, p: Any, indent_chars: float = 0) -> None:
        pos = Mm(self.text_width_mm) - Pt(indent_chars * self.size("body"))
        p.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.RIGHT)

    # ---- page numbers

    def _page_field(self, p: Any, font: str, size: float, prefix: str, suffix: str) -> None:
        if prefix:
            self.run(p, prefix, font, size)
        r = self.run(p, "", font, size)
        for tag, attrs in (
            ("w:fldChar", {"w:fldCharType": "begin"}),
            ("w:instrText", {}),
            ("w:fldChar", {"w:fldCharType": "end"}),
        ):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            if tag == "w:instrText":
                el.set(qn("xml:space"), "preserve")
                el.text = " PAGE "
            r._r.append(el)
        if suffix:
            self.run(p, suffix, font, size)

    def page_numbers(self) -> None:
        sec = self.doc.sections[0]
        font = self.fonts["page"]
        if self.s["page_number"] == "gov":
            # 4 号半角宋体，"— 1 —"；单页居右空一字，双页居左空一字。
            self.doc.settings.odd_and_even_pages_header_footer = True
            size = 14.0
            for footer, align, side in (
                (sec.footer, WD_ALIGN_PARAGRAPH.RIGHT, "right"),
                (sec.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT, "left"),
            ):
                p = footer.paragraphs[0]
                p.alignment = align
                if side == "right":
                    p.paragraph_format.right_indent = Pt(size)
                else:
                    p.paragraph_format.left_indent = Pt(size)
                self._page_field(p, font, size, "— ", " —")
        else:
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._page_field(p, font, 10.5, "— ", " —")

    # ---- body blocks

    def heading_level(self, text: str) -> int:
        for level, pattern in LEVEL_PATTERNS:
            if pattern.match(text):
                return level
        return 0

    def body_para(self, text: str, level: int = 0) -> None:
        """Body paragraph; gov-style numbered headings get their level font."""
        f = self.fonts
        size = self.size("body")
        indent = self.s["indent_chars"]
        if level == 1:
            self.para(text, font=f["h1"], first_indent=indent)
            return
        if level in (2, 3):
            # "（一）小标题。正文……" — only the lead phrase takes the heading font.
            head, sep, rest = text.partition("。")
            lead = head + sep if sep and rest else text
            tail = rest if sep and rest else ""
            p = self.para(first_indent=indent)
            font = f["h2"] if level == 2 else f["h3"]
            for chunk, strong in _inline(lead):
                self.run(p, chunk, font, size, bold=strong or level == 3)
            for chunk, strong in _inline(tail):
                self.run(p, chunk, f["body"], size, bold=strong)
            return
        if re.fullmatch(r"[^。！？]{1,15}[：:]", text):
            # 称呼语（"同志们："、"尊敬的各位来宾："）顶格。
            self.para(text, align=WD_ALIGN_PARAGRAPH.LEFT)
            return
        self.para(text, first_indent=indent)

    def table(self, rows: list[list[str]]) -> None:
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        t = self.doc.add_table(rows=len(rows), cols=width)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = self.size("table")
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                font = self.fonts["h1"] if i == 0 else self.fonts["body"]
                for chunk, strong in _inline(value):
                    self.run(p, chunk, font, size, bold=strong or (i == 0 and self.name != "gov"))
                if i == 0 and self.name != "gov":
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "E7E6E6")
                    cell._tc.get_or_add_tcPr().append(shd)
        self.blank()

    def render_blocks(self, blocks: list[Block], *, skip_title: bool) -> None:
        f = self.fonts
        gov = self.name == "gov"
        for b in blocks:
            if b.kind == "title":
                if not skip_title:
                    self.title(b.text)
            elif b.kind == "pagebreak":
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif b.kind == "h1":
                if gov:
                    self.body_para(b.text, 1)
                else:
                    self.para(b.text, font=f["h1"], size=self.size("h1"), before=12, after=6)
            elif b.kind == "h2":
                if gov:
                    self.body_para(b.text, 2)
                else:
                    self.para(b.text, font=f["h2"], size=self.size("h2"), before=8, after=4)
            elif b.kind == "h3":
                if gov:
                    self.body_para(b.text, 3)
                else:
                    self.para(b.text, font=f["h3"], size=self.size("h3"), bold=True, before=6)
            elif b.kind == "table":
                self.table(b.rows)
            elif b.kind == "image":
                self.image(b.rows[0][0], b.text)
            elif b.kind == "quote":
                p = self.para(b.text, font=f["h2"], first_indent=self.s["indent_chars"])
                if not gov:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "F2F2F2")
                    p._p.get_or_add_pPr().insert_element_before(
                        shd, "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr"
                    )
            elif b.kind == "bullet":
                if gov:
                    self.body_para(b.text)
                else:
                    p = self.para("• " + b.text)
                    self._indent_chars(p, first=-1, left=self.s["indent_chars"] + 1)
            else:  # para, number
                numbered = gov or self.name == "speech"
                self.body_para(b.text, self.heading_level(b.text) if numbered else 0)

    def title(self, text: str, subtitle: str = "") -> None:
        size = self.size("title")
        for line in text.split("\\n"):
            self.para(
                line,
                font=self.fonts["title"],
                size=size,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        if subtitle:
            self.para(subtitle, font=self.fonts["h2"], align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank()

    # ---- gov-only parts

    def gov_header(self, meta: dict[str, str]) -> bool:
        """版头。Returns True when a 版头 was written (title then follows 2 blank lines)."""
        f = self.fonts
        top_lines = [meta[k] for k in ("份号", "密级", "紧急程度") if meta.get(k)]
        for value in top_lines:
            self.para(value, font=f["h1"], align=WD_ALIGN_PARAGRAPH.LEFT)
        org = meta.get("发文机关标志", "")
        if not org:
            return False
        # 发文机关标志：红色小标宋，上边缘距版心上边缘 35mm，宽度不超过版心。
        width_pt = self.text_width_mm / 25.4 * 72
        size = min(42.0, int(width_pt / max(_em_width(org), 1) * 0.95))
        used = len(top_lines) * GOV_LINE
        p = self.para(
            org,
            font=f["header"],
            size=size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            before=max(0.0, 35 / 25.4 * 72 - used - GOV_LINE),
            color=RED,
        )
        p.paragraph_format.line_spacing = Pt(size * 1.2)
        self.blank(2)
        number = meta.get("发文字号", "")
        signer = meta.get("签发人", "")
        p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT if signer else WD_ALIGN_PARAGRAPH.CENTER)
        if signer:
            # 上行文：发文字号居左空一字，签发人居右空一字。
            self._right_tab(p, 1)
            self.run(p, "　" + number + "\t", f["body"], self.size("body"))
            self.run(p, "签发人：", f["body"], self.size("body"))
            self.run(p, signer, f["h2"], self.size("body"))
        else:
            self.run(p, number, f["body"], self.size("body"))
        # 版头分隔线：发文字号之下 4mm，红色，与版心等宽。
        self._border(p, "bottom", color="FF0000", size_eighths=12, space_pt=11)
        self.blank(1)
        return True

    def gov_tail(self, meta: dict[str, str]) -> None:
        body = self.size("body")
        attachments = [a.strip() for a in re.split(r"[；;\n]", meta.get("附件", "")) if a.strip()]
        if attachments:
            self.blank()
            for i, name in enumerate(attachments, start=1):
                label = f"{i}．{name}" if len(attachments) > 1 else name
                if i == 1:
                    p = self.para("附件：" + label, first_indent=self.s["indent_chars"])
                else:
                    p = self.para(label)
                    self._indent_chars(p, left=self.s["indent_chars"] + 3)
        signer = meta.get("署名", "")
        date = meta.get("日期", "")
        if signer or date:
            self.blank(2)  # 留出用印位置
            # 成文日期右空四字；发文机关署名以成文日期为准居中。
            date_w = _em_width(date)
            if signer:
                offset = 4 + max(0.0, (date_w - _em_width(signer)) / 2)
                p = self.para(signer, align=WD_ALIGN_PARAGRAPH.RIGHT)
                p.paragraph_format.right_indent = Pt(offset * body)
            if date:
                p = self.para(date, align=WD_ALIGN_PARAGRAPH.RIGHT)
                p.paragraph_format.right_indent = Pt(4 * body)
        if meta.get("附注"):
            note = meta["附注"].strip("（）()")
            self.para(f"（{note}）", first_indent=self.s["indent_chars"])
        self.gov_footer_block(meta)

    def gov_footer_block(self, meta: dict[str, str]) -> None:
        """版记：抄送、印发机关和印发日期。"""
        cc = meta.get("抄送", "")
        issuer = meta.get("印发机关", "")
        issued = meta.get("印发日期", "")
        if not (cc or issuer):
            return
        f = self.fonts
        size = self.size("note")
        lines = (["cc"] if cc else []) + (["issuer"] if issuer else [])
        # 单列表格，边框即版记分隔线。紧随正文排在文末（浮动定位在各软件中表现不一致），
        # 需要贴到版心底部时在版记前补空行。
        self.blank()
        t = self.doc.add_table(rows=len(lines), cols=1)
        tbl_pr = t._tbl.tblPr
        width = OxmlElement("w:tblW")
        width.set(qn("w:w"), str(int(self.text_width_mm / 25.4 * 1440)))
        width.set(qn("w:type"), "dxa")
        old = tbl_pr.find(qn("w:tblW"))
        if old is not None:
            tbl_pr.replace(old, width)
        borders = OxmlElement("w:tblBorders")
        # 首条、末条分隔线 0.35mm 粗线；中间分隔线 0.25mm 细线。
        for side, sz in (
            ("top", 8),
            ("left", 0),
            ("bottom", 8),
            ("right", 0),
            ("insideH", 6),
            ("insideV", 0),
        ):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), "single" if sz else "nil")
            if sz:
                edge.set(qn("w:sz"), str(sz))
                edge.set(qn("w:color"), "000000")
            borders.append(edge)
        tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
        for row, kind in zip(t.rows, lines, strict=True):
            # 版记各条不跨页、彼此同页。
            cant = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant)
            p = row.cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self._spacing(p)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.left_indent = Pt(size)
            p.paragraph_format.right_indent = Pt(size)
            if kind == "cc":
                self.run(p, "抄送：" + cc.rstrip("。") + "。", f["body"], size)
            else:
                self._right_tab(p, 1)
                self.run(p, issuer + "\t" + (issued + "印发" if issued else ""), f["body"], size)

    # ---- entry

    def build(self, meta: dict[str, str], blocks: list[Block], output: Path) -> None:
        gov = self.name == "gov"
        title_block = next((b for b in blocks if b.kind == "title"), None)
        if gov and self.gov_header(meta):
            self.blank(1)
        if title_block is not None:
            self.title(title_block.text, meta.get("副标题", ""))
            if not gov and meta.get("署名行"):
                self.para(meta["署名行"], align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
        if gov and meta.get("主送"):
            p = self.para(meta["主送"].rstrip(":：") + "：", align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.first_line_indent = None
        self.render_blocks(blocks, skip_title=True)
        if gov:
            self.gov_tail(meta)
        elif meta.get("署名") or meta.get("日期"):
            self.blank()
            for key in ("署名", "日期"):
                if meta.get(key):
                    p = self.para(meta[key], align=WD_ALIGN_PARAGRAPH.RIGHT)
                    p.paragraph_format.right_indent = Pt(2 * self.size("body"))
        self.page_numbers()
        if title_block is not None:
            self.doc.core_properties.title = title_block.text.replace("\\n", "")
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input", help="Markdown 文件")
    parser.add_argument("-o", "--output", required=True, help="输出 .docx 路径")
    parser.add_argument("--style", choices=sorted(STYLES), required=True)
    parser.add_argument("--fonts", help='覆盖字体，JSON，如 \'{"body": "仿宋", "title": "黑体"}\'')
    args = parser.parse_args(argv)
    try:
        source = Path(args.input).read_text(encoding="utf-8")
        fonts = json.loads(args.fonts) if args.fonts else None
        meta, blocks = parse_markdown(source)
        if not blocks:
            raise InputError("正文为空")
        builder = Builder(args.style, fonts, base_dir=Path(args.input).resolve().parent)
        builder.build(meta, blocks, Path(args.output))
    except (OSError, json.JSONDecodeError, InputError) as exc:
        print(f"生成失败：{exc}", file=sys.stderr)
        return 1
    print(f"已生成 {args.output}（版式 {args.style}）")
    for warning in builder.warnings:
        print(f"注意：{warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
