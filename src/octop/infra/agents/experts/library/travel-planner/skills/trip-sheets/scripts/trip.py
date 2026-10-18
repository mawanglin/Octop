#!/usr/bin/env python3
"""旅行规划表格：行程 Markdown → Excel（行程表、每日费用、分类预算），以及按条件生成行李清单。

Usage:
    python scripts/trip.py sheet 行程.md -o 行程.xlsx
    python scripts/trip.py packing --days 5 --climate cold,rain --type city,abroad [--people 2] -o 行李清单.xlsx

行程 Markdown：可选 front matter（行程、出发、人数、预算）+ 一张表格：
    | 日期 | 时间 | 地点 | 活动 | 类别 | 交通 | 费用 | 备注 |
"日期"写 2026-11-01 或 D1/D2；"时间"写 09:00 或 09:00-11:30；"类别"为 交通/住宿/餐饮/门票/购物/其他；
"费用"写数字（总费用，不是人均）。sheet 会检查：同一天时间重叠或倒序、多日行程缺住宿的晚上、预算超支。
Exit codes: 0 ok（可能有提示）, 1 输入错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.chart import PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
DAY_FILLS = ("F0FDFA", "FFFFFF")
CATEGORIES = ("交通", "住宿", "餐饮", "门票", "购物", "其他")
CAT_WORDS = {
    "交通": (
        "机票",
        "高铁",
        "火车",
        "动车",
        "打车",
        "地铁",
        "公交",
        "包车",
        "租车",
        "加油",
        "停车",
        "船",
        "航班",
    ),
    "住宿": ("酒店", "民宿", "住宿", "入住", "客栈"),
    "餐饮": ("早餐", "午餐", "晚餐", "午饭", "晚饭", "火锅", "小吃", "餐", "咖啡"),
    "门票": ("门票", "景区", "博物馆", "公园", "乐园", "演出", "索道"),
    "购物": ("购物", "特产", "伴手礼", "商场"),
}


class TripError(ValueError):
    pass


def parse_doc(text: str) -> tuple[dict[str, str], list[dict[str, str]]]:
    lines = text.replace("\r\n", "\n").split("\n")
    meta: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                for raw in lines[1:i]:
                    key, sep, value = raw.replace("：", ":", 1).partition(":")
                    if sep:
                        meta[key.strip()] = value.strip()
                lines = lines[i + 1 :]
                break
    rows = [
        [c.strip() for c in ln.strip().strip("|").split("|")]
        for ln in lines
        if ln.strip().startswith("|")
    ]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    if len(rows) < 2:
        raise TripError("没有找到行程表格")
    header = rows[0]
    for need in ("日期", "活动"):
        if need not in header:
            raise TripError(f"行程表缺少「{need}」列")
    items = []
    for i, r in enumerate(rows[1:], start=1):
        if len(r) != len(header):
            raise TripError(f"第 {i} 行有 {len(r)} 列，表头有 {len(header)} 列")
        items.append(dict(zip(header, r, strict=True)))
    return meta, items


def money(text: str) -> float:
    cleaned = re.sub(r"[¥￥,，元\s]", "", text or "")
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError as exc:
        raise TripError(f"费用「{text}」不是数字") from exc


def day_key(text: str, start: dt.date | None) -> tuple[str, dt.date | None]:
    m = re.fullmatch(r"[Dd](?:ay)?\s*(\d+)", text.strip())
    if m:
        n = int(m.group(1))
        when = start + dt.timedelta(days=n - 1) if start else None
        return (when.isoformat() if when else f"D{n}"), when
    try:
        when = dt.date.fromisoformat(text.strip().replace("/", "-"))
        return when.isoformat(), when
    except ValueError:
        return text.strip(), None


def minutes(text: str) -> tuple[int, int] | None:
    m = re.match(r"(\d{1,2})[:：](\d{2})(?:\s*[-~至]\s*(\d{1,2})[:：](\d{2}))?", text or "")
    if not m:
        return None
    start = int(m.group(1)) * 60 + int(m.group(2))
    end = int(m.group(3)) * 60 + int(m.group(4)) if m.group(3) else start
    return start, end


def guess_category(item: dict[str, str]) -> str:
    given = item.get("类别", "").strip()
    if given in CATEGORIES:
        return given
    # 只看活动和备注；“交通”列是怎么去那儿，不代表这笔费用的类别
    text = f"{item.get('活动', '')} {item.get('备注', '')}"
    for cat, words in CAT_WORDS.items():
        if any(w in text for w in words):
            return cat
    return "其他"


def cmd_sheet(src: Path, out: Path) -> int:
    meta, items = parse_doc(src.read_text(encoding="utf-8"))
    start = None
    if meta.get("出发"):
        try:
            start = dt.date.fromisoformat(meta["出发"].strip())
        except ValueError as exc:
            raise TripError(f"出发日期「{meta['出发']}」应写成 2026-11-01") from exc
    people = int(meta.get("人数", "1") or 1)
    budget = money(meta.get("预算", "")) if meta.get("预算") else 0.0
    warnings: list[str] = []
    days: dict[str, list[dict[str, str]]] = defaultdict(list)
    order: list[str] = []
    for it in items:
        key, _ = day_key(it["日期"], start)
        it["_day"], it["_cat"], it["_cost"] = (
            key,
            guess_category(it),
            str(money(it.get("费用", ""))),
        )
        if key not in days:
            order.append(key)
        days[key].append(it)
    for key in order:
        spans = [(minutes(it.get("时间", "")), it) for it in days[key]]
        timed = [(s, it) for s, it in spans if s]
        for (a, ia), (b, ib) in zip(timed, timed[1:], strict=False):
            if b[0] < a[0]:
                warnings.append(
                    f"{key}：「{ib['活动']}」的时间早于前一项「{ia['活动']}」，请检查顺序"
                )
            elif b[0] < a[1]:
                warnings.append(f"{key}：「{ia['活动']}」与「{ib['活动']}」时间重叠")
    if len(order) > 1:
        # “两晚”“3 晚”的住宿覆盖之后的夜晚
        covered: set[int] = set()
        for idx, key in enumerate(order):
            for it in days[key]:
                if it["_cat"] == "住宿" or "住" in it.get("活动", ""):
                    m = re.search(
                        r"([0-9]+|[一二两三四五六七八九十])\s*晚",
                        f"{it.get('活动', '')} {it.get('备注', '')}",
                    )
                    nights = 1
                    if m:
                        nights = (
                            int(m.group(1))
                            if m.group(1).isdigit()
                            else "一二三四五六七八九十".find(m.group(1).replace("两", "二")) + 1
                        )
                    covered.update(range(idx, idx + max(nights, 1)))
        for idx, key in enumerate(order[:-1]):
            if idx not in covered:
                warnings.append(f"{key}：当晚没有安排住宿")

    wb = Workbook()
    ws = wb.active
    ws.title = "行程"
    title = meta.get("行程", "旅行行程")
    ws.cell(1, 1, title).font = Font(size=14, bold=True)
    info = "　".join(f"{k}：{meta[k]}" for k in ("出发", "人数", "预算") if meta.get(k))
    ws.cell(2, 1, info).font = Font(color="6B7280")
    cols = ["日期", "时间", "地点", "活动", "类别", "交通", "费用", "备注"]
    for j, name in enumerate(cols, start=1):
        c = ws.cell(4, j, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    r = 5
    for d_index, key in enumerate(order):
        fill = PatternFill("solid", fgColor=DAY_FILLS[d_index % 2])
        for it in days[key]:
            values = [key, it.get("时间", ""), it.get("地点", ""), it.get("活动", ""), it["_cat"], it.get("交通", ""),
                      float(it["_cost"]), it.get("备注", "")]  # fmt: skip
            for j, v in enumerate(values, start=1):
                c = ws.cell(r, j, v)
                c.fill = fill
                c.alignment = Alignment(wrap_text=j in (3, 4, 8), vertical="top")
            ws.cell(r, 7).number_format = "#,##0"
            r += 1
    ws.cell(r, 6, "合计").font = Font(bold=True)
    ws.cell(r, 7, f"=SUM(G5:G{r - 1})").font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (12, 12, 16, 30, 8, 14, 10, 24), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    daily = wb.create_sheet("每日费用")
    daily.append(["日期", "项目数", "费用", "人均"])
    for key in order:
        cost = sum(float(it["_cost"]) for it in days[key])
        daily.append([key, len(days[key]), cost, round(cost / people, 2)])

    cat_ws = wb.create_sheet("分类预算")
    cat_ws.append(["类别", "费用", "占比"])
    totals = {c: sum(float(it["_cost"]) for it in items if it["_cat"] == c) for c in CATEGORIES}
    grand = sum(totals.values())
    for c in CATEGORIES:
        if totals[c]:
            cat_ws.append([c, totals[c], totals[c] / grand if grand else None])
            cat_ws.cell(cat_ws.max_row, 3).number_format = "0.0%"
    cat_ws.append(["合计", grand, None])
    if budget:
        cat_ws.append(["预算", budget, None])
        cat_ws.append(["剩余", budget - grand, None])
        if grand > budget:
            warnings.append(
                f"预计花费 {grand:,.0f} 元，超出预算 {budget:,.0f} 元 {grand - budget:,.0f} 元"
            )
    n_cat = sum(1 for c in CATEGORIES if totals[c])
    if n_cat:
        pie = PieChart()
        pie.title = "费用构成"
        pie.add_data(
            Reference(cat_ws, min_col=2, min_row=1, max_row=n_cat + 1), titles_from_data=True
        )
        pie.set_categories(Reference(cat_ws, min_col=1, min_row=2, max_row=n_cat + 1))
        cat_ws.add_chart(pie, "E2")
    for sheet in (daily, cat_ws):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(
        f"已生成 {out}：{len(order)} 天 {len(items)} 项，预计花费 {grand:,.0f} 元（人均 {grand / people:,.0f} 元）"
    )
    for w in warnings:
        print(f"提示：{w}")
    return 0


# ------------------------------------------------------------------ packing

BASE = {
    "证件与钱": ["身份证", "手机", "银行卡/少量现金", "交通和住宿订单截图（离线保存）"],
    "电子": ["充电器和数据线", "充电宝（≤20000mAh，登机需随身）", "耳机"],
    "洗漱护理": [
        "牙刷牙膏",
        "洗面奶",
        "护肤品小样",
        "毛巾（可选）",
        "常用药（感冒、肠胃、创可贴）",
    ],
}
CLOTHES = {  # (item, per_day ratio, max)
    "hot": [("短袖", 1.0, 6), ("短裤/裙子", 0.5, 3), ("凉鞋或透气鞋", 0, 1), ("防晒霜", 0, 1), ("遮阳帽/墨镜", 0, 1)],
    "mild": [("长袖上衣", 0.7, 5), ("长裤", 0.4, 3), ("薄外套", 0, 1), ("舒适步行鞋", 0, 1)],
    "cold": [("保暖内衣", 0.4, 3), ("毛衣/抓绒", 0.4, 3), ("羽绒服/厚外套", 0, 1), ("围巾手套帽子", 0, 1), ("保暖袜", 0.7, 5), ("润唇膏", 0, 1)],
    "rain": [("折叠伞", 0, 1), ("防水外套", 0, 1), ("防水鞋套/备用鞋", 0, 1), ("防水袋", 0, 2)],
}  # fmt: skip
EXTRA = {
    "city": ["城市交通卡/乘车码", "小背包"],
    "beach": ["泳衣", "沙滩拖鞋", "高倍防晒", "晒后修复", "防水手机袋"],
    "hiking": ["登山鞋", "速干衣裤", "登山杖", "头灯", "饮用水和能量补给", "急救包", "离线地图"],
    "business": ["正装/商务衬衫", "皮鞋", "笔记本电脑及转换头", "名片", "会议资料"],
    "family": ["儿童证件（户口本/出生证明）", "儿童常用药和体温计", "湿巾", "零食", "儿童安抚玩具", "备用衣物"],
    "abroad": ["护照（有效期≥6 个月）", "签证/入境材料", "转换插头", "境外流量卡/漫游", "外币或可用的信用卡", "旅行保险单", "证件复印件（与原件分开放）"],
}  # fmt: skip
CLIMATE_ZH = {"hot": "炎热", "mild": "温和", "cold": "寒冷", "rain": "多雨"}
TYPE_ZH = {
    "city": "城市观光",
    "beach": "海岛海边",
    "hiking": "徒步户外",
    "business": "商务出差",
    "family": "亲子出游",
    "abroad": "出境",
}


def cmd_packing(
    days: int, climates: list[str], types: list[str], people: int, out: Path | None
) -> int:
    for c in climates:
        if c not in CLOTHES:
            raise TripError(f"气候「{c}」不支持，可选：{', '.join(CLOTHES)}")
    for t in types:
        if t not in EXTRA:
            raise TripError(f"出行类型「{t}」不支持，可选：{', '.join(EXTRA)}")
    sections: dict[str, list[str]] = {k: list(v) for k, v in BASE.items()}
    clothes: list[str] = [f"内衣袜子 × {min(days + 1, 7)}", "睡衣 × 1"]
    for c in climates:
        for item, ratio, cap in CLOTHES[c]:
            n = max(1, min(cap, round(days * ratio))) if ratio else cap
            clothes.append(f"{item} × {n}" if n > 1 else item)
    sections["衣物"] = clothes
    for t in types:
        sections[TYPE_ZH[t]] = list(EXTRA[t])
    if days > 7:
        sections["衣物"].append("洗衣液小包（超过一周建议中途洗衣）")
    title = f"行李清单：{days} 天，{'/'.join(CLIMATE_ZH[c] for c in climates)}，{'/'.join(TYPE_ZH[t] for t in types)}"
    if people > 1:
        title += f"，{people} 人（衣物和洗漱按每人一份准备）"
    print(f"# {title}\n")
    for name, items in sections.items():
        print(f"## {name}")
        for it in items:
            print(f"- [ ] {it}")
        print()
    if out:
        wb = Workbook()
        ws = wb.active
        ws.title = "行李清单"
        ws.append(["分类", "物品", "已准备"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
        for name, items in sections.items():
            for it in items:
                ws.append([name, it, "☐"])
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 40
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        print(f"已生成 {out}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sheet")
    s.add_argument("input")
    s.add_argument("-o", "--output", required=True)
    p = sub.add_parser("packing")
    p.add_argument("--days", type=int, required=True)
    p.add_argument("--climate", default="mild", help="hot,mild,cold,rain，可多选用逗号分隔")
    p.add_argument(
        "--type", default="city", help="city,beach,hiking,business,family,abroad，可多选"
    )
    p.add_argument("--people", type=int, default=1)
    p.add_argument("-o", "--output")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "sheet":
            return cmd_sheet(Path(args.input), Path(args.output))
        if args.days < 1:
            raise TripError("--days 至少为 1")
        split = lambda v: [x.strip() for x in v.split(",") if x.strip()]  # noqa: E731
        return cmd_packing(
            args.days,
            split(args.climate),
            split(args.type),
            args.people,
            Path(args.output) if args.output else None,
        )
    except (OSError, TripError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
