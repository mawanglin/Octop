#!/usr/bin/env python3
"""简历初筛：按评分标准给一批简历打分排序，输出 Excel（每项命中了哪些词，便于人工复核）。

Usage:
    python scripts/screen.py 评分标准.csv 简历目录/ -o 初筛结果.xlsx [--today 2026-10-18]

评分标准 CSV（UTF-8 或 GBK），列：维度,权重,关键词,必备
    - 关键词用 | 分隔，命中任意一个即得该维度满分；
    - 维度名为「学历」时关键词写最低学历（大专/本科/硕士/博士），为「工作年限」时写最少年数；
    - 必备 填 是/否；必备项未满足的简历标红，排在后面。
简历支持 .txt / .md / .docx；有 pdftotext 时也支持 .pdf（图片型 PDF 读不到文字，会标注）。
初筛结果只是参考，必须人工复核；不得以性别、年龄、婚育、民族、籍贯等作为评分维度（脚本会拒绝）。
Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

DEGREES = ["中专", "大专", "本科", "硕士", "博士"]
DEGREE_ALIAS = {
    "专科": "大专",
    "学士": "本科",
    "研究生": "硕士",
    "master": "硕士",
    "bachelor": "本科",
    "phd": "博士",
}
FORBIDDEN = ("性别", "男性", "女性", "年龄", "婚姻", "婚育", "已婚", "未婚", "生育", "民族", "籍贯", "户籍", "宗教", "信仰",
             "身高", "体重", "相貌", "长相", "星座", "生肖", "政治面貌")  # fmt: skip
HEADER_FILL = PatternFill("solid", fgColor="0F766E")
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")
TOP_FILL = PatternFill("solid", fgColor="DCFCE7")


class ScreenError(ValueError):
    pass


@dataclass
class Criterion:
    name: str
    weight: float
    keywords: list[str]
    required: bool


def read_csv(path: Path) -> list[dict[str, str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ScreenError(f"{path} 编码无法识别")
    return [
        {k.strip(): (v or "").strip() for k, v in row.items() if k}
        for row in csv.DictReader(io.StringIO(text))
    ]


def load_criteria(path: Path) -> list[Criterion]:
    rows = read_csv(path)
    if not rows or not {"维度", "权重", "关键词"} <= set(rows[0]):
        raise ScreenError("评分标准需要列：维度,权重,关键词,必备")
    out = []
    for r in rows:
        name = r["维度"]
        if any(w in name or w in r["关键词"] for w in FORBIDDEN):
            raise ScreenError(
                f"评分维度「{name}」涉及与岗位无关的个人特征，可能构成就业歧视，请删除后再筛选"
            )
        try:
            weight = float(r["权重"])
        except ValueError:
            raise ScreenError(f"维度「{name}」的权重「{r['权重']}」不是数字") from None
        kws = [k.strip() for k in r["关键词"].split("|") if k.strip()]
        if not kws:
            raise ScreenError(f"维度「{name}」没有关键词")
        if name == "学历" and DEGREE_ALIAS.get(kws[0], kws[0]) not in DEGREES:
            raise ScreenError(f"学历的关键词请写 {'/'.join(DEGREES)} 之一")
        if name == "工作年限" and not re.fullmatch(r"\d+(\.\d+)?", kws[0]):
            raise ScreenError("工作年限的关键词请写最少年数，如 3")
        out.append(Criterion(name, weight, kws, r.get("必备", "") in ("是", "Y", "y", "必备", "1")))
    return out


def read_resume(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        return "\n".join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for row in t.rows for c in row.cells]
        )
    if suffix == ".pdf":
        exe = shutil.which("pdftotext")
        if not exe:
            raise ScreenError("读取 PDF 需要 pdftotext（poppler-utils）")
        return subprocess.run(
            [exe, "-layout", str(path), "-"], capture_output=True, text=True, timeout=60
        ).stdout
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def highest_degree(text: str) -> str | None:
    low = text.lower()
    found = [d for d in DEGREES if d in text] + [v for k, v in DEGREE_ALIAS.items() if k in low]
    return max(found, key=DEGREES.index) if found else None


def work_years(text: str, today: dt.date) -> float | None:
    m = re.search(r"(\d+(?:\.\d+)?)\s*年(?:以上)?(?:的)?(?:\S{0,8})?(?:工作)?经验", text)
    if m:
        return float(m.group(1))
    seg = re.split(r"教育|学历", text.split("工作", 1)[1] if "工作" in text else text, maxsplit=1)[
        0
    ]
    ds = [dt.date(int(y), int(mo), 1) for y, mo in re.findall(r"((?:19|20)\d{2})\s*[.\-/年]\s*(\d{1,2})", seg)
          if 1 <= int(mo) <= 12]  # fmt: skip
    ds = [d for d in ds if d <= today]
    return round((today - min(ds)).days / 365.25, 1) if ds else None


def score(
    text: str, criteria: list[Criterion], today: dt.date
) -> tuple[float, bool, list[str], dict[str, str]]:
    low = text.lower().replace(" ", "")
    total, ok, misses, detail = 0.0, True, [], {}
    for c in criteria:
        if c.name == "学历":
            need = DEGREE_ALIAS.get(c.keywords[0], c.keywords[0])
            have = highest_degree(text)
            hit = have is not None and DEGREES.index(have) >= DEGREES.index(need)
            detail[c.name] = have or "未写明"
        elif c.name == "工作年限":
            years = work_years(text, today)
            hit = years is not None and years >= float(c.keywords[0])
            detail[c.name] = f"约 {years} 年" if years is not None else "未写明"
        else:
            matched = [k for k in c.keywords if k.lower().replace(" ", "") in low]
            hit = bool(matched)
            detail[c.name] = "、".join(matched)
        if hit:
            total += c.weight
        else:
            if c.required:
                ok = False
            misses.append(c.name)
    return total, ok, misses, detail


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("criteria")
    parser.add_argument("folder")
    parser.add_argument("-o", "--output", required=True)
    parser.add_argument("--today")
    args = parser.parse_args(argv)
    try:
        today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
        criteria = load_criteria(Path(args.criteria))
        files = sorted(p for p in Path(args.folder).iterdir()
                       if p.suffix.lower() in (".txt", ".md", ".docx", ".pdf") and not p.name.startswith(("~$", ".")))  # fmt: skip
        if not files:
            raise ScreenError(f"{args.folder} 里没有 .txt/.md/.docx/.pdf 简历")
    except (OSError, ScreenError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1

    full = sum(c.weight for c in criteria)
    rows = []
    for f in files:
        try:
            text = read_resume(f)
        except (ScreenError, OSError, ValueError) as exc:
            rows.append((f.name, 0.0, False, [], {}, f"读取失败：{exc}"))
            continue
        if len(text.strip()) < 30:
            rows.append((f.name, 0.0, False, [], {}, "读不到文字（可能是图片型 PDF），请人工查看"))
            continue
        total, ok, misses, detail = score(text, criteria, today)
        rows.append((f.name, total, ok, misses, detail, ""))
    rows.sort(key=lambda r: (not r[2], -r[1], r[0]))

    wb = Workbook()
    ws = wb.active
    ws.title = "初筛结果"
    header = [
        "排名",
        "简历文件",
        f"总分（满分 {full:g}）",
        "必备项",
        *[c.name for c in criteria],
        "未满足",
        "备注",
    ]
    ws.append(header)
    for i, (name, total, ok, misses, detail, note) in enumerate(rows, start=1):
        ws.append([i, name, round(total, 1), "满足" if ok else "不满足", *[detail.get(c.name, "") for c in criteria],
                   "、".join(misses), note])  # fmt: skip
        fill = TOP_FILL if ok and total >= full * 0.8 else FAIL_FILL if not ok else None
        if fill:
            for c in ws[ws.max_row]:
                c.fill = fill
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate([6, 28, 12, 8, *([16] * len(criteria)), 20, 24], start=1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    ws.freeze_panes = "C2"
    st = wb.create_sheet("评分标准")
    st.append(["维度", "权重", "关键词（任一命中）", "必备"])
    for c in criteria:
        st.append([c.name, c.weight, " | ".join(c.keywords), "是" if c.required else "否"])
    st.append([])
    st.append(["说明：关键词匹配只用于初筛排序，结果需人工复核；未命中可能只是表述不同。"])
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))

    passed = sum(1 for r in rows if r[2])
    print(f"已生成 {out}：{len(rows)} 份简历，必备项全部满足 {passed} 份")
    for i, (name, total, ok, misses, _d, note) in enumerate(rows[:10], start=1):
        flag = "" if ok else "（必备项未满足）"
        print(
            f"{i}. {name}：{total:g}/{full:g}{flag}"
            + (f"，缺：{'、'.join(misses)}" if misses else "")
            + (f"，{note}" if note else "")
        )
    print("提示：关键词匹配只是初筛参考，请人工复核排名靠后和未满足必备项的简历。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
