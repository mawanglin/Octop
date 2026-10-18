#!/usr/bin/env python3
"""结构化面试：题目 Markdown → 面试评估表 Excel；多位面试官填好的评估表 → 汇总对比。

Usage:
    python scripts/interview_sheet.py form 面试题.md -o 评估表.xlsx [--position 数据分析师]
    python scripts/interview_sheet.py aggregate 评估表1.xlsx 评估表2.xlsx ... -o 面试汇总.xlsx

面试题 Markdown 表格：
    | 维度 | 问题 | 考察点 | 权重 |
    |------|------|--------|------|
    | 数据分析能力 | 讲一个你用数据推动业务决策的例子 | 问题拆解、方法、结果量化 | 30 |
权重可以不写（平均分配）。评估表每题按 1–5 分打分，总分按权重换算为百分制（公式自动计算）。
aggregate：按候选人汇总各面试官的加权得分和各维度均分；同一维度面试官打分相差 ≥2 分时标出，建议复盘讨论。
Exit codes: 0 ok, 1 输入错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
INPUT_FILL = PatternFill("solid", fgColor="FEF9C3")
WARN_FILL = PatternFill("solid", fgColor="FEE2E2")
THIN = Side(style="thin", color="CBD5E1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FIRST_ROW = 8  # first question row in the form
ANCHORS = [
    ("5 分", "远超预期：能举出多个具体例子，方法清晰，有可验证的结果，能举一反三"),
    ("4 分", "超出预期：例子具体，方法和结果清楚，个别细节需追问"),
    ("3 分", "符合要求：有相关经历，能说清做法，结果不够量化"),
    ("2 分", "部分符合：例子笼统，主要是参与，说不清自己的贡献"),
    ("1 分", "不符合：没有相关经历或回答与问题无关"),
]


class SheetError(ValueError):
    pass


def read_questions(path: Path) -> list[tuple[str, str, str, float]]:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")]
            for ln in path.read_text(encoding="utf-8").splitlines() if ln.strip().startswith("|")]  # fmt: skip
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    if len(rows) < 2 or "问题" not in rows[0]:
        raise SheetError(
            "面试题需要一张表格，表头至少包含「问题」，建议：| 维度 | 问题 | 考察点 | 权重 |"
        )
    h = rows[0]
    idx = {k: h.index(k) if k in h else None for k in ("维度", "问题", "考察点", "权重")}
    out = []
    for r in rows[1:]:
        if len(r) != len(h):
            raise SheetError(f"「{' | '.join(r)}」的列数和表头不一致")

        def get(k: str, row: list[str] = r) -> str:
            i = idx[k]
            return row[i] if i is not None else ""

        w = get("权重").rstrip("%")
        try:
            weight = float(w) if w else 0.0
        except ValueError:
            raise SheetError(f"权重「{w}」不是数字") from None
        out.append((get("维度") or "综合", get("问题"), get("考察点"), weight))
    if all(q[3] == 0 for q in out):
        out = [(d, q, p, 1.0) for d, q, p, _ in out]
    return out


def cmd_form(src: Path, out: Path, position: str) -> int:
    qs = read_questions(src)
    wb = Workbook()
    ws = wb.active
    ws.title = "评估表"
    ws["A1"] = f"结构化面试评估表{('（' + position + '）') if position else ''}"
    ws["A1"].font = Font(bold=True, size=14)
    for i, label in enumerate(("候选人", "面试官", "面试日期", "面试轮次"), start=2):
        ws.cell(i, 1, label).font = Font(bold=True)
        ws.cell(i, 2).fill = INPUT_FILL
        ws.cell(i, 2).border = BOX
    header = [
        "序号",
        "维度",
        "问题",
        "考察点",
        "权重",
        "评分(1-5)",
        "记录（行为事实，不写主观印象）",
    ]
    for j, v in enumerate(header, start=1):
        c = ws.cell(FIRST_ROW - 1, j, v)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        error="请填写 1–5 的整数", showErrorMessage=True)  # fmt: skip
    ws.add_data_validation(dv)
    for i, (dim, q, point, w) in enumerate(qs):
        r = FIRST_ROW + i
        for j, v in enumerate([i + 1, dim, q, point, w, None, None], start=1):
            c = ws.cell(r, j, v)
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 6).fill = INPUT_FILL
        ws.cell(r, 7).fill = INPUT_FILL
        dv.add(ws.cell(r, 6))
        ws.row_dimensions[r].height = 48
    last = FIRST_ROW + len(qs) - 1
    tot = last + 2
    ws.cell(tot, 4, "加权总分（百分制）").font = Font(bold=True)
    ws.cell(
        tot,
        6,
        f'=IFERROR(SUMPRODUCT(E{FIRST_ROW}:E{last},F{FIRST_ROW}:F{last})/SUM(E{FIRST_ROW}:E{last})/5*100,"")',
    )
    ws.cell(tot, 6).number_format = "0.0"
    ws.cell(tot, 6).font = Font(bold=True, size=12)
    ws.cell(tot + 1, 4, "结论（录用 / 待定 / 不录用）").font = Font(bold=True)
    ws.cell(tot + 1, 6).fill = INPUT_FILL
    ws.cell(tot + 2, 4, "主要优势 / 主要顾虑").font = Font(bold=True)
    ws.cell(tot + 2, 6).fill = INPUT_FILL
    ws.cell(tot + 4, 1, "评分锚点").font = Font(bold=True)
    for k, (score, desc) in enumerate(ANCHORS):
        ws.cell(tot + 5 + k, 1, score)
        ws.cell(tot + 5 + k, 2, desc)
    for col, w in zip("ABCDEFG", (8, 14, 40, 26, 7, 10, 40), strict=True):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"已生成 {out}：{len(qs)} 道题，{len({q[0] for q in qs})} 个维度；黄色格子由面试官填写")
    return 0


def read_form(path: Path) -> tuple[str, str, list[tuple[str, float, int]]]:
    ws = load_workbook(path, data_only=False)["评估表"]
    cand, interviewer = str(ws["B2"].value or "").strip(), str(ws["B3"].value or "").strip()
    if not cand:
        raise SheetError(f"{path.name}：没有填写候选人")
    items = []
    r = FIRST_ROW
    while ws.cell(r, 3).value:
        s = ws.cell(r, 6).value
        if s not in (None, ""):
            try:
                score = int(s)
            except (TypeError, ValueError):
                raise SheetError(f"{path.name} 第 {r} 行评分「{s}」不是 1–5 的整数") from None
            if not 1 <= score <= 5:
                raise SheetError(f"{path.name} 第 {r} 行评分 {score} 超出 1–5")
            items.append((str(ws.cell(r, 2).value), float(ws.cell(r, 5).value or 1), score))
        r += 1
    return cand, interviewer or path.stem, items


def cmd_aggregate(files: list[Path], out: Path) -> int:
    per: dict[str, dict[str, dict[str, list[tuple[float, int]]]]] = defaultdict(
        lambda: defaultdict(lambda: defaultdict(list))
    )
    for f in files:
        cand, who, items = read_form(f)
        if not items:
            print(f"提示：{f.name} 没有任何评分，已跳过")
            continue
        for dim, w, s in items:
            per[cand][who][dim].append((w, s))
    if not per:
        raise SheetError("没有可汇总的评分")
    wb = Workbook()
    ws = wb.active
    ws.title = "候选人汇总"
    ws.append(["候选人", "面试官人数", "平均加权分（百分制）", "最高", "最低", "分歧维度"])
    detail = wb.create_sheet("维度明细")
    detail.append(["候选人", "维度", "面试官", "平均分(1-5)"])
    summary = []
    for cand, by_who in per.items():
        totals = []
        dim_scores: dict[str, list[float]] = defaultdict(list)
        for who, dims in by_who.items():
            wsum = sum(w for d in dims.values() for w, _ in d)
            totals.append(sum(w * s for d in dims.values() for w, s in d) / wsum / 5 * 100)
            for dim, vals in dims.items():
                avg = sum(s for _, s in vals) / len(vals)
                dim_scores[dim].append(avg)
                detail.append([cand, dim, who, round(avg, 2)])
        split = [d for d, v in dim_scores.items() if len(v) > 1 and max(v) - min(v) >= 2]
        summary.append(
            (cand, len(by_who), sum(totals) / len(totals), max(totals), min(totals), split)
        )
    summary.sort(key=lambda x: -x[2])
    for cand, n, avg, hi, lo, split in summary:
        ws.append([cand, n, round(avg, 1), round(hi, 1), round(lo, 1), "、".join(split)])
        if split:
            ws.cell(ws.max_row, 6).fill = WARN_FILL
    for sheet in (ws, detail):
        for c in sheet[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        for col in "ABCDEF":
            sheet.column_dimensions[col].width = 16
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"已生成 {out}：{len(summary)} 位候选人")
    for cand, n, avg, _hi, _lo, split in summary:
        print(
            f"- {cand}：{n} 位面试官，平均 {avg:.1f} 分"
            + (f"；分歧较大：{'、'.join(split)}（建议面试官复盘讨论）" if split else "")
        )
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    f = sub.add_parser("form")
    f.add_argument("input")
    f.add_argument("-o", "--output", required=True)
    f.add_argument("--position", default="")
    a = sub.add_parser("aggregate")
    a.add_argument("files", nargs="+")
    a.add_argument("-o", "--output", required=True)
    args = parser.parse_args(argv)
    try:
        if args.cmd == "form":
            return cmd_form(Path(args.input), Path(args.output), args.position)
        return cmd_aggregate([Path(p) for p in args.files], Path(args.output))
    except (OSError, SheetError, KeyError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
