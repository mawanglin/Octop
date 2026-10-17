#!/usr/bin/env python3
"""透视汇总 + Excel 原生图表：分组求和/计数/平均，按月/季/年汇总并算环比、同比，排名和占比。

Usage:
    python scripts/pivot_chart.py data.xlsx -o result.xlsx --rows 区域 --values 销售额 [--agg sum]
    python scripts/pivot_chart.py data.xlsx -o result.xlsx --date 日期 --period month --values 销售额 --chart line
    python scripts/pivot_chart.py data.xlsx -o result.xlsx --rows 产品 --columns 区域 --values 销售额

参数：
    --rows      分组字段（类别）；与 --date 二选一或同时使用（同时使用时 --date 作为列）
    --date      日期字段，配合 --period month|quarter|year 按时间汇总，自动计算环比（以及跨年数据的同比）
    --columns   交叉分组字段（二维透视）
    --values    数值字段；--agg 为 count 时可省略
    --agg       sum（默认）| count | mean
    --top N     只保留前 N 个分组，其余合并为“其他”
    --chart     auto（默认）| column | bar | line | pie | none
输出 xlsx：汇总表（带合计、占比/环比/同比列）+ 图表 + 原始数据副本；同时在终端打印 Markdown 汇总，便于写结论。
Exit codes: 0 ok, 1 参数或数据错误, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from profile_data import read_table, to_date, to_number

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
TOTAL_FILL = PatternFill("solid", fgColor="FEF3C7")


class PivotError(ValueError):
    pass


def period_key(value: Any, period: str) -> str | None:
    d = to_date(value)
    if d is None:
        return None
    if period == "year":
        return f"{d.year}"
    if period == "quarter":
        return f"{d.year}-Q{(d.month - 1) // 3 + 1}"
    return f"{d.year}-{d.month:02d}"


def shift_year(key: str) -> str:
    year, _, rest = key.partition("-")
    return f"{int(year) - 1}" + (f"-{rest}" if rest else "")


def aggregate(values: list[float], agg: str) -> float:
    if agg == "count":
        return float(len(values))
    if not values:
        return 0.0
    return sum(values) / len(values) if agg == "mean" else sum(values)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input")
    parser.add_argument("-o", "--output", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--rows")
    parser.add_argument("--date")
    parser.add_argument("--period", choices=["month", "quarter", "year"], default="month")
    parser.add_argument("--columns")
    parser.add_argument("--values")
    parser.add_argument("--agg", choices=["sum", "count", "mean"], default="sum")
    parser.add_argument("--top", type=int)
    parser.add_argument(
        "--chart", choices=["auto", "column", "bar", "line", "pie", "none"], default="auto"
    )
    args = parser.parse_args(argv)
    try:
        header, body = read_table(Path(args.input), args.sheet, args.header_row, 1_000_000)
        result = run(args, header, body)
    except (OSError, ValueError, KeyError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    print(result)
    return 0


def run(args: argparse.Namespace, header: list[str], body: list[list[Any]]) -> str:
    def col(name: str | None) -> int | None:
        if name is None:
            return None
        if name not in header:
            raise PivotError(f"找不到字段「{name}」，可用字段：{'、'.join(header)}")
        return header.index(name)

    if not args.rows and not args.date:
        raise PivotError("需要 --rows（分组字段）或 --date（日期字段）")
    if args.agg != "count" and not args.values:
        raise PivotError("--agg 为 sum/mean 时需要 --values（数值字段）")
    by_time = bool(args.date) and not args.rows
    row_i = col(args.rows) if args.rows else col(args.date)
    col_name = args.columns or (args.date if args.rows and args.date else None)
    col_i = col(col_name)
    val_i = col(args.values)

    cells: dict[tuple[str, str], list[float]] = defaultdict(list)
    skipped = 0
    for r in body:
        key = (
            period_key(r[row_i], args.period)
            if by_time
            else (str(r[row_i]).strip() if r[row_i] not in (None, "") else "（空）")
        )
        if key is None:
            skipped += 1
            continue
        sub = ""
        if col_i is not None:
            sub = (
                period_key(r[col_i], args.period)
                if col_name == args.date
                else str(r[col_i]).strip()
            )
            if sub is None:
                skipped += 1
                continue
        if args.agg == "count":
            cells[(key, sub)].append(1.0)
        else:
            n = to_number(r[val_i])  # type: ignore[index]
            if n is None:
                skipped += 1
                continue
            cells[(key, sub)].append(n)
    if not cells:
        raise PivotError("没有可汇总的数据（检查字段名和数值格式）")

    row_keys = sorted({k for k, _ in cells}) if by_time else None
    col_keys = sorted({s for _, s in cells}) if col_i is not None else [""]
    totals = {
        k: aggregate([v for (rk, _), vs in cells.items() if rk == k for v in vs], args.agg)
        for k, _ in cells
    }
    if row_keys is None:
        row_keys = sorted(totals, key=lambda k: -totals[k])
        if args.top and len(row_keys) > args.top:
            keep, rest = row_keys[: args.top], row_keys[args.top :]
            for (rk, sub), vs in list(cells.items()):
                if rk in rest:
                    cells[("其他", sub)].extend(vs)
                    del cells[(rk, sub)]
            totals["其他"] = (
                sum(totals.pop(k) for k in rest)
                if args.agg != "mean"
                else aggregate(
                    [v for (rk, _), vs in cells.items() if rk == "其他" for v in vs], "mean"
                )
            )
            row_keys = [*keep, "其他"]

    agg_label = {"sum": "合计", "count": "计数", "mean": "平均"}[args.agg]
    value_label = f"{args.values or '记录'}（{agg_label}）"
    first_label = args.date if by_time else args.rows
    headers = [first_label]
    if col_i is not None:
        headers += col_keys + ["总计"]
    else:
        headers.append(value_label)
    extra: list[str] = []
    grand = sum(totals[k] for k in row_keys)
    if not by_time and args.agg != "mean":
        extra.append("占比")
    if by_time:
        extra.append("环比")
        if any(shift_year(k) in totals for k in row_keys):
            extra.append("同比")
    headers += extra

    table: list[list[Any]] = []
    for i, k in enumerate(row_keys):
        row: list[Any] = [k]
        if col_i is not None:
            row += [
                round(aggregate(cells.get((k, s), []), args.agg), 4) if (k, s) in cells else None
                for s in col_keys
            ]
        row.append(round(totals[k], 4))
        if "占比" in extra:
            row.append(totals[k] / grand if grand else None)
        if "环比" in extra:
            prev = totals.get(row_keys[i - 1]) if i > 0 else None
            row.append((totals[k] - prev) / prev if prev else None)
        if "同比" in extra:
            last = totals.get(shift_year(k))
            row.append((totals[k] - last) / last if last else None)
        table.append(row)

    out = Path(args.output)
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(24, len(str(h)) * 2 + 4))
    for i, row in enumerate(table, start=2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(i, j, v)
            if headers[j - 1] in ("占比", "环比", "同比"):
                c.number_format = "0.0%"
            elif isinstance(v, float):
                c.number_format = "#,##0.##"
    n_value_cols = (len(col_keys) + 1) if col_i is not None else 1
    if args.agg != "mean":
        total_row = len(table) + 2
        ws.cell(total_row, 1, "合计").font = Font(bold=True)
        for j in range(2, 2 + n_value_cols):
            letter = get_column_letter(j)
            c = ws.cell(total_row, j, f"=SUM({letter}2:{letter}{total_row - 1})")
            c.font = Font(bold=True)
            c.number_format = "#,##0.##"
        for j in range(1, len(headers) + 1):
            ws.cell(total_row, j).fill = TOTAL_FILL
    ws.freeze_panes = "B2"

    kind = args.chart
    if kind == "auto":
        kind = (
            "line"
            if by_time
            else (
                "column"
                if col_i is not None or len(table) > 6
                else "pie"
                if args.agg != "mean" and len(table) <= 6
                else "column"
            )
        )
    if kind != "none":
        series_cols = range(2, 2 + len(col_keys)) if col_i is not None else range(2, 3)
        if kind == "pie":
            chart: Any = PieChart()
        elif kind == "line":
            chart = LineChart()
        else:
            chart = BarChart()
            chart.type = "bar" if kind == "bar" else "col"
            if col_i is not None:
                chart.grouping = "clustered"
        chart.title = f"{first_label} · {value_label}"
        chart.height, chart.width = 9, 18
        for j in series_cols:
            data = Reference(ws, min_col=j, min_row=1, max_row=len(table) + 1)
            chart.add_data(data, titles_from_data=True)
            if kind == "pie":
                break
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(table) + 1))
        if kind == "line":
            for s in chart.series:
                s.smooth = False
        ws.add_chart(chart, f"A{len(table) + 5}")  # below the table: visible and printable

    raw = wb.create_sheet("原始数据")
    raw.append(header)
    for r in body:
        raw.append([v if not isinstance(v, str) or len(v) < 32000 else v[:32000] for v in r])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))

    lines = [f"已生成 {out}（{len(table)} 组，图表：{kind}）", ""]
    lines.append("| " + " | ".join(map(str, headers)) + " |")
    lines.append("|" + "---|" * len(headers))
    for row in table[:30]:
        shown = []
        for h, v in zip(headers, row, strict=True):
            if v is None:
                shown.append("")
            elif h in ("占比", "环比", "同比"):
                shown.append(f"{v:+.1%}" if h != "占比" else f"{v:.1%}")
            elif isinstance(v, float):
                shown.append(f"{v:,.2f}".rstrip("0").rstrip("."))
            else:
                shown.append(str(v))
        lines.append("| " + " | ".join(shown) + " |")
    if len(table) > 30:
        lines.append(f"（仅显示前 30 组，共 {len(table)} 组）")
    if args.agg != "mean":
        lines.append(f"\n总计：{grand:,.2f}")
    if skipped:
        lines.append(f"提示：{skipped} 行因字段为空、日期或数值无法识别而未计入，结论中需要说明")
    return "\n".join(lines)


if __name__ == "__main__":
    sys.exit(main())
