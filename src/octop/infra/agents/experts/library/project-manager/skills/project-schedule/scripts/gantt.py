#!/usr/bin/env python3
"""项目排期：任务表 → 按工作日排期、计算关键路径和浮动时间，生成甘特图 Excel，并对照今天检查进度。

Usage:
    python scripts/gantt.py 任务.md --start 2026-11-02 -o 项目排期.xlsx [--today 2026-11-20] [--holidays 节假日.txt] [--weekends]

任务表（Markdown 表格）：
    | 编号 | 任务 | 负责人 | 工期 | 前置任务 | 进度 |
    | 1 | 需求调研 | 张三 | 5 | | 100% |
    | 2 | 方案设计 | 李四 | 3 | 1 | 50% |
"工期"为工作日数（0 表示里程碑）；"前置任务"写编号，多个用逗号分隔，默认完成-开始（FS）关系，
可写 2+1（前置完成后再等 1 天）；"进度"可选。默认跳过周六周日，--holidays 每行一个日期（法定节假日），
--weekends 表示周末也工作。
输出：每个任务的最早 / 最晚开始结束、总浮动时间、是否关键任务；甘特图（关键任务红色、里程碑◆、已完成部分深色）；
终端汇总项目完成日期、关键路径和按计划应完成但进度落后的任务。
Exit codes: 0 ok, 1 输入错误（如循环依赖、前置任务不存在）, 2 缺少 openpyxl。
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - environment check
    print("缺少依赖 openpyxl：请运行 `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)


class PlanError(ValueError):
    pass


@dataclass
class Task:
    tid: str
    name: str
    owner: str
    days: int
    preds: list[tuple[str, int]]
    progress: float | None
    es: int = 0
    ef: int = 0
    ls: int = 0
    lf: int = 0
    succ: list[str] = field(default_factory=list)

    @property
    def slack(self) -> int:
        return self.ls - self.es

    @property
    def critical(self) -> bool:
        return self.slack == 0


def read_tasks(path: Path) -> list[Task]:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")]
            for ln in path.read_text(encoding="utf-8").splitlines() if ln.strip().startswith("|")]  # fmt: skip
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    if len(rows) < 2:
        raise PlanError("没有找到任务表格")
    h = rows[0]
    need = ("编号", "任务", "工期")
    for k in need:
        if k not in h:
            raise PlanError(
                f"任务表缺少「{k}」列（需要：编号、任务、工期，可选：负责人、前置任务、进度）"
            )
    tasks: list[Task] = []
    for r in rows[1:]:
        if len(r) != len(h):
            raise PlanError(f"「{' | '.join(r)}」的列数和表头不一致")
        rec = dict(zip(h, r, strict=True))
        try:
            days = int(float(rec["工期"].rstrip("天d ")))
        except ValueError:
            raise PlanError(f"任务 {rec['编号']} 的工期「{rec['工期']}」不是数字") from None
        preds = []
        for p in re.split(r"[,，、\s]+", rec.get("前置任务", "")):
            if not p:
                continue
            m = re.fullmatch(r"([^+]+?)(?:\+(\d+))?", p)
            assert m
            preds.append((m.group(1), int(m.group(2) or 0)))
        prog = rec.get("进度", "").rstrip("%")
        tasks.append(
            Task(
                rec["编号"],
                rec["任务"],
                rec.get("负责人", ""),
                days,
                preds,
                float(prog) / 100 if prog else None,
            )
        )
    ids = [t.tid for t in tasks]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        raise PlanError(f"任务编号重复：{'、'.join(sorted(dup))}")
    return tasks


def schedule(tasks: list[Task]) -> list[Task]:
    by_id = {t.tid: t for t in tasks}
    for t in tasks:
        for p, _lag in t.preds:
            if p not in by_id:
                raise PlanError(f"任务 {t.tid}「{t.name}」的前置任务 {p} 不存在")
            by_id[p].succ.append(t.tid)
    order: list[Task] = []
    state: dict[str, int] = {}

    def visit(t: Task, path: list[str]) -> None:
        if state.get(t.tid) == 2:
            return
        if state.get(t.tid) == 1:
            raise PlanError("存在循环依赖：" + " → ".join([*path[path.index(t.tid) :], t.tid]))
        state[t.tid] = 1
        for p, _lag in t.preds:
            visit(by_id[p], [*path, t.tid])
        state[t.tid] = 2
        order.append(t)

    for t in tasks:
        visit(t, [])
    for t in order:  # forward pass (day offsets, end exclusive)
        t.es = max((by_id[p].ef + lag for p, lag in t.preds), default=0)
        t.ef = t.es + t.days
    end = max(t.ef for t in tasks)
    for t in reversed(order):  # backward pass
        succ_ls = [
            by_id[s].ls - next(lag for p, lag in by_id[s].preds if p == t.tid) for s in t.succ
        ]
        t.lf = min(succ_ls, default=end)
        t.ls = t.lf - t.days
    return order


class Calendar:
    def __init__(self, start: dt.date, holidays: set[dt.date], weekends: bool) -> None:
        self.days: list[dt.date] = []
        d = start
        self._hol, self._we = holidays, weekends
        while len(self.days) < 2000:
            if self.is_work(d):
                self.days.append(d)
            d += dt.timedelta(days=1)

    def is_work(self, d: dt.date) -> bool:
        return d not in self._hol and (self._we or d.weekday() < 5)

    def date(self, offset: int) -> dt.date:
        return self.days[offset]

    def span(self, t: Task) -> tuple[dt.date, dt.date]:
        """Calendar start/finish; a milestone sits on the day its predecessors finish."""
        if t.days == 0:
            d = self.date(max(t.es - 1, 0))
            return d, d
        return self.date(t.es), self.date(t.ef - 1)

    def late_span(self, t: Task) -> tuple[dt.date, dt.date]:
        if t.days == 0:
            d = self.date(max(t.ls - 1, 0))
            return d, d
        return self.date(t.ls), self.date(t.lf - 1)

    def offset_of(self, d: dt.date) -> int:
        """Number of working days elapsed before *d* (i.e. completed by the start of d)."""
        return sum(1 for x in self.days if x < d)


def write_xlsx(tasks: list[Task], cal: Calendar, out: Path, today: dt.date | None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "排期"
    header = [
        "编号",
        "任务",
        "负责人",
        "工期",
        "前置",
        "开始",
        "结束",
        "最晚开始",
        "最晚结束",
        "浮动(天)",
        "关键",
        "进度",
    ]
    ws.append(header)
    red = PatternFill("solid", fgColor="FECACA")
    for t in tasks:
        start, finish = cal.span(t)
        lstart, lfinish = cal.late_span(t)
        ws.append([t.tid, t.name, t.owner, t.days, ",".join(p + (f"+{lag}" if lag else "") for p, lag in t.preds), start, finish,
                   lstart, lfinish, t.slack, "是" if t.critical else "", t.progress])  # fmt: skip
        for j in (6, 7, 8, 9):
            ws.cell(ws.max_row, j).number_format = "mm-dd"
        ws.cell(ws.max_row, 12).number_format = "0%"
        if t.critical:
            ws.cell(ws.max_row, 11).fill = red
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E40AF")
    for col, w in zip("ABCDEFGHIJKL", (6, 24, 10, 6, 8, 8, 8, 9, 9, 8, 6, 7), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"

    g = wb.create_sheet("甘特图")
    first = cal.date(0)
    last = max(cal.span(t)[1] for t in tasks)
    span = [first + dt.timedelta(days=i) for i in range((last - first).days + 1)]
    g.append(["编号", "任务", "负责人", *[f"{d:%m-%d}" for d in span]])
    g.append(["", "", "", *["一二三四五六日"[d.weekday()] for d in span]])
    fills = {"crit": PatternFill("solid", fgColor="EF4444"), "norm": PatternFill("solid", fgColor="60A5FA"),
             "done": PatternFill("solid", fgColor="1E3A8A"), "off": PatternFill("solid", fgColor="F3F4F6"),
             "today": PatternFill("solid", fgColor="FEF08A")}  # fmt: skip
    for t in tasks:
        g.append([t.tid, t.name, t.owner, *[""] * len(span)])
        row = g.max_row
        work = [cal.date(i) for i in range(t.es, t.ef)]
        done_n = round(len(work) * (t.progress or 0))
        for j, d in enumerate(span, start=4):
            cell = g.cell(row, j)
            if not cal.is_work(d):
                cell.fill = fills["off"]
            if d in work:
                k = work.index(d)
                cell.fill = fills["done"] if k < done_n else fills["crit" if t.critical else "norm"]
            if t.days == 0 and d == cal.span(t)[0]:
                cell.value = "◆"
                cell.alignment = Alignment(horizontal="center")
            if (
                today
                and d == today
                and not cell.fill.fgColor.rgb.endswith(("EF4444", "60A5FA", "1E3A8A"))
            ):
                cell.fill = fills["today"]
    for j in range(4, len(span) + 4):
        g.column_dimensions[g.cell(1, j).column_letter].width = 4
        g.cell(1, j).alignment = Alignment(text_rotation=90, horizontal="center")
    g.column_dimensions["B"].width = 24
    g.freeze_panes = "D3"
    g.cell(
        g.max_row + 2,
        2,
        "红色：关键任务（延误会推迟整个项目）；蓝色：非关键任务；深蓝：已完成部分；◆：里程碑；灰色：非工作日；黄色：今天",
    )
    for sheet in (ws, g):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input")
    parser.add_argument("--start", required=True)
    parser.add_argument("-o", "--output")
    parser.add_argument("--today")
    parser.add_argument("--holidays")
    parser.add_argument("--weekends", action="store_true")
    args = parser.parse_args(argv)
    try:
        start = dt.date.fromisoformat(args.start)
        holidays: set[dt.date] = set()
        if args.holidays:
            for ln in Path(args.holidays).read_text(encoding="utf-8").splitlines():
                if ln.strip() and not ln.startswith("#"):
                    holidays.add(dt.date.fromisoformat(ln.strip()[:10]))
        tasks = read_tasks(Path(args.input))
        order = schedule(tasks)
    except (OSError, PlanError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    cal = Calendar(start, holidays, args.weekends)
    end_off = max(t.ef for t in tasks)
    end_date = cal.date(end_off - 1) if end_off else cal.date(0)
    print(f"共 {len(tasks)} 个任务，{end_off} 个工作日，{cal.date(0)} 开始，预计 {end_date} 完成")
    crit = [t for t in order if t.critical]
    print("关键路径：" + " → ".join(f"{t.tid} {t.name}" for t in crit))
    slack = [t for t in tasks if not t.critical]
    if slack:
        print(
            "有浮动时间的任务："
            + "、".join(
                f"{t.tid} {t.name}（{t.slack} 天）" for t in sorted(slack, key=lambda x: x.slack)
            )
        )
    if args.today:
        today = dt.date.fromisoformat(args.today)
        elapsed = cal.offset_of(today)
        late = []
        for t in tasks:
            if t.days == 0 or t.progress is None:
                continue
            planned = min(1.0, max(0.0, (elapsed - t.es) / t.days))
            if planned - t.progress >= 0.2:
                late.append((t, planned))
        print(
            f"截至 {today}（第 {elapsed} 个工作日）："
            + ("进度正常" if not late else f"{len(late)} 个任务落后")
        )
        for t, planned in late:
            flag = (
                "，关键任务，会推迟项目完成日期" if t.critical else f"，还有 {t.slack} 天浮动时间"
            )
            print(
                f"- {t.tid} {t.name}（{t.owner or '未指定'}）：计划应完成 {planned:.0%}，实际 {t.progress:.0%}{flag}"
            )
    if args.output:
        write_xlsx(
            tasks, cal, Path(args.output), dt.date.fromisoformat(args.today) if args.today else None
        )
        print(f"已生成 {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
