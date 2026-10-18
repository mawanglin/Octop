#!/usr/bin/env python3
"""简历工具：Markdown 简历 → Word（3 种版式），以及简历体检。

Usage:
    python scripts/resume_docx.py build 简历.md -o 简历.docx [--style classic|modern|compact]
    python scripts/resume_docx.py check 简历.md [--today 2026-10-18]

简历 Markdown：
    # 姓名
    求职意向：数据分析师 | 上海            （标题下的普通行作为联系方式/意向，原样保留）
    电话：138-0000-0000 · 邮箱：name@example.com
    ## 工作经历                           （二级标题 = 模块）
    ### 某某科技有限公司 | 数据分析师 | 2021.07 – 至今   （三级标题 = 经历，用 | 分隔，最后一段是时间）
    - 搭建销售看板，周报制作时间从 4 小时缩短到 30 分钟
check：联系方式、敏感信息（身份证号等）、没有数字的经历条目、弱动词开头（负责/参与/协助）、条目过长、
       时间倒序、经历之间超过 6 个月的空档和重叠、预计页数。
Exit codes: 0 ok, 1 输入错误, 2 缺少 python-docx（仅 build）。
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

STYLES = {
    #          正文字体          字号  标题色     页边距(mm) 姓名字号
    "classic": ("SimSun", "宋体", 10.5, "000000", 20, 22),
    "modern": ("Microsoft YaHei", "微软雅黑", 10, "1D4ED8", 18, 24),
    "compact": ("Microsoft YaHei", "微软雅黑", 9.5, "374151", 12, 18),
}
WEAK_STARTS = ("负责", "参与", "协助", "帮助", "配合", "做了", "主要做", "日常")
DATE = r"(\d{4})\s*[.\-/年]\s*(\d{1,2})\s*月?"
RANGE = re.compile(
    DATE + r"\s*(?:[–—~\-至到]+)\s*(?:" + DATE + r"|(至今|今|现在|present|now))", re.I
)


class ResumeError(ValueError):
    pass


@dataclass
class Entry:
    title: str
    bullets: list[str] = field(default_factory=list)
    lines: list[str] = field(default_factory=list)


@dataclass
class Section:
    name: str
    entries: list[Entry] = field(default_factory=list)


@dataclass
class Resume:
    name: str
    header: list[str]
    sections: list[Section]


def parse(text: str) -> Resume:
    name = ""
    header: list[str] = []
    sections: list[Section] = []
    cur_sec: Section | None = None
    cur: Entry | None = None
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            name = line[2:].strip()
        elif line.startswith("## "):
            cur_sec = Section(line[3:].strip())
            sections.append(cur_sec)
            cur = None
        elif line.startswith("### "):
            if cur_sec is None:
                cur_sec = Section("经历")
                sections.append(cur_sec)
            cur = Entry(line[4:].strip())
            cur_sec.entries.append(cur)
        elif re.match(r"\s*[-*•]\s+", line):
            if cur_sec is None:
                raise ResumeError(f"条目「{line.strip()}」前面缺少 ## 模块标题")
            if cur is None:
                cur = Entry("")
                cur_sec.entries.append(cur)
            cur.bullets.append(re.sub(r"^\s*[-*•]\s+", "", line))
        elif cur_sec is None:
            header.append(line.strip())
        else:
            if cur is None:
                cur = Entry("")
                cur_sec.entries.append(cur)
            cur.lines.append(line.strip())
    if not name:
        raise ResumeError("第一行请用「# 姓名」")
    if not sections:
        raise ResumeError("没有找到「## 模块」（如 ## 工作经历）")
    return Resume(name, header, sections)


def _month(y: str, m: str) -> dt.date:
    return dt.date(int(y), max(1, min(12, int(m))), 1)


def date_range(title: str, today: dt.date) -> tuple[dt.date, dt.date] | None:
    m = RANGE.search(title)
    if not m:
        return None
    start = _month(m.group(1), m.group(2))
    end = today.replace(day=1) if m.group(5) else _month(m.group(3), m.group(4))
    return start, end


def months_between(a: dt.date, b: dt.date) -> int:
    return (b.year - a.year) * 12 + b.month - a.month


def check(resume: Resume, today: dt.date) -> list[str]:
    tips: list[str] = []
    head = " ".join(resume.header)
    full = (
        head
        + "\n"
        + "\n".join(
            e.title + "\n" + "\n".join(e.bullets + e.lines)
            for s in resume.sections
            for e in s.entries
        )
    )
    if not re.search(r"(?<!\d)1[3-9]\d[\s-]?\d{4}[\s-]?\d{4}(?!\d)", full):
        tips.append("联系方式：没有找到手机号")
    if not re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", full):
        tips.append("联系方式：没有找到邮箱")
    if re.search(r"\b\d{17}[\dXx]\b", full):
        tips.append("敏感信息：简历中出现身份证号，建议删除")
    for word in ("婚姻", "已婚", "未婚", "生育", "政治面貌", "籍贯", "身高", "体重"):
        if word in full:
            tips.append(f"敏感信息：出现「{word}」，除非岗位明确要求，建议删除")
            break

    total_lines = len(resume.header) + 2
    for sec in resume.sections:
        total_lines += 2
        is_exp = any(k in sec.name for k in ("工作", "项目", "实习", "经历", "经验"))
        ranges: list[tuple[dt.date, dt.date, str]] = []
        for e in sec.entries:
            total_lines += 1 + len(e.lines)
            r = date_range(e.title, today)
            if r:
                ranges.append((*r, e.title.split("|")[0].strip()))
                if r[1] < r[0]:
                    tips.append(f"时间：「{e.title}」结束早于开始")
            for b in e.bullets:
                total_lines += max(1, -(-len(b) // 48))
                short = b if len(b) <= 24 else b[:24] + "…"
                if is_exp and not re.search(r"\d", b):
                    tips.append(f"量化：「{short}」没有数字，补充规模、频次、提升比例或结果")
                if b.startswith(WEAK_STARTS):
                    tips.append(
                        f"动词：「{short}」以弱动词开头，改成主导/搭建/优化/推动等具体动作，并写出结果"
                    )
                if len(b) > 80:
                    tips.append(f"长度：「{short}」超过 80 字，拆成两条或删去背景描述")
        if "工作" in sec.name or "实习" in sec.name:
            starts = [r[0] for r in ranges]
            if starts != sorted(starts, reverse=True):
                tips.append(f"顺序：「{sec.name}」请按时间倒序排列（最近的在最前）")
            ordered = sorted(ranges)
            for (_s1, e1, n1), (s2, _e2, n2) in zip(ordered, ordered[1:], strict=False):
                gap = months_between(e1, s2)
                if gap > 6:
                    tips.append(
                        f"空档：「{n1}」到「{n2}」之间有 {gap} 个月空档，面试时可能被问到，可简要说明"
                    )
                elif gap < 0:
                    tips.append(
                        f"重叠：「{n1}」与「{n2}」时间重叠 {-gap} 个月，确认是否为兼职或写错"
                    )
    pages = total_lines / 46
    tips.append(
        f"篇幅：预计约 {pages:.1f} 页（A4）"
        + ("，建议 3 年以下经验控制在 1 页" if pages > 1.1 else "")
    )
    return tips


def build(resume: Resume, out: Path, style: str) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError:  # pragma: no cover - environment check
        print("缺少依赖 python-docx：请运行 `pip install python-docx`", file=sys.stderr)
        sys.exit(2)
    latin, east, size, color, margin, name_size = STYLES[style]
    accent = RGBColor.from_string(color)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Mm(margin))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = latin, Pt(size)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    def para(
        text: str = "",
        *,
        bold: bool = False,
        pt: float | None = None,
        rgb: RGBColor | None = None,
        align: int | None = None,
        before: float = 0,
        after: float = 0,
    ):  # type: ignore[no-untyped-def]
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if align is not None:
            p.alignment = align
        if text:
            run = p.add_run(text)
            run.bold = bold
            if pt:
                run.font.size = Pt(pt)
            if rgb is not None:
                run.font.color.rgb = rgb
        return p

    def bottom_border(p) -> None:  # type: ignore[no-untyped-def]
        pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", color)):
            b.set(qn(f"w:{k}"), v)
        pbdr.append(b)
        p._p.get_or_add_pPr().append(pbdr)

    center = WD_ALIGN_PARAGRAPH.CENTER if style == "classic" else None
    para(
        resume.name,
        bold=True,
        pt=name_size,
        rgb=accent if style != "classic" else None,
        align=center,
        after=2,
    )
    for line in resume.header:
        para(line, pt=size - 0.5, rgb=RGBColor(0x4B, 0x55, 0x63), align=center)
    gap = 4 if style == "compact" else 8
    for s in resume.sections:
        h = para(s.name, bold=True, pt=size + 1.5, rgb=accent, before=gap, after=2)
        bottom_border(h)
        for e in s.entries:
            if e.title:
                parts = [x.strip() for x in e.title.split("|")]
                when = parts.pop() if len(parts) > 1 and RANGE.search(parts[-1]) else ""
                p = para(before=3)
                p.add_run("　".join(parts)).bold = True
                if when:
                    # Right-aligned date via a right tab stop at the text width.
                    width = sec.page_width - sec.left_margin - sec.right_margin
                    p.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)  # 2 = RIGHT
                    r = p.add_run("\t" + when)
                    r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
            for line in e.lines:
                para(line)
            for b in e.bullets:
                p = para("• " + b)
                p.paragraph_format.left_indent = Mm(4)
                p.paragraph_format.first_line_indent = Mm(-3)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("input")
    b.add_argument("-o", "--output", required=True)
    b.add_argument("--style", choices=sorted(STYLES), default="modern")
    c = sub.add_parser("check")
    c.add_argument("input")
    c.add_argument("--today")
    args = parser.parse_args(argv)
    try:
        resume = parse(Path(args.input).read_text(encoding="utf-8"))
        if args.cmd == "build":
            build(resume, Path(args.output), args.style)
            n = sum(len(e.bullets) for s in resume.sections for e in s.entries)
            print(
                f"已生成 {args.output}（{args.style} 版式）：{len(resume.sections)} 个模块，{n} 条经历要点"
            )
            return 0
        today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
        tips = check(resume, today)
        print(f"简历体检：{resume.name}，共 {len(tips) - 1} 条建议")
        for t in tips:
            print(f"- {t}")
        return 0
    except (OSError, ResumeError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
