#!/usr/bin/env python3
"""岗位匹配：从岗位描述（JD）中提取技能和要求关键词，检查简历覆盖了哪些、缺了哪些。

Usage:
    python scripts/jd_match.py JD.txt 简历.md [--json]

- JD 和简历都可以是 .txt / .md / .docx；
- 关键词来源：英文技术词（SQL、Python、Power BI…）、内置常见中文技能词、"熟悉/掌握/精通/具备…"后面的短语；
- 含"优先/加分/者优先"的行算加分项，其余算必备项；同时比对学历和工作年限要求；
- 只做字面匹配（含少量同义词），结果用于提示"简历里是否写到了"，不代表真实能力。
Exit codes: 0 ok, 1 输入错误。
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

CN_TERMS = [
    "数据分析",
    "数据挖掘",
    "数据可视化",
    "数据仓库",
    "数据治理",
    "数据建模",
    "指标体系",
    "机器学习",
    "深度学习",
    "自然语言处理",
    "计算机视觉",
    "大模型",
    "推荐系统",
    "统计学",
    "回归分析",
    "假设检验",
    "AB测试",
    "用户画像",
    "用户增长",
    "用户运营",
    "内容运营",
    "活动运营",
    "社群运营",
    "新媒体运营",
    "短视频",
    "直播",
    "私域",
    "产品设计",
    "需求分析",
    "竞品分析",
    "原型设计",
    "用户研究",
    "交互设计",
    "项目管理",
    "敏捷开发",
    "跨部门沟通",
    "沟通协调",
    "团队管理",
    "团队协作",
    "带团队",
    "商业分析",
    "经营分析",
    "财务分析",
    "预算管理",
    "成本控制",
    "财务报表",
    "税务",
    "审计",
    "内控",
    "风险控制",
    "合规",
    "招聘",
    "培训",
    "绩效管理",
    "薪酬福利",
    "员工关系",
    "组织发展",
    "人力资源",
    "供应链",
    "采购",
    "仓储",
    "物流",
    "质量管理",
    "生产管理",
    "精益生产",
    "六西格玛",
    "销售",
    "大客户",
    "渠道管理",
    "商务谈判",
    "客户关系",
    "市场推广",
    "品牌策划",
    "文案",
    "广告投放",
    "搜索引擎优化",
    "前端",
    "后端",
    "微服务",
    "分布式",
    "高并发",
    "数据库",
    "云计算",
    "容器",
    "自动化测试",
    "性能优化",
    "网络安全",
    "运维",
    "英语",
    "日语",
    "粤语",
    "驾照",
]
SYNONYMS = {
    "ab测试": ["a/b测试", "a/b test", "ab test", "ab实验", "a/b实验"],
    "机器学习": ["machine learning", "ml"],
    "数据可视化": ["tableau", "power bi", "powerbi", "可视化", "看板", "bi"],
    "跨部门沟通": ["跨部门", "沟通协调"],
    "团队管理": ["带团队", "管理团队", "带领", "团队负责人"],
    "英语": ["cet-6", "cet-4", "cet6", "cet4", "ielts", "雅思", "托福", "toefl", "english"],
    "excel": ["数据透视表", "power query", "vba"],
    "python": ["pandas", "numpy"],
    "数据库": ["sql", "mysql", "postgresql", "oracle"],
}
EN_STOP = {"and", "or", "the", "of", "to", "in", "with", "for", "a", "an", "is", "are", "be", "on", "as", "we", "you",
           "our", "at", "by", "etc", "app", "jd", "hr"}  # fmt: skip
EN_PHRASES = ["power bi", "power query", "machine learning", "deep learning", "a/b", "spring boot", "react native",
              "google analytics", "unreal engine", "after effects", "premiere pro"]  # fmt: skip
DEGREES = ["中专", "大专", "本科", "硕士", "博士"]
LEAD = re.compile(
    r"(?:熟悉|掌握|精通|熟练(?:使用|掌握|运用)?|具备|具有|了解|擅长|有)(.{2,30}?)(?:经验|能力|背景|基础|者优先|优先|[，,。；;：:！!\n]|$)"
)


class MatchError(ValueError):
    pass


def read_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError:  # pragma: no cover - environment check
            raise MatchError("读取 .docx 需要 python-docx：pip install python-docx") from None
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
        return "\n".join(parts)
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    raise MatchError(f"{path} 编码无法识别")


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.lower().replace("Ａ", "a").replace("／", "/")).strip()


def extract_terms(line: str) -> list[str]:
    low = norm(line)
    found: list[str] = []
    compact = low.replace(" ", "")
    for term in CN_TERMS:  # Chinese terms first so "AB测试" doesn't also yield a stray "ab"
        t = term.lower()
        if t in compact:
            found.append(t)
            low = low.replace(t, " ")
    for ph in EN_PHRASES:
        if ph in low:
            found.append(ph)
            low = low.replace(ph, " ")
    for tok in re.findall(r"[a-z][a-z0-9+#.\-]*[a-z0-9+#]|[a-z]", low):
        if (tok not in EN_STOP and len(tok) > 1) or tok in ("r", "c"):
            found.append(tok)
    for m in LEAD.finditer(line):
        for piece in re.split(r"[、/和及与或,，]|以及", m.group(1)):
            piece = re.sub(r"^(使用|运用|常用|相关|一定的?|良好的?|较强的?|丰富的?|\d+年以上)|(等|相关|工具|软件|技能|方法|的)+$", "",
                           piece.strip())  # fmt: skip
            if not 2 <= len(piece) <= 12 or re.search(r"[a-zA-Z]", piece):
                continue
            if any(t in piece or piece in t for t in found):
                continue  # already covered by a known term ("机器学习常用算法" → 机器学习)
            found.append(piece.lower())
    out: list[str] = []
    for t in found:
        if t not in out:
            out.append(t)
    return out


def analyze(jd: str, resume: str, today: dt.date) -> dict[str, object]:
    must: list[str] = []
    plus: list[str] = []
    for line in jd.splitlines():
        if not line.strip():
            continue
        bucket = plus if re.search(r"优先|加分|更佳|plus|preferred", line, re.I) else must
        for t in extract_terms(line):
            if t not in must and t not in plus:
                bucket.append(t)
    body = norm(resume).replace(" ", "")

    def hit(term: str) -> bool:
        cands = [term, *SYNONYMS.get(term, [])]
        return any(c.replace(" ", "") in body for c in cands)

    result: dict[str, object] = {
        "必备_已覆盖": [t for t in must if hit(t)],
        "必备_未覆盖": [t for t in must if not hit(t)],
        "加分_已覆盖": [t for t in plus if hit(t)],
        "加分_未覆盖": [t for t in plus if not hit(t)],
    }
    covered = len(result["必备_已覆盖"])  # type: ignore[arg-type]
    result["必备覆盖率"] = round(covered / len(must), 3) if must else None

    need_deg = next(
        (d for d in reversed(DEGREES) if re.search(d + r"(?:及|或)?以上|" + d + "学历", jd)), None
    )
    have_deg = next((d for d in reversed(DEGREES) if d in resume), None)
    if need_deg:
        ok = have_deg is not None and DEGREES.index(have_deg) >= DEGREES.index(need_deg)
        result["学历"] = (
            f"要求{need_deg}及以上，简历{'为' + have_deg if have_deg else '未写明学历'}"
            + ("" if ok else "（不满足或未写明）")
        )
    m = re.search(r"(\d+)\s*(?:年|\+\s*年|年以上)\S{0,6}?(?:经验|工作)", jd)
    if m:
        need = int(m.group(1))
        # Only count dates in the work section so education years don't inflate experience.
        seg = resume.split("工作", 1)[1] if "工作" in resume else resume
        seg = re.split(r"教育|学历", seg, maxsplit=1)[0]
        ds = [
            dt.date(int(y), int(mo), 1)
            for y, mo in re.findall(r"(\d{4})\s*[.\-/年]\s*(\d{1,2})", seg)
        ]
        ds = [d for d in ds if d <= today]
        years = round((today - min(ds)).days / 365.25, 1) if ds else None
        result["工作年限"] = (
            f"要求 {need} 年以上，简历约 {years} 年"
            if years is not None
            else f"要求 {need} 年以上，简历中未找到工作时间"
        )
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("jd")
    parser.add_argument("resume")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--today")
    args = parser.parse_args(argv)
    try:
        today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
        res = analyze(read_text(Path(args.jd)), read_text(Path(args.resume)), today)
    except (OSError, MatchError, ValueError) as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1
    if args.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
        return 0
    rate = res["必备覆盖率"]
    print(
        f"必备关键词覆盖率：{rate:.0%}" if isinstance(rate, float) else "JD 中没有识别到必备关键词"
    )
    for key in ("必备_已覆盖", "必备_未覆盖", "加分_已覆盖", "加分_未覆盖"):
        items = res[key]
        assert isinstance(items, list)
        print(f"- {key.replace('_', '·')}（{len(items)}）：{'、'.join(items) or '无'}")
    for key in ("学历", "工作年限"):
        if key in res:
            print(f"- {key}：{res[key]}")
    if res["必备_未覆盖"]:
        print(
            "建议：确实具备的能力，写进对应经历的条目里（场景 + 做法 + 结果），不要只堆在技能栏；不具备的不要写。"
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
